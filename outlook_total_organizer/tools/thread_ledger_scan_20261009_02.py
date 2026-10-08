# -*- coding: utf-8 -*-
"""
過去スレッド台帳 S1: 走査エンジン + 台帳CSV（画面なし・読み取り専用）

【目的】
3年分（既定 2023-07〜2026-09）の輸出入トラブルのメールスレッドを漏れなく拾うための台帳を作る。
  第1段（軽量・COM）: 全ストア → メールフォルダ（ルート直下を含む）を、ISO書式の Restrict で1回だけ走査し、
                      件名・日時・送信者名・宛先表示名・会話IDなど軽い項目だけをフォルダ単位でキャッシュする。
  評価（純粋関数・COM不要）: ストア横断の重複除去 → スレッド化 → 候補判定 → スコア。
  第2段（候補のみ・COM）: 候補スレッドの各メールを開き、SMTPアドレス・添付名・本文のキーワード照合
                      （語ごとの件数のみ保持。本文は保存しない）を取得し、再評価に反映する。

【候補スレッドの定義（越智さん）】 スレッド内の全メールの送信者/To/CCの和集合で評価する。
  条件A: アンカー（PM・事務担当等）の誰かが関与
  条件B: 越智本人（theme の owner）が From/To/CC のいずれかに関与
  条件C: (i) 関係者（related_parties）の誰かが関与  または  (ii) 件名/本文のキーワード信号
         (ii) は「strong の語が1語でも当たる」または「weak の語が異なる2語以上当たる」
  候補 = A かつ B かつ C

【使い方】（Windows + Outlook 起動済み + pywin32。--evaluate-only はCOM不要）
    python tools\\thread_ledger_scan_20261009_02.py --theme json\\thread_ledger\\themes\\trade_import.json
    python tools\\thread_ledger_scan_20261009_02.py --evaluate-only --suggest --suggest-terms
    python tools\\thread_ledger_scan_20261009_02.py --check-subjects
    python tools\\thread_ledger_scan_20261009_02.py --open <thread_id>      （件名で絞り込んで Outlook に表示）
    python tools\\thread_ledger_scan_20261009_02.py --register-protocol    （Excelの件名リンク ledger:<thread_id> を有効にする。HKCUのみ）
    python tools\\thread_ledger_scan_20261009_02.py --import-judgments <台帳.xlsx|csv>  （Excelで編集した判定を取り込む）
  主なオプション:
    --theme            テーマJSON（既定 json/thread_ledger/themes/trade_import.json。無ければ example を案内）
    --stores           走査/評価するストア名の一部（カンマ区切り）。既定は全ストア
    --skip-stores      除外するストア名の一部（カンマ区切り）
    --from / --to      期間（YYYY-MM）。既定は theme.period、無ければ 2023-07〜2026-09
    --rescan           完了済みフォルダも再取得する
    --evaluate-only    COMを使わず、キャッシュだけから再評価する
    --no-stage2        第2段（COM）を省略する（既にある第2段キャッシュは評価に使う）
    --no-body          第2段で本文照合を省略する
    --stage2-max       第2段の対象通数の上限（既定5000。超えると --yes が無ければ中止）。--yes で上限を無視
    --mask-names       コンソール/ログのストア名・フォルダ名を連番に置換（CSVには元の名前）
    --stage2-scope     第2段の対象: ab(既定: 条件AかつB) / candidate(候補のみ) / either(AまたはB)
    --suggest          新メンバー候補（themeに無い参加者）のCSVを出す
    --suggest-terms    辞書に足す語の候補ランキングCSVを出す
    --check-subjects   正解件名の再現率チェック（既定 json/thread_ledger/calibration_subjects.txt）
    --open             thread_id を指定し、最新メールのあるフォルダを件名で絞り込んで Outlook に表示（失敗時は Display()）
    --entry-id/--store-id  --open の代わりに EntryID/StoreID（16進）を直接指定
    --open-url         ledger:<thread_id> を受け取って開く（プロトコル登録後に Excel から呼ばれる）
    --register-protocol / --unregister-protocol  HKCU の ledger: プロトコルの登録/削除（実行前に内容を表示して確認。--yes で省略）
    --import-judgments Excelで編集した確認結果・メモ・問題の種類を判定JSONに取り込む（許可リスト検証・空は上書きしない・
                       台帳に存在するIDのみ・取込前に .bak_ へ退避）。注意: 台帳CSVをExcelで開くと thread_id が数値化
                       （例 12345e678901）されることがあるため、取込は xlsx で編集したファイルを推奨
    --include-no-anchor  アンカー無し（B∧C成立でA不成立）の参考候補も台帳に出す（列「参考(アンカー無し)」）
    --output-dir       CSV/ログの出力先（既定 mail_reports/）
    --data-dir         台帳データの置き場（既定 json/thread_ledger/）

【初回の試走（推奨）】
  初回は --from/--to を1か月（例: --from 2024-04 --to 2024-04）、--stores で1ストア（例: --stores 2024_Q2）に絞って
  試走し、キャッシュと台帳CSVができること・所要時間を確認してから、全期間・全ストアを夜間に実行する。
  第2段（COM）は対象が多いと時間がかかるため、最初は --stage2-scope candidate で試す。

【安全方針】
  * 読み取り専用。Delete / Move / Save / AddStore / MarkAs / UnRead代入 / Categories代入などは一切呼ばない。
    Display() は --open のときだけ。
  * 第1段では Recipients・Body・Attachments・PropertyAccessor を読まない（遅いため）。
  * コンソール/ログには件数・所要時間・エラー種別だけを出す。件名・氏名・アドレスは出さない
    （例外: CSV、--check-subjects の出力、--suggest の CSV）。
  * 1フォルダ・1項目が失敗しても全体は止めず、エラー種別のみ集計して続行。Ctrl+C では保存して終了。
  * キャッシュ・判定・CSVは .gitignore 済みの json/ と mail_reports/ に置く。本文はキャッシュしない。

【Linuxでのテスト】
  win32com / pythoncom はCOM部分でだけ遅延importする。COM部分はフェイクを注入して（com 引数）テストできる。
"""
import argparse
import bisect
import csv
import gc
import glob
import hashlib
import json
import math
import os
import re
import shutil
import sys
import time
import unicodedata
from datetime import datetime, timedelta
from functools import lru_cache

SCRIPT_VERSION = "20261009_02"

# ============================================================
# 定数
# ============================================================
DEFAULT_FROM = "2023-07"
DEFAULT_TO = "2026-09"
CHECKPOINT_EVERY = 5000          # 第1段: このアイテム数ごとにフォルダのチェックポイントを保存
STAGE2_SAVE_EVERY = 1000         # 第2段: このメール数ごとにキャッシュを保存（終了時・Ctrl+C時にも保存）
STAGE2_EST_SEC = 0.8             # 第2段の1通あたり概算秒（実機試走の実測。開始前の見積り用。実行中は実測で補正して表示）
STAGE1_EST_SEC_PER_ITEM = 0.06   # 第1段の1アイテムあたり概算秒（実機試走の実測 60ms）
STAGE2_MAX_DEFAULT = 5000        # 第2段の対象がこれを超えると --yes が無ければ中止
MAX_ITEM_SKIP = 50               # 読めないアイテムを飛ばす回数の上限（フォルダごと）
ROOT_FOLDER_LABEL = "(ルート直下)"
MAX_FOLDER_DEPTH = 30
OL_MAIL_CLASS = 43               # olMail
MAX_BODY_CHARS = 300000          # 本文照合の上限（長すぎる引用連鎖で遅くならないように）
MAX_SUBJECT_CHARS = 300          # 件名照合の上限
MAX_PATTERN_LEN = 200            # patterns の正規表現の長さ上限
_NESTED_QUANT_RE = re.compile(r"\([^()]*[+*][^()]*\)\s*[+*{]")   # (a+)+ のような極端なネスト量指定子
DEFAULT_CALIBRATION_NAME = "calibration_subjects.txt"
CHECK_SUBJECTS_DEFAULT = "__default__"
THEME_DEFAULT_NAME = "trade_import.json"
THEME_EXAMPLE_NAME = "thread_ledger_theme.example.json"
MAX_PARTICIPANTS_IN_CSV = 60

# PropertyAccessor のプロパティURL（第2段のみ）
PR_SENDER_SMTP = "http://schemas.microsoft.com/mapi/proptag/0x5D01001F"
PR_RECIP_SMTP = "http://schemas.microsoft.com/mapi/proptag/0x39FE001E"
PR_INTERNET_MESSAGE_ID = "http://schemas.microsoft.com/mapi/proptag/0x1035001F"

# themeに無いときの既定（除外フォルダ）
DEFAULT_EXCLUDE_FOLDER_NAMES = [
    "削除済みアイテム", "Deleted Items", "迷惑メール", "Junk Email", "Junk E-mail",
    "RSS フィード", "RSS Feeds", "RSS Subscriptions", "検索フォルダー", "Search Folders",
    "同期の問題", "Sync Issues", "会話の履歴", "Conversation History",
    "下書き", "Drafts", "送信トレイ", "Outbox",
]
DEFAULT_EXCLUDE_FOLDER_PREFIXES = ["note-", "zenn", "qiita", "newsletter", "ニュースレター", "sync"]
DEFAULT_EXCLUDE_STORES = []

DEFAULT_SCORE_WEIGHTS = {
    "subject_multiplier": 2.0,   # 件名ヒットは本文ヒットの何倍重いか
    "weak_multiplier": 0.5,      # weak の語の重み係数
    "anchor_each": 3.0,          # アンカー1人あたり
    "related_each": 2.0,         # 関係者1件あたり（関係者側に weight があればそちら）
    "reply_every": 5,            # 何通ごとに加点するか
    "reply_points": 1.0,         # 上記ごとの加点
    "reply_cap": 5.0,            # 返信加点の上限
}

# 件名の前置語（NFKC・小文字化後に先頭から繰り返し除去）
_PREFIX_RE = re.compile(
    r"^\s*(?:"
    r"(?:re|fw|fwd|aw|wg|sv|rv|antw)(?:\[\d+\])?\s*:"
    r"|転送\s*:|返信\s*:|回覆\s*:"
    r"|\[(?:外部|external|ext|caution|注意)\]"
    r"|【外部】"
    r"|\*\*\s*internal only\s*\*\*"
    r")\s*",
    re.IGNORECASE,
)
EMAIL_RE = re.compile(r"[\w.+\-']+@[\w\-]+(?:\.[\w\-]+)+", re.UNICODE)
MAIL_ADDR_RE = EMAIL_RE
_HONORIFIC_RE = re.compile(r"(?:さん|様|殿)$")
_TOKEN_RE = re.compile(r"[ァ-ヴー]{3,}|[一-龥々]{2,}|[a-z][a-z0-9\-/]{2,}")
STOPWORDS = {
    "the", "and", "for", "with", "from", "please", "thank", "thanks", "regarding", "this", "that",
    "you", "your", "are", "not", "have", "will", "can", "may", "info", "mail", "email", "message",
    "お願い", "ください", "いたします", "よろしく", "お世話", "ありがとう", "ございます", "確認",
    "連絡", "件名", "添付", "資料", "対応", "依頼", "共有", "について",
}

LEDGER_THREAD_ID_RE = re.compile(r"^[A-Za-z0-9_-]{4,64}\Z")     # $ は末尾の改行を許すので \Z を使う
HEX_ID_RE = re.compile(r"^[0-9A-Fa-f]{16,2048}\Z")
ALLOWED_VERDICTS = ("本物", "違う", "保留", "")
PROTOCOL_NAME = "ledger"
NO_ANCHOR_COLUMN = "参考(アンカー無し)"
EXPLORER_SETTLE_SEC = 0.4        # CurrentFolder 切替後、Search の前に待つ秒数
SELECT_RETRY_SEC = 0.3           # 選択（IsItemSelectableInView/AddToSelection）のリトライ間隔
SELECT_RETRIES = 5               # 選択のリトライ回数の上限
IMPORT_MAX_BYTES = 50 * 1024 * 1024   # --import-judgments のファイルサイズ上限
IMPORT_MAX_ROWS = 200000              # --import-judgments の行数上限
LEDGER_COLUMNS = [
    "thread_id", "件名", "最初の日時", "最後の日時", "メール数", "参加者", "アンカー", "関係者(ラベル)",
    "キーワードカテゴリ(語数)", "条件A/B/C内訳", "スコア", "スコア内訳", "取得元(ストア/フォルダ)",
    "最新メールEntryID", "最新メールStoreID", "確認結果(本物/違う/保留)", "メモ", "問題の種類", "theme_version", NO_ANCHOR_COLUMN,
]
JUDGMENT_KEYS = ["確認結果", "メモ", "問題の種類"]


# ============================================================
# 純粋関数: 日付・月
# ============================================================
def parse_ym(text):
    """'YYYY-MM'（または 'YYYY/MM'）を (年, 月) に変換する。不正なら ValueError。"""
    if not isinstance(text, str):
        raise ValueError(f"年月の形式が不正です: {text!r}")
    m = re.fullmatch(r"\s*(\d{4})[-/](\d{1,2})\s*", text)
    if not m:
        raise ValueError(f"年月は YYYY-MM の形式で指定してください: {text!r}")
    y, mo = int(m.group(1)), int(m.group(2))
    if not (1 <= mo <= 12) or not (1900 <= y <= 2200):
        raise ValueError(f"年月の範囲が不正です: {text!r}")
    return (y, mo)


def ym_label(ym):
    return f"{ym[0]}-{ym[1]:02d}"


def month_range(start, end):
    """(年,月) の start から end まで（両端含む）を古い順のリストで返す。start>end は ValueError。"""
    if tuple(start) > tuple(end):
        raise ValueError("開始月が終了月より後になっています")
    out = []
    y, m = start
    while (y, m) <= tuple(end):
        out.append((y, m))
        m += 1
        if m == 13:
            m = 1
            y += 1
    return out


def month_bounds(year, month):
    """暦月の [開始, 翌月開始) を datetime で返す。"""
    start = datetime(year, month, 1)
    end = datetime(year + 1, 1, 1) if month == 12 else datetime(year, month + 1, 1)
    return start, end


def period_bounds(from_ym, to_ym):
    """期間全体の [開始, 終了(翌月1日)) を返す。"""
    return month_bounds(*from_ym)[0], month_bounds(*to_ym)[1]


def build_or_filter(start_dt, end_dt):
    """([ReceivedTime] 範囲) OR ([SentOn] 範囲)。送信済み系（ReceivedTimeが未設定）の取りこぼし防止。"""
    s = start_dt.strftime("%Y-%m-%d %H:%M")
    e = end_dt.strftime("%Y-%m-%d %H:%M")
    return (f"([ReceivedTime] >= '{s}' AND [ReceivedTime] < '{e}') OR "
            f"([SentOn] >= '{s}' AND [SentOn] < '{e}')")


def build_restrict_filter(field, start_dt, end_dt):
    """Restrict用: ISO書式 'YYYY-MM-DD HH:MM' のみ（実機診断で安全と確認済み。DASLの日付条件は使わない）。"""
    s = start_dt.strftime("%Y-%m-%d %H:%M")
    e = end_dt.strftime("%Y-%m-%d %H:%M")
    return f"[{field}] >= '{s}' AND [{field}] < '{e}'"


def dt_to_str(dt):
    return dt.strftime("%Y-%m-%dT%H:%M:%S") if dt else ""


def parse_dt(text):
    """キャッシュの日時文字列 -> datetime。不正は None。"""
    try:
        return datetime.strptime(str(text)[:19], "%Y-%m-%dT%H:%M:%S")
    except Exception:
        return None


def to_naive_dt(value):
    """COMの日時（pywintypes.datetime等）を naive datetime にする。未設定（年4500等）・不正は None。"""
    try:
        if value is None:
            return None
        if isinstance(value, datetime):
            dt = datetime(value.year, value.month, value.day, value.hour, value.minute, value.second)
        elif isinstance(value, str):
            dt = parse_dt(value)
            if dt is None:
                return None
        else:
            return None
        if dt.year < 1900 or dt.year > 3000:
            return None
        return dt
    except Exception:
        return None


def format_duration(seconds):
    if seconds is None:
        return "-"
    if seconds < 10:
        return f"{seconds:.3f}秒"
    if seconds < 60:
        return f"{seconds:.1f}秒"
    if seconds < 3600:
        m, s = divmod(int(round(seconds)), 60)
        return f"{m}分{s:02d}秒"
    h, rest = divmod(int(round(seconds)), 3600)
    return f"{h}時間{rest // 60:02d}分"


def error_kind(exc):
    """例外から「種別」だけを返す（メッセージは件名等を含みうるので出さない）。"""
    name = type(exc).__name__
    try:
        args = getattr(exc, "args", ())
        if args and isinstance(args[0], int):
            return f"{name}(0x{args[0] & 0xFFFFFFFF:08X})"
    except Exception:
        pass
    return name


def summarize_errors(counts):
    """エラー種別 -> 件数 の辞書を、多い順の文字列リストにする。"""
    return [f"{k}: {v}件" for k, v in sorted(counts.items(), key=lambda kv: (-kv[1], kv[0]))]


def add_error(counts, kind):
    counts[kind] = counts.get(kind, 0) + 1


def mask_emails(text):
    """メールアドレス形式を <MAIL> に置換する（ストア名がアドレスのため、コンソール表示用）。"""
    if not text:
        return text or ""
    return MAIL_ADDR_RE.sub("<MAIL>", str(text))


# ============================================================
# 純粋関数: 文字列の正規化
# ============================================================
def nfkc(text):
    return unicodedata.normalize("NFKC", text if isinstance(text, str) else ("" if text is None else str(text)))


def norm_space(text):
    return re.sub(r"\s+", " ", nfkc(text)).strip()


def normalize_subject(subject):
    """件名正規化: 全角半角統一・空白圧縮・大文字小文字無視、RE:/FW:/FWD:/AW:/WG:/転送:/返信:/[外部]/** internal only ** 等の
    接頭語を繰り返し除去する。"""
    t = norm_space(subject).lower()
    for _ in range(60):
        new = _PREFIX_RE.sub("", t, count=1)
        if new == t:
            break
        t = new
    return re.sub(r"\s+", " ", t).strip()


def _name_tokens(name):
    t = nfkc(name).lower()
    t = t.replace(".", "")
    t = re.sub(r"[,、，;/()\[\]<>\"'・·]", " ", t)
    tokens = []
    for tok in t.split():
        if tok in ("san", "-san"):
            continue
        tok = re.sub(r"-san$", "", tok)
        if len(tok) > 1:
            tok = _HONORIFIC_RE.sub("", tok)
        if tok:
            tokens.append(tok)
    return tokens


@lru_cache(maxsize=None)
def name_key(name):
    """氏名の正規化キー: (トークン列, 順方向文字列, 逆順文字列, 順方向の空白なし, 逆順の空白なし)。
    大文字小文字・全角半角・空白・「姓, 名」「名 姓」・敬称（さん/様/殿/san）を吸収する。"""
    tokens = tuple(_name_tokens(name))
    fwd = " ".join(tokens)
    rev = " ".join(reversed(tokens))
    return tokens, fwd, rev, fwd.replace(" ", ""), rev.replace(" ", "")


def _is_ascii(s):
    return all(ord(c) < 128 for c in s)


def _all_non_ascii(s):
    return bool(s) and all(ord(c) >= 128 for c in s)


def _ascii_boundary_find(needle, hay):
    """ASCIIの語を、語頭・語末の両方の境界つきで探す（区切り・敬称は英数字でないので境界扱い）。"""
    pat = r"(?<![a-z0-9])" + re.escape(needle) + r"(?![a-z0-9])"
    return re.search(pat, hay) is not None


@lru_cache(maxsize=500000)
def name_matches(alias, name):
    """氏名エイリアスが表示名に一致するか（正規化後。メモ化）。
    ASCIIの別名: 長さに関わらず語頭・語末の境界を要求する（'Ono' が 'Onozuka'、'Yuki' が 'Yukiko' に当たらない）。
      別名が2語以上なら、語順が違っても全ての語がトークンとして含まれれば一致。
    日本語などの別名: トークン単位の一致が基本（連結文字列の部分一致にはしない）。
      1語の別名は表示名のトークンと完全一致（'田中' は '田中 花子'・'田中さん' に一致、'山田中' には不一致）。
      加えて、表示名が全て非ASCIIで連結が別名で『始まり』残りが0〜4文字なら一致（姓だけの別名がスペース無しの
      フルネームに当たる。'越智'→'越智太郎'）。後ろが1〜4文字続く別姓（'木村'→'木村戸'）も一致してしまうリスクがある。
      2語以上は、連結が表示名の連結（順方向/逆方向）と一致するか、全トークンが含まれれば一致。"""
    ak, nk = name_key(alias), name_key(name)
    if not ak[0] or not nk[0]:
        return False
    nts = set(nk[0])
    if _is_ascii(ak[1]):
        for av in {ak[1], ak[2]}:
            for hay in (nk[1], nk[2]):
                if _ascii_boundary_find(av, hay):
                    return True
        return len(ak[0]) >= 2 and all(at in nts for at in ak[0])
    if len(ak[0]) == 1:
        a = ak[0][0]
        if a in nts or a in (nk[3], nk[4]):       # トークン一致、またはフルネーム全体の一致
            return True
        # 姓だけの別名: 表示名が全て非ASCIIで、連結(空白除去)が別名で「始まり」、残りが0〜4文字（名の長さ）なら一致。
        # 例: '越智'→'越智太郎'、'木村'→'木村美穂'。'山田中'・'山木村'(別名'木村')のように別名で始まらないものは不一致。
        # 既知のリスク: 後ろに1〜4文字続く別姓（'木村'→'木村戸'）は姓+名と区別できず一致してしまう。
        return (len(a) >= 2 and _all_non_ascii(nk[3]) and nk[3].startswith(a) and len(nk[3]) - len(a) <= 4)
    if ak[3] in (nk[3], nk[4]):
        return True
    return all(at in nts for at in ak[0])


def split_names(field):
    """To/CC の表示名文字列（; 区切り）を氏名リストにする。<...> のアドレス部分は除く。"""
    out = []
    for part in re.split(r"[;\n]", field or ""):
        nm = re.sub(r"<[^>]*>", "", part).strip().strip("'\"").strip()
        if nm:
            out.append(nm)
    return out


def extract_emails(text):
    return [m.lower() for m in EMAIL_RE.findall(text or "")]


def smtp_like(raw):
    """SenderEmailAddress の生値から SMTP アドレスだけを返す（EXのX500形式は空文字）。"""
    if not raw or not isinstance(raw, str):
        return ""
    if raw.startswith("/") or "@" not in raw:
        return ""
    m = EMAIL_RE.search(raw)
    return m.group(0).lower() if m else ""


def email_domain(addr):
    return addr.rsplit("@", 1)[-1].lower()


def domain_matches(domain, pattern):
    d = (pattern or "").lower().strip().lstrip("@").lstrip(".")
    return bool(d) and (domain == d or domain.endswith("." + d))


# ============================================================
# 純粋関数: テーマ
# ============================================================
class ThemeError(Exception):
    pass


def _as_list(v):
    if v is None:
        return []
    if isinstance(v, (list, tuple)):
        return [x for x in v if isinstance(x, str) and x.strip()]
    if isinstance(v, str) and v.strip():
        return [v]
    return []


def _norm_entity(d, kind_default="person"):
    d = d if isinstance(d, dict) else {}
    return {
        "label": str(d.get("label") or "").strip(),
        "kind": str(d.get("kind") or kind_default),
        "name_aliases": _as_list(d.get("name_aliases")),
        "emails": [e.lower().strip() for e in _as_list(d.get("emails"))],
        "domains": [x.lower().strip() for x in _as_list(d.get("domains"))],
        "weight": d.get("weight"),
    }


def build_term_regex(term_norm):
    """語の照合用正規表現（正規化・小文字化済みの語に対して）。
    ASCIIの語は語境界を要求（右側は複数形の s/es を許容）。日本語などは単純な部分一致。"""
    esc = re.escape(term_norm)
    if term_norm and _is_ascii(term_norm):
        left = r"(?<![a-z0-9])" if term_norm[0].isalnum() else ""
        right = r"(?:s|es)?(?![a-z0-9])" if term_norm[-1].isalnum() else ""
        return re.compile(left + esc + right)
    return re.compile(esc)


def normalize_theme(data, name="theme"):
    """テーマJSON(dict)を検証し、既定値を補い、照合用テーブルを付けて返す。不正は ThemeError。"""
    if not isinstance(data, dict):
        raise ThemeError("テーマJSONの最上位がオブジェクトではありません")
    owner = _norm_entity(data.get("owner"))
    owner["label"] = owner["label"] or "owner"
    if not (owner["name_aliases"] or owner["emails"]):
        raise ThemeError("owner に name_aliases または emails が必要です")
    anchors = [_norm_entity(a) for a in (data.get("anchors") or [])]
    anchors = [a for a in anchors if a["name_aliases"] or a["emails"] or a["domains"]]
    if not anchors:
        raise ThemeError("anchors が空です（条件Aを判定できません）")
    for i, a in enumerate(anchors, 1):
        a["label"] = a["label"] or f"anchor{i}"
    related = [_norm_entity(r) for r in (data.get("related_parties") or [])]
    related = [r for r in related if r["name_aliases"] or r["emails"] or r["domains"]]
    for i, r in enumerate(related, 1):
        r["label"] = r["label"] or f"related{i}"

    hit_table = {}      # キー(語 / 'pat:'+正規表現) -> {strong, category, weight, regex, kind}
    keywords = data.get("keywords") or {}
    if not isinstance(keywords, dict):
        raise ThemeError("keywords はオブジェクトで指定してください")
    cat_weights = {}
    for cat, spec in keywords.items():
        if str(cat).startswith("_") or not isinstance(spec, dict):
            continue
        try:
            weight = float(spec.get("weight", 1))
        except Exception:
            raise ThemeError(f"keywords.{cat}.weight が数値ではありません")
        cat_weights[cat] = weight
        for strong, key in ((True, "strong"), (False, "weak")):
            for term in _as_list(spec.get(key)):
                tn = norm_space(term).lower()
                if not tn:
                    continue
                old = hit_table.get(tn)
                if old is not None and (old["strong"] or not strong):
                    continue   # 同じ語が複数ある場合は strong を優先
                hit_table[tn] = {"strong": strong, "category": cat, "weight": weight,
                                 "regex": build_term_regex(tn), "kind": "term", "label": tn}
    patterns = []
    for p in data.get("patterns") or []:
        if not isinstance(p, dict) or not p.get("regex"):
            continue
        if len(str(p["regex"])) > MAX_PATTERN_LEN:
            raise ThemeError(f"patterns の正規表現が長すぎます（{MAX_PATTERN_LEN}字まで）")
        if _NESTED_QUANT_RE.search(str(p["regex"])):
            raise ThemeError("patterns の正規表現にネストした量指定子（(a+)+ 等）は使えません")
        try:
            rx = re.compile(p["regex"], re.IGNORECASE)
        except re.error:
            raise ThemeError(f"patterns の正規表現が不正です: {p.get('regex')!r}")
        cat = str(p.get("category") or "パターン")
        key = "pat:" + p["regex"]
        hit_table[key] = {"strong": bool(p.get("strong")), "category": cat,
                          "weight": cat_weights.get(cat, 1.0), "regex": rx, "kind": "pattern", "label": key}
        patterns.append({"category": cat, "regex": p["regex"], "strong": bool(p.get("strong"))})

    sw = dict(DEFAULT_SCORE_WEIGHTS)
    for k, v in (data.get("score_weights") or {}).items():
        if k.startswith("_"):
            continue
        try:
            sw[k] = float(v)
        except Exception:
            raise ThemeError(f"score_weights.{k} が数値ではありません")
    period = data.get("period") or {}
    p_from = p_to = None
    try:
        if period.get("from"):
            p_from = parse_ym(period["from"])
        if period.get("to"):
            p_to = parse_ym(period["to"])
    except ValueError as e:
        raise ThemeError(f"period が不正です: {e}")

    def lst(key, default):
        v = data.get(key)
        return _as_list(v) if v is not None else list(default)

    dict_src = json.dumps({"k": {c: keywords[c] for c in keywords if not str(c).startswith("_")},
                           "p": patterns}, sort_keys=True, ensure_ascii=False)
    return {
        "name": name,
        "theme_version": str(data.get("theme_version") or ""),
        "owner": owner, "anchors": anchors, "related_parties": related,
        "hit_table": hit_table, "patterns": patterns,
        "exclude_folder_names": lst("exclude_folder_names", DEFAULT_EXCLUDE_FOLDER_NAMES),
        "exclude_folder_prefixes": lst("exclude_folder_prefixes", DEFAULT_EXCLUDE_FOLDER_PREFIXES),
        "exclude_stores": lst("exclude_stores", DEFAULT_EXCLUDE_STORES),
        "score_weights": sw,
        "period_from": p_from, "period_to": p_to,
        "dict_hash": hashlib.sha1(dict_src.encode("utf-8")).hexdigest()[:12],
    }


def load_theme(path):
    """テーマJSONファイルを読み込む。読めない/不正は ThemeError。"""
    try:
        with open(path, "r", encoding="utf-8-sig") as fh:
            data = json.load(fh)
    except FileNotFoundError:
        raise ThemeError("テーマファイルがありません")
    except Exception as e:
        raise ThemeError(f"テーマJSONを読めません: {error_kind(e)}")
    name = os.path.splitext(os.path.basename(path))[0]
    return normalize_theme(data, name)


def folder_excluded(name, exclude_names, exclude_prefixes):
    """フォルダ名が除外対象か（名前の完全一致 または 前置詞。大文字小文字・全角半角を無視）。"""
    low = norm_space(name).lower()
    if not low:
        return False
    names = {norm_space(n).lower() for n in exclude_names}
    if low in names:
        return True
    return any(low.startswith(norm_space(p).lower()) for p in exclude_prefixes if p)


def store_selected(display_name, only_words, skip_words, exclude_stores=()):
    """ストア名（部分一致・小文字）で対象を選ぶ。only_words が空なら全ストア。"""
    low = (display_name or "").lower()
    if any(w and w in low for w in skip_words):
        return False
    if any(norm_space(w).lower() in low for w in exclude_stores if w):
        return False
    if only_words and not any(w and w in low for w in only_words):
        return False
    return True


def split_words(text):
    return [w.strip().lower() for w in (text or "").split(",") if w.strip()]


# ============================================================
# 純粋関数: 照合・候補判定・スコア
# ============================================================
def entity_matches(entity, names, emails, domains=None):
    """entity（owner/anchor/related）が、表示名の集合・アドレス集合に関与するか。"""
    if domains is None:
        domains = {email_domain(e) for e in emails}
    for a in entity["name_aliases"]:
        for n in names:
            if name_matches(a, n):
                return True
    for e in entity["emails"]:
        if e in emails:
            return True
    for d in entity["domains"]:
        if any(domain_matches(x, d) for x in domains):
            return True
    return False


def scan_text_hits(text, hit_table):
    """テキストに対する辞書照合。キー -> 件数 の辞書（0件は含めない）。本文自体は返さない。"""
    out = {}
    if not text:
        return out
    norm = norm_space(text)[:MAX_BODY_CHARS]
    low = norm.lower()
    for key, ent in hit_table.items():
        target = norm if ent["kind"] == "pattern" else low
        n = len(ent["regex"].findall(target))
        if n:
            out[key] = n
    return out


def keyword_signal(hit_keys, hit_table):
    """トピック信号 C(ii): strong が1語以上、または weak が異なる2語以上。
    戻り値: {'ok': bool, 'strong': [キー], 'weak': [キー]}"""
    strong = sorted(k for k in hit_keys if k in hit_table and hit_table[k]["strong"])
    weak = sorted(k for k in hit_keys if k in hit_table and not hit_table[k]["strong"])
    return {"ok": len(strong) >= 1 or len(weak) >= 2, "strong": strong, "weak": weak}


def fmt_num(x):
    x = round(float(x), 1)
    return str(int(x)) if x == int(x) else f"{x:.1f}"


def compute_score(anchor_labels, related, subj_hits, all_hits, hit_table, mail_count, sw):
    """単純加算スコアと内訳文字列。
    related: [(ラベル, 点)]。subj_hits: 件名で当たったキー集合。all_hits: 件名+本文で当たったキー集合。"""
    parts = []
    total = 0.0
    if anchor_labels:
        pts = sw["anchor_each"] * len(anchor_labels)
        total += pts
        parts.append(f"アンカー+{fmt_num(pts)}" if len(anchor_labels) == 1
                     else f"アンカー({len(anchor_labels)}名)+{fmt_num(pts)}")
    for label, pts in related:
        total += pts
        parts.append(f"{label}+{fmt_num(pts)}")
    cat_pts, cat_n = {}, {}
    for key in all_hits:
        ent = hit_table.get(key)
        if not ent:
            continue
        pts = ent["weight"]
        if key in subj_hits:
            pts *= sw["subject_multiplier"]
        if not ent["strong"]:
            pts *= sw["weak_multiplier"]
        cat_pts[ent["category"]] = cat_pts.get(ent["category"], 0.0) + pts
        cat_n[ent["category"]] = cat_n.get(ent["category"], 0) + 1
    for cat in sorted(cat_pts, key=lambda c: (-cat_pts[c], c)):
        total += cat_pts[cat]
        parts.append(f"{cat}({cat_n[cat]}語)+{fmt_num(cat_pts[cat])}")
    if sw["reply_every"] > 0 and mail_count >= sw["reply_every"]:
        pts = min(sw["reply_cap"], (mail_count // int(sw["reply_every"])) * sw["reply_points"])
        if pts > 0:
            total += pts
            parts.append(f"返信{mail_count}通+{fmt_num(pts)}")
    return round(total, 1), ", ".join(parts)


# ============================================================
# 純粋関数: メール・重複除去・スレッド化
# ============================================================
def mail_people(mail, st2):
    """1通の関与者。戻り値: (表示名リスト, アドレス集合, [(表示名, アドレス)])。
    第1段は表示名とSMTP形式のSenderEmailAddressのみ。第2段があればSMTPを加える。"""
    names, emails, pairs = [], set(), []
    sn = mail.get("sender") or ""
    se = smtp_like(mail.get("sender_email"))
    if sn:
        names.append(sn)
    if se:
        emails.add(se)
    for nm in split_names(mail.get("to")) + split_names(mail.get("cc")):
        names.append(nm)
        pairs.append((nm, ""))
    emails.update(extract_emails(mail.get("to")))
    emails.update(extract_emails(mail.get("cc")))
    sender_addr = se
    if st2:
        s = (st2.get("s") or "").lower().strip()
        if s:
            emails.add(s)
            sender_addr = sender_addr or s
        for r in (st2.get("to") or []) + (st2.get("cc") or []):
            n, a = r.get("n") or "", (r.get("a") or "").lower().strip()
            if n:
                names.append(n)
            if a:
                emails.add(a)
            if n or a:
                pairs.append((n, a))
    if sn or sender_addr:
        pairs.append((sn, sender_addr))
    return names, emails, pairs


def stage2_key(sid, eid):
    return hashlib.sha1(f"{sid}|{eid}".encode("utf-8", "replace")).hexdigest()[:20]


def mail_stage2(mail, stage2):
    if not stage2:
        return None
    loc = mail["locs"][0]
    return stage2.get(stage2_key(loc["sid"], loc["eid"]))


def cache_ident(c):
    """ストアの識別子。StoreID（PSTはFilePathも併用）。無ければ表示名。store_index は使わない。"""
    sid = c.get("store_id") or ""
    fp = (c.get("file_path") or "").lower()
    if sid or fp:
        return f"{sid}|{fp}"
    return "name:" + str(c.get("store_name", ""))


def assign_store_labels(entries):
    """表示名が重複するストアに ' #2' 等を付けて一意化する。entries: [(識別子, 表示名, 順序)] -> {識別子: ラベル}"""
    seen, labels = {}, {}
    for ident, name, _order in sorted(entries, key=lambda e: (e[2], e[0])):
        if ident in labels:
            continue
        n = seen.get(name, 0) + 1
        seen[name] = n
        labels[ident] = name if n == 1 else f"{name} #{n}"
    return labels


def store_ok(name, label, only_words, skip_words, exclude_stores=()):
    """ストアの選択（表示名またはラベルの部分一致）。--stores / --skip-stores / exclude_stores 用。"""
    if not store_selected(name, [], skip_words, exclude_stores) or not store_selected(label, [], skip_words, []):
        return False
    if only_words:
        low = (name + "\n" + label).lower()
        return any(w and w in low for w in only_words)
    return True


def select_caches(caches, only_words, skip_words, exclude_stores, exclude_names, exclude_prefixes):
    """キャッシュ(フォルダ単位)を選ぶ。(ストア識別子,フォルダパス)が同じものは scanned_at が新しい方だけ残す。
    表示名が同じで別StoreIDのストアは両方残し、ラベルで区別する（store_label を付けたコピーを返す）。
    除外ストア/除外フォルダ名（パスの各階層）/ストア選択を適用する。"""
    best = {}
    for c in caches:
        key = (cache_ident(c), c.get("folder_path", ""))
        old = best.get(key)
        if old is None or str(c.get("scanned_at", "")) >= str(old.get("scanned_at", "")):
            best[key] = c
    ents = {}
    for c in best.values():
        i, o = cache_ident(c), c.get("store_index", 0)
        if i not in ents or o < ents[i][1]:
            ents[i] = (c.get("store_name", ""), o)
    labels = assign_store_labels([(i, n, o) for i, (n, o) in ents.items()])
    out = []
    for c in best.values():
        label = labels[cache_ident(c)]
        if not store_ok(c.get("store_name", ""), label, only_words, skip_words, exclude_stores):
            continue
        segs = [x for x in (c.get("folder_path") or "").split("\\") if x]
        if any(folder_excluded(x, exclude_names, exclude_prefixes) for x in segs):
            continue
        c2 = dict(c)
        c2["store_label"] = label
        out.append(c2)
    out.sort(key=lambda c: (c.get("store_index", 0), c["store_label"], c.get("folder_path", "")))
    return out


def union_names(a, b):
    """To/CC の表示名文字列を和集合にする（'; ' 連結）。"""
    out = split_names(a)
    seen = {name_key(x)[3] or x.lower() for x in out}
    for x in split_names(b):
        k = name_key(x)[3] or x.lower()
        if k not in seen:
            seen.add(k)
            out.append(x)
    return "; ".join(out)


def merge_caches_to_mails(caches, from_ym, to_ym):
    """フォルダ単位のキャッシュを、ストア横断で重複除去した1通ごとのメールにまとめる。
    重複除去キー: (ReceivedTimeを秒まで, 正規化件名, 正規化SenderName)。統合するのはストア間だけで、
    同一ストア内でEntryIDが異なるメールは別メールとして残す。統合時の To/CC は和集合にする。
    取得元(ストア/フォルダ)は全て記録する。
    戻り値: (メールのリスト(時刻順), 統計)。統計: raw_by_store_month, raw_total, unique_total, cross_store, same_store, dup_by_month"""
    start, end = period_bounds(from_ym, to_ym)
    by_key = {}
    order = []
    raw_by_sm = {}
    stats = {"raw_total": 0, "cross_store": 0, "same_store": 0, "dup_by_month": {}, "no_time": 0}
    for c in caches:
        ident = cache_ident(c)
        sname = c.get("store_label") or c.get("store_name", "")
        sidx, fpath = c.get("store_index", 0), c.get("folder_path", "")
        sid = c.get("store_id", "")
        for r in c.get("records") or []:
            t = parse_dt(r.get("t"))
            if t is None:
                stats["no_time"] += 1
                continue
            if t < start or t >= end:
                continue
            stats["raw_total"] += 1
            ym = (t.year, t.month)
            raw_by_sm[(sname, ym)] = raw_by_sm.get((sname, ym), 0) + 1
            nsubj = normalize_subject(r.get("s"))
            key = (r.get("t", "")[:19], nsubj, norm_space(r.get("n")).lower())
            loc = {"store": sname, "ident": ident, "sidx": sidx, "folder": fpath, "eid": r.get("e", ""), "sid": sid}
            cands = by_key.setdefault(key, [])
            target = None
            for m in cands:
                if not any(l["ident"] == ident for l in m["locs"]):
                    target = m
                    break
            if target is None:
                if cands:
                    stats["same_store"] += 1    # 同一ストアの似たメール: 別メールとして残す
                m = {"key": key, "t": t, "subj": r.get("s") or "", "nsubj": nsubj,
                     "sender": r.get("n") or "", "sender_email": r.get("a") or "",
                     "to": r.get("to") or "", "cc": r.get("cc") or "", "cid": r.get("c") or "",
                     "locs": [loc]}
                cands.append(m)
                order.append(m)
            else:
                stats["cross_store"] += 1
                stats["dup_by_month"][ym] = stats["dup_by_month"].get(ym, 0) + 1
                target["locs"].append(loc)
                for fld, src in (("sender_email", "a"), ("cid", "c")):
                    if not target[fld] and r.get(src):
                        target[fld] = r.get(src)
                for fld in ("to", "cc"):
                    if r.get(fld) and r.get(fld) != target[fld]:
                        target[fld] = union_names(target[fld], r.get(fld))
    order.sort(key=lambda m: m["t"])
    stats["unique_total"] = len(order)
    stats["raw_by_store_month"] = raw_by_sm
    return order, stats


GENERIC_SUBJECTS = {
    "お疲れ様です", "お疲れさまです", "ご連絡", "お願い", "確認", "ありがとうございます", "hello", "thanks",
    "thank you", "update", "meeting", "fyi", "test", "報告", "質問", "お世話になっております", "ご確認ください",
    "ご相談", "相談", "連絡", "情報共有", "please check", "question", "request", "hi", "no subject",
    "(件名なし)", "件名なし",
}
GENERIC_MIN_LEN = 8


def is_generic_subject(nsubj):
    """短い（8文字未満）/一般語だけの正規化件名。スレッド束ね・部分一致に使わない。"""
    n = (nsubj or "").strip()
    return len(n) < GENERIC_MIN_LEN or n in GENERIC_SUBJECTS


def thread_key_of(mail):
    """ConversationID。空なら正規化件名。ただし短い/汎用件名は束ねず単独スレッドにする。"""
    cid = (mail.get("cid") or "").strip()
    if cid:
        return "c:" + cid
    ns = mail.get("nsubj")
    if ns and not is_generic_subject(ns):
        return "s:" + ns
    loc = mail["locs"][0]
    return "m:" + hashlib.sha1(repr((mail["key"], loc["sid"], loc["eid"])).encode("utf-8", "replace")).hexdigest()


def thread_id_of(key):
    return hashlib.sha1(key.encode("utf-8", "replace")).hexdigest()[:12]


def build_threads(mails):
    """ConversationID（空なら正規化件名）でグループ化する。mails は時刻順であること。"""
    groups = {}
    for m in mails:
        groups.setdefault(thread_key_of(m), []).append(m)
    threads = []
    for key, ms in groups.items():
        ms.sort(key=lambda m: m["t"])
        threads.append({"id": thread_id_of(key), "key": key, "mails": ms,
                        "subject": ms[-1]["subj"], "first": ms[0]["t"], "last": ms[-1]["t"], "count": len(ms)})
    threads.sort(key=lambda t: t["last"])
    return threads


def thread_people(thread, stage2):
    """スレッド内の全メールの送信者/To/CCの和集合。戻り値: (表示名リスト(重複除去), アドレス集合, ペア集合)。"""
    names, seen, emails, pairs = [], set(), set(), set()
    for m in thread["mails"]:
        n, e, p = mail_people(m, mail_stage2(m, stage2))
        for x in n:
            k = name_key(x)[3]
            if k and k not in seen:
                seen.add(k)
                names.append(x)
        emails |= e
        pairs.update(p)
    return names, emails, pairs


def evaluate_thread(thread, theme, stage2, memo=None):
    """スレッド1件の候補判定とスコア。"""
    memo = memo if memo is not None else {}
    table = theme["hit_table"]
    names, emails, _pairs = thread_people(thread, stage2)
    domains = {email_domain(e) for e in emails}
    anchors = [a["label"] for a in theme["anchors"] if entity_matches(a, names, emails, domains)]
    owner_in = entity_matches(theme["owner"], names, emails, domains)
    rel_hit = [r for r in theme["related_parties"] if entity_matches(r, names, emails, domains)]

    subj_hits, body_hits = set(), set()
    seen_subj = set()
    for m in thread["mails"]:
        ns = m["nsubj"]
        if ns not in seen_subj:
            seen_subj.add(ns)
            if ns not in memo:
                memo[ns] = set(scan_text_hits(ns[:MAX_SUBJECT_CHARS], table))
            subj_hits |= memo[ns]
        st2 = mail_stage2(m, stage2)
        if st2 and st2.get("bh") == theme["dict_hash"]:    # 辞書が変わった後の古い本文ヒットは使わない
            body_hits |= {k for k in (st2.get("hits") or {}) if k in table}
    all_hits = subj_hits | body_hits
    sig = keyword_signal(all_hits, table)
    a_ok, b_ok = bool(anchors), bool(owner_in)
    c_i, c_ii = bool(rel_hit), sig["ok"]
    sw = theme["score_weights"]
    related_pts = [(r["label"], float(r["weight"]) if r.get("weight") is not None else sw["related_each"])
                   for r in rel_hit]
    score, breakdown = compute_score(anchors, related_pts, subj_hits, all_hits, table, thread["count"], sw)
    cats = {}
    for k in all_hits:
        cat = table[k]["category"]
        cats[cat] = cats.get(cat, 0) + 1
    return {
        "id": thread["id"], "A": a_ok, "B": b_ok, "C_i": c_i, "C_ii": c_ii, "C": c_i or c_ii,
        "candidate": a_ok and b_ok and (c_i or c_ii),
        "anchors": anchors, "related": [r["label"] for r in rel_hit],
        "strong_n": len(sig["strong"]), "weak_n": len(sig["weak"]),
        "cats": cats, "score": score, "breakdown": breakdown,
        "participants": names, "hits": sorted(all_hits),
    }


def cond_text(ev):
    mark = lambda b: "○" if b else "×"
    a = f"A:{mark(ev['A'])}" + (f"({'/'.join(ev['anchors'])})" if ev["anchors"] else "")
    b = f"B:{mark(ev['B'])}"
    ci = f"C(i):{mark(ev['C_i'])}" + (f"({'/'.join(ev['related'])})" if ev["related"] else "")
    cii = f"C(ii):{mark(ev['C_ii'])}(強{ev['strong_n']}/弱{ev['weak_n']})"
    return f"{a} {b} {ci} {cii}"


def failed_conditions(ev):
    out = []
    if not ev["A"]:
        out.append("A")
    if not ev["B"]:
        out.append("B")
    if not ev["C"]:
        out.append("C")
    return out


def summarize_conditions(evals):
    """条件別の件数。単独で不成立のスレッド数（再現率の検証・調整用）も数える。"""
    s = {"total": len(evals), "A": 0, "B": 0, "C": 0, "C_i": 0, "C_ii": 0, "AB": 0, "candidate": 0,
         "only_A_fails": 0, "only_B_fails": 0, "only_C_fails": 0}
    for e in evals:
        s["A"] += e["A"]
        s["B"] += e["B"]
        s["C"] += e["C"]
        s["C_i"] += e["C_i"]
        s["C_ii"] += e["C_ii"]
        s["AB"] += e["A"] and e["B"]
        s["candidate"] += e["candidate"]
        f = failed_conditions(e)
        if f == ["A"]:
            s["only_A_fails"] += 1
        elif f == ["B"]:
            s["only_B_fails"] += 1
        elif f == ["C"]:
            s["only_C_fails"] += 1
    return s


def count_stale_stage2(stage2, theme):
    """本文照合済みだが、辞書ハッシュが現在のテーマと異なる第2段レコードの件数。"""
    return sum(1 for v in (stage2 or {}).values()
               if isinstance(v, dict) and v.get("bh") is not None and v.get("bh") != theme["dict_hash"])


def select_stage2_targets(threads, evals, scope):
    """第2段で開くメールの選択。scope: candidate(候補のみ) / ab(AかつB) / either(AまたはB)。"""
    def ok(e):
        if scope == "candidate":
            return e["candidate"]
        if scope == "either":
            return e["A"] or e["B"]
        return e["A"] and e["B"]
    out = []
    for t, e in zip(threads, evals):
        if ok(e):
            out.extend(t["mails"])
    return out


# ============================================================
# 純粋関数: CSV・judgments
# ============================================================
def csv_safe(value):
    """CSVインジェクション対策: 先頭が = + - @ タブ CR の文字列は先頭に ' を付ける（空白を除いた先頭も見る）。"""
    if value is None:
        return ""
    s = value if isinstance(value, str) else str(value)
    if s and (s[0] in "=+-@\t\r" or s.lstrip()[:1] in ("=", "+", "-", "@")):
        return "'" + s
    return s


def write_csv(path, header, rows):
    """utf-8-sig で書く。全セルを csv_safe に通す。"""
    d = os.path.dirname(os.path.abspath(path))
    os.makedirs(d, exist_ok=True)
    with open(path, "w", encoding="utf-8-sig", newline="") as fh:
        w = csv.writer(fh)
        w.writerow([csv_safe(h) for h in header])
        for r in rows:
            w.writerow([csv_safe(c) for c in r])


def merge_judgments(existing, thread_ids):
    """判定の保持: 既存の判定は上書きせず、無い thread_id の空行だけ追加する（不正な形の値もそのまま残す）。
    戻り値: (新しい辞書, 追加件数)。"""
    out = {k: (dict(v) if isinstance(v, dict) else v) for k, v in (existing or {}).items()}
    added = 0
    for tid in thread_ids:
        if tid not in out:
            out[tid] = {k: "" for k in JUDGMENT_KEYS}
            added += 1
    for tid, v in out.items():
        if isinstance(v, dict):
            for k in JUDGMENT_KEYS:
                v.setdefault(k, "")   # 欠けたキーは空で補うだけ（既存の値は変えない）
    return out, added


def load_judgments(path):
    """判定ファイルの読み込み。戻り値: (辞書, 状態)。状態: 'new'(無い) / 'ok' / 'broken'(存在するが読めない・空・形が不正)。
    broken のとき呼び出し側は上書きしてはいけない。"""
    if not os.path.exists(path):
        return {}, "new"
    try:
        with open(path, "r", encoding="utf-8-sig") as fh:
            raw = fh.read()
        if not raw.strip():
            return {}, "broken"
        data = json.loads(raw)
        if not isinstance(data, dict):
            return {}, "broken"
        return data, "ok"
    except Exception:
        return {}, "broken"


def backup_file(path):
    """壊れたファイルをタイムスタンプ付きで退避する（元のファイルは触らない）。退避先を返す。失敗は None。"""
    dst = f"{path}.bak_{datetime.now().strftime('%Y%m%d_%H%M%S')}"
    try:
        shutil.copy2(path, dst)
        return dst
    except Exception:
        return None


def load_json(path, default=None):
    try:
        with open(path, "r", encoding="utf-8-sig") as fh:
            return json.load(fh)
    except Exception:
        return default


def write_json_atomic(path, obj):
    d = os.path.dirname(os.path.abspath(path))
    os.makedirs(d, exist_ok=True)
    tmp = path + ".tmp"
    with open(tmp, "w", encoding="utf-8") as fh:
        json.dump(obj, fh, ensure_ascii=False)
        fh.flush()
        try:
            os.fsync(fh.fileno())
        except OSError:
            pass
    os.replace(tmp, path)


def sources_text(thread):
    seen = []
    for m in thread["mails"]:
        for l in m["locs"]:
            s = f"{l['store']}:{l['folder']}"
            if s not in seen:
                seen.append(s)
    return "; ".join(seen)


def is_reference_thread(ev):
    """参考候補: B∧C が成立し A だけが不成立（アンカー無し）。"""
    return (not ev["A"]) and ev["B"] and ev["C"]


def build_ledger_rows(threads, evals, judgments, theme, include_no_anchor=False):
    """候補スレッドを、スコア降順（同点は新しい順）で台帳の行にする。
    include_no_anchor=True のときは、アンカー無しの参考候補（B∧C成立でA不成立）を通常候補の後ろに加え、
    列「参考(アンカー無し)」に '参考' を入れる。"""
    def order(sel):
        pairs = [(t, e) for t, e in zip(threads, evals) if sel(e)]
        pairs.sort(key=lambda te: (-te[1]["score"], -te[0]["last"].timestamp()))
        return pairs
    groups = [(order(lambda e: e["candidate"]), "")]
    if include_no_anchor:
        groups.append((order(is_reference_thread), "参考"))
    rows = []
    for pairs, flag in groups:
        for t, e in pairs:
            j = judgments.get(t["id"], {}) if judgments else {}
            if not isinstance(j, dict):
                j = {}
            names = e["participants"]
            ptxt = " / ".join(names[:MAX_PARTICIPANTS_IN_CSV])
            if len(names) > MAX_PARTICIPANTS_IN_CSV:
                ptxt += f" …他{len(names) - MAX_PARTICIPANTS_IN_CSV}名"
            cats = " / ".join(f"{c}({n})" for c, n in sorted(e["cats"].items(), key=lambda kv: (-kv[1], kv[0])))
            last = t["mails"][-1]["locs"][0]
            rows.append([
                t["id"], t["subject"], t["first"].strftime("%Y-%m-%d %H:%M"), t["last"].strftime("%Y-%m-%d %H:%M"),
                t["count"], ptxt, " / ".join(e["anchors"]), " / ".join(e["related"]), cats, cond_text(e),
                fmt_num(e["score"]), e["breakdown"], sources_text(t), last["eid"], last["sid"],
                j.get("確認結果", ""), j.get("メモ", ""), j.get("問題の種類", ""), theme["theme_version"], flag,
            ])
    return rows


def import_openpyxl():
    """openpyxl を遅延import。無ければ None（その場合はCSVのみ出力する）。"""
    try:
        import openpyxl  # noqa: WPS433
        return openpyxl
    except Exception:
        return None


def write_ledger_xlsx(path, header, rows, openpyxl_mod=None):
    """台帳を .xlsx で書く。件名セルに hyperlink（`ledger:<thread_id>`）を設定する。
    リンクのスキームと目標文字列は固定の自前生成のみ（メール由来の文字列を混ぜない。thread_id は形式検証済みのもののみ）。
    メール由来の文字列は csv_safe を通し、さらに文字列型を強制して数式化を防ぐ。
    戻り値: 書けたら True、openpyxl が無ければ False。"""
    ox = openpyxl_mod or import_openpyxl()
    if ox is None:
        return False
    from openpyxl.styles import Font
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = ox.Workbook()
    ws = wb.active
    ws.title = "台帳"
    ws.append([csv_safe(h) for h in header])
    for c in ws[1]:
        c.font = Font(bold=True)
    numeric = {"メール数": int, "スコア": float}
    for ri, row in enumerate(rows, start=2):
        for ci, val in enumerate(row, start=1):
            name = header[ci - 1]
            cell = ws.cell(row=ri, column=ci)
            conv = numeric.get(name)
            if conv is not None:
                try:
                    cell.value = conv(val)
                    continue
                except (TypeError, ValueError):
                    pass
            cell.value = csv_safe(val)
            cell.data_type = "s"        # 文字列型を強制（先頭が = のセルを数式にしない）
        tid = str(row[0])
        if LEDGER_THREAD_ID_RE.match(tid):
            link = ws.cell(row=ri, column=2)
            link.hyperlink = f"{PROTOCOL_NAME}:{tid}"
            link.font = Font(color="0563C1", underline="single")
    ws.freeze_panes = "C2"
    if rows:
        ws.auto_filter.ref = ws.dimensions
        col = header.index("確認結果(本物/違う/保留)") + 1
        dv = DataValidation(type="list", formula1='"本物,違う,保留"', allow_blank=True)
        ws.add_data_validation(dv)
        letter = ws.cell(row=1, column=col).column_letter
        dv.add(f"{letter}2:{letter}{len(rows) + 1}")
    widths = {"thread_id": 14, "件名": 60, "参加者": 40, "条件A/B/C内訳": 44, "スコア内訳": 40}
    for ci, name in enumerate(header, start=1):
        ws.column_dimensions[ws.cell(row=1, column=ci).column_letter].width = widths.get(name, 16)
    d = os.path.dirname(os.path.abspath(path))
    os.makedirs(d, exist_ok=True)
    wb.save(path)
    return True


# ---------- 件名で絞り込んで Outlook で開く ----------
def safe_search_subject(subject):
    """Outlook の Instant Search（subject:"..."）に渡す件名の安全化。
    RE:/FW:/** internal only ** 等の前置を除去し、末尾の括弧書きを除去し、検索を壊す記号を空白にする。"""
    t = norm_space(subject)
    for _ in range(60):
        n = _PREFIX_RE.sub("", t, count=1)
        if n == t:
            break
        t = n
    t = re.sub(r"\s*[\(（][^\)）]*[\)）]\s*$", "", t)
    t = re.sub(r"[\":;\(\)（）\-\[\]\{\}<>'\\*%]", " ", t)
    return re.sub(r"\s+", " ", t).strip()[:200]


def build_subject_query(subject):
    safe = safe_search_subject(subject)
    return f'subject:"{safe}"' if safe else ""


def open_thread_in_outlook(client, eid, sid, subject_hint, rep):
    """最新メールを開く。メールのあるフォルダ（PSTルート直下を含む）へ切り替え、件名で絞り込み、
    該当メールを選択状態にする。選択できない/失敗したときは Display() にフォールバック。読み取りのみ。
    戻り値: 'explorer'(絞り込み表示) / 'display'(メール単体表示)"""
    outlook = client.Dispatch("Outlook.Application")
    ns = outlook.GetNamespace("MAPI")
    item = call_by_id(ns.GetItemFromID, eid, sid)
    subject = safe_attr(item, "Subject", "") or subject_hint or ""
    mode = "display"
    try:
        query = build_subject_query(subject)
        if not query:
            raise LookupError("件名が空")
        try:
            explorer = outlook.ActiveExplorer()
        except Exception:
            explorer = None
        folder = item.Parent
        if folder is None:
            raise LookupError("親フォルダを取得できません")
        if explorer is None:
            explorer = folder.GetExplorer()
            explorer.Display()
        else:
            explorer.CurrentFolder = folder
            try:
                explorer.Activate()
            except Exception:
                pass
        time.sleep(EXPLORER_SETTLE_SEC)     # フォルダ切替の反映を待ってから検索する
        explorer.Search(query, 0)       # 0 = 現在のフォルダ
        selected = False
        for attempt in range(SELECT_RETRIES):
            time.sleep(SELECT_RETRY_SEC)    # 絞り込みの反映を待ってから選択する（最大 SELECT_RETRIES 回）
            try:
                if explorer.IsItemSelectableInView(item):
                    explorer.ClearSelection()
                    explorer.AddToSelection(item)
                    selected = True
                    break
            except Exception:
                selected = False
        if selected:
            mode = "explorer"
    except Exception as e:
        rep.log(f"⚠️ 絞り込み表示に失敗: {error_kind(e)}（メール単体の表示に切り替えます）")
    if mode != "explorer":
        item.Display()      # 表示するだけ（変更はしない）
    return mode


# ---------- ledger: プロトコル（HKCU のみ・オプトイン） ----------
def parse_ledger_url(url):
    """`ledger:<thread_id>` から thread_id だけを取り出す。形式が違えば None（任意のパスやコマンドは受け付けない）。"""
    m = re.fullmatch(r"\s*(?i:ledger):(?://)?([A-Za-z0-9_-]{4,64})/?\s*", url or "")
    return m.group(1) if m else None


def protocol_command(pythonw, script):
    return f'"{pythonw}" "{script}" --open-url "%1"'


def pythonw_path(executable=None):
    exe = executable or sys.executable
    d, base = os.path.split(exe)
    if base.lower() == "python.exe":
        cand = os.path.join(d, "pythonw.exe")
        if os.path.exists(cand):
            return cand
    return exe


PROTOCOL_KEY = "Software\\Classes\\" + PROTOCOL_NAME


def protocol_plan(pythonw, script):
    """登録内容 [(サブキー, 値名, 値)]。HKCU 配下のみ。"""
    if '"' in pythonw or '"' in script:
        raise ValueError("パスに二重引用符が含まれています")
    return [
        (PROTOCOL_KEY, "", "URL:Ledger Protocol"),
        (PROTOCOL_KEY, "URL Protocol", ""),
        (PROTOCOL_KEY + "\\shell\\open\\command", "", protocol_command(pythonw, script)),
    ]


def register_protocol(winreg_mod, plan):
    for sub, name, val in plan:
        key = winreg_mod.CreateKey(winreg_mod.HKEY_CURRENT_USER, sub)
        try:
            winreg_mod.SetValueEx(key, name, 0, winreg_mod.REG_SZ, val)
        finally:
            winreg_mod.CloseKey(key)


def unregister_protocol(winreg_mod):
    """HKCU\\Software\\Classes\\ledger を削除する（深い方から）。戻り値: 削除したキー数。"""
    n = 0
    for sub in (PROTOCOL_KEY + "\\shell\\open\\command", PROTOCOL_KEY + "\\shell\\open",
                PROTOCOL_KEY + "\\shell", PROTOCOL_KEY):
        try:
            winreg_mod.DeleteKey(winreg_mod.HKEY_CURRENT_USER, sub)
            n += 1
        except FileNotFoundError:
            pass
    return n


def read_existing_command(winreg_mod):
    """既存の ledger: 登録（開くコマンド）を読む。無ければ None。"""
    try:
        key = winreg_mod.OpenKey(winreg_mod.HKEY_CURRENT_USER, PROTOCOL_KEY + "\\shell\\open\\command")
        try:
            return winreg_mod.QueryValueEx(key, "")[0]
        finally:
            winreg_mod.CloseKey(key)
    except Exception:
        return None


def run_protocol(args, rep, winreg_mod=None):
    if winreg_mod is None:
        try:
            import winreg as winreg_mod  # noqa: WPS433
        except ImportError:
            rep.log("❌ winreg が使えません（Windows 専用の機能です）")
            return 2
    script = os.path.abspath(__file__)
    try:
        plan = protocol_plan(pythonw_path(), script)
    except ValueError as e:
        rep.log(f"❌ {e}")
        return 2
    action = "登録" if args.register_protocol else "削除"
    rep.log(f"🔗 ledger: プロトコルを{action}します（HKEY_CURRENT_USER のみ・管理者権限は不要）")
    if args.register_protocol:
        for sub, name, val in plan:
            rep.log(f"   HKCU\\{sub}  [{name or '(既定)'}] = {val!r}")
    else:
        rep.log(f"   HKCU\\{PROTOCOL_KEY} の次の4キーを、深い方から1本ずつ削除します（非再帰）: "
                "shell\\open\\command → shell\\open → shell → ledger")
    existing_cmd = read_existing_command(winreg_mod)
    if existing_cmd is not None:
        rep.log(f"   ⚠️ 既存の ledger: 登録があります。{'上書きします' if args.register_protocol else '削除対象です'}: {existing_cmd!r}")
    if not args.yes:
        try:
            ans = input("実行しますか？ [y/N]: ").strip().lower()
        except (EOFError, OSError):
            ans = ""
        if ans not in ("y", "yes"):
            rep.log("⏭️ 中止しました（何も変更していません。--yes で確認を省略できます）")
            return 1
    try:
        if args.register_protocol:
            register_protocol(winreg_mod, plan)
            rep.log("✅ 登録しました。Excel の件名リンク（ledger:<thread_id>）から Outlook が開きます")
        else:
            n = unregister_protocol(winreg_mod)
            rep.log(f"✅ 削除しました（{n}/4キー）")
    except Exception as e:
        rep.log(f"❌ レジストリ操作に失敗: {error_kind(e)}")
        return 2
    return 0


# ---------- 判定のExcel取込 ----------
def _unescape_cell(v):
    """セルの値を文字列にし、csv_safe が付けた先頭の ' を戻す。"""
    if v is None:
        return ""
    s = str(v).strip()
    if s.startswith("'") and s[1:2] in ("=", "+", "-", "@", "\t", "\r"):
        s = s[1:]
    return s


def read_table_rows(path, openpyxl_mod=None):
    """CSV または xlsx を、見出し行つきの辞書のリストとして読む。サイズ上限（50MB・20万行）あり。"""
    if os.path.getsize(path) > IMPORT_MAX_BYTES:
        raise ValueError("ファイルが大きすぎます")
    ext = os.path.splitext(path)[1].lower()
    if ext == ".xlsx":
        ox = openpyxl_mod or import_openpyxl()
        if ox is None:
            raise RuntimeError("openpyxl が無いため xlsx を読めません（CSVで保存し直してください）")
        wb = ox.load_workbook(path, read_only=True, data_only=True)
        try:
            ws = wb.worksheets[0]
            it = ws.iter_rows(values_only=True)
            header = [str(h).strip() if h is not None else "" for h in next(it, ())]
            rows = []
            for r in it:
                if len(rows) >= IMPORT_MAX_ROWS:
                    raise ValueError("行数が多すぎます")
                rows.append(dict(zip(header, r)))
            return rows
        finally:
            wb.close()
    for enc in ("utf-8-sig", "cp932"):
        try:
            with open(path, "r", encoding=enc, newline="") as fh:
                rows = []
                for r in csv.DictReader(fh):
                    if len(rows) >= IMPORT_MAX_ROWS:
                        raise ValueError("行数が多すぎます")
                    rows.append(r)
                return rows
        except UnicodeDecodeError:
            continue
    raise ValueError("CSVの文字コードを判定できません")


def import_judgment_rows(rows, existing):
    """台帳（CSV/xlsx）の判定列を judgments に取り込む。thread_id で突合する（台帳に存在する既知IDのみ）。
    確認結果は許可リスト（本物/違う/保留/空）で検証し、不正値はスキップ。空のセルは既存を上書きしない。
    judgments に無い thread_id は取り込まずスキップ（件数を数える）。
    戻り値: (新しい辞書, 統計dict)"""
    out = {k: (dict(v) if isinstance(v, dict) else v) for k, v in (existing or {}).items()}
    st = {"rows": 0, "matched": 0, "unknown": 0, "updated": 0, "invalid_value": 0, "bad_id": 0, "bad_entry": 0}
    for r in rows:
        st["rows"] += 1
        tid = _unescape_cell(r.get("thread_id"))
        if not LEDGER_THREAD_ID_RE.match(tid):
            st["bad_id"] += 1
            continue
        if tid not in out:
            st["unknown"] += 1          # 台帳に存在しない thread_id は取り込まない
            continue
        if not isinstance(out[tid], dict):
            st["bad_entry"] += 1       # 形が不正な既存行は触らない
            continue
        st["matched"] += 1
        verdict = _unescape_cell(r.get("確認結果(本物/違う/保留)", r.get("確認結果")))
        memo = _unescape_cell(r.get("メモ"))[:2000]
        kind = _unescape_cell(r.get("問題の種類"))[:200]
        entry = out[tid]
        for key, val in (("確認結果", verdict), ("メモ", memo), ("問題の種類", kind)):
            if val == "":
                continue                # 空は変更なし
            if key == "確認結果" and val not in ALLOWED_VERDICTS:
                st["invalid_value"] += 1
                continue
            if entry.get(key) != val:
                entry[key] = val
                st["updated"] += 1
    return out, st


def run_import_judgments(args, theme, rep, paths):
    jpath = judgments_path(paths, theme["name"])
    existing, status = load_judgments(jpath)
    if status == "broken":
        bak = backup_file(jpath)
        rep.log("❌ 既存の判定ファイルを読めません（壊れている・空・形が不正）。上書きせず取込を中止しました。"
                + (f" 退避コピー: {bak}" if bak else ""))
        return 2
    try:
        rows = read_table_rows(args.import_judgments)
    except Exception as e:
        rep.log(f"❌ 取込ファイルを読めません: {error_kind(e)}")
        return 2
    if rows and "thread_id" not in rows[0]:
        rep.log("❌ 取込ファイルに thread_id 列がありません")
        return 2
    merged, st = import_judgment_rows(rows, existing)
    if status == "ok":
        bak = backup_file(jpath)         # 正常な既存ファイルでも、上書きの前に必ず退避コピーを作る
        if not bak:
            rep.log("❌ 判定ファイルの退避コピーを作れないため、取込を中止しました（何も変更していません）")
            return 2
        rep.log(f"💾 取込前の判定ファイルを退避しました: {bak}")
    try:
        write_json_atomic(jpath, merged)
    except Exception as e:
        rep.log(f"❌ 判定ファイルを保存できません: {error_kind(e)}")
        return 2
    rep.log(f"📥 判定の取込: 行 {st['rows']} / 突合 {st['matched']} / 台帳に無いIDをスキップ {st['unknown']} / 更新した項目 {st['updated']}"
            f" / 不正な確認結果をスキップ {st['invalid_value']} / thread_id不正 {st['bad_id']} / 形が不正な既存行 {st['bad_entry']}")
    return 0


def build_coverage_rows(stats, mails, threads, evals, stage2, theme, months, store_names):
    """ストア×月の走査件数・重複除去後件数・アンカー関与件数・候補スレッド数。
    重複除去後件数 = そのストアが代表の取得元になっているメール数。0件の月は「未取得の恐れ」。"""
    raw = stats["raw_by_store_month"]
    uniq, anch = {}, {}
    for m in mails:
        loc = m["locs"][0]
        k = (loc["store"], (m["t"].year, m["t"].month))
        uniq[k] = uniq.get(k, 0) + 1
        names, emails, _p = mail_people(m, mail_stage2(m, stage2))
        domains = {email_domain(e) for e in emails}
        if any(entity_matches(a, names, emails, domains) for a in theme["anchors"]):
            anch[k] = anch.get(k, 0) + 1
    cand = {}
    for t, e in zip(threads, evals):
        if not e["candidate"]:
            continue
        seen = set()
        for m in t["mails"]:
            for l in m["locs"]:
                seen.add((l["store"], (m["t"].year, m["t"].month)))
        for k in seen:
            cand[k] = cand.get(k, 0) + 1
    rows = []
    for ym in months:
        tot_raw = sum(raw.get((s, ym), 0) for s in store_names)
        flag_all = "未取得の恐れ(全ストア0件・取得元ゼロ・要確認)" if tot_raw == 0 else ""
        for s in store_names:
            r = raw.get((s, ym), 0)
            rows.append([s, ym_label(ym), r, uniq.get((s, ym), 0), anch.get((s, ym), 0), cand.get((s, ym), 0),
                         flag_all or ("0件(取得元ゼロ・要確認)" if r == 0 else "")])
        # 全ストア合計の行（候補スレッド数は重複を数えないよう月で数え直す）
        month_cand = sum(1 for t, e in zip(threads, evals) if e["candidate"]
                         and any((m["t"].year, m["t"].month) == ym for m in t["mails"]))
        rows.append(["(全ストア)", ym_label(ym), tot_raw,
                     sum(v for (s, y), v in uniq.items() if y == ym),
                     sum(v for (s, y), v in anch.items() if y == ym), month_cand,
                     "未取得の恐れ(取得元ゼロ・要確認)" if tot_raw == 0 else ""])
    return rows


def store_zero_months(stats, months, store_names):
    """ストア別に0件の月の一覧。戻り値: {ストアラベル: [(年,月), ...]}（期間内に1件も無いストアは全月）。"""
    raw = stats["raw_by_store_month"]
    return {s: [ym for ym in months if raw.get((s, ym), 0) == 0] for s in store_names}


COVERAGE_COLUMNS = ["ストア", "年月", "走査件数", "重複除去後件数", "アンカー関与件数", "候補スレッド数", "備考"]


def suggest_participants(threads, evals, theme, stage2):
    """候補スレッドに登場するが theme に無い人（表示名/SMTP）の出現スレッド数ランキング。
    戻り値: [(表示名, SMTP(; 区切り), スレッド数)]"""
    ents = [theme["owner"]] + theme["anchors"] + theme["related_parties"]
    count, disp, addrs = {}, {}, {}
    for t, e in zip(threads, evals):
        if not e["candidate"]:
            continue
        _n, _e, pairs = thread_people(t, stage2)
        seen = set()
        for name, addr in pairs:
            key = name_key(name)[3] if name_key(name)[0] else addr
            if not key or key in seen:
                continue
            ems = {addr} if addr else set()
            if any(entity_matches(en, [name] if name else [], ems) for en in ents):
                continue
            seen.add(key)
            disp.setdefault(key, name or addr)
            if addr:
                addrs.setdefault(key, set()).add(addr)
        for k in seen:
            count[k] = count.get(k, 0) + 1
    out = [(disp[k], "; ".join(sorted(addrs.get(k, ()))), n) for k, n in count.items()]
    out.sort(key=lambda r: (-r[2], r[0]))
    return out


def tokenize_for_terms(text):
    """辞書候補語の抽出用: カタカナ3字以上・漢字2字以上・英字始まり3字以上の連続を語とする。"""
    low = norm_space(text).lower()
    return {t for t in _TOKEN_RE.findall(low) if t not in STOPWORDS}


def rank_candidate_terms(fg_docs, bg_docs, known_terms=(), min_fg=2):
    """候補スレッド(fg)で偏って出る語のランキング。各 doc は語集合。
    全体(bg)での出現率が高い一般語は lift が下がるので順位が下がる。
    戻り値: [(語, fg件数, bg件数, fg率, bg率, lift, スコア, 既に辞書にあるか)]"""
    fg_n, bg_n = len(fg_docs), len(bg_docs)
    if fg_n == 0 or bg_n == 0:
        return []
    fg_c, bg_c = {}, {}
    for d in fg_docs:
        for w in d:
            fg_c[w] = fg_c.get(w, 0) + 1
    for d in bg_docs:
        for w in d:
            bg_c[w] = bg_c.get(w, 0) + 1
    known = {norm_space(k).lower() for k in known_terms}
    out = []
    for w, f in fg_c.items():
        if f < min_fg:
            continue
        b = max(bg_c.get(w, 0), f)
        fr, br = f / fg_n, (b + 1) / (bg_n + 1)
        lift = fr / br
        if lift <= 1.0:
            continue
        out.append((w, f, b, round(fr, 4), round(b / bg_n, 4), round(lift, 2), round(f * math.log(lift), 2),
                    "はい" if w in known else ""))
    out.sort(key=lambda r: (-r[6], r[0]))
    return out


def candidate_term_docs(threads, evals, theme, stage2):
    """候補語ランキング用の語集合。件名(正規化)と第2段の添付名から作る（本文は保存していないので使えない）。
    fg = 関係者が絡む候補スレッド、bg = 全スレッド。"""
    fg, bg = [], []
    for t, e in zip(threads, evals):
        words = set()
        for m in t["mails"]:
            words |= tokenize_for_terms(m["nsubj"])
            st2 = mail_stage2(m, stage2)
            if st2:
                for a in st2.get("att") or []:
                    words |= tokenize_for_terms(os.path.splitext(a)[0])
        bg.append(words)
        if e["candidate"] and e["C_i"]:
            fg.append(words)
    return fg, bg


# ============================================================
# 純粋関数: 正解件名の再現率チェック
# ============================================================
def load_calibration_subjects(path):
    """1行1件名。# で始まる行と空行は無視。"""
    out = []
    with open(path, "r", encoding="utf-8-sig") as fh:
        for line in fh:
            s = line.strip()
            if s and not s.startswith("#"):
                out.append(s)
    return out


def classify_check_subjects(queries, threads, evals, min_partial_len=GENERIC_MIN_LEN):
    """各件名を正規化し、スレッドの件名（スレッド内の全メールの正規化件名）と完全一致/部分一致で探す。
    完全一致を優先し、完全一致があるときは部分一致を見ない（完全一致が候補外なら、部分一致に候補があっても
    found にしない）。部分一致は汎用件名（短い・一般語）を除外する。
    戻り値: query ごとの dict。status は 'found'(候補) / 'found_not_candidate' / 'missing'。"""
    total = len(threads)
    neg = sorted(-e["score"] for e in evals)   # 昇順
    results = []
    subj_sets = [{m["nsubj"] for m in t["mails"]} for t in threads]

    def make(t, e, kind):
        rank = bisect.bisect_left(neg, -e["score"]) + 1   # 自分より高いスコアの数 + 1
        return {"thread_id": t["id"], "kind": kind, "mails": t["count"],
                "last": t["last"].strftime("%Y-%m-%d"), "A": e["A"], "B": e["B"],
                "C": e["C"], "C_i": e["C_i"], "C_ii": e["C_ii"],
                "candidate": e["candidate"], "score": e["score"], "rank": rank,
                "total": total, "failed": failed_conditions(e)}

    for q in queries:
        nq = normalize_subject(q)
        matches = []
        if nq:
            for t, e, ss in zip(threads, evals, subj_sets):
                if nq in ss:
                    matches.append(make(t, e, "完全一致"))
            if not matches and len(nq) >= min_partial_len and not is_generic_subject(nq):
                for t, e, ss in zip(threads, evals, subj_sets):
                    if any(nq in x or (not is_generic_subject(x) and x in nq) for x in ss):
                        matches.append(make(t, e, "部分一致"))
            matches.sort(key=lambda x: -x["score"])
        if not matches:
            status = "missing"
        elif any(m["candidate"] for m in matches):
            status = "found"
        else:
            status = "found_not_candidate"
        results.append({"query": q, "norm": nq, "status": status, "matches": matches[:5]})
    return results


def format_check_result(i, n, res):
    """--check-subjects の出力行（件名の表示が許可されている唯一の出力）。"""
    mk = lambda b: "○" if b else "×"
    lines = [f"🔎 [{i}/{n}] {res['query']}"]
    if res["status"] == "missing":
        lines.append("   ❌ 見つからない: キャッシュに無い（期間・ストア・フォルダの問題の可能性）")
        return lines
    for m in res["matches"]:
        verdict = "候補" if m["candidate"] else "候補外"
        lines.append(f"   ✅ 見つかった({m['kind']}): thread_id={m['thread_id']} / メール{m['mails']}通 / 最終 {m['last']}")
        lines.append(f"      判定: {verdict} | 条件A{mk(m['A'])} B{mk(m['B'])} C{mk(m['C'])}"
                     f"(i{mk(m['C_i'])} ii{mk(m['C_ii'])}) | スコア {fmt_num(m['score'])} | 順位 {m['rank']}/{m['total']}")
        if not m["candidate"]:
            lines.append(f"      ⚠️ あるが候補外: 不成立の条件 = {','.join(m['failed'])}")
    return lines


# ============================================================
# パス・キャッシュ
# ============================================================
def tool_root():
    """outlook_total_organizer/（tools の1つ上）"""
    return os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def make_paths(data_dir=None, output_dir=None):
    root = tool_root()
    data = data_dir or os.path.join(root, "json", "thread_ledger")
    return {
        "data": data,
        "scan_cache": os.path.join(data, "scan_cache"),
        "stage2_cache": os.path.join(data, "candidate_cache", "stage2.json"),
        "themes": os.path.join(data, "themes"),
        "output": output_dir or os.path.join(root, "mail_reports"),
    }


def judgments_path(paths, theme_name):
    return os.path.join(paths["data"], f"judgments_{theme_name}.json")


def folder_cache_path(paths, store_index, store_name, folder_path):
    h = hashlib.sha1(f"{store_name}\n{folder_path}".encode("utf-8", "replace")).hexdigest()[:12]
    return os.path.join(paths["scan_cache"], f"{store_index:02d}_{h}.json")


def load_all_caches(scan_dir):
    out = []
    for p in sorted(glob.glob(os.path.join(scan_dir, "*.json"))):
        c = load_json(p)
        if isinstance(c, dict) and isinstance(c.get("records"), list) and c.get("schema") == 1:
            out.append(c)
    return out


def decide_cache_action(cache, from_ym, to_ym, rescan):
    """フォルダ単位キャッシュの扱い: 'skip'(完了済みで期間を含む) / 'resume'(チェックポイントから再開) / 'fresh'(取り直し)"""
    if rescan or not isinstance(cache, dict):
        return "fresh"
    try:
        cf, ct = parse_ym(cache["range"][0]), parse_ym(cache["range"][1])
    except Exception:
        return "fresh"
    if cache.get("complete"):
        return "skip" if (cf <= from_ym and ct >= to_ym) else "fresh"
    if (cf, ct) == (from_ym, to_ym) and cache.get("last_received"):
        return "resume"
    return "fresh"


def make_cache_payload(task, state, from_ym, to_ym, complete):
    errs = dict(state.get("prev_errors") or {})
    for k, v in state["errors"].items():
        errs[k] = errs.get(k, 0) + v
    return {
        "schema": 1, "scanner_version": SCRIPT_VERSION,
        "store_index": task["sidx"], "store_name": task["store_name"], "store_id": task["store_id"],
        "file_path": task.get("file_path", ""),
        "folder_path": task["folder_path"], "date_field": state.get("date_field") or "or",
        "range": [ym_label(from_ym), ym_label(to_ym)], "complete": bool(complete),
        "last_received": dt_to_str(state.get("last")), "scanned_at": datetime.now().strftime("%Y-%m-%dT%H:%M:%S"),
        "counts": {"mail": len(state["records"]), "non_mail": state["non_mail"], "errors": errs},
        "records": state["records"],
    }


# ============================================================
# 出力（コンソール + 追記保存ファイル）
# ============================================================
class NameMasker:
    """--mask-names 用。同じ名前には同じ連番を割り当てる（Store-01 / Folder-001）。"""

    def __init__(self, prefix, width):
        self.prefix, self.width, self._map = prefix, width, {}

    def mask(self, name):
        if name not in self._map:
            self._map[name] = f"{self.prefix}-{len(self._map) + 1:0{self.width}d}"
        return self._map[name]


_MASK = {"on": False, "store": NameMasker("Store", 2), "folder": NameMasker("Folder", 3)}


def configure_masking(enabled):
    """コンソール/ログのストア名・フォルダ名を連番に置換するか（CSVには元の名前を出す）。"""
    _MASK["on"] = bool(enabled)
    _MASK["store"] = NameMasker("Store", 2)
    _MASK["folder"] = NameMasker("Folder", 3)


def show_store(label):
    """コンソール/ログ用のストア表示名（メールアドレス形式は <MAIL>、--mask-names なら連番）。"""
    return _MASK["store"].mask(label) if _MASK["on"] else mask_emails(label)


def show_folder(path):
    """コンソール/ログ用のフォルダパス（--mask-names なら階層ごとに連番）。"""
    if not _MASK["on"]:
        return mask_emails(path)
    return "\\".join(_MASK["folder"].mask(seg) if seg else seg for seg in str(path).split("\\"))


class Reporter:
    """コンソールに出しつつ、1行ごとにファイルへ追記・flushする。
    渡されるのは件数・所要時間・エラー種別・フォルダ名・マスク済みストア名のみ。
    console() はコンソールだけに出す（--check-subjects の件名はファイルに残さない）。"""

    def __init__(self, path=None):
        self.lines = []
        self.path = path
        self.fh = None
        if path:
            try:
                os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
                self.fh = open(path, "w", encoding="utf-8-sig", newline="\n")
            except Exception:
                self.fh = None
                self.path = None

    def _print(self, text):
        try:
            print(text, flush=True)
        except UnicodeEncodeError:
            try:
                print(text.encode("ascii", "replace").decode("ascii"), flush=True)
            except Exception:
                pass
        except Exception:
            pass     # pythonw.exe など標準出力が無い環境ではログファイルだけに出す

    def log(self, text=""):
        self.lines.append(text)
        self._print(text)
        if self.fh:
            try:
                self.fh.write(text + "\n")
                self.fh.flush()
            except Exception:
                self.fh = None

    def console(self, text=""):
        self._print(text)

    def close(self):
        try:
            if self.fh:
                self.fh.close()
        except Exception:
            pass
        self.fh = None


# ============================================================
# COM 部分（Windows + Outlook 専用。importは遅延）
# ============================================================
def import_com():
    """win32com / pythoncom を遅延import。無ければ (None, None)。"""
    try:
        import win32com.client as client  # noqa: WPS433
        import pythoncom  # noqa: WPS433
        return client, pythoncom
    except Exception:
        return None, None


def safe_attr(obj, name, default=None):
    try:
        return getattr(obj, name)
    except Exception:
        return default


def call_by_id(fn, eid, sid):
    """StoreID が空のときは空文字を渡さず、第2引数を省略して呼ぶ（GetFolderFromID / GetItemFromID）。"""
    return fn(eid, sid) if sid else fn(eid)


def run_with_com(com, func, rep):
    """COMの初期化・解放（gc.collect → CoUninitialize）を一括で行う。func(client) の戻り値を返す。
    COM参照は func の中だけで持ち、関数を抜けてから解放される。"""
    client, pythoncom = com if com is not None else import_com()
    if client is None:
        rep.log("❌ win32com / pythoncom が import できません。Windows + Outlook + pywin32 の環境で実行してください。")
        return 2
    pythoncom.CoInitialize()
    exc = None
    code = 1
    try:
        code = func(client)
    except BaseException as e:  # noqa: BLE001  解放後に再送出する
        exc = e
    if exc is not None:
        exc.__traceback__ = None
    gc.collect()
    try:
        pythoncom.CoUninitialize()
    except Exception:
        pass
    if exc is not None:
        raise exc
    return code


def store_kind_label(display_name, file_path):
    low = (display_name or "").lower()
    if any(p in low for p in ("オンライン アーカイブ", "online archive", "in-place archive", "archive -")):
        return "アーカイブ"
    if (file_path or "").lower().endswith(".pst"):
        return "PST"
    return "メールボックス/PST"


def enumerate_mail_folders(root, store_id, exclude_names, exclude_prefixes, errors, stats):
    """ストアのルート配下を再帰列挙し、メールフォルダ(DefaultItemType==0)の記録リストを返す。
    ルート直下にメールがあれば「(ルート直下)」を加える。除外フォルダは配下ごと対象外。
    COMオブジェクトは保持せず、EntryID と StoreID だけを記録する（走査時に GetFolderFromID で取得する）。
    記録: {'path','name','total','eid','sid'}"""
    result = []

    def walk(folder, path, depth):
        if depth > MAX_FOLDER_DEPTH:
            return
        try:
            subs = folder.Folders
            n = int(subs.Count)
        except Exception as e:
            add_error(errors, f"フォルダ列挙:{error_kind(e)}")
            return
        for i in range(1, n + 1):
            try:
                sub = subs.Item(i)
                raw = safe_attr(sub, "Name", "") or "(名前なし)"
                if folder_excluded(raw, exclude_names, exclude_prefixes):
                    stats["excluded_folders"] = stats.get("excluded_folders", 0) + 1
                    continue
                sub_path = f"{path}\\{raw}"
                if safe_attr(sub, "DefaultItemType", None) == 0:
                    try:
                        total = int(sub.Items.Count)
                    except Exception as e:
                        total = 0
                        add_error(errors, f"件数取得:{error_kind(e)}")
                    result.append({"path": sub_path, "name": raw, "total": total,
                                   "eid": safe_attr(sub, "EntryID", "") or "", "sid": store_id})
                walk(sub, sub_path, depth + 1)
                sub = None
            except Exception as e:
                add_error(errors, f"フォルダ取得:{error_kind(e)}")

    try:
        root_total = int(root.Items.Count)
    except Exception:
        root_total = 0
    if root_total > 0:
        result.append({"path": "\\" + ROOT_FOLDER_LABEL, "name": ROOT_FOLDER_LABEL, "total": root_total,
                       "eid": safe_attr(root, "EntryID", "") or "", "sid": store_id, "is_root": True})
    walk(root, "", 1)
    return result


def _get(item, name, errs):
    try:
        return getattr(item, name)
    except Exception as e:
        add_error(errs, f"{name}:{error_kind(e)}")
        return None


def read_light_item(item, errs, known=None):
    """軽量項目だけを読む（Recipients・Body・Attachments・PropertyAccessor は読まない）。
    known(取得済みEntryIDの集合)に含まれるときは、残りの項目を読まずに 'dup' を返す。
    戻り値: ('mail', record) / ('nonmail', None) / ('dup', None) / ('error', None)"""
    cls = _get(item, "Class", errs)
    if cls is None:
        return "error", None
    try:
        if int(cls) != OL_MAIL_CLASS:
            return "nonmail", None
    except Exception:
        return "error", None
    eid = _get(item, "EntryID", errs)
    if not eid:
        return "error", None
    if known is not None and eid in known:
        return "dup", None
    t = to_naive_dt(_get(item, "ReceivedTime", errs))
    if t is None:
        t = to_naive_dt(_get(item, "SentOn", errs))
    if t is None:
        add_error(errs, "日時なし")
        return "error", None
    rec = {
        "e": str(eid), "s": _get(item, "Subject", errs) or "", "t": dt_to_str(t),
        "n": _get(item, "SenderName", errs) or "", "a": _get(item, "SenderEmailAddress", errs) or "",
        "to": _get(item, "To", errs) or "", "cc": _get(item, "CC", errs) or "",
        "c": _get(item, "ConversationID", errs) or "",
    }
    for k in ("s", "n", "a", "to", "cc", "c"):
        if not isinstance(rec[k], str):
            rec[k] = str(rec[k])
    return "mail", rec


def new_scan_state(cache=None):
    """errors はこの実行分だけ（過去分は prev_errors に持ち、保存時に合算する。再開時の二重計上を避ける）。"""
    if cache:
        recs = list(cache.get("records") or [])
        return {"records": recs, "known": {r.get("e") for r in recs},
                "non_mail": int(cache.get("counts", {}).get("non_mail", 0)),
                "errors": {}, "prev_errors": dict(cache.get("counts", {}).get("errors", {})),
                "processed": 0, "since_cp": 0, "last": parse_dt(cache.get("last_received")),
                "date_field": cache.get("date_field") or None, "iter_error": False, "complete": False}
    return {"records": [], "known": set(), "non_mail": 0, "errors": {}, "prev_errors": {}, "processed": 0,
            "since_cp": 0, "last": None, "date_field": None, "iter_error": False, "complete": False}


def scan_folder_items(folder, start_dt, end_dt, state, resume=False, checkpoint_cb=None,
                      checkpoint_every=CHECKPOINT_EVERY, total_hint=0):
    """1フォルダを期間全体で1回だけ Restrict(ISO) → Sort昇順 → 走査する。state を更新する（Ctrl+Cでも途中結果が残る）。
    Restrict は ([ReceivedTime]範囲) OR ([SentOn]範囲)。失敗する環境では ReceivedTime → 0件なら SentOn にフォールバック。
    再開: OR方式は並び順が保証できないため期間全体を取り直し、取得済みEntryIDを早期に飛ばす。
    ReceivedTime/SentOn 方式は state['last'] 以降だけを Restrict する。
    GetNext が失敗したら、その位置を飛ばして Item(i) の添字アクセスで先へ進む（最大 MAX_ITEM_SKIP 回）。"""
    if resume and state.get("date_field") in ("ReceivedTime", "SentOn"):
        mode0 = state["date_field"]
        b0 = max(start_dt, state["last"].replace(second=0)) if state.get("last") else start_dt
        attempts = [(mode0, b0)]
    else:
        attempts = [("or", start_dt), ("ReceivedTime", start_dt), ("SentOn", start_dt)]
    items = folder.Items
    restricted = first = chosen = chosen_begin = None
    last_exc = None
    for idx, (mode, begin) in enumerate(attempts):
        try:
            flt = build_or_filter(begin, end_dt) if mode == "or" else build_restrict_filter(mode, begin, end_dt)
            cand = items.Restrict(flt)
        except Exception as e:
            last_exc = e
            add_error(state["errors"], f"Restrict({mode}):{error_kind(e)}")
            continue
        try:
            cand.Sort("[ReceivedTime]" if mode == "or" else f"[{mode}]")
        except Exception as e:
            add_error(state["errors"], f"Sort:{error_kind(e)}")
        try:
            f = cand.GetFirst()
        except Exception as e:
            last_exc = e
            add_error(state["errors"], f"GetFirst:{error_kind(e)}")
            continue
        if f is None and mode == "ReceivedTime" and total_hint > 0 and idx + 1 < len(attempts):
            continue   # 送信済み系の可能性: 次に SentOn で試す
        restricted, first, chosen, chosen_begin = cand, f, mode, begin
        break
    if chosen is None:
        raise last_exc
    state["date_field"] = chosen
    if resume and chosen_begin == start_dt:
        state["non_mail"] = 0      # 期間全体を取り直すので、メール以外の件数も数え直す
    pos, indexed, fails = 1, False, 0
    item = first
    while item is not None:
        kind, rec = read_light_item(item, state["errors"], state["known"])
        if kind != "dup":
            state["processed"] += 1
            state["since_cp"] += 1
        if kind == "nonmail":
            state["non_mail"] += 1
        elif kind == "mail":
            state["known"].add(rec["e"])
            state["records"].append(rec)
            t = parse_dt(rec["t"])
            if t is not None and (state["last"] is None or t > state["last"]):
                state["last"] = t
        item = None
        if not indexed:
            try:
                item = restricted.GetNext()
                pos += 1
            except Exception as e:
                add_error(state["errors"], f"GetNext:{error_kind(e)}")
                indexed = True
                pos += 1             # 読めなかった位置
        if indexed and item is None:
            while True:
                if fails >= MAX_ITEM_SKIP:
                    state["iter_error"] = True
                    break
                pos += 1
                try:
                    total = int(restricted.Count)
                except Exception as e:
                    add_error(state["errors"], f"Count:{error_kind(e)}")
                    state["iter_error"] = True      # 件数が分からないまま完了扱いにしない（取りこぼし検知）
                    break
                if pos > total:
                    break
                try:
                    item = restricted.Item(pos)
                    break
                except Exception as e:
                    add_error(state["errors"], f"Item:{error_kind(e)}")
                    fails += 1
        if checkpoint_cb and state["since_cp"] >= checkpoint_every:
            state["since_cp"] = 0
            checkpoint_cb(state)
    state["complete"] = not state["iter_error"]
    return state


def run_scan(args, theme, rep, paths, from_ym, to_ym, com=None):
    """第1段: 全ストア（またはfilter）→ メールフォルダ → 期間のRestrict。戻り値: 終了コード(0=完了, 130=中断)"""
    def core(client):
        return _scan_core(args, theme, rep, paths, from_ym, to_ym, client)
    return run_with_com(com, core, rep)


def open_task_folder(namespace, task):
    """走査対象フォルダを取得する。通常は GetFolderFromID(EntryID, StoreID)。
    ルートのEntryIDが空のときは GetStoreFromID(StoreID).GetRootFolder()、取れなければ列挙時のストア参照を使う。"""
    if task["feid"]:
        return call_by_id(namespace.GetFolderFromID, task["feid"], task["fsid"])
    folder = None
    if task["fsid"]:
        try:
            folder = namespace.GetStoreFromID(task["fsid"]).GetRootFolder()
        except Exception:
            folder = None
    if folder is None and task.get("store_ref") is not None:
        folder = task["store_ref"].GetRootFolder()
    if folder is None:
        raise LookupError("ルートフォルダを取得できません")
    return folder


def _scan_core(args, theme, rep, paths, from_ym, to_ym, client):
    t_all = time.perf_counter()
    errors, stats = {}, {"excluded_folders": 0}
    start_dt, end_dt = period_bounds(from_ym, to_ym)
    only_words, skip_words = split_words(args.stores), split_words(args.skip_stores)
    rep.log("=" * 78)
    rep.log(f"📚 過去スレッド台帳 S1 第1段: 走査 (v{SCRIPT_VERSION}) 読み取り専用")
    rep.log(f"   実行日時: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    rep.log(f"   期間: {ym_label(from_ym)} 〜 {ym_label(to_ym)} / テーマ: {theme['name']} (版 {theme['theme_version'] or '-'})")
    rep.log("   🔒 件名・氏名・アドレスはコンソール/ログに出しません。Recipients・Body・Attachmentsは第1段では読みません。")
    rep.log("=" * 78)
    try:
        outlook = client.Dispatch("Outlook.Application")
        namespace = outlook.GetNamespace("MAPI")
        com_stores = namespace.Stores
        n_stores = int(com_stores.Count)
    except Exception as e:
        rep.log(f"❌ Outlookに接続できません: {error_kind(e)}  (Outlookを起動してから再実行してください)")
        return 2
    rep.log(f"📦 検出ストア数: {n_stores}")
    # 1) ストア情報（COM参照は保持しない。識別は StoreID(+FilePath)）
    infos = []
    for i in range(1, n_stores + 1):
        try:
            st = com_stores.Item(i)
            display = safe_attr(st, "DisplayName", "") or "(名前取得不可)"
            store_id = safe_attr(st, "StoreID", "") or ""
            fpath = safe_attr(st, "FilePath", "") or ""
        except Exception as e:
            add_error(errors, f"ストア取得:{error_kind(e)}")
            rep.log(f"   ⚠️ [{i}] ストア取得エラー: {error_kind(e)}")
            continue
        st = None
        ident = cache_ident({"store_id": store_id, "file_path": fpath, "store_name": display})
        infos.append({"i": i, "display": display, "store_id": store_id, "fpath": fpath, "ident": ident})
    labels = assign_store_labels([(x["ident"], x["display"], x["i"]) for x in infos])
    # 2) ストアごとにフォルダを列挙（フォルダは EntryID/StoreID だけ保存）
    tasks = []
    for info in infos:
        i, display, label = info["i"], info["display"], labels[info["ident"]]
        kind_label = store_kind_label(display, info["fpath"])
        if not store_ok(display, label, only_words, skip_words, theme["exclude_stores"]):
            rep.log(f"   ⏭️ [{i}] {show_store(label)} ({kind_label}): 対象外（--stores/--skip-stores/exclude_stores）")
            continue
        t0 = time.perf_counter()
        folders = None
        st_keep = None
        try:
            st = com_stores.Item(i)
            root = st.GetRootFolder()
            folders = enumerate_mail_folders(root, info["store_id"], theme["exclude_folder_names"],
                                             theme["exclude_folder_prefixes"], errors, stats)
            if any(f.get("is_root") and not f["eid"] for f in folders):
                st_keep = st      # ルートのEntryIDが空のときだけ、走査時のフォールバック用にストア参照を保持する
        except Exception as e:
            add_error(errors, f"ルート取得:{error_kind(e)}")
            rep.log(f"   ❌ [{i}] {show_store(label)}: ルートフォルダ取得エラー {error_kind(e)}")
            continue
        finally:
            st = root = None
        rep.log(f"   📬 [{i}] {show_store(label)} ({kind_label}): メールフォルダ {len(folders)}個 "
                f"(総{sum(f['total'] for f in folders)}件) 列挙 {format_duration(time.perf_counter() - t0)}")
        for f in folders:
            if not f["eid"] and not f.get("is_root"):
                add_error(errors, "フォルダEntryIDなし")
                continue
            tasks.append({"store_ref": (st_keep if not f["eid"] else None), "sidx": i, "store_name": display, "store_label": label, "store_id": info["store_id"],
                          "file_path": info["fpath"], "folder_path": f["path"], "total": f["total"],
                          "feid": f["eid"], "fsid": f["sid"], "label": kind_label})
    com_stores = None
    # 3) 計画: キャッシュの有無で skip / resume / fresh を決める（再取得は別ファイル .part に書いて完了時に置換）
    req = [ym_label(from_ym), ym_label(to_ym)]
    for task in tasks:
        key = (f"{task['store_id']}|{task['file_path']}" if (task["store_id"] or task["file_path"])
               else task["store_name"])
        task["cache_path"] = folder_cache_path(paths, task["sidx"], key, task["folder_path"])
        task["part_path"] = task["cache_path"] + ".part"
        main = load_json(task["cache_path"])
        part = load_json(task["part_path"])
        if (isinstance(part, dict) and not part.get("complete") and part.get("last_received")
                and list(part.get("range") or []) == req):
            task.update(action="resume", cache=part, write_path=task["part_path"])
        else:
            action = decide_cache_action(main, from_ym, to_ym, args.rescan)
            task["action"] = action
            task["cache"] = main if action == "resume" else None
            replacing = action == "fresh" and isinstance(main, dict)
            task["write_path"] = task["part_path"] if replacing else task["cache_path"]
    todo_items = sum(t["total"] for t in tasks if t["action"] != "skip")
    n_skip = sum(1 for t in tasks if t["action"] == "skip")
    rep.log(f"📋 走査計画: フォルダ {len(tasks)}個（完了済みスキップ {n_skip} / 再開 "
            f"{sum(1 for t in tasks if t['action'] == 'resume')} / 新規 {sum(1 for t in tasks if t['action'] == 'fresh')}）"
            f" 対象アイテム約{todo_items}件")
    rep.log(f"⏱️ 第1段の見積り: 約{todo_items}件 × {int(STAGE1_EST_SEC_PER_ITEM * 1000)}ms ≈ "
            f"{format_duration(todo_items * STAGE1_EST_SEC_PER_ITEM)}（実機の実測値。実行中は実測で補正して残り時間を表示）")
    done_items, scan_secs, interrupted = 0, 0.0, False
    summary = {"scanned": 0, "failed": 0, "mail": 0, "non_mail": 0}
    k = 0
    for task in tasks:
        k += 1
        if task["action"] == "skip":
            continue
        rate = (scan_secs / done_items) if done_items else None
        remain_items = sum(t["total"] for t in tasks[k - 1:] if t["action"] != "skip")
        eta = f" / 残り約{format_duration(rate * remain_items)}" if rate else ""
        rep.log(f"⏳ [{k}/{len(tasks)}] [{task['sidx']}] {show_folder(task['folder_path'])} (総{task['total']}件) "
                f"{'再開' if task['action'] == 'resume' else '新規'}{eta}")
        t0 = time.perf_counter()
        state = new_scan_state(task["cache"] if task["action"] == "resume" else None)
        folder = None
        try:
            folder = open_task_folder(namespace, task)

            def cp(s_, _task=task):
                write_json_atomic(_task["write_path"], make_cache_payload(_task, s_, from_ym, to_ym, False))
                rep.log(f"   💾 チェックポイント保存 ({len(s_['records'])}件)")
            scan_folder_items(folder, start_dt, end_dt, state, resume=(task["action"] == "resume"),
                              checkpoint_cb=cp, checkpoint_every=CHECKPOINT_EVERY, total_hint=task["total"])
            if state["complete"]:
                write_json_atomic(task["cache_path"], make_cache_payload(task, state, from_ym, to_ym, True))
                if os.path.exists(task["part_path"]):
                    try:
                        os.remove(task["part_path"])
                    except OSError:
                        pass
            else:
                # 途中で止めたときは、完了済みの旧キャッシュを置換しない（書き込み先は .part または未完了の新規キャッシュ）
                write_json_atomic(task["write_path"], make_cache_payload(task, state, from_ym, to_ym, False))
            for kind, n in state["errors"].items():
                errors[kind] = errors.get(kind, 0) + n
            summary["scanned"] += 1
            summary["mail"] += len(state["records"])
            summary["non_mail"] += state["non_mail"]
            sec = time.perf_counter() - t0
            done_items += max(state["processed"], 1)
            scan_secs += sec
            rep.log(f"   {'✅' if state['complete'] else '⚠️'} メール{len(state['records'])}件 "
                    f"(メール以外スキップ {state['non_mail']}件, エラー {sum(state['errors'].values())}件) "
                    f"所要 {format_duration(sec)}"
                    + ("" if state["complete"] else " ※読めないアイテムが多く途中で止めました（再実行で再開します）"))
        except KeyboardInterrupt:
            try:
                write_json_atomic(task["write_path"], make_cache_payload(task, state, from_ym, to_ym, False))
            except Exception:
                pass
            for kind, n in state["errors"].items():
                errors[kind] = errors.get(kind, 0) + n
            interrupted = True
            rep.log("")
            rep.log("⛔ Ctrl+C を受け付けました。ここまでの取得分を保存して終了します（再実行で続きから再開します）。")
            break
        except Exception as e:
            summary["failed"] += 1
            add_error(errors, f"フォルダ走査:{error_kind(e)}")
            rep.log(f"   ❌ フォルダ走査エラー: {error_kind(e)}（このフォルダはスキップして続行）")
        finally:
            folder = None     # フォルダごとにCOM参照を解放する
            task["store_ref"] = None
    namespace = outlook = None
    rep.log("")
    rep.log("🧾 [第1段 総括]" + ("（途中まで）" if interrupted else ""))
    rep.log(f"   フォルダ: 走査 {summary['scanned']} / 完了済みスキップ {n_skip} / 失敗 {summary['failed']}"
            f" / 除外フォルダ {stats['excluded_folders']}個")
    rep.log(f"   取得メール {summary['mail']}件 / メール以外スキップ {summary['non_mail']}件"
            f" / 所要 {format_duration(time.perf_counter() - t_all)}")
    if errors:
        rep.log(f"   ❌ エラー種別（{sum(errors.values())}件）:")
        for line in summarize_errors(errors):
            rep.log(f"      {line}")
    else:
        rep.log("   ✅ エラーなし")
    return 130 if interrupted else 0


# ---------- 第2段 ----------
STAGE2_HARD_ERRORS = ("Recipients:", "Recipient取得:", "Attachments:", "Body:")


def read_stage2_item(item, theme, no_body, errs_out):
    """メール1通の第2段情報。各項目は独立に try/except（失敗は種別だけ集計）。本文は返さない。
    戻り値: {'s','to','cc','att','imid','bh','hits'}"""
    errs = {}
    out = {"s": "", "to": [], "cc": [], "att": [], "imid": "", "bh": None, "hits": {}, "retry": False}
    # 送信者SMTP
    try:
        v = item.PropertyAccessor.GetProperty(PR_SENDER_SMTP)
        if isinstance(v, str) and "@" in v:
            out["s"] = v.strip().lower()
    except Exception as e:
        add_error(errs, f"送信者SMTP(PA):{error_kind(e)}")
    if not out["s"]:
        try:
            out["s"] = smtp_like(item.SenderEmailAddress)
        except Exception as e:
            add_error(errs, f"SenderEmailAddress:{error_kind(e)}")
    if not out["s"]:
        try:
            if str(item.SenderEmailType).upper() == "EX":
                out["s"] = (item.Sender.GetExchangeUser().PrimarySmtpAddress or "").strip().lower()
        except Exception as e:
            add_error(errs, f"送信者SMTP(EX):{error_kind(e)}")
    # 宛先
    try:
        recs = item.Recipients
        n = int(recs.Count)
    except Exception as e:
        add_error(errs, f"Recipients:{error_kind(e)}")
        recs, n = None, 0
    for i in range(1, n + 1):
        try:
            r = recs.Item(i)
        except Exception as e:
            add_error(errs, f"Recipient取得:{error_kind(e)}")
            continue
        rtype = safe_attr(r, "Type", 1)
        name = safe_attr(r, "Name", "") or ""
        addr = ""
        try:
            v = r.PropertyAccessor.GetProperty(PR_RECIP_SMTP)
            if isinstance(v, str) and "@" in v:
                addr = v.strip().lower()
        except Exception as e:
            add_error(errs, f"宛先SMTP(PA):{error_kind(e)}")
        if not addr:
            addr = smtp_like(safe_attr(r, "Address", ""))
        if not addr:
            try:
                addr = (r.AddressEntry.GetExchangeUser().PrimarySmtpAddress or "").strip().lower()
            except Exception as e:
                add_error(errs, f"宛先SMTP(EX):{error_kind(e)}")
        entry = {"n": name, "a": addr}
        if rtype == 1:
            out["to"].append(entry)
        elif rtype == 2:
            out["cc"].append(entry)
    # 添付ファイル名
    try:
        atts = item.Attachments
        for j in range(1, int(atts.Count) + 1):
            try:
                fn = atts.Item(j).FileName
                if fn:
                    out["att"].append(str(fn))
            except Exception as e:
                add_error(errs, f"添付名:{error_kind(e)}")
    except Exception as e:
        add_error(errs, f"Attachments:{error_kind(e)}")
    # インターネットメッセージID
    try:
        v = item.PropertyAccessor.GetProperty(PR_INTERNET_MESSAGE_ID)
        out["imid"] = v if isinstance(v, str) else ""
    except Exception as e:
        add_error(errs, f"MessageID:{error_kind(e)}")
    # 本文照合（語ごとの件数のみ保持）
    if not no_body:
        try:
            body = item.Body
            if isinstance(body, str):
                out["hits"] = scan_text_hits(body[:MAX_BODY_CHARS], theme["hit_table"])
                out["bh"] = theme["dict_hash"]
        except Exception as e:
            add_error(errs, f"Body:{error_kind(e)}")
    # 開けたが一部の重要項目が読めなかったレコードは、次回の実行で再取得する
    out["retry"] = any(k.startswith(STAGE2_HARD_ERRORS) for k in errs)
    for k, v in errs.items():
        errs_out[k] = errs_out.get(k, 0) + v
    return out


def run_stage2(args, theme, rep, paths, targets, stage2, com=None):
    """第2段: 対象メールを GetItemFromID で開き、結果を stage2(dict) とキャッシュに追加する。
    開始前に対象件数と概算時間を表示し、--stage2-max を超える場合は --yes が無ければ中止する（戻り値 3）。"""
    todo = [m for m in targets if stage2_needed(m, stage2, theme, args.no_body)]
    est = len(todo) * STAGE2_EST_SEC
    rep.log("")
    rep.log(f"🔬 第2段の対象: {len(todo)}通（スコープ {args.stage2_scope}）/ 概算 {format_duration(est)}"
            f"（1通{STAGE2_EST_SEC}秒の仮置き。実行中は実測で補正して表示）")
    limit = getattr(args, "stage2_max", STAGE2_MAX_DEFAULT)
    if len(todo) > limit and not getattr(args, "yes", False):
        rep.log(f"⛔ 対象が --stage2-max ({limit}通) を超えるため第2段を中止しました。")
        rep.log("   ・まず試走: --stage2-scope candidate（候補スレッドだけ開く）")
        rep.log("   ・全部実行する場合: --yes を付ける（または --stage2-max を増やす）")
        rep.log("   ・第2段なしで進める場合: --no-stage2")
        return 3

    def core(client):
        return _stage2_core(args, theme, rep, paths, targets, stage2, client)
    return run_with_com(com, core, rep)


def stage2_needed(mail, stage2, theme, no_body):
    rec = stage2.get(stage2_key(mail["locs"][0]["sid"], mail["locs"][0]["eid"]))
    if rec is None or rec.get("retry"):
        return True            # 未取得、または前回エラーがあったレコードは取り直す
    if no_body:
        return False
    return rec.get("bh") != theme["dict_hash"]


def _stage2_core(args, theme, rep, paths, targets, stage2, client):
    t0 = time.perf_counter()
    errors = {}
    todo = [m for m in targets if stage2_needed(m, stage2, theme, args.no_body)]
    rep.log("")
    rep.log(f"🔬 第2段: 対象メール {len(targets)}通（取得済みスキップ {len(targets) - len(todo)} / 今回 {len(todo)}）"
            f"{' / 本文照合なし' if args.no_body else ''}")
    if not todo:
        return 0
    try:
        namespace = client.Dispatch("Outlook.Application").GetNamespace("MAPI")
    except Exception as e:
        rep.log(f"❌ Outlookに接続できません: {error_kind(e)}")
        return 2
    done, failed, interrupted = 0, 0, False

    def save():
        write_json_atomic(paths["stage2_cache"], stage2)

    try:
        for m in todo:
            item, ok = None, False
            for loc in m["locs"]:
                try:
                    item = call_by_id(namespace.GetItemFromID, loc["eid"], loc["sid"])
                    ok = item is not None
                    if ok:
                        break
                except Exception as e:
                    add_error(errors, f"GetItemFromID:{error_kind(e)}")
            if not ok:
                failed += 1
                continue
            rec = read_stage2_item(item, theme, args.no_body, errors)
            item = None
            stage2[stage2_key(m["locs"][0]["sid"], m["locs"][0]["eid"])] = rec
            done += 1
            if done % STAGE2_SAVE_EVERY == 0:
                try:
                    save()
                except Exception as e:
                    rep.log(f"⚠️ 第2段キャッシュの途中保存に失敗: {error_kind(e)}（続行します）")
                rate = (time.perf_counter() - t0) / done
                rep.log(f"   ⏳ 第2段 {done}/{len(todo)}通 残り約{format_duration(rate * (len(todo) - done))}")
    except KeyboardInterrupt:
        interrupted = True
        rep.log("⛔ Ctrl+C を受け付けました。第2段のここまでの結果を保存して終了します。")
    finally:
        try:
            save()      # 途中保存や取得で例外が出ても、最終保存は必ず試みる
        except Exception as e:
            rep.log(f"❌ 第2段キャッシュを保存できません: {error_kind(e)}")
    rep.log(f"   第2段: 取得 {done}通 / 開けず {failed}通 / 所要 {format_duration(time.perf_counter() - t0)}"
            + ("（途中まで）" if interrupted else ""))
    if errors:
        rep.log(f"   ❌ 第2段エラー種別（{sum(errors.values())}件）:")
        for line in summarize_errors(errors):
            rep.log(f"      {line}")
    return 130 if interrupted else 0


# ---------- --open ----------
def find_thread_in_ledger(output_dir, thread_id):
    """最新の台帳CSVから thread_id の (EntryID, StoreID) を引く。見つからなければ None。"""
    files = sorted(glob.glob(os.path.join(output_dir, "thread_ledger_*.csv")), key=os.path.getmtime, reverse=True)
    for p in files:
        try:
            with open(p, "r", encoding="utf-8-sig", newline="") as fh:
                for row in csv.DictReader(fh):
                    if (row.get("thread_id") or "").lstrip("'") == thread_id:
                        eid = (row.get("最新メールEntryID") or "").lstrip("'")
                        sid = (row.get("最新メールStoreID") or "").lstrip("'")
                        if eid and sid:
                            return eid, sid
        except Exception:
            continue
    return None


def find_thread_in_cache(paths, thread_id, from_ym, to_ym):
    caches = select_caches(load_all_caches(paths["scan_cache"]), [], [], [], DEFAULT_EXCLUDE_FOLDER_NAMES,
                           DEFAULT_EXCLUDE_FOLDER_PREFIXES)
    mails, _s = merge_caches_to_mails(caches, from_ym, to_ym)
    for t in build_threads(mails):
        if t["id"] == thread_id:
            loc = t["mails"][-1]["locs"][0]
            return loc["eid"], loc["sid"]
    return None


def run_open(args, rep, paths, from_ym, to_ym, com=None):
    """最新メールを Outlook で開く（件名で絞り込み＋選択。失敗時は Display）。
    指定は thread_id（--open / --open-url）または entry_id+store_id（--entry-id/--store-id）。読み取りのみ。"""
    if getattr(args, "entry_id", ""):
        eid, sid = args.entry_id, getattr(args, "store_id", "") or ""
        if not HEX_ID_RE.match(eid) or (sid and not HEX_ID_RE.match(sid)):
            rep.log("❌ --entry-id/--store-id の形式が不正です（16進文字列のみ）")
            return 2
    else:
        tid = args.open
        if not LEDGER_THREAD_ID_RE.match(tid or ""):
            rep.log("❌ thread_id の形式が不正です（英数字・_・- の4〜64文字のみ）")
            return 2
        found = find_thread_in_ledger(paths["output"], tid) or find_thread_in_cache(paths, tid, from_ym, to_ym)
        if not found:
            rep.log("❌ 指定した thread_id が台帳CSV・キャッシュに見つかりません")
            return 2
        eid, sid = found
        if not HEX_ID_RE.match(eid or "") or (sid and not HEX_ID_RE.match(sid)):
            rep.log("❌ 台帳/キャッシュ由来の EntryID/StoreID の形式が不正です（16進文字列のみ）")
            return 2

    def core(client):
        try:
            mode = open_thread_in_outlook(client, eid, sid, "", rep)
            rep.log("✅ Outlookで件名絞り込み表示にしました" if mode == "explorer" else "✅ Outlookでメールを開きました")
            return 0
        except Exception as e:
            rep.log(f"❌ メールを開けません: {error_kind(e)}")
            return 2
    return run_with_com(com, core, rep)


# ============================================================
# 評価〜出力
# ============================================================
def load_stage2(paths):
    d = load_json(paths["stage2_cache"], {})
    return d if isinstance(d, dict) else {}


def evaluate_all(caches, theme, stage2, from_ym, to_ym, only_words=(), skip_words=()):
    """キャッシュ → 重複除去 → スレッド化 → 評価。戻り値 dict(mails, threads, evals, stats, caches)"""
    sel = select_caches(caches, list(only_words), list(skip_words), theme["exclude_stores"],
                        theme["exclude_folder_names"], theme["exclude_folder_prefixes"])
    mails, stats = merge_caches_to_mails(sel, from_ym, to_ym)
    threads = build_threads(mails)
    memo = {}
    evals = [evaluate_thread(t, theme, stage2, memo) for t in threads]
    return {"mails": mails, "threads": threads, "evals": evals, "stats": stats, "caches": sel}


def log_evaluation_summary(rep, res, label):
    s = summarize_conditions(res["evals"])
    st = res["stats"]
    rep.log(f"📊 [{label}] 取得メール(重複除去前) {st['raw_total']}通 → 重複除去後 {st['unique_total']}通 → スレッド {s['total']}件")
    rep.log(f"   条件A(アンカー関与) {s['A']} / 条件B(本人関与) {s['B']} / A∧B {s['AB']} / "
            f"条件C {s['C']} (C(i)関係者 {s['C_i']} / C(ii)キーワード {s['C_ii']})")
    rep.log(f"   ✅ 候補スレッド(A∧B∧C): {s['candidate']}件")
    rep.log(f"   🔧 調整用: Aのみ不成立 {s['only_A_fails']} / Bのみ不成立 {s['only_B_fails']} / Cのみ不成立 {s['only_C_fails']}")
    return s


def log_dup_summary(rep, st):
    rep.log(f"🔁 重複除去: ストア間の重複 {st['cross_store']}件 / 同一ストア内の重複 {st['same_store']}件")
    if st["dup_by_month"]:
        rep.log("   ストア間重複の月別: " + ", ".join(f"{ym_label(ym)}={n}" for ym, n in sorted(st["dup_by_month"].items())))
    else:
        rep.log("   ストア間の重複はありません")


def warn_stale_stage2(rep, stage2, theme):
    n = count_stale_stage2(stage2, theme)
    if n:
        rep.log(f"⚠️ 辞書が変わる前に本文照合した第2段レコードが {n}通あります。評価では本文ヒットを使っていません"
                "（--no-body なしで第2段を再実行すると取り直します）")


def run_pipeline(args, theme, rep, paths, from_ym, to_ym, com=None):
    """評価 → 第2段 → 再評価 → 出力。戻り値: 終了コード"""
    caches = load_all_caches(paths["scan_cache"])
    if not caches:
        rep.log("❌ 走査キャッシュがありません。先に走査（--evaluate-only なし）を実行してください。")
        return 2
    only_words, skip_words = split_words(args.stores), split_words(args.skip_stores)
    stage2 = load_stage2(paths)
    warn_stale_stage2(rep, stage2, theme)
    res = evaluate_all(caches, theme, stage2, from_ym, to_ym, only_words, skip_words)
    if args.check_subjects:
        # 再現率チェックは評価済みキャッシュを読むだけ（走査・第2段・CSV出力はしない）
        return run_check_subjects(args, rep, paths, res)
    rep.log("")
    log_evaluation_summary(rep, res, "第1段のみ")
    log_dup_summary(rep, res["stats"])
    code = 0
    if not args.evaluate_only and not args.no_stage2:
        targets = select_stage2_targets(res["threads"], res["evals"], args.stage2_scope)
        code = run_stage2(args, theme, rep, paths, targets, stage2, com)
        if code in (0, 130):
            warn_stale_stage2(rep, stage2, theme)
            res = evaluate_all(caches, theme, stage2, from_ym, to_ym, only_words, skip_words)
            log_evaluation_summary(rep, res, "第2段反映後")
    elif stage2:
        rep.log(f"   ℹ️ 既存の第2段キャッシュ {len(stage2)}通分を評価に使用しました")
    write_outputs(args, theme, rep, paths, res, stage2, from_ym, to_ym)
    return code


def write_outputs(args, theme, rep, paths, res, stage2, from_ym, to_ym):
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    out = paths["output"]
    # 判定の保持（既存は上書きしない。新規行だけ追加）
    jpath = judgments_path(paths, theme["name"])
    existing, jstatus = load_judgments(jpath)
    cand_ids = [t["id"] for t, e in zip(res["threads"], res["evals"]) if e["candidate"]]
    if getattr(args, "include_no_anchor", False):
        cand_ids += [t["id"] for t, e in zip(res["threads"], res["evals"]) if is_reference_thread(e)]
    if jstatus == "broken":
        # 存在するのに読めない（壊れている/空/形が不正）: 上書きして判定を全消去しない
        bak = backup_file(jpath)
        rep.log("⚠️ 判定ファイルを読めません（壊れている・空・形が不正）。上書きせず、判定の書き出しをスキップします。"
                + (f" 退避コピー: {bak}" if bak else " ※退避コピーも作れませんでした")
                + " / 今回のCSVの確認結果・メモ・問題の種類は空欄です。ファイルを直してから再評価してください。")
        merged = {}
    else:
        merged, added = merge_judgments(existing, cand_ids)
        if added or jstatus == "new":
            try:
                write_json_atomic(jpath, merged)
            except Exception as e:
                rep.log(f"⚠️ 判定ファイルを保存できません: {error_kind(e)}")
    path = os.path.join(out, f"thread_ledger_{theme['name']}_{ts}.csv")
    ledger_rows = build_ledger_rows(res["threads"], res["evals"], merged, theme,
                                    include_no_anchor=getattr(args, "include_no_anchor", False))
    write_csv(path, LEDGER_COLUMNS, ledger_rows)
    rep.log(f"💾 台帳CSV: {len(ledger_rows)}スレッド → {path}")
    ox = import_openpyxl()
    if ox is None:
        rep.log("ℹ️ openpyxl が無いため .xlsx は出力しません（CSVのみ）。Excelでリンクを使うには `pip install openpyxl` を実行してください")
    else:
        xpath = os.path.splitext(path)[0] + ".xlsx"
        try:
            write_ledger_xlsx(xpath, LEDGER_COLUMNS, ledger_rows, ox)
            rep.log(f"💾 台帳Excel（件名リンク ledger:<thread_id>）→ {xpath}")
        except Exception as e:
            rep.log(f"⚠️ .xlsx を出力できません: {error_kind(e)}（CSVは出力済み）")
    months = month_range(from_ym, to_ym)
    store_names = []
    for c in res["caches"]:
        lab = c.get("store_label") or c.get("store_name")
        if lab not in store_names:
            store_names.append(lab)
    cpath = os.path.join(out, f"coverage_{ts}.csv")
    write_csv(cpath, COVERAGE_COLUMNS,
              build_coverage_rows(res["stats"], res["mails"], res["threads"], res["evals"], stage2, theme, months,
                                  store_names))
    zero = [ym_label(ym) for ym in months if not any(res["stats"]["raw_by_store_month"].get((s, ym), 0) for s in store_names)]
    rep.log(f"💾 カバレッジCSV → {cpath}")
    if zero:
        rep.log(f"   ⚠️ 全ストア0件の月（未取得の恐れ・取得元ゼロ・要確認）: {', '.join(zero)}")
    for sname, zs in store_zero_months(res["stats"], months, store_names).items():
        rep.log(f"   ℹ️ {show_store(sname)}: 0件の月 {len(zs)}/{len(months)}"
                + (f"（{ym_label(zs[0])}〜{ym_label(zs[-1])}の範囲に分布）" if 0 < len(zs) < len(months) else ""))
    if args.suggest:
        rows = suggest_participants(res["threads"], res["evals"], theme, stage2)
        spath = os.path.join(out, f"suggest_participants_{ts}.csv")
        write_csv(spath, ["表示名", "SMTP", "出現スレッド数"], rows)
        rep.log(f"💾 新メンバー候補CSV: {len(rows)}人 → {spath}")
    if args.suggest_terms:
        fg, bg = candidate_term_docs(res["threads"], res["evals"], theme, stage2)
        rows = rank_candidate_terms(fg, bg, theme["hit_table"].keys())
        tpath = os.path.join(out, f"candidate_terms_{ts}.csv")
        write_csv(tpath, ["語", "候補スレッド数", "全体スレッド数", "候補での出現率", "全体での出現率", "lift", "スコア", "既に辞書にある"],
                  rows)
        rep.log(f"💾 辞書候補語CSV: {len(rows)}語 → {tpath}（件名・添付名から抽出。本文は保存していないため対象外）")


def run_check_subjects(args, rep, paths, res):
    path = (os.path.join(paths["data"], DEFAULT_CALIBRATION_NAME)
            if args.check_subjects == CHECK_SUBJECTS_DEFAULT else args.check_subjects)
    try:
        queries = load_calibration_subjects(path)
    except Exception as e:
        rep.log(f"❌ 正解件名ファイルを読めません: {error_kind(e)}")
        return 2
    results = classify_check_subjects(queries, res["threads"], res["evals"])
    rep.console("")
    rep.console(f"🎯 正解件名の再現率チェック（{len(queries)}件）※この出力のみ件名を表示します（ログには残しません）")
    for i, r in enumerate(results, 1):
        for line in format_check_result(i, len(results), r):
            rep.console(line)
    n_found = sum(1 for r in results if r["status"] == "found")
    n_not = sum(1 for r in results if r["status"] == "found_not_candidate")
    n_miss = sum(1 for r in results if r["status"] == "missing")
    rep.log(f"🎯 再現率チェック: 候補 {n_found} / あるが候補外 {n_not} / キャッシュに無い {n_miss} （全{len(results)}件）")
    return 0


# ============================================================
# エントリポイント
# ============================================================
def parse_args(argv=None):
    ap = argparse.ArgumentParser(
        description="過去スレッド台帳 S1: 走査エンジン + 台帳CSV（読み取り専用）",
        epilog="注意: 台帳CSVをExcelで開くと、thread_id が数値（例 12345e678901 → 指数表記）に変換されることがあります。"
               "判定の取込は xlsx（台帳と同時に出力される .xlsx）を編集して --import-judgments に渡すことを推奨します。",
        allow_abbrev=False)
    ap.add_argument("--theme", default="", help="テーマJSON（既定 json/thread_ledger/themes/trade_import.json）")
    ap.add_argument("--stores", default="", help="対象ストア名の一部（カンマ区切り）")
    ap.add_argument("--skip-stores", default="", help="除外するストア名の一部（カンマ区切り）")
    ap.add_argument("--from", dest="from_ym", default="", help="開始月 YYYY-MM")
    ap.add_argument("--to", dest="to_ym", default="", help="終了月 YYYY-MM")
    ap.add_argument("--rescan", action="store_true", help="完了済みフォルダも再取得する")
    ap.add_argument("--evaluate-only", action="store_true", help="COMを使わずキャッシュだけから再評価する")
    ap.add_argument("--no-stage2", action="store_true", help="第2段（COM）を省略する")
    ap.add_argument("--no-body", action="store_true", help="第2段で本文照合を省略する")
    ap.add_argument("--stage2-scope", default="ab", choices=["ab", "candidate", "either"],
                    help="第2段の対象: ab(条件AかつB) / candidate(候補のみ) / either(AまたはB)")
    ap.add_argument("--suggest", action="store_true", help="新メンバー候補CSVを出す")
    ap.add_argument("--suggest-terms", "--candidate-terms", dest="suggest_terms", action="store_true",
                    help="辞書候補語ランキングCSVを出す")
    ap.add_argument("--stage2-max", type=int, default=STAGE2_MAX_DEFAULT,
                    help="第2段の対象がこの通数を超えると --yes が無ければ中止する（既定 5000）")
    ap.add_argument("--yes", action="store_true", help="第2段の通数上限を超えても実行する")
    ap.add_argument("--mask-names", action="store_true",
                    help="コンソール/ログのストア名・フォルダ名を連番に置換する（CSVには元の名前を出す）")
    ap.add_argument("--check-subjects", nargs="?", const=CHECK_SUBJECTS_DEFAULT, default="",
                    help="正解件名ファイル（1行1件名）で再現率をチェックする")
    ap.add_argument("--open", default="", help="thread_id を指定して最新メールをOutlookで件名絞り込み表示にする")
    ap.add_argument("--entry-id", default="", help="--open の代わりに EntryID を直接指定する（16進文字列）")
    ap.add_argument("--store-id", default="", help="--entry-id と組み合わせる StoreID（16進文字列）")
    ap.add_argument("--open-url", default=None, help="ledger:<thread_id> を受け取って開く（プロトコル登録後に Excel から呼ばれる）")
    ap.add_argument("--register-protocol", action="store_true",
                    help="HKCU に ledger: プロトコルを登録する（オプトイン。管理者権限不要）")
    ap.add_argument("--unregister-protocol", action="store_true", help="登録した ledger: プロトコルを削除する")
    ap.add_argument("--import-judgments", default="",
                    help="Excelで編集した台帳（xlsx推奨。CSVはExcelでthread_idが数値化されることがある）の確認結果・メモ・問題の種類を "
                         "judgments JSON に取り込む（台帳に存在するIDのみ。取込前に判定JSONを .bak_ に退避）")
    ap.add_argument("--include-no-anchor", action="store_true",
                    help="アンカー無し（B∧C成立でA不成立）の参考候補も台帳に出す")
    ap.add_argument("--output-dir", default="", help="CSV/ログの出力先（既定 mail_reports/）")
    ap.add_argument("--data-dir", default="", help="台帳データの置き場（既定 json/thread_ledger/）")
    return ap.parse_args(argv)


def resolve_period(args, theme):
    """期間の決定: コマンドライン > theme.period > 既定。"""
    f = parse_ym(args.from_ym) if args.from_ym else (theme["period_from"] if theme else None) or parse_ym(DEFAULT_FROM)
    t = parse_ym(args.to_ym) if args.to_ym else (theme["period_to"] if theme else None) or parse_ym(DEFAULT_TO)
    month_range(f, t)   # 順序の検証
    return f, t


def show_message_box(text, title="過去スレッド台帳"):
    """Windows のメッセージボックスを出す（pythonw では print が見えないため）。非Windows・失敗時は何もしない。"""
    try:
        import ctypes  # noqa: WPS433
        ctypes.windll.user32.MessageBoxW(0, str(text), str(title), 0x10)
    except Exception:
        pass


def has_open_url_token(argv):
    """argv に '--open-url' または '--open-url=' で始まる要素が1つでもあるか。"""
    return any(a == "--open-url" or str(a).startswith("--open-url=") for a in argv)


def exact_open_url_argv(argv, open_url):
    """--open-url の呼び出しは ['--open-url', <値>] だけを許す（値に " が入って他のオプションが注入されるのを防ぐ）。"""
    return isinstance(open_url, str) and list(argv) == ["--open-url", open_url]


def write_failure_log(path, lines):
    """--open-url が失敗したときだけログを書く（成功時はクリックのたびにログが増えないよう書かない）。"""
    try:
        os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
        with open(path, "w", encoding="utf-8-sig", newline="\n") as fh:
            fh.write("\n".join(lines) + "\n")
        return path
    except Exception:
        return None


def main_open_url(args, argv_used, com):
    """Excel のリンク（ledger:<thread_id>）から呼ばれる入口。pythonw では画面が出ないので、失敗時はメッセージボックスを出す。
    - 呼び出しは ['--open-url', <値>] だけを許す（他のオプションが注入されていれば何もせず終了コード2）。
    - --data-dir / --output-dir は無視して既定パスを使う。成功時はログを作らず、失敗時だけログを書く。"""
    if not exact_open_url_argv(argv_used, args.open_url):
        show_message_box("不正な呼び出しのため、何もせず終了しました。")
        return 2
    paths = make_paths(None, None)
    rep = Reporter(None)
    code = 2
    try:
        tid = parse_ledger_url(args.open_url)
        if tid is None:
            rep.log("❌ URLの形式が不正です（ledger:<thread_id> のみ受け付けます）")
        else:
            args.open = tid
            from_ym, to_ym = resolve_period(args, None)
            code = run_open(args, rep, paths, from_ym, to_ym, com)
    except Exception as e:
        rep.log(f"❌ 想定外のエラー: {error_kind(e)}")
        code = 1
    if code != 0:
        ts = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
        logp = write_failure_log(os.path.join(paths["output"], f"thread_ledger_open_{ts}.log"), rep.lines)
        reason = next((l.lstrip("❌ ").strip() for l in rep.lines if l.startswith("❌")), "Outlookで開けませんでした")
        show_message_box(f"台帳のメールを開けませんでした。\n{reason}" + (f"\n\n詳細ログ: {logp}" if logp else ""))
    return code


def main(argv=None, com=None, winreg_mod=None):
    try:
        sys.stdout.reconfigure(errors="replace")
    except Exception:
        pass
    argv_used = list(sys.argv[1:] if argv is None else argv)
    # --open-url を含む呼び出しは、値が空文字でも厳密検証に入る（空値＋他オプションの迂回を防ぐ二重の判定）
    url_call = has_open_url_token(argv_used)
    try:
        args = parse_args(argv_used)
    except SystemExit as e:
        if url_call:
            show_message_box("呼び出しの引数が不正なため、何もせず終了しました。")     # pythonw では無反応に見えないように知らせる
        raise
    configure_masking(args.mask_names)
    if url_call or args.open_url is not None:
        return main_open_url(args, argv_used, com)
    paths = make_paths(args.data_dir or None, args.output_dir or None)
    theme_path = args.theme or os.path.join(paths["themes"], THEME_DEFAULT_NAME)
    # 開く・プロトコル登録はテーマ不要
    no_theme = bool(args.open or args.entry_id or args.register_protocol or args.unregister_protocol)
    theme = None
    if not no_theme:
        try:
            theme = load_theme(theme_path)
        except ThemeError as e:
            print(f"❌ テーマを読み込めません: {e}")
            if not os.path.exists(theme_path):
                ex = os.path.join(os.path.dirname(os.path.abspath(__file__)), THEME_EXAMPLE_NAME)
                print(f"   example をコピーして編集してください: {ex}")
                print(f"   コピー先: {theme_path}")
            return 2
    try:
        from_ym, to_ym = resolve_period(args, theme)
    except ValueError as e:
        print(f"❌ 引数エラー: {e}")
        return 2
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    rep = Reporter(os.path.join(paths["output"], f"thread_ledger_scan_{ts}.log"))
    code = 1
    try:
        if args.register_protocol or args.unregister_protocol:
            code = run_protocol(args, rep, winreg_mod)
        elif args.open or args.entry_id:
            code = run_open(args, rep, paths, from_ym, to_ym, com)
        elif args.import_judgments:
            code = run_import_judgments(args, theme, rep, paths)
        else:
            code = 0
            if args.check_subjects:
                args.evaluate_only = True
            if not args.evaluate_only:
                code = run_scan(args, theme, rep, paths, from_ym, to_ym, com)
            if code == 0 or args.evaluate_only:
                code = run_pipeline(args, theme, rep, paths, from_ym, to_ym, com)
    except KeyboardInterrupt:
        rep.log("⛔ Ctrl+C を受け付けました。途中結果は保存済みです。")
        code = 130
    except Exception as e:
        rep.log(f"❌ 想定外のエラー: {error_kind(e)}（途中結果は保存済みです）")
        code = 1
    finally:
        rep.close()
        if rep.path:
            try:
                print(f"📝 ログ: {rep.path}")
            except Exception:
                pass
    return code


if __name__ == "__main__":
    sys.exit(main())
