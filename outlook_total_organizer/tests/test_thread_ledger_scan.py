# -*- coding: utf-8 -*-
"""過去スレッド台帳 S1 (tools/thread_ledger_scan_*.py) のテスト。

app/ のモジュールではないので importlib で tools/ から読み込む。
純粋関数（正規化・照合・重複除去・スレッド化・候補判定・スコア・CSV・判定保持・再現率チェック・候補語）と、
COM部分（フェイクを注入）を確認する。ダミーデータの氏名・アドレス・番号はすべて架空。
"""
import contextlib
import csv
import glob
import importlib.util
import io
import json
import os
import re
import sys
import tempfile
import unittest
from datetime import datetime

sys.dont_write_bytecode = True

TESTS_DIR = os.path.dirname(os.path.abspath(__file__))
TOOLS_DIR = os.path.join(os.path.dirname(TESTS_DIR), "tools")


def _load_tool():
    paths = sorted(glob.glob(os.path.join(TOOLS_DIR, "thread_ledger_scan_*.py")))
    if not paths:
        raise FileNotFoundError("tools/thread_ledger_scan_*.py がありません")
    spec = importlib.util.spec_from_file_location("thread_ledger_scan_under_test", paths[-1])
    mod = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = mod
    spec.loader.exec_module(mod)
    return mod


T = _load_tool()
ORIG_WAITS = (T.EXPLORER_SETTLE_SEC, T.SELECT_RETRY_SEC, T.SELECT_RETRIES)
ORIG_RETRY_WAITS = T.RETRY_WAITS
T.RETRY_WAITS = (0, 0)          # リトライの待ちはテストでは無し（既定値は test_resilience_defaults で確認）
T.EXPLORER_SETTLE_SEC = T.SELECT_RETRY_SEC = 0      # テストでは待たない（既定値は test_explorer_wait_defaults で確認）


def _rd(path, enc="utf-8"):
    with open(path, "r", encoding=enc) as fh:
        return fh.read()


def _rb(path):
    with open(path, "rb") as fh:
        return fh.read()


def _wr(path, text, enc="utf-8"):
    with open(path, "w", encoding=enc) as fh:
        fh.write(text)


def _rj(path):
    return json.loads(_rd(path))


def _wj(path, obj):
    _wr(path, json.dumps(obj, ensure_ascii=False))


def _rcsv(path):
    with open(path, "r", encoding="utf-8-sig", newline="") as fh:
        return list(csv.reader(fh))


def _rdict(path):
    with open(path, "r", encoding="utf-8-sig", newline="") as fh:
        return list(csv.DictReader(fh))
SECRET = "SECRET-SUBJECT-ZZ9"
SECRET_ERR = "SECRET-ERR-MESSAGE-QQ7"


# ============================================================
# テーマ・データ作成ヘルパ
# ============================================================
def make_theme(**over):
    data = {
        "theme_version": "t1",
        "owner": {"name_aliases": ["Owner Taro"], "emails": ["owner@example.com"]},
        "anchors": [
            {"label": "PM", "name_aliases": ["Sato Pm", "佐藤"], "emails": ["pm@example.com"]},
            {"label": "Admin", "name_aliases": ["Suzuki Admin", "鈴木"], "emails": []},
        ],
        "related_parties": [
            {"label": "Sachi", "kind": "person", "name_aliases": ["Tanaka Partner"], "emails": [], "domains": []},
            {"label": "Courier", "kind": "org", "name_aliases": [], "emails": [], "domains": ["courier.example.com"]},
        ],
        "keywords": {
            "通関": {"weight": 3, "strong": ["通関", "customs clearance", "HS code"], "weak": ["hold", "保留"]},
            "トラブル語": {"weight": 2, "strong": [], "weak": ["delay", "urgent", "approval request"]},
        },
        "patterns": [
            {"category": "物流", "regex": r"#\s?\d{10}", "strong": True},
            {"category": "関税", "regex": r"\b\d{4}\.\d{2}\b", "strong": False},
        ],
        "score_weights": {"subject_multiplier": 2, "weak_multiplier": 0.5, "anchor_each": 3, "related_each": 2,
                          "reply_every": 5, "reply_points": 1, "reply_cap": 5},
        "period": {"from": "2024-01", "to": "2024-12"},
    }
    data.update(over)
    return T.normalize_theme(data, "unit")


def rec(eid, subj, t, sender="", to="", cc="", cid="", addr=""):
    return {"e": eid, "s": subj, "t": t, "n": sender, "a": addr, "to": to, "cc": cc, "c": cid}


def cache(store, sidx, folder, records, sid=None, scanned_at="2024-12-31T00:00:00"):
    return {"schema": T.SCHEMA_VERSION, "store_index": sidx, "store_name": store, "store_id": sid or f"SID{sidx}",
            "folder_path": folder, "range": ["2024-01", "2024-12"], "complete": True, "records": records,
            "scanned_at": scanned_at, "counts": {"mail": len(records), "non_mail": 0, "errors": {}}}


def build(caches, theme=None, stage2=None, from_ym=(2024, 1), to_ym=(2024, 12)):
    theme = theme or make_theme()
    return T.evaluate_all(caches, theme, stage2 or {}, from_ym, to_ym), theme


def ev_of(res, i=0):
    return res["evals"][i]


def one_thread(records, theme=None, stage2=None):
    res, th = build([cache("S", 1, "\\Inbox", records)], theme, stage2)
    assert len(res["evals"]) == 1, len(res["evals"])
    return res, ev_of(res)


# ============================================================
# 期間・Restrict
# ============================================================
class TestPeriod(unittest.TestCase):
    def test_parse_ym(self):
        self.assertEqual(T.parse_ym("2023-07"), (2023, 7))
        self.assertEqual(T.parse_ym(" 2026/9 "), (2026, 9))
        for bad in ("2023", "2023-13", "2023-00", "x", None, ""):
            with self.assertRaises(ValueError):
                T.parse_ym(bad)

    def test_month_range_boundaries(self):
        self.assertEqual(len(T.month_range((2023, 7), (2026, 9))), 39)
        self.assertEqual(T.month_range((2024, 12), (2025, 1)), [(2024, 12), (2025, 1)])
        with self.assertRaises(ValueError):
            T.month_range((2025, 2), (2025, 1))

    def test_period_bounds_and_iso_filter(self):
        s, e = T.period_bounds((2023, 7), (2026, 9))
        self.assertEqual(s, datetime(2023, 7, 1))
        self.assertEqual(e, datetime(2026, 10, 1))
        self.assertEqual(T.period_bounds((2025, 12), (2025, 12))[1], datetime(2026, 1, 1))
        self.assertEqual(T.build_restrict_filter("ReceivedTime", s, e),
                         "[ReceivedTime] >= '2023-07-01 00:00' AND [ReceivedTime] < '2026-10-01 00:00'")

    def test_to_naive_dt(self):
        self.assertEqual(T.to_naive_dt(datetime(2024, 1, 2, 3, 4, 5)), datetime(2024, 1, 2, 3, 4, 5))
        self.assertIsNone(T.to_naive_dt(datetime(4501, 1, 1)))   # Outlookの「未設定」
        self.assertIsNone(T.to_naive_dt(None))
        self.assertIsNone(T.to_naive_dt(12345))


# ============================================================
# 正規化
# ============================================================
class TestNormalize(unittest.TestCase):
    def test_subject_prefix_chain(self):
        self.assertEqual(T.normalize_subject("RE: FW: Re: Fwd:  Hello World"), "hello world")
        self.assertEqual(T.normalize_subject("** internal only ** RE: FW: x"), "x")
        self.assertEqual(T.normalize_subject("FW: ** Internal Only ** Re: x"), "x")
        self.assertEqual(T.normalize_subject("AW: WG: 転送: 返信: [外部] x"), "x")
        self.assertEqual(T.normalize_subject("ＲＥ：　Ｈｅｌｌｏ　　Ｗｏｒｌｄ"), "hello world")

    def test_subject_keeps_bracket_tags_and_empty(self):
        self.assertEqual(T.normalize_subject("RE: [Export request]2024-01-01 10:00"), "[export request]2024-01-01 10:00")
        self.assertEqual(T.normalize_subject(""), "")
        self.assertEqual(T.normalize_subject(None), "")
        self.assertEqual(T.normalize_subject("RE: "), "")

    def test_subject_does_not_strip_inner_re(self):
        self.assertEqual(T.normalize_subject("Report re: status"), "report re: status")

    def test_name_matches_order_case_width(self):
        self.assertTrue(T.name_matches("Taro Sato", "SATO, Taro"))
        self.assertTrue(T.name_matches("Taro Sato", "ｔａｒｏ　ｓａｔｏ"))
        self.assertTrue(T.name_matches("Sato Taro", "Taro Sato (Dept)"))
        self.assertTrue(T.name_matches("Sato", "Sato, Hana"))
        self.assertTrue(T.name_matches("Sato", "Sato-san"))
        self.assertTrue(T.name_matches("Sato", "Mr. Sato"))
        self.assertTrue(T.name_matches("佐藤", "佐藤 太郎"))
        self.assertTrue(T.name_matches("佐藤", "太郎 佐藤さん"))
        self.assertTrue(T.name_matches("佐藤さん", "佐藤 美穂"))
        self.assertTrue(T.name_matches("佐藤", "Sato, 佐藤"))
        self.assertTrue(T.name_matches("太郎 佐藤", "佐藤, 太郎"))
        self.assertTrue(T.name_matches("佐藤太郎", "佐藤 太郎"))      # 連結が一致
        self.assertTrue(T.name_matches("佐藤 太郎", "佐藤太郎"))

    def test_name_matches_ascii_always_requires_word_end(self):
        self.assertFalse(T.name_matches("Jay", "Jaycee Smith"))
        self.assertTrue(T.name_matches("Jay", "Jay Smith"))
        self.assertFalse(T.name_matches("Chris", "Michristo"))                # 語頭境界
        self.assertFalse(T.name_matches("Chris", "Christopher Lee"))          # 4文字以上でも語末境界
        self.assertFalse(T.name_matches("Ono", "Onozuka Hana"))
        self.assertFalse(T.name_matches("Yuki", "Yukiko Tanaka"))
        self.assertFalse(T.name_matches("Sato", "Satomura Ken"))
        self.assertTrue(T.name_matches("Ono", "Ono Hana"))
        self.assertTrue(T.name_matches("Yuki", "Tanaka, Yuki"))
        self.assertFalse(T.name_matches("Taro Sato", "Taro Satomura"))

    def test_name_matches_japanese_surname_only_alias(self):
        # 姓だけの別名は、スペースの有無にかかわらずフルネームに一致する（表示名が全て非ASCIIで、残りが0〜4文字）
        self.assertTrue(T.name_matches("越智", "越智太郎"))
        self.assertTrue(T.name_matches("越智", "越智 太郎"))
        self.assertTrue(T.name_matches("木村", "木村美穂"))
        self.assertTrue(T.name_matches("木村", "木村 美穂"))
        self.assertTrue(T.name_matches("木村さん", "木村美穂"))
        self.assertTrue(T.name_matches("木村さん", "木村さん"))
        self.assertTrue(T.name_matches("田中", "田中 花子"))
        self.assertTrue(T.name_matches("田中さん", "田中"))
        self.assertTrue(T.name_matches("田中", "Tanaka, 田中さん"))

    def test_name_matches_japanese_other_surnames_do_not_match(self):
        self.assertFalse(T.name_matches("田中", "山田中"))          # 別名で始まらない
        self.assertFalse(T.name_matches("木村", "山木村"))
        self.assertFalse(T.name_matches("木村", "山木村太郎"))
        self.assertFalse(T.name_matches("佐藤", "田佐藤"))
        self.assertFalse(T.name_matches("木村", "木村ふじさわたろう"))   # 残りが5文字以上は姓+名とみなさない
        self.assertFalse(T.name_matches("木村", "木村Miho"))            # 非ASCIIだけの表示名に限る
        self.assertFalse(T.name_matches("田", "田中"))                  # 1文字の別名は完全トークン一致のみ

    def test_name_matches_known_risk_trailing_chars(self):
        # 既知のリスク（仕様）: 後ろに1〜4文字続く別姓は、姓+名と区別できず一致する
        self.assertTrue(T.name_matches("木村", "木村戸"))
        self.assertTrue(T.name_matches("田中", "田中村"))

    def test_name_matches_negative_basic(self):
        self.assertFalse(T.name_matches("", "Anyone"))
        self.assertFalse(T.name_matches("Taro Sato", ""))
        self.assertFalse(T.name_matches("Taro Sato", "Hanako Suzuki"))
        self.assertFalse(T.name_matches("田中", "佐藤"))

    def test_split_names_and_emails(self):
        self.assertEqual(T.split_names("A One; B Two <b@example.com>;;"), ["A One", "B Two"])
        self.assertEqual(T.split_names(""), [])
        self.assertEqual(T.extract_emails("X <A@Example.com>; y@b.co"), ["a@example.com", "y@b.co"])
        self.assertEqual(T.smtp_like("/O=EXCH/CN=ABC"), "")
        self.assertEqual(T.smtp_like("User@Example.com"), "user@example.com")
        self.assertTrue(T.domain_matches("mail.courier.example.com", "courier.example.com"))
        self.assertFalse(T.domain_matches("xcourier.example.com", "courier.example.com"))


# ============================================================
# テーマ・キーワード強弱
# ============================================================
class TestTheme(unittest.TestCase):
    def test_example_theme_loads_and_has_no_real_data(self):
        p = os.path.join(TOOLS_DIR, "thread_ledger_theme.example.json")
        th = T.load_theme(p)
        self.assertGreater(len(th["hit_table"]), 30)
        raw = _rd(p)
        import re
        for addr in re.findall(r"[\w.+-]+@[\w.-]+", raw):
            self.assertTrue(addr.endswith("example.com"), addr)

    def test_errors(self):
        with self.assertRaises(T.ThemeError):
            T.normalize_theme([], "x")
        with self.assertRaises(T.ThemeError):
            T.normalize_theme({"owner": {}, "anchors": [{"name_aliases": ["a"]}]}, "x")
        with self.assertRaises(T.ThemeError):
            T.normalize_theme({"owner": {"emails": ["a@b.c"]}, "anchors": []}, "x")
        with self.assertRaises(T.ThemeError):
            make_theme(patterns=[{"category": "c", "regex": "(", "strong": True}])
        with self.assertRaises(T.ThemeError):
            make_theme(period={"from": "bad"})
        with self.assertRaises(T.ThemeError):
            T.load_theme("/nonexistent/theme.json")

    def test_strong_wins_when_term_in_both(self):
        th = make_theme(keywords={"a": {"weight": 1, "weak": ["foo"]}, "b": {"weight": 1, "strong": ["foo"]}})
        self.assertTrue(th["hit_table"]["foo"]["strong"])

    def test_defaults_for_excludes(self):
        th = make_theme()
        self.assertIn("削除済みアイテム", th["exclude_folder_names"])
        self.assertIn("note-", th["exclude_folder_prefixes"])
        th2 = make_theme(exclude_folder_names=["X"])
        self.assertEqual(th2["exclude_folder_names"], ["X"])

    def test_dict_hash_changes_with_keywords(self):
        self.assertNotEqual(make_theme()["dict_hash"],
                            make_theme(keywords={"c": {"weight": 1, "strong": ["zzz"]}})["dict_hash"])

    def test_term_boundaries(self):
        tbl = make_theme()["hit_table"]
        self.assertIn("hold", T.scan_text_hits("please hold the shipment", tbl))
        self.assertNotIn("hold", T.scan_text_hits("the holder of the card", tbl))
        self.assertIn("delay", T.scan_text_hits("many delays happened", tbl))      # 複数形は許容
        self.assertNotIn("delay", T.scan_text_hits("delayed again", tbl))
        self.assertIn("通関", T.scan_text_hits("通関業者へ", tbl))                   # 日本語は部分一致
        self.assertIn("hs code", T.scan_text_hits("HS Code is wrong", tbl))
        self.assertEqual(T.scan_text_hits("", tbl), {})

    def test_pattern_hits(self):
        tbl = make_theme()["hit_table"]
        h = T.scan_text_hits("DHL #1234567890 ok", tbl)
        self.assertIn("pat:#\\s?\\d{10}", h)
        h2 = T.scan_text_hits("code 8541.10 listed", tbl)
        self.assertIn("pat:\\b\\d{4}\\.\\d{2}\\b", h2)
        self.assertEqual(T.scan_text_hits("#12345", tbl), {})


class TestKeywordSignal(unittest.TestCase):
    def setUp(self):
        self.tbl = make_theme()["hit_table"]

    def test_weak_one_not_enough(self):
        self.assertFalse(T.keyword_signal({"approval request"}, self.tbl)["ok"])
        self.assertFalse(T.keyword_signal({"hold"}, self.tbl)["ok"])

    def test_weak_two_distinct_enough(self):
        s = T.keyword_signal({"approval request", "hold"}, self.tbl)
        self.assertTrue(s["ok"])
        self.assertEqual(s["strong"], [])

    def test_strong_one_enough(self):
        self.assertTrue(T.keyword_signal({"通関"}, self.tbl)["ok"])

    def test_weak_pattern_counts_as_weak_and_strong_pattern_as_strong(self):
        self.assertFalse(T.keyword_signal({"pat:\\b\\d{4}\\.\\d{2}\\b"}, self.tbl)["ok"])
        self.assertTrue(T.keyword_signal({"pat:\\b\\d{4}\\.\\d{2}\\b", "hold"}, self.tbl)["ok"])
        self.assertTrue(T.keyword_signal({"pat:#\\s?\\d{10}"}, self.tbl)["ok"])

    def test_unknown_keys_ignored(self):
        self.assertFalse(T.keyword_signal({"nothing"}, self.tbl)["ok"])
        self.assertFalse(T.keyword_signal(set(), self.tbl)["ok"])


# ============================================================
# 重複除去・スレッド化
# ============================================================
class TestMergeAndThreads(unittest.TestCase):
    def test_cross_store_dedupe_records_all_sources_and_month_stats(self):
        a = cache("MB", 1, "\\Inbox", [rec("E1", "RE: Topic", "2024-04-10T09:00:00", "Pm", "Owner Taro")])
        b = cache("PST", 2, "\\(ルート直下)", [rec("E9", "Topic", "2024-04-10T09:00:00", "Pm", "Owner Taro")])
        mails, st = T.merge_caches_to_mails([a, b], (2024, 1), (2024, 12))
        self.assertEqual(len(mails), 1)
        self.assertEqual(st["raw_total"], 2)
        self.assertEqual(st["unique_total"], 1)
        self.assertEqual(st["cross_store"], 1)
        self.assertEqual(st["dup_by_month"], {(2024, 4): 1})
        self.assertEqual({l["store"] for l in mails[0]["locs"]}, {"MB", "PST"})
        self.assertEqual(st["raw_by_store_month"][("MB", (2024, 4))], 1)

    def test_not_merged_when_seconds_or_sender_differ(self):
        a = cache("MB", 1, "\\Inbox", [rec("E1", "Topic", "2024-04-10T09:00:00", "Pm"),
                                       rec("E2", "Topic", "2024-04-10T09:00:01", "Pm"),
                                       rec("E3", "Topic", "2024-04-10T09:00:00", "Other")])
        mails, st = T.merge_caches_to_mails([a], (2024, 1), (2024, 12))
        self.assertEqual(len(mails), 3)
        self.assertEqual(st["cross_store"], 0)

    def test_same_store_dup_counted_separately(self):
        a = cache("MB", 1, "\\Inbox", [rec("E1", "Topic", "2024-04-10T09:00:00", "Pm")])
        b = cache("MB", 1, "\\Sub", [rec("E2", "Topic", "2024-04-10T09:00:00", "Pm")])
        mails, st = T.merge_caches_to_mails([a, b], (2024, 1), (2024, 12))
        self.assertEqual((len(mails), st["same_store"], st["cross_store"]), (2, 1, 0))   # 同一ストア内は別メール

    def test_period_filter_and_bad_time(self):
        a = cache("MB", 1, "\\Inbox", [rec("E1", "A", "2023-12-31T23:59:59"), rec("E2", "B", "2024-01-01T00:00:00"),
                                       rec("E3", "C", "2024-12-31T23:59:59"), rec("E4", "D", "2025-01-01T00:00:00"),
                                       rec("E5", "E", "garbage")])
        mails, st = T.merge_caches_to_mails([a], (2024, 1), (2024, 12))
        self.assertEqual([m["subj"] for m in mails], ["B", "C"])
        self.assertEqual(st["no_time"], 1)

    def test_thread_by_conversation_id_then_subject(self):
        a = cache("MB", 1, "\\Inbox", [
            rec("E1", "Topic one", "2024-04-10T09:00:00", "A", cid="C1"),
            rec("E2", "Different subject", "2024-04-11T09:00:00", "B", cid="C1"),
            rec("E3", "RE: Longer shared subject", "2024-04-12T09:00:00", "A"),
            rec("E4", "FW: Longer shared subject", "2024-04-13T09:00:00", "B"),
            rec("E5", "Other", "2024-04-14T09:00:00", "A"),
        ])
        mails, _ = T.merge_caches_to_mails([a], (2024, 1), (2024, 12))
        threads = T.build_threads(mails)
        self.assertEqual(sorted(t["count"] for t in threads), [1, 2, 2])
        t1 = [t for t in threads if t["key"] == "c:C1"][0]
        self.assertEqual(t1["subject"], "Different subject")       # 最新の件名
        self.assertEqual(t1["first"], datetime(2024, 4, 10, 9))
        self.assertEqual(t1["last"], datetime(2024, 4, 11, 9))
        self.assertEqual(len({t["id"] for t in threads}), 3)

    def test_empty_subject_without_cid_not_lumped(self):
        a = cache("MB", 1, "\\Inbox", [rec("E1", "", "2024-04-10T09:00:00", "A"),
                                       rec("E2", "RE:", "2024-04-11T09:00:00", "B")])
        mails, _ = T.merge_caches_to_mails([a], (2024, 1), (2024, 12))
        self.assertEqual(len(T.build_threads(mails)), 2)

    def test_select_caches_exclusions_and_latest(self):
        old = cache("MB", 1, "\\Inbox", [], scanned_at="2024-01-01T00:00:00")
        new = cache("MB", 1, "\\Inbox", [rec("E1", "x", "2024-04-10T09:00:00")], scanned_at="2024-02-01T00:00:00")
        deleted = cache("MB", 1, "\\削除済みアイテム\\sub", [rec("E2", "y", "2024-04-10T09:00:00")])
        news = cache("MB", 1, "\\Zenn-weekly", [rec("E3", "z", "2024-04-10T09:00:00")])
        other = cache("Other", 2, "\\Inbox", [rec("E4", "w", "2024-04-10T09:00:00")])
        th = make_theme()
        sel = T.select_caches([old, new, deleted, news, other], [], [], [], th["exclude_folder_names"],
                              th["exclude_folder_prefixes"])
        self.assertEqual(sorted((c["store_name"], c["folder_path"]) for c in sel),
                         [("MB", "\\Inbox"), ("Other", "\\Inbox")])
        self.assertEqual(len([c for c in sel if c["store_name"] == "MB"][0]["records"]), 1)
        sel2 = T.select_caches([new, other], ["other"], [], [], [], [])
        self.assertEqual([c["store_name"] for c in sel2], ["Other"])
        sel3 = T.select_caches([new, other], [], ["other"], [], [], [])
        self.assertEqual([c["store_name"] for c in sel3], ["MB"])

    def test_folder_excluded_and_store_selected(self):
        self.assertTrue(T.folder_excluded("Deleted Items", ["deleted items"], []))
        self.assertTrue(T.folder_excluded("ＮＯＴＥ-foo", [], ["note-"]))
        self.assertFalse(T.folder_excluded("Inbox", ["deleted items"], ["note-"]))
        self.assertFalse(T.folder_excluded("", ["x"], ["y"]))
        self.assertTrue(T.store_selected("User@Example.com", [], [], []))
        self.assertFalse(T.store_selected("User@Example.com", ["pst"], [], []))
        self.assertFalse(T.store_selected("2023_Q3", [], ["2023"], []))
        self.assertFalse(T.store_selected("Archive X", [], [], ["archive"]))


# ============================================================
# 候補判定
# ============================================================
class TestCandidate(unittest.TestCase):
    def test_all_three_conditions_via_related_party(self):
        _, e = one_thread([rec("E1", "Hello", "2024-04-10T09:00:00", "Sato Pm", "Owner Taro; Tanaka Partner")])
        self.assertTrue(e["A"] and e["B"] and e["C_i"])
        self.assertFalse(e["C_ii"])
        self.assertTrue(e["candidate"])
        self.assertEqual(e["anchors"], ["PM"])
        self.assertEqual(e["related"], ["Sachi"])

    def test_candidate_via_strong_subject_keyword(self):
        _, e = one_thread([rec("E1", "RE: 通関の件", "2024-04-10T09:00:00", "Sato Pm", "Owner Taro")])
        self.assertTrue(e["candidate"])
        self.assertFalse(e["C_i"])
        self.assertTrue(e["C_ii"])

    def test_no_anchor_fails(self):
        _, e = one_thread([rec("E1", "通関", "2024-04-10T09:00:00", "Someone", "Owner Taro; Tanaka Partner")])
        self.assertFalse(e["A"])
        self.assertFalse(e["candidate"])
        self.assertEqual(T.failed_conditions(e), ["A"])

    def test_no_owner_fails(self):
        _, e = one_thread([rec("E1", "通関", "2024-04-10T09:00:00", "Sato Pm", "Tanaka Partner")])
        self.assertFalse(e["B"])
        self.assertFalse(e["candidate"])
        self.assertEqual(T.failed_conditions(e), ["B"])

    def test_no_signal_fails(self):
        _, e = one_thread([rec("E1", "Lunch plan", "2024-04-10T09:00:00", "Sato Pm", "Owner Taro")])
        self.assertTrue(e["A"] and e["B"])
        self.assertFalse(e["C"])
        self.assertEqual(T.failed_conditions(e), ["C"])

    def test_weak_one_word_not_enough_but_two_is(self):
        _, e1 = one_thread([rec("E1", "[Approval request] Purchase lab parts", "2024-04-10T09:00:00",
                                "Sato Pm", "Owner Taro")])
        self.assertTrue(e1["A"] and e1["B"])
        self.assertFalse(e1["C_ii"])
        self.assertFalse(e1["candidate"])
        _, e2 = one_thread([rec("E1", "Approval request: urgent", "2024-04-10T09:00:00", "Sato Pm", "Owner Taro")])
        self.assertTrue(e2["C_ii"])
        self.assertTrue(e2["candidate"])

    def test_weak_one_but_related_party_rescues(self):
        _, e = one_thread([rec("E1", "[Approval request] Purchase", "2024-04-10T09:00:00", "Sato Pm",
                               "Owner Taro; Tanaka Partner")])
        self.assertFalse(e["C_ii"])
        self.assertTrue(e["candidate"])

    def test_union_over_thread_mails(self):
        # アンカーは1通目、越智は2通目、信号は3通目の件名 -> スレッド単位では全て満たす
        _, e = one_thread([
            rec("E1", "Topic", "2024-04-10T09:00:00", "Sato Pm", "X Other", cid="C1"),
            rec("E2", "Topic", "2024-04-11T09:00:00", "X Other", "Owner Taro", cid="C1"),
            rec("E3", "Topic customs clearance", "2024-04-12T09:00:00", "X Other", "Y", cid="C1"),
        ])
        self.assertTrue(e["candidate"])

    def test_weak_terms_across_subject_and_body_union(self):
        th = make_theme()
        recs = [rec("E1", "hold on", "2024-04-10T09:00:00", "Sato Pm", "Owner Taro")]
        res, _ = build([cache("S", 1, "\\Inbox", recs)], th)
        m = res["mails"][0]
        k = T.stage2_key(m["locs"][0]["sid"], m["locs"][0]["eid"])
        st2 = {k: {"s": "", "to": [], "cc": [], "att": [], "bh": th["dict_hash"], "hits": {"delay": 2}}}
        res2, _ = build([cache("S", 1, "\\Inbox", recs)], th, st2)
        self.assertFalse(ev_of(res)["candidate"])
        self.assertTrue(ev_of(res2)["candidate"])        # 件名の hold + 本文の delay = 異なる弱語2語

    def test_body_strong_from_stage2_and_stale_terms_ignored(self):
        th = make_theme()
        recs = [rec("E1", "Hello", "2024-04-10T09:00:00", "Sato Pm", "Owner Taro")]
        res, _ = build([cache("S", 1, "\\Inbox", recs)], th)
        m = res["mails"][0]
        k = T.stage2_key(m["locs"][0]["sid"], m["locs"][0]["eid"])
        st2 = {k: {"bh": th["dict_hash"], "hits": {"通関": 3, "removed-term": 9}}}
        res2, _ = build([cache("S", 1, "\\Inbox", recs)], th, st2)
        self.assertTrue(ev_of(res2)["candidate"])
        st3 = {k: {"bh": th["dict_hash"], "hits": {"removed-term": 9}}}
        res3, _ = build([cache("S", 1, "\\Inbox", recs)], th, st3)
        self.assertFalse(ev_of(res3)["candidate"])

    def test_address_only_match_via_stage2_and_domain(self):
        th = make_theme()
        # 表示名が別名と一致しない（'N. P.'）。SMTP で anchors(pm@example.com)・owner を拾う
        recs = [rec("E1", "通関", "2024-04-10T09:00:00", "N. P.", "O. P.")]
        res, _ = build([cache("S", 1, "\\Inbox", recs)], th)
        self.assertFalse(ev_of(res)["candidate"])
        m = res["mails"][0]
        k = T.stage2_key(m["locs"][0]["sid"], m["locs"][0]["eid"])
        st2 = {k: {"s": "pm@example.com", "to": [{"n": "O. P.", "a": "owner@example.com"}], "cc": [], "hits": {}}}
        res2, _ = build([cache("S", 1, "\\Inbox", recs)], th, st2)
        self.assertTrue(ev_of(res2)["candidate"])
        # ドメイン一致の関係者
        st3 = {k: {"s": "pm@example.com", "to": [{"n": "O. P.", "a": "owner@example.com"},
                                                  {"n": "Driver", "a": "x@sub.courier.example.com"}], "cc": []}}
        recs2 = [rec("E1", "Hello", "2024-04-10T09:00:00", "N. P.", "O. P.")]
        res3, _ = build([cache("S", 1, "\\Inbox", recs2)], th, st3)
        self.assertEqual(ev_of(res3)["related"], ["Courier"])
        self.assertTrue(ev_of(res3)["candidate"])

    def test_sender_smtp_in_stage1_used(self):
        recs = [rec("E1", "通関", "2024-04-10T09:00:00", "Zed", "Owner Taro", addr="pm@example.com")]
        _, e = one_thread(recs)
        self.assertTrue(e["candidate"])

    def test_exchange_x500_sender_not_used_as_address(self):
        recs = [rec("E1", "通関", "2024-04-10T09:00:00", "Zed", "Owner Taro", addr="/O=EXCH/CN=PM@EXAMPLE.COM")]
        _, e = one_thread(recs)
        self.assertFalse(e["A"])

    def test_summary_counts(self):
        recs = [
            rec("E1", "通関", "2024-04-10T09:00:00", "Sato Pm", "Owner Taro", cid="a"),   # 候補
            rec("E2", "通関", "2024-04-10T09:00:01", "Zed", "Owner Taro", cid="b"),        # Aのみ不成立
            rec("E3", "通関", "2024-04-10T09:00:02", "Sato Pm", "Zed", cid="c"),            # Bのみ不成立
            rec("E4", "Lunch", "2024-04-10T09:00:03", "Sato Pm", "Owner Taro", cid="d"),  # Cのみ不成立
            rec("E5", "Lunch", "2024-04-10T09:00:04", "Zed", "Zed", cid="e"),                # 全部不成立
        ]
        res, _ = build([cache("S", 1, "\\Inbox", recs)])
        s = T.summarize_conditions(res["evals"])
        self.assertEqual((s["total"], s["candidate"]), (5, 1))
        self.assertEqual((s["only_A_fails"], s["only_B_fails"], s["only_C_fails"]), (1, 1, 1))
        self.assertEqual((s["A"], s["B"], s["AB"]), (3, 3, 2))

    def test_select_stage2_targets_scope(self):
        recs = [
            rec("E1", "通関", "2024-04-10T09:00:00", "Sato Pm", "Owner Taro", cid="a"),
            rec("E4", "Lunch", "2024-04-10T09:00:03", "Sato Pm", "Owner Taro", cid="d"),
            rec("E2", "通関", "2024-04-10T09:00:01", "Zed", "Owner Taro", cid="b"),
        ]
        res, _ = build([cache("S", 1, "\\Inbox", recs)])
        n = lambda scope: len(T.select_stage2_targets(res["threads"], res["evals"], scope))
        self.assertEqual((n("candidate"), n("ab"), n("either")), (1, 2, 3))


class TestAnchorException(unittest.TestCase):
    """案②: 候補 = B ∧ C ∧ (A ∨ アンカー例外)。例外は anchor_exception_categories の strong 語が『件名』に当たるときだけ。"""
    EXC = ["通関"]

    def test_exception_makes_candidate_without_anchor(self):
        th = make_theme(anchor_exception_categories=self.EXC)
        _, e = one_thread([rec("E1", "RE: 通関の件", "2024-04-10T09:00:00", "Zed", "Owner Taro")], th)
        self.assertFalse(e["A"])
        self.assertTrue(e["anchor_exception"])
        self.assertTrue(e["candidate"])
        self.assertEqual(T.failed_conditions(e), [])
        self.assertIn(T.ANCHOR_EXCEPTION_LABEL, T.cond_text(e))

    def test_exception_marked_in_ledger_row_without_changing_columns(self):
        th = make_theme(anchor_exception_categories=self.EXC)
        res, _ = build([cache("S", 1, "\\Inbox", [
            rec("E1", "通関の件", "2024-04-10T09:00:00", "Zed", "Owner Taro", cid="a"),
            rec("E2", "通関の件2", "2024-04-10T09:00:01", "Sato Pm", "Owner Taro", cid="b")])], th)
        rows = T.build_ledger_rows(res["threads"], res["evals"], {}, th)
        self.assertEqual(len(rows), 2)
        for r in rows:
            self.assertEqual(len(r), len(T.LEDGER_COLUMNS))
        idx = T.LEDGER_COLUMNS.index("条件A/B/C内訳")
        marked = [r for r in rows if T.ANCHOR_EXCEPTION_LABEL in r[idx]]
        self.assertEqual(len(marked), 1)
        self.assertIn("A:×", marked[0][idx])
        # アンカー有りの行には付かない
        s = T.summarize_conditions(res["evals"])
        self.assertEqual((s["candidate"], s["anchor_exception"]), (2, 1))

    def test_exception_is_subject_only_not_body(self):
        th = make_theme(anchor_exception_categories=self.EXC)
        recs = [rec("E1", "Hello", "2024-04-10T09:00:00", "Zed", "Owner Taro; Tanaka Partner")]
        res, _ = build([cache("S", 1, "\\Inbox", recs)], th)
        m = res["mails"][0]
        k = T.stage2_key(m["locs"][0]["sid"], m["locs"][0]["eid"])
        st2 = {k: {"bh": th["dict_hash"], "hits": {"通関": 3}}}
        res2, _ = build([cache("S", 1, "\\Inbox", recs)], th, st2)
        e = ev_of(res2)
        self.assertTrue(e["C"])                # 本文の強い語でCは成立するが…
        self.assertFalse(e["anchor_exception"])  # 例外は認めない
        self.assertFalse(e["candidate"])
        self.assertEqual(T.failed_conditions(e), ["A"])

    def test_exception_not_for_other_category(self):
        th = make_theme(anchor_exception_categories=self.EXC)
        # 「物流」カテゴリの strong パターンは例外カテゴリ外
        _, e = one_thread([rec("E1", "Parcel #1234567890", "2024-04-10T09:00:00", "Zed", "Owner Taro")], th)
        self.assertTrue(e["C_ii"])
        self.assertFalse(e["anchor_exception"])
        self.assertFalse(e["candidate"])
        # 例外カテゴリでも weak 語は例外にならない
        th2 = make_theme(anchor_exception_categories=["通関"])
        _, e2 = one_thread([rec("E1", "hold", "2024-04-10T09:00:00", "Zed", "Owner Taro; Tanaka Partner")], th2)
        self.assertFalse(e2["anchor_exception"])
        self.assertFalse(e2["candidate"])

    def test_empty_or_missing_keeps_legacy_behavior(self):
        for th in (make_theme(), make_theme(anchor_exception_categories=[])):
            self.assertEqual(th["anchor_exception_categories"], [])
            _, e = one_thread([rec("E1", "通関", "2024-04-10T09:00:00", "Zed", "Owner Taro")], th)
            self.assertFalse(e["anchor_exception"])
            self.assertFalse(e["candidate"])
            self.assertEqual(T.failed_conditions(e), ["A"])

    def test_b_or_c_missing_still_fails(self):
        th = make_theme(anchor_exception_categories=self.EXC)
        _, e = one_thread([rec("E1", "通関", "2024-04-10T09:00:00", "Zed", "Someone")], th)   # B不成立
        self.assertFalse(e["B"])
        self.assertFalse(e["candidate"])
        self.assertEqual(T.failed_conditions(e), ["B"])
        # C不成立 = 件名に語が無い以上、例外も成立しない
        _, e2 = one_thread([rec("E1", "Lunch", "2024-04-10T09:00:00", "Zed", "Owner Taro")], th)
        self.assertFalse(e2["candidate"])
        self.assertFalse(e2["anchor_exception"])
        self.assertEqual(T.failed_conditions(e2), ["A", "C"])

    def test_reference_thread_excludes_exception_candidates(self):
        th = make_theme(anchor_exception_categories=self.EXC)
        _, e = one_thread([rec("E1", "通関", "2024-04-10T09:00:00", "Zed", "Owner Taro")], th)
        self.assertFalse(T.is_reference_thread(e))
        _, e2 = one_thread([rec("E1", "customs delay", "2024-04-10T09:00:00", "Zed", "Owner Taro")], th)
        self.assertFalse(T.is_reference_thread(e2))   # weak2語でC成立。例外ではない → 参考候補

    def test_check_subjects_follows_new_rule(self):
        th = make_theme(anchor_exception_categories=self.EXC)
        res, _ = build([cache("S", 1, "\\Inbox", [
            rec("E1", "通関 exception thread", "2024-04-10T09:00:00", "Zed", "Owner Taro", cid="a"),
            rec("E2", "通関 no owner thread", "2024-04-10T09:00:01", "Zed", "Someone", cid="b")])], th)
        r = T.classify_check_subjects(["通関 exception thread"], res["threads"], res["evals"])[0]
        self.assertEqual(r["status"], "found")
        self.assertTrue(r["matches"][0]["anchor_exception"])
        self.assertTrue(any(T.ANCHOR_EXCEPTION_LABEL in ln for ln in T.format_check_result(1, 1, r)))
        r2 = T.classify_check_subjects(["通関 no owner thread"], res["threads"], res["evals"])[0]
        self.assertEqual(r2["status"], "found_not_candidate")
        self.assertEqual(r2["matches"][0]["failed"], ["B"])

    def test_schema_validation(self):
        with self.assertRaises(T.ThemeError):
            make_theme(anchor_exception_categories="通関")            # 配列でない
        with self.assertRaises(T.ThemeError):
            make_theme(anchor_exception_categories=[1])               # 文字列でない
        with self.assertRaises(T.ThemeError):
            make_theme(anchor_exception_categories=[""])              # 空文字
        with self.assertRaises(T.ThemeError):
            make_theme(anchor_exception_categories=["存在しないカテゴリ"])   # 未知カテゴリ名
        th = make_theme(anchor_exception_categories=["通関", "通関", "物流"])   # patterns のカテゴリも可・重複は除く
        self.assertEqual(th["anchor_exception_categories"], ["通関", "物流"])

    def test_no_label_when_b_or_c_fails(self):
        th = make_theme(anchor_exception_categories=self.EXC)
        res, _ = build([cache("S", 1, "\\Inbox", [
            rec("E1", "通関 noB", "2024-04-10T09:00:00", "Zed", "Someone", cid="a"),
            rec("E2", "Lunch noC", "2024-04-10T09:00:01", "Zed", "Owner Taro", cid="b")])], th)
        for e in res["evals"]:
            self.assertFalse(e["candidate"])
            self.assertFalse(e["anchor_exception"])
            self.assertNotIn(T.ANCHOR_EXCEPTION_LABEL, T.cond_text(e))
        for q in ("通関 noB", "Lunch noC"):
            r = T.classify_check_subjects([q], res["threads"], res["evals"])[0]
            self.assertFalse(r["matches"][0]["anchor_exception"])
            self.assertFalse(any(T.ANCHOR_EXCEPTION_LABEL in ln or "で候補" in ln for ln in T.format_check_result(1, 1, r)))

    def test_attachment_name_only_does_not_make_exception(self):
        th = make_theme(anchor_exception_categories=self.EXC)
        recs = [rec("E1", "Hello", "2024-04-10T09:00:00", "Zed", "Owner Taro; Tanaka Partner")]
        res, _ = build([cache("S", 1, "\\Inbox", recs)], th)
        m = res["mails"][0]
        k = T.stage2_key(m["locs"][0]["sid"], m["locs"][0]["eid"])
        st2 = {k: {"bh": th["dict_hash"], "hits": {}, "att": ["通関 customs clearance.pdf"]}}
        res2, _ = build([cache("S", 1, "\\Inbox", recs)], th, st2)
        e = ev_of(res2)
        self.assertFalse(e["anchor_exception"])
        self.assertFalse(e["candidate"])

    def test_anchor_thread_is_not_marked_exception(self):
        th = make_theme(anchor_exception_categories=self.EXC)
        _, e = one_thread([rec("E1", "通関", "2024-04-10T09:00:00", "Sato Pm", "Owner Taro")], th)
        self.assertTrue(e["A"] and e["candidate"])
        self.assertFalse(e["anchor_exception"])
        self.assertNotIn(T.ANCHOR_EXCEPTION_LABEL, T.cond_text(e))

    def test_strong_pattern_category_exception(self):
        th = make_theme(anchor_exception_categories=["物流"])
        _, e = one_thread([rec("E1", "Parcel #1234567890", "2024-04-10T09:00:00", "Zed", "Owner Taro")], th)
        self.assertTrue(e["anchor_exception"])
        self.assertTrue(e["candidate"])

    def test_category_must_have_strong_terms(self):
        # トラブル語 = weak のみ、関税 = weak パターンのみ、空カテゴリ、dict でない spec
        kw = {"通関": {"weight": 3, "strong": ["通関"], "weak": []},
              "トラブル語": {"weight": 2, "strong": [], "weak": ["delay"]},
              "空": {"weight": 1, "strong": [], "weak": []},
              "壊れ": "string-spec"}
        for bad in ("トラブル語", "関税", "空", "壊れ"):
            with self.assertRaises(T.ThemeError, msg=bad):
                make_theme(keywords=kw, anchor_exception_categories=[bad])
        with self.assertRaises(T.ThemeError):
            make_theme(anchor_exception_categories=[" 通関"])   # 前後空白はエラー
        self.assertEqual(make_theme(keywords=kw, anchor_exception_categories=["通関"])["anchor_exception_categories"], ["通関"])

    def test_stage2_scope_ab_includes_exception_candidates(self):
        th = make_theme(anchor_exception_categories=self.EXC)
        res, _ = build([cache("S", 1, "\\Inbox", [
            rec("E1", "通関", "2024-04-10T09:00:00", "Zed", "Owner Taro", cid="a"),
            rec("E2", "Lunch", "2024-04-10T09:00:01", "Zed", "Owner Taro", cid="b")])], th)
        n = lambda scope: len(T.select_stage2_targets(res["threads"], res["evals"], scope))
        self.assertEqual((n("ab"), n("candidate"), n("either")), (1, 1, 2))

    def test_dict_hash_and_evaluate_only_unaffected(self):
        # 例外設定は辞書ハッシュ(第2段の本文ヒットの有効性)を変えない。キャッシュschemaも不変 = 再評価のみで反映できる
        self.assertEqual(make_theme()["dict_hash"], make_theme(anchor_exception_categories=self.EXC)["dict_hash"])
        self.assertEqual(T.SCHEMA_VERSION, 2)
        recs = [rec("E1", "通関", "2024-04-10T09:00:00", "Zed", "Owner Taro")]
        c = cache("S", 1, "\\Inbox", recs)
        r_old, _ = build([c], make_theme())
        r_new, _ = build([c], make_theme(anchor_exception_categories=self.EXC))
        self.assertFalse(ev_of(r_old)["candidate"])
        self.assertTrue(ev_of(r_new)["candidate"])

    def test_example_theme_has_exception_key_with_known_category(self):
        th = T.load_theme(os.path.join(TOOLS_DIR, "thread_ledger_theme.example.json"))
        self.assertTrue(th["anchor_exception_categories"])


# ============================================================
# スコア
# ============================================================
class TestScore(unittest.TestCase):
    def test_breakdown_text_and_subject_weighting(self):
        th = make_theme()
        recs = [rec(f"E{i}", "Re: customs clearance delay", f"2024-04-1{i}T09:00:00", "Sato Pm",
                    "Owner Taro; Tanaka Partner", cid="C") for i in range(5)]
        _, e = one_thread(recs, th)
        # アンカー1名+3, 関係者+2, 通関(1語) 3*2(件名)=6, トラブル語(1語) 2*2*0.5=2, 返信5通+1
        self.assertEqual(e["score"], 14.0)
        self.assertEqual(e["breakdown"], "アンカー+3, Sachi+2, 通関(1語)+6, トラブル語(1語)+2, 返信5通+1")

    def test_body_hit_lighter_than_subject_and_multiple_anchors(self):
        th = make_theme()
        recs = [rec("E1", "Hello", "2024-04-10T09:00:00", "Sato Pm", "Owner Taro; Suzuki Admin")]
        res, _ = build([cache("S", 1, "\\Inbox", recs)], th)
        m = res["mails"][0]
        k = T.stage2_key(m["locs"][0]["sid"], m["locs"][0]["eid"])
        res2, _ = build([cache("S", 1, "\\Inbox", recs)], th, {k: {"bh": th["dict_hash"], "hits": {"通関": 1}}})
        e = ev_of(res2)
        self.assertEqual(e["breakdown"], "アンカー(2名)+6, 通関(1語)+3")
        self.assertEqual(e["score"], 9.0)

    def test_reply_cap_and_fmt(self):
        sw = make_theme()["score_weights"]
        s, b = T.compute_score([], [], set(), set(), {}, 100, sw)
        self.assertEqual((s, b), (5.0, "返信100通+5"))
        s2, _ = T.compute_score([], [], set(), set(), {}, 4, sw)
        self.assertEqual(s2, 0.0)
        self.assertEqual(T.fmt_num(2.0), "2")
        self.assertEqual(T.fmt_num(2.25), "2.2")

    def test_related_weight_override(self):
        th = make_theme(related_parties=[{"label": "Big", "name_aliases": ["Tanaka Partner"], "weight": 7}])
        _, e = one_thread([rec("E1", "Hello", "2024-04-10T09:00:00", "Sato Pm", "Owner Taro; Tanaka Partner")], th)
        self.assertIn("Big+7", e["breakdown"])


# ============================================================
# CSV・判定
# ============================================================
class TestCsv(unittest.TestCase):
    def test_csv_safe(self):
        for ch in "=+-@":
            self.assertEqual(T.csv_safe(ch + "cmd"), "'" + ch + "cmd")
        self.assertEqual(T.csv_safe("\tabc"), "'\tabc")
        self.assertEqual(T.csv_safe("\rabc"), "'\rabc")
        self.assertEqual(T.csv_safe(" =1+1"), "' =1+1")
        self.assertEqual(T.csv_safe("normal"), "normal")
        self.assertEqual(T.csv_safe(""), "")
        self.assertEqual(T.csv_safe(None), "")
        self.assertEqual(T.csv_safe(12), "12")
        self.assertEqual(T.csv_safe("a=b"), "a=b")

    def test_write_csv_escapes_all_cells_and_utf8_sig(self):
        with tempfile.TemporaryDirectory() as d:
            p = os.path.join(d, "x.csv")
            T.write_csv(p, ["=h", "b"], [["=SUM(A1)", "@x"], ["日本語", 3]])
            raw = _rb(p)
            self.assertTrue(raw.startswith(b"\xef\xbb\xbf"))
            rows = _rcsv(p)
            self.assertEqual(rows[0], ["'=h", "b"])
            self.assertEqual(rows[1], ["'=SUM(A1)", "'@x"])
            self.assertEqual(rows[2], ["日本語", "3"])

    def test_ledger_rows_escape_mail_derived_strings_and_columns(self):
        th = make_theme()
        recs = [rec("E1", "=HYPERLINK(\"x\") 通関", "2024-04-10T09:00:00", "-Sato Pm", "@Owner Taro")]
        res, _ = build([cache("S", 1, "\\Inbox", recs)], th)
        # B は表示名の先頭記号を含んでも name_key が吸収するか確認（先頭 '@' は記号として残る）
        rows = T.build_ledger_rows(res["threads"], res["evals"], {}, th)
        for r in rows:
            self.assertEqual(len(r), len(T.LEDGER_COLUMNS))
        with tempfile.TemporaryDirectory() as d:
            p = os.path.join(d, "l.csv")
            T.write_csv(p, T.LEDGER_COLUMNS, rows)
            for line in _rcsv(p):
                for cell in line:
                    self.assertFalse(cell[:1] in ("=", "+", "@") or (cell[:1] == "-" and False), cell)

    def test_ledger_has_judgment_columns_and_version(self):
        th = make_theme()
        recs = [rec("E1", "通関", "2024-04-10T09:00:00", "Sato Pm", "Owner Taro")]
        res, _ = build([cache("S", 1, "\\Inbox", recs)], th)
        tid = res["threads"][0]["id"]
        rows = T.build_ledger_rows(res["threads"], res["evals"], {tid: {"確認結果": "本物", "メモ": "m", "問題の種類": "通関"}}, th)
        d = dict(zip(T.LEDGER_COLUMNS, rows[0]))
        self.assertEqual((d["確認結果(本物/違う/保留)"], d["メモ"], d["問題の種類"], d["theme_version"]), ("本物", "m", "通関", "t1"))
        self.assertEqual(d["最新メールEntryID"], "E1")
        self.assertEqual(d["最新メールStoreID"], "SID1")
        self.assertIn("A:○(PM)", d["条件A/B/C内訳"])

    def test_non_candidates_not_in_ledger(self):
        th = make_theme()
        recs = [rec("E1", "Lunch", "2024-04-10T09:00:00", "Sato Pm", "Owner Taro")]
        res, _ = build([cache("S", 1, "\\Inbox", recs)], th)
        self.assertEqual(T.build_ledger_rows(res["threads"], res["evals"], {}, th), [])

    def test_ledger_sorted_by_score(self):
        th = make_theme()
        recs = [rec("E1", "通関", "2024-04-10T09:00:00", "Sato Pm", "Owner Taro", cid="a"),
                rec("E2", "通関 customs clearance", "2024-04-11T09:00:00", "Sato Pm", "Owner Taro; Tanaka Partner", cid="b")]
        res, _ = build([cache("S", 1, "\\Inbox", recs)], th)
        rows = T.build_ledger_rows(res["threads"], res["evals"], {}, th)
        self.assertGreater(float(rows[0][10]), float(rows[1][10]))


class TestJudgments(unittest.TestCase):
    def test_existing_never_overwritten_new_rows_added(self):
        existing = {"t1": {"確認結果": "本物", "メモ": "keep", "問題の種類": "通関"}, "old": {"確認結果": "違う"}}
        merged, added = T.merge_judgments(existing, ["t1", "t2"])
        self.assertEqual(added, 1)
        self.assertEqual(merged["t1"], existing["t1"])
        self.assertEqual(merged["t2"], {"確認結果": "", "メモ": "", "問題の種類": ""})
        self.assertIn("old", merged)                       # 候補から外れた行も消さない
        self.assertEqual(merged["old"]["確認結果"], "違う")
        self.assertEqual(merged["old"]["メモ"], "")        # 欠けたキーは空で補うだけ
        self.assertIsNot(merged["t1"], existing["t1"])

    def test_none_existing(self):
        merged, added = T.merge_judgments(None, ["a"])
        self.assertEqual((len(merged), added), (1, 1))


# ============================================================
# カバレッジ・新メンバー・候補語
# ============================================================
class TestReports(unittest.TestCase):
    def test_coverage_flags_and_counts(self):
        th = make_theme()
        a = cache("MB", 1, "\\Inbox", [rec("E1", "通関", "2024-04-10T09:00:00", "Sato Pm", "Owner Taro"),
                                       rec("E2", "Lunch", "2024-04-11T09:00:00", "Zed", "Zed")])
        b = cache("PST", 2, "\\(ルート直下)", [rec("E9", "通関", "2024-04-10T09:00:00", "Sato Pm", "Owner Taro")])
        res, _ = build([a, b], th)
        rows = T.build_coverage_rows(res["stats"], res["mails"], res["threads"], res["evals"], {}, th,
                                     [(2024, 3), (2024, 4)], ["MB", "PST"])
        d = {(r[0], r[1]): r for r in rows}
        self.assertEqual(d[("MB", "2024-04")][2:6], [2, 2, 1, 1])
        self.assertEqual(d[("PST", "2024-04")][2:6], [1, 0, 0, 1])      # 重複除去後は代表(MB)側に数える
        self.assertEqual(d[("(全ストア)", "2024-04")][2:6], [3, 2, 1, 1])
        self.assertIn("未取得の恐れ", d[("MB", "2024-03")][6])
        self.assertIn("取得元ゼロ・要確認", d[("(全ストア)", "2024-03")][6])
        self.assertIn("未取得の恐れ", d[("(全ストア)", "2024-03")][6])
        self.assertEqual(d[("MB", "2024-04")][6], "")

    def test_suggest_participants_excludes_known_and_ranks(self):
        th = make_theme()
        recs = [
            rec("E1", "通関", "2024-04-10T09:00:00", "Sato Pm", "Owner Taro; New Person", cid="a"),
            rec("E2", "通関", "2024-04-11T09:00:00", "Sato Pm", "Owner Taro; New Person; Rare One", cid="b"),
            rec("E3", "Lunch", "2024-04-12T09:00:00", "Sato Pm", "Owner Taro; Not Candidate", cid="c"),
        ]
        res, _ = build([cache("S", 1, "\\Inbox", recs)], th)
        rows = T.suggest_participants(res["threads"], res["evals"], th, {})
        names = [r[0] for r in rows]
        self.assertEqual(names[0], "New Person")
        self.assertEqual(rows[0][2], 2)
        self.assertIn("Rare One", names)
        for known in ("Sato Pm", "Owner Taro", "Not Candidate"):
            self.assertNotIn(known, names)

    def test_rank_candidate_terms_down_weights_common_words(self):
        fg = [{"customs", "meeting"}, {"customs", "meeting"}, {"customs", "report"}]
        bg = fg + [{"meeting"}] * 50 + [{"other"}] * 10
        rows = T.rank_candidate_terms(fg, bg, known_terms=["customs"])
        terms = [r[0] for r in rows]
        self.assertEqual(terms[0], "customs")
        self.assertEqual(rows[0][7], "はい")
        self.assertNotIn("report", terms)               # fg 1件のみ(min_fg=2未満)
        if "meeting" in terms:
            self.assertLess(terms.index("customs"), terms.index("meeting"))
        self.assertEqual(T.rank_candidate_terms([], bg), [])
        self.assertEqual(T.rank_candidate_terms(fg, []), [])

    def test_tokenize_for_terms(self):
        w = T.tokenize_for_terms("Re: パラメータシート提出 and customs-entry 2024 XY")
        self.assertIn("パラメータシート", w)
        self.assertIn("提出", w)
        self.assertIn("customs-entry", w)
        self.assertNotIn("and", w)
        self.assertNotIn("2024", w)
        self.assertNotIn("xy", w)

    def test_candidate_term_docs_only_related_in_fg(self):
        th = make_theme()
        recs = [rec("E1", "foobar baz", "2024-04-10T09:00:00", "Sato Pm", "Owner Taro; Tanaka Partner", cid="a"),
                rec("E2", "通関 qux", "2024-04-10T09:00:01", "Sato Pm", "Owner Taro", cid="b")]
        res, _ = build([cache("S", 1, "\\Inbox", recs)], th)
        fg, bg = T.candidate_term_docs(res["threads"], res["evals"], th, {})
        self.assertEqual((len(fg), len(bg)), (1, 2))


# ============================================================
# 正解件名チェック
# ============================================================
class TestCheckSubjects(unittest.TestCase):
    def setUp(self):
        th = make_theme()
        recs = [
            rec("E1", "** internal only ** RE: [Export request]2024-04-01 10:00", "2024-04-10T09:00:00",
                "Sato Pm", "Owner Taro; Tanaka Partner", cid="a"),
            rec("E2", "FW: Weekly lunch plan arrangement", "2024-04-11T09:00:00", "Sato Pm", "Owner Taro", cid="b"),
            rec("E3", "customs clearance for other", "2024-04-12T09:00:00", "Zed", "Owner Taro", cid="c"),
        ]
        self.res, _ = build([cache("S", 1, "\\Inbox", recs)], th)

    def test_found_candidate_exact_with_prefix_stripped(self):
        r = T.classify_check_subjects(["RE: [Export request]2024-04-01 10:00"], self.res["threads"], self.res["evals"])[0]
        self.assertEqual(r["status"], "found")
        m = r["matches"][0]
        self.assertEqual((m["kind"], m["A"], m["B"], m["C"], m["candidate"]), ("完全一致", True, True, True, True))
        self.assertEqual(m["total"], 3)
        self.assertGreaterEqual(m["rank"], 1)

    def test_found_but_not_candidate_reports_failed_conditions(self):
        r = T.classify_check_subjects(["Weekly lunch plan arrangement"], self.res["threads"], self.res["evals"])[0]
        self.assertEqual(r["status"], "found_not_candidate")
        self.assertEqual(r["matches"][0]["failed"], ["C"])
        r2 = T.classify_check_subjects(["customs clearance for other"], self.res["threads"], self.res["evals"])[0]
        self.assertEqual(r2["matches"][0]["failed"], ["A"])

    def test_partial_match(self):
        r = T.classify_check_subjects(["Export request"], self.res["threads"], self.res["evals"])[0]
        self.assertEqual(r["status"], "found")
        self.assertEqual(r["matches"][0]["kind"], "部分一致")

    def test_short_query_no_partial_and_missing(self):
        r = T.classify_check_subjects(["lunch"], self.res["threads"], self.res["evals"])[0]
        self.assertEqual(r["status"], "missing")
        r2 = T.classify_check_subjects(["totally unknown subject here"], self.res["threads"], self.res["evals"])[0]
        self.assertEqual(r2["status"], "missing")
        r3 = T.classify_check_subjects([""], self.res["threads"], self.res["evals"])[0]
        self.assertEqual(r3["status"], "missing")

    def test_rank_is_by_score(self):
        ranks = {}
        for q in ("[Export request]2024-04-01 10:00", "customs clearance for other", "Weekly lunch plan arrangement"):
            ranks[q] = T.classify_check_subjects([q], self.res["threads"], self.res["evals"])[0]["matches"][0]["rank"]
        self.assertEqual(ranks["Weekly lunch plan arrangement"], 3)
        self.assertLess(ranks["[Export request]2024-04-01 10:00"], 3)

    def test_format_lines_and_load_file(self):
        rs = T.classify_check_subjects(["Weekly lunch plan arrangement", "unknown unknown unknown"],
                                       self.res["threads"], self.res["evals"])
        text = "\n".join(T.format_check_result(1, 2, rs[0]) + T.format_check_result(2, 2, rs[1]))
        self.assertIn("あるが候補外", text)
        self.assertIn("キャッシュに無い", text)
        self.assertIn("順位", text)
        with tempfile.TemporaryDirectory() as d:
            p = os.path.join(d, "c.txt")
            _wr(p, "# comment\n\nSubject A\n  Subject B  \n#x\n")
            self.assertEqual(T.load_calibration_subjects(p), ["Subject A", "Subject B"])
        ex = os.path.join(TOOLS_DIR, "calibration_subjects.example.txt")
        self.assertGreaterEqual(len(T.load_calibration_subjects(ex)), 3)


# ============================================================
# キャッシュ方針
# ============================================================
class TestCachePolicy(unittest.TestCase):
    def test_decide_cache_action(self):
        f, t = (2024, 1), (2024, 12)
        done = {"schema": T.SCHEMA_VERSION, "complete": True, "range": ["2023-07", "2026-09"]}
        self.assertEqual(T.decide_cache_action(done, f, t, False), "skip")
        self.assertEqual(T.decide_cache_action(done, f, t, True), "fresh")
        self.assertEqual(T.decide_cache_action(None, f, t, False), "fresh")
        self.assertEqual(T.decide_cache_action({"schema": T.SCHEMA_VERSION, "complete": True, "range": ["2024-03", "2024-12"]}, f, t, False), "fresh")
        self.assertEqual(T.decide_cache_action({"schema": T.SCHEMA_VERSION, "complete": True, "range": ["2024-01", "2024-11"]}, f, t, False), "fresh")
        cp = {"schema": T.SCHEMA_VERSION, "complete": False, "range": ["2024-01", "2024-12"], "last_received": "2024-05-01T00:00:00"}
        self.assertEqual(T.decide_cache_action(cp, f, t, False), "resume")
        self.assertEqual(T.decide_cache_action(cp, (2024, 2), t, False), "fresh")
        self.assertEqual(T.decide_cache_action({"schema": T.SCHEMA_VERSION, "complete": False, "range": ["2024-01", "2024-12"]}, f, t, False), "fresh")
        self.assertEqual(T.decide_cache_action({"range": "bad"}, f, t, False), "fresh")
        old = {"schema": 1, "complete": True, "range": ["2023-07", "2026-09"]}
        self.assertEqual(T.decide_cache_action(old, f, t, False), "fresh")        # 旧スキーマは再走査

    def test_folder_cache_path_stable_and_distinct(self):
        p = T.make_paths("/x/data", "/x/out")
        a = T.folder_cache_path(p, 1, "S", "\\Inbox")
        self.assertEqual(a, T.folder_cache_path(p, 1, "S", "\\Inbox"))
        self.assertNotEqual(a, T.folder_cache_path(p, 1, "S", "\\Other"))
        self.assertTrue(os.path.basename(a).startswith("01_"))

    def test_load_all_caches_skips_broken(self):
        with tempfile.TemporaryDirectory() as d:
            T.write_json_atomic(os.path.join(d, "a.json"), cache("S", 1, "\\Inbox", []))
            _wr(os.path.join(d, "b.json"), "{broken")
            T.write_json_atomic(os.path.join(d, "c.json"), {"schema": 99, "records": []})
            self.assertEqual(len(T.load_all_caches(d)), 1)


# ============================================================
# COM フェイク
# ============================================================
FORBIDDEN = {"Delete", "Move", "Save", "SaveAs", "Copy", "MarkAsTask", "UnRead", "Categories", "AddStore",
             "RemoveStore", "FlagStatus", "Send"}
ACCESSED = []         # 禁止属性にアクセスされた記録
STAGE2_ONLY = {"Body", "Recipients", "Attachments"}      # PropertyAccessor は会話ID(PR_CONVERSATION_ID)のため第1段でも読む
LIGHT_ACCESSED = []   # 第1段で読んではいけない属性にアクセスされた記録


class FakeComError(Exception):
    pass


class GuardedPA:
    """第1段のアイテムの PropertyAccessor。PR_CONVERSATION_ID 以外の GetProperty は LIGHT_ACCESSED に記録する。"""

    def __init__(self, inner):
        self.inner = inner

    def GetProperty(self, url):
        if url != T.PR_CONVERSATION_ID_URL:
            LIGHT_ACCESSED.append("GetProperty:" + str(url))
        return self.inner.GetProperty(url)


class FakeItem:
    def __init__(self, fields, raises=None, stage1=True, display_log=None):
        object.__setattr__(self, "_f", fields)
        object.__setattr__(self, "_r", raises or {})
        object.__setattr__(self, "_stage1", stage1)
        object.__setattr__(self, "_display_log", display_log)

    def __setattr__(self, name, value):
        ACCESSED.append(f"set:{name}")
        raise AttributeError(name)

    def Display(self):
        if self._display_log is not None:
            self._display_log.append(self._f.get("EntryID"))

    def __getattr__(self, name):
        if name.startswith("_"):
            raise AttributeError(name)
        if name in FORBIDDEN:
            ACCESSED.append(name)
            raise AttributeError(name)
        if self._stage1 and name in STAGE2_ONLY:
            LIGHT_ACCESSED.append(name)
        if self._stage1 and name == "PropertyAccessor":
            # 第1段で許すのは PR_CONVERSATION_ID(0x30130102) の GetProperty だけ。本文系などの混入を検出する
            if name in self._r:
                raise self._r[name]
            inner = self._f.get("PropertyAccessor")
            if inner is None:
                raise AttributeError(name)
            return GuardedPA(inner)
        if name in self._r:
            raise self._r[name]
        if name in self._f:
            return self._f[name]
        raise AttributeError(name)


def mk_item(eid, subj, t, sender="", to="", cc="", cid="", cls=43, addr="", sent=None, raises=None, **extra):
    f = {"Class": cls, "EntryID": eid, "Subject": subj, "ReceivedTime": t, "SenderName": sender,
         "SenderEmailAddress": addr, "To": to, "CC": cc, "ConversationID": cid}
    if sent is not None:
        f["SentOn"] = sent
    f.update(extra)
    return FakeItem(f, raises)


class FakeRestricted:
    def __init__(self, items, field):
        self.items = list(items)
        self.field = field
        self.i = 0
        self.next_raises = None

    def Sort(self, prop):
        key = prop.strip("[]")
        self.items.sort(key=lambda it: getattr(it, key))

    @property
    def Count(self):
        return len(self.items)

    def Item(self, i):
        if getattr(self, "item_raises", None) is not None:
            raise self.item_raises
        if i < 1 or i > len(self.items):
            raise IndexError(i)
        return self.items[i - 1]

    def GetFirst(self):
        self.i = 0
        return self.items[0] if self.items else None

    def GetNext(self):
        if self.next_raises is not None and self.i + 1 == self.next_raises[0]:
            raise self.next_raises[1]
        self.i += 1
        return self.items[self.i] if self.i < len(self.items) else None


class FakeItems:
    def __init__(self, items, filters, items_raise=None, next_raises=None, fail_or=False, item_raises=None):
        self.items = items
        self.filters = filters
        self.items_raise = items_raise
        self.next_raises = next_raises
        self.fail_or = fail_or
        self.item_raises = item_raises

    @property
    def Count(self):
        return len(self.items)

    def Restrict(self, flt):
        import re
        self.filters.append(flt)
        if self.items_raise is not None:
            raise self.items_raise
        if self.fail_or and " OR " in flt:
            raise FakeComError(SECRET_ERR)
        clauses = re.findall(r"\[(\w+)\] >= '([^']+)' AND \[\w+\] < '([^']+)'", flt)
        assert clauses, flt
        out = []
        for it in self.items:
            for field, s, e in clauses:
                s_dt, e_dt = datetime.strptime(s, "%Y-%m-%d %H:%M"), datetime.strptime(e, "%Y-%m-%d %H:%M")
                try:
                    v = getattr(it, field)
                except Exception:
                    continue
                if isinstance(v, datetime) and s_dt <= v < e_dt:
                    out.append(it)
                    break
        field = clauses[0][0]
        r = FakeRestricted(out, field)
        r.next_raises = self.next_raises
        r.item_raises = self.item_raises
        return r


class FakeFolders:
    def __init__(self, folders):
        self.folders = folders

    @property
    def Count(self):
        return len(self.folders)

    def Item(self, i):
        return self.folders[i - 1]


NAME2KEY = {name: key for key, name in T.TABLE_COLUMNS}


class FakeColumns:
    def __init__(self, tbl):
        self.tbl = tbl

    def RemoveAll(self):
        self.tbl.cols = []

    def Add(self, name):
        if name in self.tbl.folder.table_opts.get("col_fail", ()):
            raise FakeComError(SECRET_ERR)
        self.tbl.cols.append(NAME2KEY[name])


class FakeTable:
    """GetArray でバルク取得できる Table のフェイク（値は列の追加順のタプル。None・バイナリ・日時を含む）"""

    def __init__(self, folder, rows):
        self.folder, self.rows, self.pos, self.cols = folder, rows, 0, []
        self.Columns = FakeColumns(self)

    @property
    def EndOfTable(self):
        return self.pos >= len(self.rows)

    def Sort(self, prop, desc):
        if self.folder.table_opts.get("sort_fail"):
            raise FakeComError(SECRET_ERR)
        self.rows = sorted(self.rows, key=lambda r: r.get("recv") or datetime.max)

    def GetArray(self, n):
        self.folder.array_calls += 1
        if self.folder.array_calls in self.folder.table_opts.get("array_fail_calls", ()):
            raise FakeComError(SECRET_ERR)
        chunk = self.rows[self.pos:self.pos + n]
        self.pos += len(chunk)
        data = tuple(tuple(r.get(k) for k in self.cols) for r in chunk)
        if self.folder.table_opts.get("transpose") and data:
            return tuple(zip(*data))          # 列×行（転置）で返す壊れたAPIのフェイク
        return data


def trow(eid, subj="subj", t=None, sent=None, sender="Sato Pm", smtp="", semail="", to="Owner Taro", cc="",
         conv=None, cls="IPM.Note"):
    return {"eid": eid, "subj": subj, "recv": t, "sent": sent, "sender": sender, "smtp": smtp, "semail": semail,
            "to": to, "cc": cc, "conv": conv, "cls": cls}


_FID = [0]


class FakeFolder:
    def __init__(self, name, items=None, subfolders=None, item_type=0, items_raise=None, next_raises=None,
                 fail_or=False, item_raises=None, table_rows=None, table_opts=None):
        self.table_rows = table_rows
        self.table_opts = table_opts or {}
        self.table_filters = []
        self.array_calls = 0
        self.Name = name
        _FID[0] += 1
        self.EntryID = f"FID{_FID[0]}"
        self.DefaultItemType = item_type
        self.filters = []
        self._items = FakeItems(items or [], self.filters, items_raise, next_raises, fail_or, item_raises)
        self.Folders = FakeFolders(subfolders or [])

    @property
    def Items(self):
        return self._items

    def GetTable(self, flt=None, contents=0):
        self.table_filters.append(flt)
        o = self.table_opts
        self.get_table_calls = getattr(self, "get_table_calls", 0) + 1
        if o.get("get_table_exc") is not None and self.get_table_calls in o.get("get_table_exc_calls", range(1, 10 ** 6)):
            raise o["get_table_exc"]
        if self.table_rows is None or o.get("get_table_fail"):
            raise FakeComError(SECRET_ERR)
        if flt and " OR " in flt and o.get("fail_or"):
            raise FakeComError(SECRET_ERR)
        if flt and o.get("fail_any_filter"):
            raise FakeComError(SECRET_ERR)
        rows = list(self.table_rows)
        if flt:
            clauses = re.findall(r"\[(\w+)\] >= '([^']+)' AND \[\w+\] < '([^']+)'", flt)
            keep = []
            for r in rows:
                for field, a, b in clauses:
                    v = r.get("recv" if field == "ReceivedTime" else "sent")
                    if isinstance(v, datetime) and datetime.strptime(a, "%Y-%m-%d %H:%M") <= v < datetime.strptime(b, "%Y-%m-%d %H:%M"):
                        keep.append(r)
                        break
            rows = keep
        return FakeTable(self, rows)


class FakeStore:
    def __init__(self, name, root, sid=None, path=""):
        self.DisplayName = name
        self.StoreID = sid if sid is not None else "SID-" + name
        self.FilePath = path
        self._root = root

    def GetRootFolder(self):
        return self._root


class FakeStores:
    def __init__(self, stores):
        self.stores = stores

    @property
    def Count(self):
        return len(self.stores)

    def Item(self, i):
        return self.stores[i - 1]


PROXIES = []     # GetFolderFromID が返したフォルダ（弱参照）。走査後に全て解放されていること


class FolderProxy:
    """GetFolderFromID が返すフォルダ。本物のフェイクに委譲する（解放確認のため弱参照できる別オブジェクト）。"""

    def __init__(self, real):
        object.__setattr__(self, "_real", real)

    def __getattr__(self, name):
        return getattr(self._real, name)


class FakeNamespace:
    def __init__(self, stores, by_id=None):
        self.Stores = FakeStores(stores)
        self.by_id = by_id or {}
        self.folder_fetches = []
        self.store_lookups = []
        self.store_lookup_fails = False

    def GetStoreFromID(self, sid):
        self.store_lookups.append(sid)
        if self.store_lookup_fails:
            raise FakeComError(SECRET_ERR)
        for st in self.Stores.stores:
            if st.StoreID == sid:
                return st
        raise FakeComError(SECRET_ERR)

    def GetFolderFromID(self, eid, sid=None):
        import weakref

        def walk(f):
            yield f
            for sub in f.Folders.folders:
                yield from walk(sub)
        for st in self.Stores.stores:
            for f in walk(st._root):
                if f.EntryID == eid and (sid is None or sid == st.StoreID):
                    self.folder_fetches.append((eid, sid))
                    px = FolderProxy(f)
                    PROXIES.append(weakref.ref(px))
                    return px
        raise FakeComError(SECRET_ERR)

    def GetItemFromID(self, eid, sid):
        if (eid, sid) not in self.by_id:
            raise FakeComError(SECRET_ERR)
        v = self.by_id[(eid, sid)]
        if isinstance(v, BaseException):
            raise v
        return v


class FakeOutlook:
    def __init__(self, ns, explorer=None):
        self.ns = ns
        self.explorer = explorer

    def GetNamespace(self, name):
        return self.ns

    def ActiveExplorer(self):
        return self.explorer


class FakeClient:
    def __init__(self, ns, explorer=None):
        self.ns = ns
        self.explorer = explorer

    def Dispatch(self, name):
        return FakeOutlook(self.ns, self.explorer)


class FakePythonCom:
    def __init__(self):
        self.init = 0
        self.uninit = 0

    def CoInitialize(self):
        self.init += 1

    def CoUninitialize(self):
        self.uninit += 1


def fake_com(ns, explorer=None):
    pc = FakePythonCom()
    return (FakeClient(ns, explorer), pc), pc


def dt(m, d, h=9, mi=0, s=0):
    return datetime(2024, m, d, h, mi, s)


def run_main(argv, com=None):
    buf = io.StringIO()
    with contextlib.redirect_stdout(buf):
        code = T.main(argv, com=com)
    return code, buf.getvalue()


class ComTestBase(unittest.TestCase):
    def setUp(self):
        ACCESSED.clear()
        LIGHT_ACCESSED.clear()
        self.tmp = tempfile.TemporaryDirectory()
        self.data = os.path.join(self.tmp.name, "data")
        self.out = os.path.join(self.tmp.name, "out")
        self.theme_path = os.path.join(self.tmp.name, "theme.json")
        self.theme_data = {
            "theme_version": "t1",
            "owner": {"name_aliases": ["Owner Taro"], "emails": ["owner@example.com"]},
            "anchors": [{"label": "PM", "name_aliases": ["Sato Pm"], "emails": ["pm@example.com"]}],
            "related_parties": [{"label": "Sachi", "name_aliases": ["Tanaka Partner"]}],
            "keywords": {"通関": {"weight": 3, "strong": ["通関"], "weak": ["hold"]}},
            "period": {"from": "2024-01", "to": "2024-12"},
        }
        _wj(self.theme_path, self.theme_data)
        self.old_cp = T.CHECKPOINT_EVERY

    def tearDown(self):
        T.CHECKPOINT_EVERY = self.old_cp
        self.tmp.cleanup()

    def argv(self, *extra):
        return ["--theme", self.theme_path, "--data-dir", self.data, "--output-dir", self.out, "--scan-mode", "items", *extra]

    def paths(self):
        return T.make_paths(self.data, self.out)

    def caches(self):
        return T.load_all_caches(self.paths()["scan_cache"])

    def standard_namespace(self):
        inbox = FakeFolder("Inbox", [
            mk_item("E1", f"{SECRET} 通関", dt(4, 10), "Sato Pm", "Owner Taro", cid="C1", addr="pm@example.com"),
            mk_item("E2", "RE: " + SECRET + " 通関", dt(4, 11), "Owner Taro", "Sato Pm", cid="C1"),
            mk_item("E3", "Lunch", dt(5, 1), "Zed", "Owner Taro"),
            mk_item("E4", "meeting request", dt(5, 2), "Zed", "Owner Taro", cls=53),   # 会議出席依頼 -> メール以外
            mk_item("E5", "out of period", datetime(2023, 11, 1, 9), "Zed", "Owner Taro"),
        ], subfolders=[
            FakeFolder("削除済みアイテム", [mk_item("D1", "deleted", dt(4, 12), "Zed", "Owner Taro")]),
            FakeFolder("Calendar", [mk_item("K1", "cal", dt(4, 12))], item_type=1),
        ])
        mb = FakeStore("user@example.com", inbox_root(inbox))
        pst_root = FakeFolder("root", [
            mk_item("P1", f"{SECRET} 通関", dt(4, 10), "Sato Pm", "Owner Taro", cid="C9"),     # E1 と重複
            mk_item("P2", "PST only 通関", dt(4, 20), "Sato Pm", "Owner Taro", cid="C2"),
        ])
        pst = FakeStore("2024_Q2", pst_root, path="C:\\x\\2024_Q2.pst")
        return FakeNamespace([mb, pst])


def inbox_root(inbox):
    root = FakeFolder("root", [], subfolders=[inbox])
    return root


# ============================================================
# 第1段スキャン（COM）
# ============================================================
class TestScanCom(ComTestBase):
    def test_scan_end_to_end_cache_and_privacy(self):
        ns = self.standard_namespace()
        com, pc = fake_com(ns)
        code, out = run_main(self.argv("--no-stage2"), com)
        self.assertEqual(code, 0)
        self.assertNotIn(SECRET, out)
        self.assertEqual((pc.init, pc.uninit), (1, 1))            # 走査1回のみCOMを使う（評価はCOM不要）
        self.assertEqual(LIGHT_ACCESSED, [])                     # Body等は第1段で読まない
        self.assertEqual(ACCESSED, [])                           # 変更系は一切呼ばない
        caches = self.caches()
        by_path = {(c["store_name"], c["folder_path"]): c for c in caches}
        self.assertIn(("user@example.com", "\\Inbox"), by_path)
        self.assertIn(("2024_Q2", "\\(ルート直下)"), by_path)
        self.assertNotIn(("user@example.com", "\\削除済みアイテム"), by_path)          # 除外フォルダは走査しない
        self.assertNotIn(("user@example.com", "\\Calendar"), by_path)                   # メール以外のフォルダ
        inbox = by_path[("user@example.com", "\\Inbox")]
        self.assertEqual(sorted(r["e"] for r in inbox["records"]), ["E1", "E2", "E3"])  # 期間外は取らない・クラス53は除く
        self.assertEqual(inbox["counts"]["non_mail"], 1)
        self.assertTrue(inbox["complete"])
        self.assertEqual(inbox["store_id"], "SID-user@example.com")
        r1 = [r for r in inbox["records"] if r["e"] == "E1"][0]
        self.assertEqual((r1["n"], r1["to"], r1["c2"], r1["a"]), ("Sato Pm", "Owner Taro", "C1", "pm@example.com"))
        self.assertEqual(r1["c"], "")            # PropertyAccessor が無いフェイクでは PR_CONVERSATION_ID は空
        self.assertEqual(r1["t"], "2024-04-10T09:00:00")
        for k in r1:
            self.assertIn(k, {"e", "s", "t", "n", "a", "to", "cc", "c", "c2"})                # 本文は含めない
        # フィルタはISO・期間全体で1回
        flt = ns.Stores.Item(1).GetRootFolder().Folders.Item(1).filters
        self.assertEqual(flt, ["([ReceivedTime] >= '2024-01-01 00:00' AND [ReceivedTime] < '2025-01-01 00:00') OR ([SentOn] >= '2024-01-01 00:00' AND [SentOn] < '2025-01-01 00:00')"])

    def test_console_has_no_mail_addresses_of_stores(self):
        com, _ = fake_com(self.standard_namespace())
        code, out = run_main(self.argv("--no-stage2"), com)
        self.assertNotIn("user@example.com", out)
        self.assertIn("<MAIL>", out)

    def test_outputs_after_scan_dedupe_and_csv(self):
        com, _ = fake_com(self.standard_namespace())
        code, out = run_main(self.argv("--no-stage2", "--suggest", "--suggest-terms"), com)
        self.assertEqual(code, 0)
        self.assertIn("ストア間の重複 1件", out)
        files = os.listdir(self.out)
        for prefix in ("thread_ledger_theme_", "coverage_", "suggest_participants_", "candidate_terms_"):
            self.assertTrue(any(f.startswith(prefix) and f.endswith(".csv") for f in files), prefix)
        led = [f for f in files if f.startswith("thread_ledger_theme_") and f.endswith(".csv")][0]
        rows = _rdict(os.path.join(self.out, led))
        self.assertEqual(len(rows), 2)        # C1(E1/E2/P1重複) と C2
        self.assertEqual(sorted(r["メール数"] for r in rows), ["1", "2"])
        multi = [r for r in rows if r["メール数"] == "2"][0]
        self.assertIn("user@example.com", multi["取得元(ストア/フォルダ)"])
        self.assertIn("2024_Q2", multi["取得元(ストア/フォルダ)"])   # 重複の取得元も全て記録
        self.assertTrue(os.path.exists(os.path.join(self.data, "judgments_theme.json")))
        self.assertNotIn(SECRET, "".join(_rd(os.path.join(self.out, f), "utf-8-sig")
                                         for f in files if f.endswith(".log")))

    def test_second_run_skips_completed_and_rescan_refetches(self):
        com, _ = fake_com(self.standard_namespace())
        run_main(self.argv("--no-stage2"), com)
        ns2 = self.standard_namespace()
        com2, _ = fake_com(ns2)
        code, out = run_main(self.argv("--no-stage2"), com2)
        self.assertIn("完了済みスキップ 2", out)
        self.assertEqual(ns2.Stores.Item(1).GetRootFolder().Folders.Item(1).filters, [])
        ns3 = self.standard_namespace()
        com3, _ = fake_com(ns3)
        run_main(self.argv("--no-stage2", "--rescan"), com3)
        self.assertEqual(len(ns3.Stores.Item(1).GetRootFolder().Folders.Item(1).filters), 1)

    def test_one_folder_exception_continues_and_error_kind_only(self):
        bad = FakeFolder("Bad", [mk_item("B1", SECRET, dt(4, 1), "Zed", "Owner Taro")],
                         items_raise=FakeComError(SECRET_ERR))
        good = FakeFolder("Good", [mk_item("G1", "ok", dt(4, 2), "Zed", "Owner Taro")])
        root = FakeFolder("root", [], subfolders=[bad, good])
        com, _ = fake_com(FakeNamespace([FakeStore("MB", root)]))
        code, out = run_main(self.argv("--no-stage2"), com)
        self.assertEqual(code, 0)
        self.assertIn("FakeComError", out)
        self.assertNotIn(SECRET_ERR, out)
        self.assertNotIn(SECRET, out)
        paths = {c["folder_path"] for c in self.caches()}
        self.assertIn("\\Good", paths)
        self.assertNotIn("\\Bad", paths)

    def test_item_field_exceptions_are_swallowed_per_field(self):
        it = mk_item("X1", SECRET, dt(4, 1), "Zed", "Owner Taro",
                     raises={"Subject": FakeComError(SECRET_ERR), "ConversationID": FakeComError(SECRET_ERR)})
        it2 = mk_item("X2", "t", dt(4, 2), raises={"EntryID": FakeComError(SECRET_ERR)})
        it3 = mk_item("X3", "t", dt(4, 3), raises={"Class": FakeComError(SECRET_ERR)})
        it4 = mk_item("X4", "ok", dt(4, 4), "Zed", "Owner Taro")
        root = FakeFolder("root", [it, it2, it3, it4])
        com, _ = fake_com(FakeNamespace([FakeStore("MB", root)]))
        code, out = run_main(self.argv("--no-stage2"), com)
        self.assertEqual(code, 0)
        self.assertNotIn(SECRET_ERR, out)
        c = self.caches()[0]
        self.assertEqual(sorted(r["e"] for r in c["records"]), ["X1", "X4"])
        x1 = [r for r in c["records"] if r["e"] == "X1"][0]
        self.assertEqual((x1["s"], x1["c"]), ("", ""))
        self.assertIn("Subject:FakeComError", c["counts"]["errors"])

    def sent_item(self):
        # ReceivedTime が未設定（4501年）の送信済みアイテム
        return FakeItem({"Class": 43, "EntryID": "S1", "Subject": "s", "ReceivedTime": datetime(4501, 1, 1),
                         "SentOn": dt(4, 5), "SenderName": "Owner Taro", "To": "Sato Pm",
                         "SenderEmailAddress": "", "CC": "", "ConversationID": ""})

    def test_or_restrict_catches_sent_and_received_together(self):
        sent = FakeFolder("Mixed", [self.sent_item(), mk_item("R1", "r", dt(4, 6), "Zed", "Owner Taro")])
        com, _ = fake_com(FakeNamespace([FakeStore("MB", FakeFolder("root", [], subfolders=[sent]))]))
        run_main(self.argv("--no-stage2"), com)
        c = self.caches()[0]
        self.assertEqual(c["date_field"], "or")
        self.assertEqual(sorted(r["e"] for r in c["records"]), ["R1", "S1"])
        self.assertEqual({r["e"]: r["t"] for r in c["records"]}["S1"], "2024-04-05T09:00:00")
        self.assertIn(" OR ", sent.filters[0])

    def test_or_restrict_failure_falls_back_to_receivedtime_then_senton(self):
        sent = FakeFolder("Sent", [self.sent_item()], fail_or=True)
        com, _ = fake_com(FakeNamespace([FakeStore("MB", FakeFolder("root", [], subfolders=[sent]))]))
        code, out = run_main(self.argv("--no-stage2"), com)
        self.assertEqual(code, 0)
        c = self.caches()[0]
        self.assertEqual(c["date_field"], "SentOn")
        self.assertEqual(c["records"][0]["t"], "2024-04-05T09:00:00")
        self.assertEqual([f.split()[0] for f in sent.filters], ["([ReceivedTime]", "[ReceivedTime]", "[SentOn]"])
        self.assertNotIn(SECRET_ERR, out)
        # 受信側のフォルダは ReceivedTime で取れる（SentOn まで行かない）
        inbox = FakeFolder("In", [mk_item("R1", "r", dt(4, 6), "Zed", "Owner Taro")], fail_or=True)
        com2, _ = fake_com(FakeNamespace([FakeStore("MB2", FakeFolder("root", [], subfolders=[inbox]), sid="S2")]))
        run_main(self.argv("--no-stage2", "--stores", "MB2"), com2)
        self.assertEqual(len(inbox.filters), 2)

    def test_resume_in_fallback_mode_restricts_from_last_received(self):
        T.CHECKPOINT_EVERY = 100
        items = [mk_item(f"M{i}", f"s{i}", dt(4, i), "Zed", "Owner Taro") for i in range(1, 8)]
        items[4] = mk_item("M5", "s5", dt(4, 5), "Zed", "Owner Taro", raises={"Subject": KeyboardInterrupt()})
        f1 = FakeFolder("F", items, fail_or=True)
        com, _ = fake_com(FakeNamespace([FakeStore("MB", FakeFolder("root", [], subfolders=[f1]))]))
        self.assertEqual(run_main(self.argv("--no-stage2"), com)[0], 130)
        items2 = [mk_item(f"M{i}", f"s{i}", dt(4, i), "Zed", "Owner Taro") for i in range(1, 8)]
        f2 = FakeFolder("F", items2, fail_or=True)
        com2, _ = fake_com(FakeNamespace([FakeStore("MB", FakeFolder("root", [], subfolders=[f2]))]))
        self.assertEqual(run_main(self.argv("--no-stage2"), com2)[0], 0)
        self.assertEqual(f2.filters, ["[ReceivedTime] >= '2024-04-04 09:00' AND [ReceivedTime] < '2025-01-01 00:00'"])
        c = self.caches()[0]
        self.assertEqual(sorted(r["e"] for r in c["records"]), [f"M{i}" for i in range(1, 8)])

    def test_keyboard_interrupt_saves_partial_and_resume_has_no_duplicates(self):
        T.CHECKPOINT_EVERY = 2
        items = [mk_item(f"M{i}", f"s{i}", dt(4, i), "Zed", "Owner Taro") for i in range(1, 8)]
        items[4] = mk_item("M5", "s5", dt(4, 5), "Zed", "Owner Taro", raises={"Subject": KeyboardInterrupt()})
        root = FakeFolder("root", items)
        com, pc = fake_com(FakeNamespace([FakeStore("MB", root)]))
        code, out = run_main(self.argv("--no-stage2"), com)
        self.assertEqual(code, 130)
        self.assertEqual(pc.uninit, 1)
        c = self.caches()[0]
        self.assertFalse(c["complete"])
        self.assertEqual(sorted(r["e"] for r in c["records"]), ["M1", "M2", "M3", "M4"])
        self.assertEqual(c["last_received"], "2024-04-04T09:00:00")
        self.assertIn("Ctrl+C", out)
        # 再開: OR方式は期間全体を取り直し、取得済みEntryIDを早期に飛ばして重複除去する
        items2 = [mk_item(f"M{i}", f"s{i}", dt(4, i), "Zed", "Owner Taro") for i in range(1, 8)]
        root2 = FakeFolder("root", items2)
        com2, _ = fake_com(FakeNamespace([FakeStore("MB", root2)]))
        code2, out2 = run_main(self.argv("--no-stage2"), com2)
        self.assertEqual(code2, 0)
        self.assertEqual(root2.filters, ["([ReceivedTime] >= '2024-01-01 00:00' AND [ReceivedTime] < '2025-01-01 00:00') OR ([SentOn] >= '2024-01-01 00:00' AND [SentOn] < '2025-01-01 00:00')"])
        c2 = self.caches()[0]
        self.assertTrue(c2["complete"])
        eids = [r["e"] for r in c2["records"]]
        self.assertEqual(sorted(eids), [f"M{i}" for i in range(1, 8)])
        self.assertEqual(len(eids), len(set(eids)))

    def test_checkpoint_saved_every_n_items(self):
        T.CHECKPOINT_EVERY = 3
        items = [mk_item(f"M{i}", f"s{i}", dt(4, i), "Zed", "Owner Taro") for i in range(1, 9)]
        saved = []
        orig = T.write_json_atomic

        def spy(path, obj):
            if "scan_cache" in path:
                saved.append((obj["complete"], len(obj["records"])))
            return orig(path, obj)
        T.write_json_atomic = spy
        try:
            com, _ = fake_com(FakeNamespace([FakeStore("MB", FakeFolder("root", items))]))
            run_main(self.argv("--no-stage2"), com)
        finally:
            T.write_json_atomic = orig
        self.assertEqual(saved, [(False, 3), (False, 6), (True, 8)])

    def test_getnext_failure_skips_bad_item_and_goes_on(self):
        items = [mk_item(f"M{i}", f"s{i}", dt(4, i), "Zed", "Owner Taro") for i in range(1, 5)]
        f = FakeFolder("F", items, next_raises=(2, FakeComError(SECRET_ERR)))
        com, _ = fake_com(FakeNamespace([FakeStore("MB", FakeFolder("root", [], subfolders=[f]))]))
        code, out = run_main(self.argv("--no-stage2"), com)
        self.assertEqual(code, 0)
        self.assertNotIn(SECRET_ERR, out)
        c = self.caches()[0]
        self.assertTrue(c["complete"])                 # 失敗した1件を飛ばして最後まで読む（再実行でも同じ所で止まらない）
        self.assertEqual(sorted(r["e"] for r in c["records"]), ["M1", "M2", "M4"])
        self.assertIn("GetNext:FakeComError", c["counts"]["errors"])

    def test_many_unreadable_items_stop_the_folder_as_incomplete(self):
        items = [mk_item(f"M{i}", f"s{i}", dt(4, 1) , "Zed", "Owner Taro") for i in range(1, 4)]
        st = T.new_scan_state()
        restricted_items = FakeItems(items, [])

        class Bad(FakeRestricted):
            def Item(self, i):
                raise FakeComError(SECRET_ERR)

            def GetNext(self):
                raise FakeComError(SECRET_ERR)

        class F:
            class Items:
                @staticmethod
                def Restrict(flt):
                    return Bad(items, "x")
        old = T.MAX_ITEM_SKIP
        T.MAX_ITEM_SKIP = 1
        try:
            T.scan_folder_items(F, datetime(2024, 1, 1), datetime(2025, 1, 1), st, total_hint=3)
        finally:
            T.MAX_ITEM_SKIP = old
        self.assertFalse(st["complete"])
        self.assertEqual([r["e"] for r in st["records"]], ["M1"])

    def test_stores_and_skip_stores_filters(self):
        ns = self.standard_namespace()
        com, _ = fake_com(ns)
        run_main(self.argv("--no-stage2", "--stores", "2024_q2"), com)
        self.assertEqual({c["store_name"] for c in self.caches()}, {"2024_Q2"})

    def test_no_com_available(self):
        orig = T.import_com
        T.import_com = lambda: (None, None)
        try:
            code, out = run_main(self.argv("--no-stage2"))
        finally:
            T.import_com = orig
        self.assertEqual(code, 2)
        self.assertIn("win32com", out)

    def test_outlook_connection_failure(self):
        class BadClient:
            def Dispatch(self, n):
                raise FakeComError(SECRET_ERR)
        pc = FakePythonCom()
        code, out = run_main(self.argv("--no-stage2"), (BadClient(), pc))
        self.assertEqual(code, 2)
        self.assertNotIn(SECRET_ERR, out)
        self.assertEqual(pc.uninit, 1)


# ============================================================
# 第2段・評価のみ・open
# ============================================================
class TestStage2AndModes(ComTestBase):
    def scan_first(self):
        com, _ = fake_com(self.standard_namespace())
        run_main(self.argv("--no-stage2"), com)

    def stage2_ns(self, **kw):
        sid_mb, sid_pst = "SID-user@example.com", "SID-2024_Q2"
        body = "please see customs and 通関 and hold " + SECRET
        full = lambda eid, s_addr, recips, **r: FakeItem(
            {"Class": 43, "EntryID": eid, "SenderEmailAddress": s_addr, "SenderEmailType": "SMTP",
             "PropertyAccessor": FakePA({T.PR_SENDER_SMTP: s_addr, T.PR_INTERNET_MESSAGE_ID: "<id@x>"}),
             "Recipients": FakeRecipients(recips), "Attachments": FakeAtts(["a.pdf", "b.xlsx"]),
             "Body": body}, raises=r, stage1=False)
        by_id = {
            ("E1", sid_mb): full("E1", "pm@example.com", [("Owner Taro", "owner@example.com", 1), ("Tanaka Partner", "", 2)], **kw),
            ("E2", sid_mb): full("E2", "owner@example.com", [("Sato Pm", "pm@example.com", 1)]),
            ("P2", sid_pst): full("P2", "pm@example.com", [("Owner Taro", "owner@example.com", 1)]),
        }
        return FakeNamespace([], by_id)

    def test_stage2_fills_cache_without_body_text_and_reevaluates(self):
        self.scan_first()
        ns = self.stage2_ns()
        com, pc = fake_com(ns)
        code, out = run_main(self.argv("--evaluate-only"), None)      # 評価のみ（COMなし）
        self.assertEqual(code, 0)
        code, out = run_main(self.argv(), com)  # 走査はスキップ → 第2段
        self.assertEqual(code, 0)
        self.assertNotIn(SECRET, out)
        self.assertNotIn("pm@example.com", out)
        self.assertEqual(ACCESSED, [])
        raw = _rd(self.paths()["stage2_cache"])
        self.assertNotIn(SECRET, raw)                                  # 本文は保存しない
        st2 = _rj(self.paths()["stage2_cache"])
        self.assertEqual(len(st2), 3)
        k = T.stage2_key("SID-user@example.com", "E1")
        r = st2[k]
        self.assertEqual(r["s"], "pm@example.com")
        self.assertEqual(r["to"], [{"n": "Owner Taro", "a": "owner@example.com"}])
        self.assertEqual(r["cc"][0]["n"], "Tanaka Partner")
        self.assertEqual(r["att"], ["a.pdf", "b.xlsx"])
        self.assertEqual(r["imid"], "<id@x>")
        self.assertEqual(r["hits"], {"通関": 1, "hold": 1})
        self.assertEqual(r["bh"], T.load_theme(self.theme_path)["dict_hash"])

    def test_stage2_no_body_skips_body_and_second_run_reuses(self):
        self.scan_first()
        ns = self.stage2_ns()
        com, _ = fake_com(ns)
        run_main(self.argv("--no-body"), com)
        st2 = _rj(self.paths()["stage2_cache"])
        self.assertTrue(all(v["hits"] == {} and v["bh"] is None for v in st2.values()))
        # --no-body の再実行は取得済みをスキップ。本文ありで再実行すると取り直す
        ns2 = self.stage2_ns()
        ns2.by_id[("E1", "SID-user@example.com")] = FakeComError("must not be opened")
        com2, _ = fake_com(ns2)
        code, out = run_main(self.argv("--no-body"), com2)
        self.assertEqual(code, 0)
        self.assertIn("取得済みスキップ 3", out)
        ns3 = self.stage2_ns()
        com3, _ = fake_com(ns3)
        run_main(self.argv(), com3)
        st3 = _rj(self.paths()["stage2_cache"])
        self.assertEqual(st3[T.stage2_key("SID-user@example.com", "E1")]["hits"], {"通関": 1, "hold": 1})

    def test_stage2_item_errors_swallowed_and_kinds_only(self):
        self.scan_first()
        ns = self.stage2_ns(Body=FakeComError(SECRET_ERR), Attachments=FakeComError(SECRET_ERR),
                            Recipients=FakeComError(SECRET_ERR))
        ns.by_id[("E2", "SID-user@example.com")] = FakeComError(SECRET_ERR)     # 開けないメール
        com, _ = fake_com(ns)
        code, out = run_main(self.argv(), com)
        self.assertEqual(code, 0)
        self.assertNotIn(SECRET_ERR, out)
        self.assertIn("FakeComError", out)
        self.assertIn("開けず 1通", out)
        st2 = _rj(self.paths()["stage2_cache"])
        self.assertEqual(len(st2), 2)

    def test_stage2_keyboard_interrupt_saves(self):
        self.scan_first()
        ns = self.stage2_ns()
        ns.by_id[("E2", "SID-user@example.com")] = KeyboardInterrupt()
        com, pc = fake_com(ns)
        code, out = run_main(self.argv(), com)
        self.assertEqual(code, 130)
        st2 = _rj(self.paths()["stage2_cache"])
        self.assertGreaterEqual(len(st2), 1)
        self.assertEqual(pc.uninit, 2)   # 走査(スキップ)と第2段で各1回

    def test_no_stage2_does_not_open_items_but_uses_existing_cache(self):
        self.scan_first()
        com, _ = fake_com(self.stage2_ns())
        run_main(self.argv(), com)
        ns = self.stage2_ns()
        ns.by_id.clear()            # 開こうとすれば FakeComError になる
        com2, pc2 = fake_com(ns)
        code, out = run_main(self.argv("--no-stage2"), com2)
        self.assertEqual(code, 0)
        self.assertIn("既存の第2段キャッシュ", out)
        self.assertNotIn("第2段: 対象メール", out)
        self.assertNotIn("FakeComError", out)

    def test_evaluate_only_never_touches_com(self):
        self.scan_first()
        orig = T.import_com
        T.import_com = lambda: (_ for _ in ()).throw(AssertionError("COMを使ってはいけない"))
        try:
            code, out = run_main(self.argv("--evaluate-only"))
        finally:
            T.import_com = orig
        self.assertEqual(code, 0)
        self.assertIn("候補スレッド", out)

    def test_evaluate_only_without_cache(self):
        code, out = run_main(self.argv("--evaluate-only"))
        self.assertEqual(code, 2)
        self.assertIn("走査キャッシュがありません", out)

    def test_judgments_preserved_across_reevaluation(self):
        self.scan_first()
        run_main(self.argv("--evaluate-only"))
        jp = os.path.join(self.data, "judgments_theme.json")
        j = _rj(jp)
        tid = sorted(j)[0]
        j[tid]["確認結果"] = "本物"
        j[tid]["メモ"] = "手編集"
        _wj(jp, j)
        run_main(self.argv("--evaluate-only"))
        run_main(self.argv("--evaluate-only", "--rescan"))
        j2 = _rj(jp)
        self.assertEqual(j2[tid]["確認結果"], "本物")
        self.assertEqual(j2[tid]["メモ"], "手編集")
        newest = sorted(f for f in os.listdir(self.out) if f.startswith("thread_ledger_theme_") and f.endswith(".csv"))[-1]
        rows = _rdict(os.path.join(self.out, newest))
        self.assertEqual([r["確認結果(本物/違う/保留)"] for r in rows if r["thread_id"] == tid], ["本物"])

    def test_theme_edit_changes_evaluation_without_rescan(self):
        self.scan_first()
        run_main(self.argv("--evaluate-only"))
        self.theme_data["owner"] = {"name_aliases": ["Nobody Here"]}
        _wj(self.theme_path, self.theme_data)
        code, out = run_main(self.argv("--evaluate-only"))
        self.assertIn("候補スレッド(B∧C∧(A∨アンカー例外)): 0件", out)

    def test_check_subjects_command_prints_subjects_only_on_console_and_no_scan(self):
        self.scan_first()
        cal = os.path.join(self.tmp.name, "cal.txt")
        _wr(cal, f"# c\nRE: {SECRET} 通関\nnot in cache subject xyz\nLunch\n")
        orig = T.import_com
        T.import_com = lambda: (_ for _ in ()).throw(AssertionError("COM禁止"))
        try:
            code, out = run_main(self.argv("--check-subjects", cal))
        finally:
            T.import_com = orig
        self.assertEqual(code, 0)
        self.assertIn(SECRET, out)                       # この出力だけは件名を表示
        self.assertIn("キャッシュに無い", out)
        self.assertIn("あるが候補外", out)
        self.assertIn("順位", out)
        log = sorted(f for f in os.listdir(self.out) if f.endswith(".log"))[-1]
        logtext = _rd(os.path.join(self.out, log), "utf-8-sig")
        self.assertNotIn(SECRET, logtext)                # ログには件名を残さない
        self.assertIn("再現率チェック", logtext)

    def test_open_uses_ledger_and_only_displays(self):
        self.scan_first()
        run_main(self.argv("--evaluate-only"))
        led = sorted(f for f in os.listdir(self.out) if f.startswith("thread_ledger_theme_") and f.endswith(".csv"))[-1]
        rows = _rdict(os.path.join(self.out, led))
        row = rows[0]
        shown = []
        item = FakeItem({"EntryID": row["最新メールEntryID"]}, display_log=shown)
        ns = FakeNamespace([], {(row["最新メールEntryID"], row["最新メールStoreID"]): item})
        com, pc = fake_com(ns)
        import unittest.mock as mock
        with mock.patch.object(T, "HEX_ID_RE", re.compile(r".+")):          # フェイクのIDは16進ではない
            code, out = run_main(["--open", row["thread_id"], "--data-dir", self.data, "--output-dir", self.out], com)
        self.assertEqual(code, 0)
        self.assertEqual(shown, [row["最新メールEntryID"]])
        self.assertEqual(ACCESSED, [])
        self.assertEqual(pc.uninit, 1)

    def test_open_unknown_thread_and_fallback_to_cache(self):
        self.scan_first()
        com, _ = fake_com(FakeNamespace([], {}))
        code, out = run_main(["--open", "zzzzzzzzzzzz", "--data-dir", self.data, "--output-dir", self.out], com)
        self.assertEqual(code, 2)
        # 台帳CSVが無くてもキャッシュから引ける
        mails, _s = T.merge_caches_to_mails(self.caches(), (2024, 1), (2024, 12))
        t = T.build_threads(mails)[0]
        eid, sid = t["mails"][-1]["locs"][0]["eid"], t["mails"][-1]["locs"][0]["sid"]
        self.assertEqual(T.find_thread_in_cache(self.paths(), t["id"], (2024, 1), (2024, 12)), (eid, sid))

    def test_missing_theme_guides_to_example(self):
        code, out = run_main(["--data-dir", self.data, "--output-dir", self.out, "--evaluate-only"])
        self.assertEqual(code, 2)
        self.assertIn("thread_ledger_theme.example.json", out)

    def test_invalid_period_arg(self):
        code, out = run_main(self.argv("--from", "2025-01", "--to", "2024-01", "--evaluate-only"))
        self.assertEqual(code, 2)

    def test_coverage_csv_flags_empty_months(self):
        self.scan_first()
        run_main(self.argv("--evaluate-only"))
        cov = sorted(f for f in os.listdir(self.out) if f.startswith("coverage_"))[-1]
        rows = _rdict(os.path.join(self.out, cov))
        allrows = {r["年月"]: r for r in rows if r["ストア"] == "(全ストア)"}
        self.assertIn("未取得の恐れ", allrows["2024-02"]["備考"])
        self.assertEqual(allrows["2024-04"]["備考"], "")


class FakePA:
    def __init__(self, props):
        self.props = props

    def GetProperty(self, url):
        if url not in self.props:
            raise FakeComError(SECRET_ERR)
        return self.props[url]


class FakeRecipient:
    def __init__(self, name, addr, typ):
        self.Name = name
        self.Type = typ
        self.Address = addr
        self.PropertyAccessor = FakePA({T.PR_RECIP_SMTP: addr} if addr else {})


class FakeRecipients:
    def __init__(self, recips):
        self.r = [FakeRecipient(*x) for x in recips]

    @property
    def Count(self):
        return len(self.r)

    def Item(self, i):
        return self.r[i - 1]


class FakeAttachment:
    def __init__(self, name):
        self.FileName = name


class FakeAtts:
    def __init__(self, names):
        self.a = [FakeAttachment(n) for n in names]

    @property
    def Count(self):
        return len(self.a)

    def Item(self, i):
        return self.a[i - 1]


class TestReadStage2Item(unittest.TestCase):
    def test_fallback_chain_for_sender_and_recipient_smtp(self):
        class Entry:
            def GetExchangeUser(self):
                class U:
                    PrimarySmtpAddress = "Ex.User@Example.com"
                return U()

        class RecipEx:
            Name = "Ex Person"
            Type = 1
            Address = "/O=EXCH/CN=X"
            AddressEntry = Entry()
            PropertyAccessor = FakePA({})

        class Recs:
            Count = 1

            def Item(self, i):
                return RecipEx()

        item = FakeItem({"SenderEmailAddress": "/O=EXCH/CN=Y", "SenderEmailType": "EX", "Sender": Entry(),
                         "PropertyAccessor": FakePA({}), "Recipients": Recs(), "Attachments": FakeAtts([]),
                         "Body": "x"}, stage1=False)
        errs = {}
        r = T.read_stage2_item(item, make_theme(), False, errs)
        self.assertEqual(r["s"], "ex.user@example.com")
        self.assertEqual(r["to"], [{"n": "Ex Person", "a": "ex.user@example.com"}])
        self.assertTrue(all(SECRET_ERR not in k for k in errs))
        self.assertIn("送信者SMTP(PA):FakeComError", errs)

    def test_bcc_ignored(self):
        item = FakeItem({"SenderEmailAddress": "a@example.com", "PropertyAccessor": FakePA({}),
                         "Recipients": FakeRecipients([("X", "x@example.com", 3), ("Y", "y@example.com", 1)]),
                         "Attachments": FakeAtts([]), "Body": ""}, stage1=False)
        r = T.read_stage2_item(item, make_theme(), True, {})
        self.assertEqual([x["n"] for x in r["to"]], ["Y"])
        self.assertEqual(r["cc"], [])
        self.assertEqual(r["s"], "a@example.com")


class TestReporterAndMisc(unittest.TestCase):
    def test_error_kind_never_includes_message(self):
        self.assertEqual(T.error_kind(FakeComError(SECRET_ERR)), "FakeComError")
        self.assertEqual(T.error_kind(FakeComError(-2147221233, SECRET_ERR)), "FakeComError(0x8004010F)")

    def test_format_duration(self):
        self.assertEqual(T.format_duration(0.5), "0.500秒")
        self.assertEqual(T.format_duration(65), "1分05秒")
        self.assertEqual(T.format_duration(3720), "1時間02分")
        self.assertEqual(T.format_duration(None), "-")

    def test_reporter_console_not_in_file(self):
        with tempfile.TemporaryDirectory() as d:
            p = os.path.join(d, "x.log")
            rep = T.Reporter(p)
            buf = io.StringIO()
            with contextlib.redirect_stdout(buf):
                rep.log("logged")
                rep.console("consoleonly")
            rep.close()
            self.assertIn("consoleonly", buf.getvalue())
            self.assertEqual(_rd(p, "utf-8-sig").strip(), "logged")

    def test_reporter_unwritable_path_does_not_crash(self):
        rep = T.Reporter("/proc/nonexistent_dir/x.log")
        with contextlib.redirect_stdout(io.StringIO()):
            rep.log("x")
        rep.close()

    def test_mask_emails(self):
        self.assertEqual(T.mask_emails("a@b.com / 2024_Q2"), "<MAIL> / 2024_Q2")


# ============================================================
# レビュー指摘への対応（M1〜M8・任意改善・安全性）
# ============================================================
class TestStoreIdentity(unittest.TestCase):
    """M1: 同名ストアのキャッシュが消えない（識別はStoreID+FilePath）"""

    def two_caches(self):
        a = cache("PST", 1, "\\(ルート直下)", [rec("A1", "通関 alpha topic", "2024-04-10T09:00:00", "Sato Pm", "Owner Taro", cid="X1")],
                  sid="SID-A")
        b = cache("PST", 2, "\\(ルート直下)", [rec("B1", "通関 beta topic", "2024-04-11T09:00:00", "Sato Pm", "Owner Taro", cid="X2")],
                  sid="SID-B")
        return a, b

    def test_both_survive_with_unique_labels(self):
        a, b = self.two_caches()
        sel = T.select_caches([a, b], [], [], [], [], [])
        self.assertEqual(sorted(c["store_label"] for c in sel), ["PST", "PST #2"])
        res, _ = build([a, b])
        self.assertEqual(len(res["mails"]), 2)
        self.assertEqual(res["stats"]["raw_by_store_month"][("PST", (2024, 4))], 1)
        self.assertEqual(res["stats"]["raw_by_store_month"][("PST #2", (2024, 4))], 1)
        self.assertEqual({m["locs"][0]["store"] for m in res["mails"]}, {"PST", "PST #2"})

    def test_same_name_same_path_not_merged_as_dup_and_filter_by_label(self):
        a, b = self.two_caches()
        b["records"][0].update({"s": a["records"][0]["s"], "t": a["records"][0]["t"], "e": "B1"})
        res, _ = build([a, b])
        self.assertEqual(res["stats"]["cross_store"], 1)      # 別ストアなので統合（重複）
        self.assertEqual(len(res["mails"]), 1)
        self.assertEqual({l["store"] for l in res["mails"][0]["locs"]}, {"PST", "PST #2"})
        sel = T.select_caches([a, b], ["#2"], [], [], [], [])
        self.assertEqual([c["store_label"] for c in sel], ["PST #2"])
        sel = T.select_caches([a, b], [], ["#2"], [], [], [])
        self.assertEqual([c["store_label"] for c in sel], ["PST"])

    def test_filepath_distinguishes_when_store_id_same(self):
        a, b = self.two_caches()
        a["store_id"] = b["store_id"] = "SAME"
        a["file_path"], b["file_path"] = "C:\\a.pst", "C:\\b.pst"
        self.assertEqual(len(T.select_caches([a, b], [], [], [], [], [])), 2)

    def test_assign_labels_order_and_third(self):
        lab = T.assign_store_labels([("i3", "X", 3), ("i1", "X", 1), ("i2", "Y", 2), ("i4", "X", 4)])
        self.assertEqual(lab, {"i1": "X", "i3": "X #2", "i2": "Y", "i4": "X #3"})


class TestDedupeKeepsDistinctMails(unittest.TestCase):
    """M3: 同一ストア内でEntryIDが異なるメールは別メール。統合時のTo/CCは和集合"""

    def test_same_second_subject_sender_different_to_same_store_kept_separate(self):
        a = cache("MB", 1, "\\Inbox", [rec("E1", "Topic long subject", "2024-04-10T09:00:00", "Sato", "Alice One"),
                                       rec("E2", "Topic long subject", "2024-04-10T09:00:00", "Sato", "Bob Two")])
        mails, st = T.merge_caches_to_mails([a], (2024, 1), (2024, 12))
        self.assertEqual(len(mails), 2)
        self.assertEqual(st["same_store"], 1)
        self.assertEqual(sorted(m["to"] for m in mails), ["Alice One", "Bob Two"])

    def test_cross_store_merge_unions_to_and_cc(self):
        a = cache("MB", 1, "\\Inbox", [rec("E1", "Topic long subject", "2024-04-10T09:00:00", "Sato", "Alice One; Bob Two", "Cc One")])
        b = cache("PST", 2, "\\(ルート直下)", [rec("P1", "RE: Topic long subject", "2024-04-10T09:00:00", "Sato", "Bob Two; Carol Three", "Cc Two")])
        mails, st = T.merge_caches_to_mails([a, b], (2024, 1), (2024, 12))
        self.assertEqual(len(mails), 1)
        self.assertEqual(mails[0]["to"], "Alice One; Bob Two; Carol Three")
        self.assertEqual(mails[0]["cc"], "Cc One; Cc Two")
        self.assertEqual(st["cross_store"], 1)

    def test_union_names_fullwidth_and_empty(self):
        self.assertEqual(T.union_names("A One", "ａ　ｏｎｅ; B Two"), "A One; B Two")
        self.assertEqual(T.union_names("", "X"), "X")
        self.assertEqual(T.union_names("", ""), "")


class TestJudgmentsFile(unittest.TestCase):
    """M2: 判定JSONの破損で全消去しない"""

    def test_load_judgments_status(self):
        with tempfile.TemporaryDirectory() as d:
            p = os.path.join(d, "j.json")
            self.assertEqual(T.load_judgments(p), ({}, "new"))
            _wr(p, "")
            self.assertEqual(T.load_judgments(p)[1], "broken")
            _wr(p, "   \n")
            self.assertEqual(T.load_judgments(p)[1], "broken")
            _wr(p, "{broken")
            self.assertEqual(T.load_judgments(p)[1], "broken")
            _wr(p, "[1, 2]")
            self.assertEqual(T.load_judgments(p)[1], "broken")
            _wr(p, '{"a": {"メモ": "x"}}')
            self.assertEqual(T.load_judgments(p), ({"a": {"メモ": "x"}}, "ok"))

    def test_backup_file(self):
        with tempfile.TemporaryDirectory() as d:
            p = os.path.join(d, "j.json")
            _wr(p, "{broken")
            bak = T.backup_file(p)
            self.assertTrue(".bak_" in bak and _rd(bak) == "{broken")
            self.assertEqual(_rd(p), "{broken")                      # 元は触らない
            self.assertIsNone(T.backup_file(os.path.join(d, "nothing.json")))

    def test_partially_invalid_entries_are_kept_untouched(self):
        existing = {"t1": "oops", "t2": {"確認結果": "本物"}, "t3": None, "t4": [1]}
        merged, added = T.merge_judgments(existing, ["t1", "t2", "new"])
        self.assertEqual(merged["t1"], "oops")
        self.assertIsNone(merged["t3"])
        self.assertEqual(merged["t4"], [1])
        self.assertEqual(merged["t2"]["確認結果"], "本物")
        self.assertEqual(added, 1)
        th = make_theme()
        recs = [rec("E1", "通関", "2024-04-10T09:00:00", "Sato Pm", "Owner Taro")]
        res, _ = build([cache("S", 1, "\\Inbox", recs)], th)
        tid = res["threads"][0]["id"]
        rows = T.build_ledger_rows(res["threads"], res["evals"], {tid: "oops"}, th)    # 形が不正でも落ちない
        self.assertEqual(rows[0][T.LEDGER_COLUMNS.index("確認結果(本物/違う/保留)")], "")


class TestSafetyAndOptions(ComTestBase):
    def scan_first(self):
        com, _ = fake_com(self.standard_namespace())
        run_main(self.argv("--no-stage2"), com)

    def test_broken_and_empty_judgments_not_overwritten(self):
        self.scan_first()
        run_main(self.argv("--evaluate-only"))
        jp = os.path.join(self.data, "judgments_theme.json")
        for content in ("{broken", ""):
            _wr(jp, content)
            code, out = run_main(self.argv("--evaluate-only"))
            self.assertEqual(code, 0)
            self.assertEqual(_rd(jp), content)                                   # 上書きしない
            self.assertIn("判定ファイルを読めません", out)
            self.assertTrue(any(".bak_" in f for f in os.listdir(self.data)))
            self.assertTrue(any(f.startswith("thread_ledger_theme_") for f in os.listdir(self.out)))

    def test_stage2_guard_blocks_large_run_and_yes_overrides(self):
        self.scan_first()
        ns = Stage2Helper.namespace()
        com, pc = fake_com(ns)
        code, out = run_main(self.argv("--stage2-max", "1"), com)
        self.assertEqual(code, 3)
        self.assertIn("--stage2-max", out)
        self.assertIn("--yes", out)
        self.assertIn("--stage2-scope candidate", out)
        self.assertIn("概算", out)
        self.assertEqual(ns.opened, [])
        self.assertFalse(os.path.exists(self.paths()["stage2_cache"]))
        self.assertTrue(any(f.startswith("thread_ledger_theme_") for f in os.listdir(self.out)))   # 第1段の結果は出力する
        code, out = run_main(self.argv("--stage2-max", "1", "--yes"), fake_com(Stage2Helper.namespace())[0])
        self.assertEqual(code, 0)
        self.assertTrue(os.path.exists(self.paths()["stage2_cache"]))

    def test_stage2_candidate_scope_is_smaller_and_allowed_under_limit(self):
        self.scan_first()
        ns = Stage2Helper.namespace()
        code, out = run_main(self.argv("--stage2-max", "5", "--stage2-scope", "candidate"), fake_com(ns)[0])
        self.assertEqual(code, 0)
        self.assertEqual(len(ns.opened), 3)        # 候補スレッド(C1: E1,E2 / C2: P2)の3通

    def test_stage2_saved_in_big_batches_only(self):
        self.assertEqual(T.STAGE2_SAVE_EVERY, 1000)
        self.scan_first()
        saved = []
        orig = T.write_json_atomic
        T.write_json_atomic = lambda path, obj: (saved.append(path) if "stage2" in path else None, orig(path, obj))[1]
        try:
            run_main(self.argv(), fake_com(Stage2Helper.namespace())[0])
        finally:
            T.write_json_atomic = orig
        self.assertEqual(len(saved), 1)            # 3通なので終了時の1回だけ

    def test_stage2_error_records_are_flagged_and_refetched(self):
        self.scan_first()
        ns = Stage2Helper.namespace(Body=FakeComError(SECRET_ERR))
        run_main(self.argv(), fake_com(ns)[0])
        st2 = _rj(self.paths()["stage2_cache"])
        self.assertTrue(st2[T.stage2_key("SID-user@example.com", "E1")]["retry"])
        self.assertFalse(st2[T.stage2_key("SID-user@example.com", "E2")]["retry"])
        ns2 = Stage2Helper.namespace()
        run_main(self.argv(), fake_com(ns2)[0])
        self.assertEqual(ns2.opened, [("E1", "SID-user@example.com")])         # エラーだった1通だけ取り直す
        st3 = _rj(self.paths()["stage2_cache"])
        self.assertFalse(st3[T.stage2_key("SID-user@example.com", "E1")]["retry"])

    def test_stale_dictionary_hits_ignored_and_warned(self):
        self.scan_first()
        run_main(self.argv(), fake_com(Stage2Helper.namespace())[0])
        self.theme_data["keywords"] = {"通関": {"weight": 3, "strong": ["通関"], "weak": ["hold", "extra"]}}
        _wj(self.theme_path, self.theme_data)
        code, out = run_main(self.argv("--evaluate-only"))
        self.assertEqual(code, 0)
        self.assertIn("辞書が変わる前に本文照合した第2段レコードが 3通", out)

    def test_folders_are_fetched_by_id_and_com_refs_released(self):
        import gc
        PROXIES.clear()
        ns = self.standard_namespace()
        com, pc = fake_com(ns)
        run_main(self.argv("--no-stage2"), com)
        self.assertGreaterEqual(len(ns.folder_fetches), 2)        # Inbox と PST の(ルート直下)（+ 接続確認の GetFolderFromID）
        gc.collect()
        self.assertTrue(PROXIES and all(r() is None for r in PROXIES))
        task_like = T.enumerate_mail_folders(ns.Stores.Item(1).GetRootFolder(), "SID", [], [], {}, {})
        self.assertTrue(task_like and all("com" not in f and f["eid"] and f["sid"] == "SID" for f in task_like))

    def test_cache_replaced_via_part_file_when_range_changes(self):
        items = [mk_item(f"M{i}", f"s{i}", dt(4, i), "Zed", "Owner Taro") for i in range(1, 6)]
        com, _ = fake_com(FakeNamespace([FakeStore("MB", FakeFolder("root", list(items)))]))
        run_main(self.argv("--no-stage2"), com)
        old = self.caches()[0]
        self.assertEqual(old["range"], ["2024-01", "2024-12"])
        # 期間を広げて再取得 -> Ctrl+C で中断しても、完了済みの旧キャッシュは壊れない
        bad = list(items)
        bad[3] = mk_item("M4", "s4", dt(4, 4), "Zed", "Owner Taro", raises={"Subject": KeyboardInterrupt()})
        com2, _ = fake_com(FakeNamespace([FakeStore("MB", FakeFolder("root", bad))]))
        self.assertEqual(run_main(self.argv("--no-stage2", "--from", "2023-01"), com2)[0], 130)
        cur = self.caches()[0]
        self.assertEqual((cur["range"], cur["complete"], len(cur["records"])), (["2024-01", "2024-12"], True, 5))
        parts = [f for f in os.listdir(self.paths()["scan_cache"]) if f.endswith(".part")]
        self.assertEqual(len(parts), 1)
        com3, _ = fake_com(FakeNamespace([FakeStore("MB", FakeFolder("root", list(items)))]))
        self.assertEqual(run_main(self.argv("--no-stage2", "--from", "2023-01"), com3)[0], 0)
        new = self.caches()[0]
        self.assertEqual((new["range"], new["complete"]), (["2023-01", "2024-12"], True))
        self.assertEqual([f for f in os.listdir(self.paths()["scan_cache"]) if f.endswith(".part")], [])

    def test_mask_names_hides_store_and_folder_names_in_console_and_log_but_not_csv(self):
        com, _ = fake_com(self.standard_namespace())
        code, out = run_main(self.argv("--no-stage2", "--mask-names"), com)
        self.assertEqual(code, 0)
        self.assertIn("Store-01", out)
        self.assertIn("Folder-001", out)
        for secret in ("2024_Q2", "Inbox", "ルート直下", "user@example.com"):
            self.assertNotIn(secret, out)
        log = sorted(f for f in os.listdir(self.out) if f.endswith(".log"))[-1]
        logtext = _rd(os.path.join(self.out, log), "utf-8-sig")
        for secret in ("2024_Q2", "Inbox", "user@example.com"):
            self.assertNotIn(secret, logtext)
        led = sorted(f for f in os.listdir(self.out) if f.startswith("thread_ledger_theme_") and f.endswith(".csv"))[-1]
        csv_text = _rd(os.path.join(self.out, led), "utf-8-sig")
        self.assertTrue("2024_Q2" in csv_text or "user@example.com" in csv_text)           # CSVには元の名前
        T.configure_masking(False)

    def test_same_display_name_pst_stores_end_to_end(self):
        r1 = FakeFolder("root", [mk_item("A1", "通関 alpha subject", dt(4, 10), "Sato Pm", "Owner Taro", cid="X1")])
        r2 = FakeFolder("root", [mk_item("B1", "通関 beta subject", dt(4, 11), "Sato Pm", "Owner Taro", cid="X2")])
        ns = FakeNamespace([FakeStore("PST", r1, sid="SID-A", path="C:\\a.pst"),
                            FakeStore("PST", r2, sid="SID-B", path="C:\\b.pst")])
        code, out = run_main(self.argv("--no-stage2"), fake_com(ns)[0])
        self.assertEqual(code, 0)
        self.assertEqual(len(self.caches()), 2)
        self.assertIn("PST #2", out)
        led = sorted(f for f in os.listdir(self.out) if f.startswith("thread_ledger_theme_") and f.endswith(".csv"))[-1]
        rows = _rdict(os.path.join(self.out, led))
        self.assertEqual(len(rows), 2)
        self.assertTrue(any("PST #2" in r["取得元(ストア/フォルダ)"] for r in rows))
        run_main(self.argv("--evaluate-only", "--stores", "#2"))
        led2 = sorted(f for f in os.listdir(self.out) if f.startswith("thread_ledger_theme_") and f.endswith(".csv"))[-1]
        self.assertEqual(len(_rdict(os.path.join(self.out, led2))), 1)

    def test_cli_alias_and_defaults(self):
        a = T.parse_args(["--candidate-terms"])
        self.assertTrue(a.suggest_terms)
        self.assertEqual((a.stage2_max, a.yes, a.mask_names), (5000, False, False))
        self.assertTrue(T.parse_args(["--suggest-terms"]).suggest_terms)


class Stage2Helper:
    """第2段のフェイク名前空間（開いたメールを記録する）"""

    @staticmethod
    def namespace(**kw):
        sid_mb, sid_pst = "SID-user@example.com", "SID-2024_Q2"
        body = "please see 通関 and hold"

        def full(eid, s_addr, recips, **r):
            return FakeItem({"Class": 43, "EntryID": eid, "SenderEmailAddress": s_addr, "SenderEmailType": "SMTP",
                             "PropertyAccessor": FakePA({T.PR_SENDER_SMTP: s_addr, T.PR_INTERNET_MESSAGE_ID: "<id@x>"}),
                             "Recipients": FakeRecipients(recips), "Attachments": FakeAtts(["a.pdf"]),
                             "Body": body}, raises=r, stage1=False)
        by_id = {
            ("E1", sid_mb): full("E1", "pm@example.com", [("Owner Taro", "owner@example.com", 1)], **kw),
            ("E2", sid_mb): full("E2", "owner@example.com", [("Sato Pm", "pm@example.com", 1)]),
            ("P2", sid_pst): full("P2", "pm@example.com", [("Owner Taro", "owner@example.com", 1)]),
        }

        class Recording(FakeNamespace):
            opened = None

            def GetItemFromID(self, eid, sid):
                self.opened.append((eid, sid))
                return FakeNamespace.GetItemFromID(self, eid, sid)
        ns = Recording([], by_id)
        ns.opened = []
        return ns


class TestOptionalImprovements(unittest.TestCase):
    def test_classify_exact_preferred_and_partial_skips_generic(self):
        th = make_theme()
        recs = [
            rec("E1", "Weekly lunch plan arrangement", "2024-04-10T09:00:00", "Zed", "Zed", cid="a"),               # 候補外(完全一致)
            rec("E2", "Weekly lunch plan arrangement and more", "2024-04-11T09:00:00", "Sato Pm", "Owner Taro; Tanaka Partner", cid="b"),  # 候補(部分一致側)
        ]
        res, _ = build([cache("S", 1, "\\Inbox", recs)], th)
        r = T.classify_check_subjects(["Weekly lunch plan arrangement"], res["threads"], res["evals"])[0]
        self.assertEqual(r["status"], "found_not_candidate")     # 完全一致が候補外なら、部分一致の候補があっても found にしない
        self.assertEqual([m["kind"] for m in r["matches"]], ["完全一致"])
        r2 = T.classify_check_subjects(["lunch"], res["threads"], res["evals"])[0]
        self.assertEqual(r2["status"], "missing")                # 短い汎用件名は部分一致しない
        r3 = T.classify_check_subjects(["Weekly lunch plan arrangement and"], res["threads"], res["evals"])[0]
        self.assertEqual(r3["status"], "found")                  # 十分に長ければ部分一致する
        self.assertEqual(r3["matches"][0]["kind"], "部分一致")

    def test_generic_subjects_not_used_for_thread_grouping(self):
        a = cache("MB", 1, "\\Inbox", [rec("E1", "FYI", "2024-04-10T09:00:00", "A"), rec("E2", "RE: FYI", "2024-04-11T09:00:00", "B"),
                                       rec("E3", "お疲れ様です", "2024-04-12T09:00:00", "A"), rec("E4", "short", "2024-04-13T09:00:00", "A"),
                                       rec("E5", "Quarterly review meeting", "2024-04-14T09:00:00", "A"),
                                       rec("E6", "RE: Quarterly review meeting", "2024-04-15T09:00:00", "B")])
        mails, _ = T.merge_caches_to_mails([a], (2024, 1), (2024, 12))
        threads = T.build_threads(mails)
        self.assertEqual(sorted(t["count"] for t in threads), [1, 1, 1, 1, 2])
        self.assertTrue(T.is_generic_subject("short"))
        self.assertTrue(T.is_generic_subject("meeting"))
        self.assertFalse(T.is_generic_subject("quarterly review meeting"))

    def test_name_matches_is_memoized(self):
        T.name_matches.cache_clear()
        T.name_matches("Zed Alias", "Zed Alias Person")
        T.name_matches("Zed Alias", "Zed Alias Person")
        self.assertGreaterEqual(T.name_matches.cache_info().hits, 1)

    def test_store_zero_months(self):
        stats = {"raw_by_store_month": {("A", (2024, 1)): 3, ("A", (2024, 3)): 1}}
        z = T.store_zero_months(stats, [(2024, 1), (2024, 2), (2024, 3)], ["A", "B"])
        self.assertEqual(z["A"], [(2024, 2)])
        self.assertEqual(len(z["B"]), 3)

    def test_resume_does_not_double_count_errors(self):
        cache_ = {"records": [{"e": "E1", "t": "2024-04-01T09:00:00"}], "last_received": "2024-04-01T09:00:00",
                  "counts": {"non_mail": 2, "errors": {"Subject:X": 2}}, "date_field": "or"}
        st = T.new_scan_state(cache_)
        self.assertEqual(st["errors"], {})
        T.add_error(st["errors"], "Subject:X")
        task = {"sidx": 1, "store_name": "S", "store_id": "I", "folder_path": "\\F"}
        p1 = T.make_cache_payload(task, st, (2024, 1), (2024, 12), False)
        p2 = T.make_cache_payload(task, st, (2024, 1), (2024, 12), False)      # 保存を繰り返しても増えない
        self.assertEqual(p1["counts"]["errors"], {"Subject:X": 3})
        self.assertEqual(p2["counts"]["errors"], {"Subject:X": 3})

    def test_known_items_skipped_without_reading_other_fields(self):
        item = mk_item("K1", SECRET, dt(4, 1), "Zed", "Owner Taro", raises={"Subject": FakeComError(SECRET_ERR)})
        errs = {}
        self.assertEqual(T.read_light_item(item, errs, {"K1"}), ("dup", None))
        self.assertEqual(errs, {})

    def test_pattern_safety_and_length_caps(self):
        with self.assertRaises(T.ThemeError):
            make_theme(patterns=[{"category": "c", "regex": "(a+)+$", "strong": True}])
        with self.assertRaises(T.ThemeError):
            make_theme(patterns=[{"category": "c", "regex": "a" * 300, "strong": True}])
        make_theme(patterns=[{"category": "c", "regex": r"\b\d{4}\.\d{2}(\.\d{2,4})?\b", "strong": False}])   # 通常の式は可
        tbl = make_theme()["hit_table"]
        self.assertEqual(T.scan_text_hits("x" * (T.MAX_BODY_CHARS + 10) + " 通関", tbl), {})            # 本文の上限
        th = make_theme()
        recs = [rec("E1", "x" * 400 + " 通関", "2024-04-10T09:00:00", "Sato Pm", "Owner Taro")]
        _, e = one_thread(recs, th)
        self.assertFalse(e["C_ii"])                                                                      # 件名は300字まで

    def test_write_json_atomic_fsync(self):
        calls = []
        orig = os.fsync
        os.fsync = lambda fd: calls.append(fd)
        try:
            with tempfile.TemporaryDirectory() as d:
                T.write_json_atomic(os.path.join(d, "a.json"), {"x": 1})
                self.assertEqual(_rj(os.path.join(d, "a.json")), {"x": 1})
        finally:
            os.fsync = orig
        self.assertEqual(len(calls), 1)

    def test_no_real_surnames_in_repo_files(self):
        import re
        words = ["中" + "井", "梶" + "川", "佐" + "治", "Na" + "kai", "Kaji" + "kawa", "Sa" + "ji", "nexp" + "eria",
                 "dhl" + r"\.com", "trade" + "win"]
        pat = re.compile("|".join(words), re.IGNORECASE)
        targets = [os.path.join(TOOLS_DIR, f) for f in ("thread_ledger_scan_20261010_05.py",
                   "thread_ledger_theme.example.json", "calibration_subjects.example.txt")] + [os.path.abspath(__file__)]
        for t in targets:
            self.assertIsNone(pat.search(_rd(t)), t)



class TestReReviewFixes(ComTestBase):
    def test_call_by_id_omits_empty_store_id(self):
        calls = []
        T.call_by_id(lambda *a: calls.append(a), "E", "")
        T.call_by_id(lambda *a: calls.append(a), "E", None)
        T.call_by_id(lambda *a: calls.append(a), "E", "S")
        self.assertEqual(calls, [("E",), ("E",), ("E", "S")])

    def test_scan_with_empty_store_id_calls_without_second_argument(self):
        root = FakeFolder("root", [mk_item("A1", "通関 x subject", dt(4, 10), "Sato Pm", "Owner Taro")])
        inbox = FakeFolder("Inbox", [mk_item("A2", "通関 y subject", dt(4, 11), "Sato Pm", "Owner Taro")])
        root.Folders.folders.append(inbox)
        ns = FakeNamespace([FakeStore("MB", root, sid="")])
        code, out = run_main(self.argv("--no-stage2"), fake_com(ns)[0])
        self.assertEqual(code, 0)
        self.assertEqual(len(self.caches()), 2)
        self.assertTrue(all(sid is None for _e, sid in ns.folder_fetches))     # 第2引数は省略された

    def root_ns(self, lookup_fails=False):
        root = FakeFolder("root", [mk_item("A1", "通関 x subject", dt(4, 10), "Sato Pm", "Owner Taro")])
        root.EntryID = ""                       # PSTのルートでEntryIDが空
        ns = FakeNamespace([FakeStore("PST", root, sid="SID-PST", path="C:\\a.pst")])
        ns.store_lookup_fails = lookup_fails
        return ns

    def test_root_without_entry_id_uses_get_store_from_id(self):
        ns = self.root_ns()
        code, out = run_main(self.argv("--no-stage2"), fake_com(ns)[0])
        self.assertEqual(code, 0)
        self.assertEqual(ns.store_lookups, ["SID-PST"])
        self.assertEqual(len(self.caches()), 1)
        self.assertEqual(self.caches()[0]["records"][0]["e"], "A1")
        self.assertNotIn("EntryIDなし", out)

    def test_root_without_entry_id_falls_back_to_enumeration_store_ref(self):
        ns = self.root_ns(lookup_fails=True)
        code, out = run_main(self.argv("--no-stage2"), fake_com(ns)[0])
        self.assertEqual(code, 0)
        self.assertEqual(ns.store_lookups, ["SID-PST"])
        self.assertEqual(self.caches()[0]["records"][0]["e"], "A1")
        self.assertNotIn(SECRET_ERR, out)

    def test_non_root_folder_without_entry_id_is_skipped_with_error_kind(self):
        sub = FakeFolder("Sub", [mk_item("A1", "x", dt(4, 10), "Zed", "Owner Taro")])
        sub.EntryID = ""
        ns = FakeNamespace([FakeStore("MB", FakeFolder("root", [], subfolders=[sub]))])
        code, out = run_main(self.argv("--no-stage2"), fake_com(ns)[0])
        self.assertEqual(code, 2)                 # 取得できたキャッシュが無いので評価は中止
        self.assertIn("フォルダEntryIDなし", out)
        self.assertEqual(self.caches(), [])

    def test_stage2_final_save_happens_even_if_periodic_save_fails(self):
        com, _ = fake_com(self.standard_namespace())
        run_main(self.argv("--no-stage2"), com)
        old_every, orig = T.STAGE2_SAVE_EVERY, T.write_json_atomic
        calls = []

        def flaky(path, obj):
            if "stage2" in path:
                calls.append(path)
                if len(calls) == 1:
                    raise OSError(SECRET_ERR)
            return orig(path, obj)
        T.STAGE2_SAVE_EVERY, T.write_json_atomic = 1, flaky
        try:
            code, out = run_main(self.argv(), fake_com(Stage2Helper.namespace())[0])
        finally:
            T.STAGE2_SAVE_EVERY, T.write_json_atomic = old_every, orig
        self.assertEqual(code, 0)
        self.assertNotIn(SECRET_ERR, out)
        self.assertIn("途中保存に失敗", out)
        self.assertEqual(len(_rj(self.paths()["stage2_cache"])), 3)          # 最終保存で全件が残る

    def test_incomplete_rescan_does_not_replace_old_complete_cache(self):
        items = [mk_item(f"M{i}", f"s{i}", dt(4, i), "Zed", "Owner Taro") for i in range(1, 4)]
        run_main(self.argv("--no-stage2"), fake_com(FakeNamespace([FakeStore("MB", FakeFolder("root", list(items)))]))[0])
        before = self.caches()[0]
        self.assertTrue(before["complete"])
        old = T.MAX_ITEM_SKIP
        T.MAX_ITEM_SKIP = 1
        try:
            f = FakeFolder("root", list(items), next_raises=(1, FakeComError(SECRET_ERR)),
                           item_raises=FakeComError(SECRET_ERR))
            code, out = run_main(self.argv("--no-stage2", "--from", "2023-01"),
                                 fake_com(FakeNamespace([FakeStore("MB", f)]))[0])
        finally:
            T.MAX_ITEM_SKIP = old
        self.assertEqual(code, 0)
        self.assertNotIn(SECRET_ERR, out)
        cur = self.caches()[0]
        self.assertEqual((cur["range"], cur["complete"], len(cur["records"])), (["2024-01", "2024-12"], True, 3))
        parts = [x for x in os.listdir(self.paths()["scan_cache"]) if x.endswith(".part")]
        self.assertEqual(len(parts), 1)
        part = _rj(os.path.join(self.paths()["scan_cache"], parts[0]))
        self.assertFalse(part["complete"])
        self.assertEqual(part["range"], ["2023-01", "2024-12"])

    def test_count_failure_does_not_mark_folder_complete(self):
        items = [mk_item(f"M{i}", f"s{i}", dt(4, i), "Zed", "Owner Taro") for i in range(1, 4)]

        class NoCount(FakeRestricted):
            @property
            def Count(self):
                raise FakeComError(SECRET_ERR)

            def GetNext(self):
                raise FakeComError(SECRET_ERR)

        class F:
            class Items:
                @staticmethod
                def Restrict(flt):
                    return NoCount(items, "x")
        st = T.new_scan_state()
        T.scan_folder_items(F, datetime(2024, 1, 1), datetime(2025, 1, 1), st, total_hint=3)
        self.assertFalse(st["complete"])
        self.assertIn("Count:FakeComError", st["errors"])
        self.assertEqual([r["e"] for r in st["records"]], ["M1"])

    def test_usage_mentions_trial_run(self):
        doc = T.__doc__
        self.assertIn("1か月", doc)
        self.assertIn("--stores", doc)

    def test_example_theme_comment_explains_aliases(self):
        txt = _rd(os.path.join(TOOLS_DIR, "thread_ledger_theme.example.json"))
        self.assertIn("別々の別名として両方書く", txt)



# ============================================================
# S1.5: 件名絞り込みで開く・xlsx・プロトコル・判定取込・参考候補・見積り
# ============================================================
class FakeExplorer:
    def __init__(self, selectable=True, search_raises=None, folder_explorer=None):
        self.selectable = selectable
        self.search_raises = search_raises
        self.CurrentFolder = None
        self.queries = []
        self.selected = []
        self.activated = 0
        self.displayed = 0
        self.cleared = 0

    def Activate(self):
        self.activated += 1

    def Display(self):
        self.displayed += 1

    def Search(self, query, scope):
        if self.search_raises is not None:
            raise self.search_raises
        self.queries.append((query, scope))

    def IsItemSelectableInView(self, item):
        return self.selectable

    def ClearSelection(self):
        self.cleared += 1

    def AddToSelection(self, item):
        self.selected.append(item)


class FakeParentFolder:
    def __init__(self, explorer=None):
        self.explorer = explorer

    def GetExplorer(self):
        return self.explorer


class TestSearchSubject(unittest.TestCase):
    def test_safe_search_subject(self):
        f = T.safe_search_subject
        self.assertEqual(f("RE: FW: ** internal only ** [Export request]2024-04-01 10:00"), "Export request 2024 04 01 10 00")
        self.assertEqual(f("Re: Import to Japan（PO 8210238665）"), "Import to Japan")
        self.assertEqual(f("FW: 【DHL】Import \"x\" (a) to Japan from US#0000000000"), "【DHL】Import x a to Japan from US#0000000000")
        self.assertEqual(f('a"b:c;d<e>{f}\'g\\h*i%j'), "a b c d e f g h i j")
        self.assertEqual(f("  ＲＥ：  Ｈｅｌｌｏ　　Ｗｏｒｌｄ  "), "Hello World")
        self.assertEqual(f("RE:"), "")
        self.assertEqual(f(None), "")
        self.assertEqual(len(f("x" * 500)), 200)

    def test_build_subject_query(self):
        self.assertEqual(T.build_subject_query("RE: Topic (A)"), 'subject:"Topic"')
        self.assertEqual(T.build_subject_query("RE: ()"), "")
        q = T.build_subject_query('x" OR from:evil')
        self.assertNotIn('"', q[len('subject:"'):-1])


class TestOpenInOutlook(unittest.TestCase):
    def setUp(self):
        ACCESSED.clear()

    def run_open(self, explorer, item_fields=None, parent=None, **kw):
        shown = []
        fields = {"EntryID": "E" * 20, "Subject": "RE: Topic about shipment (PO 1)", "Parent": parent or FakeParentFolder()}
        fields.update(item_fields or {})
        item = FakeItem(fields, display_log=shown, **kw)
        ns = FakeNamespace([], {("E" * 20, "S" * 20): item})
        rep = T.Reporter(None)
        with contextlib.redirect_stdout(io.StringIO()):
            mode = T.open_thread_in_outlook(FakeClient(ns, explorer), "E" * 20, "S" * 20, "", rep)
        return mode, shown, item, rep

    def test_active_explorer_switches_folder_searches_and_selects(self):
        ex = FakeExplorer()
        parent = FakeParentFolder()
        mode, shown, item, _ = self.run_open(ex, parent=parent)
        self.assertEqual(mode, "explorer")
        self.assertIs(ex.CurrentFolder, parent)                 # メールのあるフォルダに切り替え
        self.assertEqual(ex.queries, [('subject:"Topic about shipment"', 0)])
        self.assertEqual(ex.selected, [item])
        self.assertEqual(shown, [])
        self.assertEqual(ACCESSED, [])

    def test_no_active_explorer_creates_one_from_parent_folder(self):
        ex = FakeExplorer()
        mode, shown, item, _ = self.run_open(None, parent=FakeParentFolder(ex))
        self.assertEqual(mode, "explorer")
        self.assertEqual(ex.displayed, 1)
        self.assertEqual(len(ex.queries), 1)

    def test_falls_back_to_display_when_not_selectable_or_error(self):
        mode, shown, _i, _r = self.run_open(FakeExplorer(selectable=False))
        self.assertEqual((mode, len(shown)), ("display", 1))
        mode, shown, _i, rep = self.run_open(FakeExplorer(search_raises=FakeComError(SECRET_ERR)))
        self.assertEqual((mode, len(shown)), ("display", 1))
        self.assertNotIn(SECRET_ERR, "\n".join(rep.lines))
        mode, shown, _i, _r = self.run_open(FakeExplorer(), item_fields={"Subject": "RE:"})
        self.assertEqual((mode, len(shown)), ("display", 1))             # 件名が空なら絞り込まない
        mode, shown, _i, _r = self.run_open(FakeExplorer(), item_fields={"Parent": None})
        self.assertEqual(mode, "display")

    def test_run_open_by_entry_id_validates_hex(self):
        args = T.parse_args(["--entry-id", "../x", "--store-id", "ab"])
        rep = T.Reporter(None)
        with contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(T.run_open(args, rep, T.make_paths("/x", "/y"), (2024, 1), (2024, 12)), 2)
        ns = FakeNamespace([], {})
        args2 = T.parse_args(["--entry-id", "ab" * 10, "--store-id", "cd" * 10])
        item = FakeItem({"EntryID": "ab" * 10, "Subject": "Topic long subject", "Parent": FakeParentFolder()})
        ns.by_id[("ab" * 10, "cd" * 10)] = item
        ex = FakeExplorer()
        with contextlib.redirect_stdout(io.StringIO()):
            code = T.run_open(args2, rep, T.make_paths("/x", "/y"), (2024, 1), (2024, 12), (FakeClient(ns, ex), FakePythonCom()))
        self.assertEqual((code, len(ex.queries)), (0, 1))


class TestThreadIdAndUrlValidation(unittest.TestCase):
    def test_thread_id_regex(self):
        ok = ["abcd", "a" * 64, "0123456789ab", "A_b-C"]
        bad = ["", "abc", "a" * 65, "../etc", "a b c d", "abc;rm", "ab\tcd", "a/b/c/d", "abcd\n", "ａｂｃｄ", "ab..cd", "a&b|cd"]
        for t in ok:
            self.assertTrue(T.LEDGER_THREAD_ID_RE.match(t), t)
        for t in bad:
            self.assertIsNone(T.LEDGER_THREAD_ID_RE.match(t), repr(t))

    def test_parse_ledger_url(self):
        self.assertEqual(T.parse_ledger_url("ledger:abcd1234"), "abcd1234")
        self.assertEqual(T.parse_ledger_url("LEDGER:abcd1234"), "abcd1234")
        self.assertEqual(T.parse_ledger_url("ledger://abcd1234/"), "abcd1234")
        for bad in ("", None, "ledger:", "ledger:../../x", "ledger:ab cd", "ledger:abcd;calc", "http://x", "ledger:" + "a" * 65,
                    "ledger:abcd&calc", "xledger:abcd1234", "ledger:abcd1234 --evil", "file:///c:/x"):
            self.assertIsNone(T.parse_ledger_url(bad), repr(bad))


class FakeWinreg:
    HKEY_CURRENT_USER = "HKCU"
    REG_SZ = 1

    def __init__(self):
        self.created, self.values, self.deleted, self.closed = [], [], [], 0

    def CreateKey(self, root, sub):
        self.created.append((root, sub))
        return (root, sub)

    def SetValueEx(self, key, name, reserved, typ, val):
        self.values.append((key, name, typ, val))

    def CloseKey(self, key):
        self.closed += 1

    def DeleteKey(self, root, sub):
        if sub.endswith("\\shell"):
            raise FileNotFoundError(sub)
        self.deleted.append((root, sub))


class TestProtocol(unittest.TestCase):
    def test_command_string_and_plan(self):
        self.assertEqual(T.protocol_command("C:\\Py\\pythonw.exe", "C:\\t\\s.py"),
                         '"C:\\Py\\pythonw.exe" "C:\\t\\s.py" --open-url "%1"')
        plan = T.protocol_plan("C:\\Py\\pythonw.exe", "C:\\t\\s.py")
        self.assertTrue(all(sub.startswith("Software\\Classes\\ledger") for sub, _n, _v in plan))     # HKCU配下のみ
        self.assertIn(("Software\\Classes\\ledger", "URL Protocol", ""), plan)
        self.assertEqual(plan[-1][0], "Software\\Classes\\ledger\\shell\\open\\command")
        with self.assertRaises(ValueError):
            T.protocol_plan('C:\\a"b\\pythonw.exe', "C:\\t\\s.py")

    def test_pythonw_path(self):
        with tempfile.TemporaryDirectory() as d:
            exe = os.path.join(d, "python.exe")
            _wr(exe, "")
            self.assertEqual(T.pythonw_path(exe), exe)                         # pythonw.exe が無ければ元のまま
            _wr(os.path.join(d, "pythonw.exe"), "")
            self.assertEqual(T.pythonw_path(exe), os.path.join(d, "pythonw.exe"))

    def run_main(self, *argv, answer=None):
        import unittest.mock as mock
        wr = FakeWinreg()
        buf = io.StringIO()
        with tempfile.TemporaryDirectory() as d, contextlib.redirect_stdout(buf):
            args = list(argv) + ["--data-dir", d, "--output-dir", d]
            if answer is None:
                code = T.main(args, winreg_mod=wr)
            else:
                with mock.patch("builtins.input", return_value=answer):
                    code = T.main(args, winreg_mod=wr)
        return code, wr, buf.getvalue()

    def test_register_with_yes_writes_hkcu_only(self):
        code, wr, out = self.run_main("--register-protocol", "--yes")
        self.assertEqual(code, 0)
        self.assertTrue(all(root == "HKCU" for root, _s in wr.created))
        cmd = [v for (_k, n, _t, v) in wr.values if n == "" and "--open-url" in v][0]
        self.assertTrue(cmd.endswith('--open-url "%1"'))
        self.assertIn(os.path.abspath(T.__file__), cmd)
        self.assertEqual(wr.closed, 3)
        self.assertIn("--open-url", out)            # 実行前に内容を表示

    def test_register_asks_confirmation(self):
        code, wr, out = self.run_main("--register-protocol", answer="n")
        self.assertEqual((code, wr.values), (1, []))
        code, wr, out = self.run_main("--register-protocol", answer="y")
        self.assertEqual(code, 0)
        self.assertTrue(wr.values)

    def test_unregister_deletes_deepest_first_and_ignores_missing(self):
        code, wr, out = self.run_main("--unregister-protocol", "--yes")
        self.assertEqual(code, 0)
        subs = [s for _r, s in wr.deleted]
        self.assertEqual(subs[0], "Software\\Classes\\ledger\\shell\\open\\command")
        self.assertEqual(subs[-1], "Software\\Classes\\ledger")

    def test_without_winreg_reports_error(self):
        import unittest.mock as mock
        buf = io.StringIO()
        with tempfile.TemporaryDirectory() as d, contextlib.redirect_stdout(buf):
            with mock.patch.dict(sys.modules, {"winreg": None}):
                code = T.main(["--register-protocol", "--yes", "--data-dir", d, "--output-dir", d])
        self.assertEqual(code, 2)


class LedgerFlowBase(ComTestBase):
    def scan_first(self):
        com, _ = fake_com(self.standard_namespace())
        run_main(self.argv("--no-stage2"), com)

    def latest(self, prefix, ext):
        fs = sorted(f for f in os.listdir(self.out) if f.startswith(prefix) and f.endswith(ext))
        return os.path.join(self.out, fs[-1])


class TestOpenUrlFlow(LedgerFlowBase):
    def patched_paths(self):
        import unittest.mock as mock
        orig = T.make_paths
        return mock.patch.object(T, "make_paths", lambda d=None, o=None: orig(self.data, self.out))

    def test_open_url_valid_opens_via_ledger_and_writes_no_log_on_success(self):
        import unittest.mock as mock
        self.scan_first()
        run_main(self.argv("--evaluate-only"))
        row = _rdict(self.latest("thread_ledger_theme_", ".csv"))[0]
        item = FakeItem({"EntryID": row["最新メールEntryID"], "Subject": "Topic long subject", "Parent": FakeParentFolder()})
        ns = FakeNamespace([], {(row["最新メールEntryID"], row["最新メールStoreID"]): item})
        ex = FakeExplorer()
        before = set(os.listdir(self.out))
        boxes = []
        with self.patched_paths(), mock.patch.object(T, "HEX_ID_RE", re.compile(r".+")), \
                mock.patch.object(T, "show_message_box", lambda *a, **k: boxes.append(a)):
            code, out = run_main(["--open-url", "ledger:" + row["thread_id"]], fake_com(ns, ex)[0])
        self.assertEqual((code, len(ex.queries)), (0, 1))
        self.assertEqual(boxes, [])
        self.assertEqual(set(os.listdir(self.out)), before)             # 成功時はログを増やさない

    def test_open_url_failures_show_message_box_and_write_log(self):
        import unittest.mock as mock
        orig = T.import_com
        T.import_com = lambda: (_ for _ in ()).throw(AssertionError("COM禁止"))
        try:
            for bad in ("ledger:../../x", "ledger:ab cd", "calc.exe", "ledger:abcd;calc"):
                boxes = []
                before = set(os.listdir(self.out)) if os.path.isdir(self.out) else set()
                with self.patched_paths(), mock.patch.object(T, "show_message_box", lambda *a, **k: boxes.append(a)):
                    code, out = run_main(["--open-url", bad])
                self.assertEqual(code, 2, bad)
                self.assertNotIn("calc", out)
                self.assertEqual(len(boxes), 1, bad)
                self.assertIn("開けませんでした", boxes[0][0])
                new = set(os.listdir(self.out)) - before
                self.assertTrue(any(f.startswith("thread_ledger_open_") for f in new), bad)    # 失敗時だけログ
        finally:
            T.import_com = orig

    def test_open_url_not_found_and_outlook_failure_show_message(self):
        import unittest.mock as mock
        boxes = []
        with self.patched_paths(), mock.patch.object(T, "show_message_box", lambda *a, **k: boxes.append(a)):
            code, _o = run_main(["--open-url", "ledger:abcd1234"], fake_com(FakeNamespace([], {}))[0])
        self.assertEqual(code, 2)
        self.assertIn("見つかりません", boxes[0][0])
        self.scan_first()
        run_main(self.argv("--evaluate-only"))
        row = _rdict(self.latest("thread_ledger_theme_", ".csv"))[0]
        boxes.clear()
        with self.patched_paths(), mock.patch.object(T, "HEX_ID_RE", re.compile(r".+")), \
                mock.patch.object(T, "show_message_box", lambda *a, **k: boxes.append(a)):
            code, _o = run_main(["--open-url", "ledger:" + row["thread_id"]], fake_com(FakeNamespace([], {}))[0])
        self.assertEqual(code, 2)
        self.assertEqual(len(boxes), 1)
        self.assertNotIn(SECRET_ERR, boxes[0][0])

    def test_non_hex_ids_from_ledger_are_rejected(self):
        self.scan_first()
        run_main(self.argv("--evaluate-only"))
        row = _rdict(self.latest("thread_ledger_theme_", ".csv"))[0]
        code, out = run_main(["--open", row["thread_id"], "--data-dir", self.data, "--output-dir", self.out],
                             fake_com(FakeNamespace([], {}))[0])
        self.assertEqual(code, 2)
        self.assertIn("形式が不正", out)


class TestOpenUrlInjection(unittest.TestCase):
    """--open-url の値に " が入って別オプションが注入されても、何も実行しない"""

    def setUp(self):
        import unittest.mock as mock
        self.boxes = []
        self.wr = FakeWinreg()
        self.patches = [mock.patch.object(T, "show_message_box", lambda *a, **k: self.boxes.append(a)),
                        mock.patch.object(T, "write_failure_log", lambda *a, **k: None),     # リポジトリ側にログを作らない
                        mock.patch.object(T, "import_com", lambda: (_ for _ in ()).throw(AssertionError("COM禁止")))]
        for p in self.patches:
            p.start()

    def tearDown(self):
        for p in self.patches:
            p.stop()

    def call(self, argv):
        buf = io.StringIO()
        with contextlib.redirect_stdout(buf), contextlib.redirect_stderr(io.StringIO()):
            try:
                code = T.main(argv, winreg_mod=self.wr)
            except SystemExit as e:
                code = e.code
        return code

    def test_extra_options_are_refused_without_side_effects(self):
        cases = [
            ["--open-url", "ledger:abcd", "--unregister-protocol", "--yes"],
            ["--open-url", "ledger:abcd", "--register-protocol", "--yes"],
            ["--open-url", "ledger:abcd", "--data-dir", "x"],
            ["--open-url", "ledger:abcd", "--output-dir", "x"],
            ["--open-url", "ledger:abcd", "--import-judgments", "x.csv"],
            ["--open-url=ledger:abcd"],
            ["--yes", "--open-url", "ledger:abcd"],
            ["--open-url", "ledger:abcd", "--open", "abcd1234"],
        ]
        for argv in cases:
            self.boxes.clear()
            code = self.call(argv)
            self.assertEqual(code, 2, argv)
            self.assertEqual((self.wr.created, self.wr.deleted, self.wr.values), ([], [], []), argv)
            self.assertEqual(len(self.boxes), 1, argv)

    def test_empty_open_url_value_cannot_bypass_the_strict_check(self):
        with tempfile.TemporaryDirectory() as d:
            combos = [
                ["--unregister-protocol", "--yes"],
                ["--register-protocol", "--yes"],
                ["--data-dir", d],
                ["--output-dir", d],
                ["--theme="],
                ["--theme", ""],
                ["--import-judgments", os.path.join(d, "x.csv")],
                ["--unregister-protocol", "--yes", "--output-dir", d, "--theme", ""],
                ["--unregister-protocol", "--yes", "--theme="],
            ]
            for extra in combos:
                for head in (["--open-url", ""], ["--open-url="]):
                    for argv in (head + extra, extra + head):
                        self.boxes.clear()
                        code = self.call(argv)
                        self.assertEqual(code, 2, argv)
                        self.assertEqual((self.wr.created, self.wr.deleted, self.wr.values), ([], [], []), argv)
                        self.assertEqual(os.listdir(d), [], argv)            # ファイルにも触れない
                        self.assertGreaterEqual(len(self.boxes), 1, argv)
            # 実攻撃形: %1 に `" --unregister-protocol --yes --theme="` が入った場合
            self.boxes.clear()
            self.assertEqual(self.call(["--open-url", "", "--unregister-protocol", "--yes", "--theme", ""]), 2)
            self.assertEqual(self.wr.deleted, [])

    def test_empty_open_url_alone_is_rejected_as_invalid_url(self):
        self.assertEqual(self.call(["--open-url", ""]), 2)
        self.assertEqual(len(self.boxes), 1)
        self.assertIsNone(T.parse_args([]).open_url)
        self.assertEqual(T.parse_args(["--open-url", ""]).open_url, "")
        self.assertTrue(T.has_open_url_token(["--x", "--open-url="]))
        self.assertFalse(T.has_open_url_token(["--open-u", "x"]))
        self.assertFalse(T.exact_open_url_argv(["--open-url", None], None))

    def test_parse_error_with_open_url_notifies_by_message_box(self):
        code = self.call(["--open-url", "ledger:abcd", "--no-such-option"])
        self.assertEqual(code, 2)                                   # argparse の SystemExit(2)
        self.assertEqual(len(self.boxes), 1)
        self.boxes.clear()
        self.assertEqual(self.call(["--no-such-option"]), 2)
        self.assertEqual(self.boxes, [])                            # --open-url が無ければ通知しない

    def test_abbreviated_options_are_refused(self):
        for argv in (["--reg", "--yes"], ["--unreg", "--yes"], ["--open-u", "ledger:abcd"], ["--open-url", "ledger:abcd", "--reg"]):
            code = self.call(argv)
            self.assertNotEqual(code, 0, argv)
            self.assertEqual((self.wr.created, self.wr.deleted), ([], []), argv)

    def test_exact_argv_helper(self):
        self.assertTrue(T.exact_open_url_argv(["--open-url", "ledger:abcd"], "ledger:abcd"))
        self.assertFalse(T.exact_open_url_argv(["--open-url", "ledger:abcd", "--yes"], "ledger:abcd"))
        self.assertFalse(T.exact_open_url_argv(["--open-url=ledger:abcd"], "ledger:abcd"))

    def test_data_and_output_dirs_ignored_for_open_url(self):
        import unittest.mock as mock
        seen = []
        orig = T.make_paths
        with mock.patch.object(T, "make_paths", lambda d=None, o=None: (seen.append((d, o)), orig(d, o))[1]):
            self.call(["--open-url", "ledger:abcd1234"])
        self.assertEqual(seen, [(None, None)])


class TestScriptEntryPoint(unittest.TestCase):
    """直接実行（python tool.py）で動くこと。末尾の main ガードが消えると何も起きなくなる。"""

    def run_tool(self, *args):
        import subprocess
        out_dir = os.path.join(os.path.dirname(TESTS_DIR), "mail_reports")
        before = set(os.listdir(out_dir)) if os.path.isdir(out_dir) else set()
        path = sorted(glob.glob(os.path.join(TOOLS_DIR, "thread_ledger_scan_*.py")))[-1]
        r = subprocess.run([sys.executable, "-I", path, *args], capture_output=True, text=True, timeout=60)
        # 失敗時ログがリポジトリ側の mail_reports に増えていたら片付ける（ごみ箱代わりに一時フォルダへ移す）
        if os.path.isdir(out_dir):
            for f in set(os.listdir(out_dir)) - before:
                shutil_move(os.path.join(out_dir, f), tempfile.gettempdir())
        return r

    def test_help_prints_usage(self):
        r = self.run_tool("--help")
        self.assertEqual(r.returncode, 0)
        self.assertIn("usage", r.stdout)
        self.assertIn("--open-url", r.stdout)
        self.assertIn("数値", r.stdout)                 # thread_id のExcel数値化の注意書き

    def test_open_url_invalid_value_is_rejected_nonzero(self):
        for bad in ("ledger:../x", "calc.exe", "ledger:ab cd"):
            r = self.run_tool("--open-url", bad)
            self.assertNotEqual(r.returncode, 0, bad)

    def test_source_ends_with_main_guard_and_newline(self):
        text = _rd(sorted(glob.glob(os.path.join(TOOLS_DIR, "thread_ledger_scan_*.py")))[-1])
        self.assertTrue(text.endswith('if __name__ == "__main__":\n    sys.exit(main())\n'))


def shutil_move(src, dst_dir):
    import shutil
    try:
        shutil.move(src, os.path.join(dst_dir, "ledger_test_" + os.path.basename(src)))
    except Exception:
        pass


class TestMessageBox(unittest.TestCase):
    def test_calls_messageboxw_and_ignores_failures(self):
        import types
        calls = []
        fake = types.SimpleNamespace(windll=types.SimpleNamespace(user32=types.SimpleNamespace(
            MessageBoxW=lambda *a: calls.append(a))))
        import unittest.mock as mock
        with mock.patch.dict(sys.modules, {"ctypes": fake}):
            T.show_message_box("本文", "題")
        self.assertEqual(calls, [(0, "本文", "題", 0x10)])
        with mock.patch.dict(sys.modules, {"ctypes": types.SimpleNamespace()}):
            T.show_message_box("x")          # windll が無い（非Windows）でも例外にならない


@unittest.skipIf(T.import_openpyxl() is None, "openpyxl が無い環境")
class TestXlsx(LedgerFlowBase):
    def test_xlsx_written_with_ledger_hyperlinks_and_safe_cells(self):
        import openpyxl
        header = ["thread_id", "件名", "メール数", "スコア", "確認結果(本物/違う/保留)", "メモ"]
        rows = [["abcd1234ef56", "=HYPERLINK(\"http://evil\",\"x\")", 3, "7.5", "", "+cmd"],
                ["0123456789ab", "-1+1 normal", 1, "2", "", "@x"],
                ["../evil", "bad id", 1, "1", "", ""]]
        with tempfile.TemporaryDirectory() as d:
            p = os.path.join(d, "l.xlsx")
            self.assertTrue(T.write_ledger_xlsx(p, header, rows))
            wb = openpyxl.load_workbook(p)
            ws = wb.active
            self.assertEqual(ws["B2"].hyperlink.target, "ledger:abcd1234ef56")
            self.assertEqual(ws["B3"].hyperlink.target, "ledger:0123456789ab")
            self.assertIsNone(ws["B4"].hyperlink)                     # 形式が不正な thread_id にはリンクを張らない
            for ref in ("B2", "F2", "B3", "F3"):
                self.assertEqual(ws[ref].data_type, "s", ref)           # 数式にならない
                self.assertTrue(str(ws[ref].value).startswith("'"), ref)
            self.assertEqual(ws["C2"].value, 3)
            self.assertEqual(ws["D2"].value, 7.5)
            self.assertEqual(ws["A1"].value, "thread_id")
            self.assertEqual(ws.freeze_panes, "C2")
            self.assertTrue(ws.data_validations.dataValidation)

    def test_pipeline_writes_xlsx_next_to_csv(self):
        import openpyxl
        self.scan_first()
        code, out = run_main(self.argv("--evaluate-only"))
        self.assertIn("台帳Excel", out)
        wb = openpyxl.load_workbook(self.latest("thread_ledger_theme_", ".xlsx"))
        ws = wb.active
        csv_rows = _rdict(self.latest("thread_ledger_theme_", ".csv"))
        self.assertEqual(ws.max_row - 1, len(csv_rows))
        self.assertEqual({ws.cell(row=r, column=1).value for r in range(2, ws.max_row + 1)}, {r["thread_id"] for r in csv_rows})
        for r in range(2, ws.max_row + 1):
            self.assertEqual(ws.cell(row=r, column=2).hyperlink.target, "ledger:" + ws.cell(row=r, column=1).value)


class TestXlsxMissing(LedgerFlowBase):
    def test_csv_only_with_guidance_when_openpyxl_missing(self):
        self.scan_first()
        for f in os.listdir(self.out):
            if f.endswith(".xlsx"):
                os.remove(os.path.join(self.out, f))      # 最初の走査（openpyxlあり）で出たものを消す
        orig = T.import_openpyxl
        T.import_openpyxl = lambda: None
        try:
            code, out = run_main(self.argv("--evaluate-only"))
        finally:
            T.import_openpyxl = orig
        self.assertEqual(code, 0)
        self.assertIn("pip install openpyxl", out)
        self.assertFalse(any(f.endswith(".xlsx") for f in os.listdir(self.out)))
        self.assertTrue(any(f.startswith("thread_ledger_theme_") and f.endswith(".csv") for f in os.listdir(self.out)))


class TestImportJudgments(LedgerFlowBase):
    def test_import_rows_allowlist_empty_not_overwrite_and_matching(self):
        existing = {"aaaa1111": {"確認結果": "本物", "メモ": "old", "問題の種類": "通関"},
                    "bbbb2222": {"確認結果": "", "メモ": "", "問題の種類": ""}, "cccc3333": "oops",
                    "ffff6666": {"確認結果": "", "メモ": "", "問題の種類": ""}}
        rows = [
            {"thread_id": "aaaa1111", "確認結果(本物/違う/保留)": "", "メモ": "", "問題の種類": ""},      # 空は変更なし
            {"thread_id": "bbbb2222", "確認結果(本物/違う/保留)": "違う", "メモ": "new memo", "問題の種類": "その他"},
            {"thread_id": "dddd4444", "確認結果(本物/違う/保留)": "保留", "メモ": "", "問題の種類": ""},    # 台帳に無い -> スキップ
            {"thread_id": "eeee5555", "確認結果(本物/違う/保留)": "たぶん本物", "メモ": "m", "問題の種類": ""},  # 台帳に無い -> スキップ
            {"thread_id": "ffff6666", "確認結果(本物/違う/保留)": "たぶん本物"},                             # 不正値のみ -> 変化なし
            {"thread_id": "../bad", "確認結果(本物/違う/保留)": "本物"},
            {"thread_id": "cccc3333", "確認結果(本物/違う/保留)": "本物"},
            {"thread_id": "'=evil01", "確認結果(本物/違う/保留)": "本物"},
        ]
        merged, st = T.import_judgment_rows(rows, existing)
        self.assertEqual(merged["aaaa1111"], existing["aaaa1111"])
        self.assertEqual(merged["bbbb2222"], {"確認結果": "違う", "メモ": "new memo", "問題の種類": "その他"})
        self.assertNotIn("dddd4444", merged)                    # 既知IDのみ
        self.assertNotIn("eeee5555", merged)                    # 空エントリも作らない
        self.assertEqual(merged["ffff6666"], {"確認結果": "", "メモ": "", "問題の種類": ""})
        self.assertEqual(merged["cccc3333"], "oops")
        self.assertNotIn("../bad", merged)
        self.assertEqual((st["invalid_value"], st["unknown"], st["bad_id"], st["bad_entry"]), (1, 2, 2, 1))
        self.assertEqual((st["matched"], st["updated"]), (3, 3))
        self.assertEqual(set(merged), set(existing))

    def test_import_unescapes_leading_apostrophe(self):
        existing = {"aaaa1111": {"確認結果": "", "メモ": "", "問題の種類": ""}}
        merged, _ = T.import_judgment_rows([{"thread_id": "aaaa1111", "メモ": "'=memo", "確認結果": "'本物"}], existing)
        self.assertEqual(merged["aaaa1111"]["メモ"], "=memo")
        self.assertEqual(merged["aaaa1111"]["確認結果"], "")        # 許可リスト外(アポストロフィ付き)はスキップ

    def test_size_limits(self):
        import unittest.mock as mock
        with tempfile.TemporaryDirectory() as d:
            p = os.path.join(d, "a.csv")
            _wr(p, "thread_id,メモ\n" + "".join("aaaa%04d,x\n" % i for i in range(10)), "utf-8-sig")
            with mock.patch.object(T, "IMPORT_MAX_ROWS", 5):
                with self.assertRaises(ValueError):
                    T.read_table_rows(p)
            with mock.patch.object(T, "IMPORT_MAX_BYTES", 10):
                with self.assertRaises(ValueError):
                    T.read_table_rows(p)
            self.assertEqual(len(T.read_table_rows(p)), 10)
            self.assertEqual((T.IMPORT_MAX_BYTES, T.IMPORT_MAX_ROWS), (50 * 1024 * 1024, 200000))

    def test_import_backs_up_existing_valid_judgments_before_writing(self):
        csv_path, rows = self.edit_ledger_csv()
        jp = os.path.join(self.data, "judgments_theme.json")
        before = _rd(jp)
        tid = rows[0]["thread_id"]
        edited = os.path.join(self.tmp.name, "e.csv")
        _wr(edited, "thread_id,確認結果(本物/違う/保留),メモ,問題の種類\n%s,本物,m,\n" % tid, "utf-8-sig")
        code, out = run_main(self.argv("--import-judgments", edited))
        self.assertEqual(code, 0)
        baks = [f for f in os.listdir(self.data) if ".bak_" in f]
        self.assertEqual(len(baks), 1)
        self.assertEqual(_rd(os.path.join(self.data, baks[0])), before)       # 取込前の内容が残っている
        self.assertEqual(_rj(jp)[tid]["確認結果"], "本物")
        self.assertIn("退避", out)

    def test_import_aborts_when_backup_cannot_be_made(self):
        import unittest.mock as mock
        csv_path, rows = self.edit_ledger_csv()
        jp = os.path.join(self.data, "judgments_theme.json")
        before = _rd(jp)
        edited = os.path.join(self.tmp.name, "e.csv")
        _wr(edited, "thread_id,確認結果(本物/違う/保留)\n%s,本物\n" % rows[0]["thread_id"], "utf-8-sig")
        with mock.patch.object(T, "backup_file", lambda p: None):
            code, out = run_main(self.argv("--import-judgments", edited))
        self.assertEqual(code, 2)
        self.assertEqual(_rd(jp), before)

    def edit_ledger_csv(self):
        self.scan_first()
        run_main(self.argv("--evaluate-only"))
        csv_path = self.latest("thread_ledger_theme_", ".csv")
        rows = _rdict(csv_path)
        return csv_path, rows

    def test_csv_import_end_to_end(self):
        csv_path, rows = self.edit_ledger_csv()
        tid = rows[0]["thread_id"]
        edited = os.path.join(self.tmp.name, "edited.csv")
        with open(edited, "w", encoding="cp932", newline="") as fh:        # Excel の「CSV(コンマ区切り)」はcp932
            w = csv.writer(fh)
            w.writerow(["thread_id", "確認結果(本物/違う/保留)", "メモ", "問題の種類"])
            w.writerow([tid, "本物", "確認済み", "通関"])
            w.writerow(["zzzz9999", "違う", "", ""])
        code, out = run_main(self.argv("--import-judgments", edited))
        self.assertEqual(code, 0)
        self.assertIn("判定の取込", out)
        j = _rj(os.path.join(self.data, "judgments_theme.json"))
        self.assertEqual(j[tid], {"確認結果": "本物", "メモ": "確認済み", "問題の種類": "通関"})
        self.assertNotIn("zzzz9999", j)                            # 台帳に無いIDは取り込まない
        self.assertIn("台帳に無いIDをスキップ 1", out)
        # 再評価しても保持される
        run_main(self.argv("--evaluate-only"))
        self.assertEqual(_rj(os.path.join(self.data, "judgments_theme.json"))[tid]["確認結果"], "本物")
        # 空セルでは上書きしない
        blank = os.path.join(self.tmp.name, "blank.csv")
        _wr(blank, "thread_id,確認結果(本物/違う/保留),メモ,問題の種類\n%s,,,\n" % tid, "utf-8-sig")
        run_main(self.argv("--import-judgments", blank))
        self.assertEqual(_rj(os.path.join(self.data, "judgments_theme.json"))[tid]["メモ"], "確認済み")

    @unittest.skipIf(T.import_openpyxl() is None, "openpyxl が無い環境")
    def test_xlsx_import_end_to_end(self):
        import openpyxl
        csv_path, rows = self.edit_ledger_csv()
        xlsx = self.latest("thread_ledger_theme_", ".xlsx")
        wb = openpyxl.load_workbook(xlsx)
        ws = wb.active
        header = [c.value for c in ws[1]]
        vi = header.index("確認結果(本物/違う/保留)") + 1
        ws.cell(row=2, column=vi).value = "保留"
        ws.cell(row=2, column=header.index("メモ") + 1).value = "Excelで編集"
        tid = ws.cell(row=2, column=1).value
        edited = os.path.join(self.tmp.name, "edited.xlsx")
        wb.save(edited)
        code, out = run_main(self.argv("--import-judgments", edited))
        self.assertEqual(code, 0)
        j = _rj(os.path.join(self.data, "judgments_theme.json"))
        self.assertEqual((j[tid]["確認結果"], j[tid]["メモ"]), ("保留", "Excelで編集"))

    def test_broken_existing_judgments_abort_import_with_backup(self):
        csv_path, rows = self.edit_ledger_csv()
        jp = os.path.join(self.data, "judgments_theme.json")
        _wr(jp, "{broken")
        code, out = run_main(self.argv("--import-judgments", csv_path))
        self.assertEqual(code, 2)
        self.assertEqual(_rd(jp), "{broken")
        self.assertTrue(any(".bak_" in f for f in os.listdir(self.data)))

    def test_missing_thread_id_column_or_file(self):
        csv_path, rows = self.edit_ledger_csv()
        bad = os.path.join(self.tmp.name, "bad.csv")
        _wr(bad, "a,b\n1,2\n", "utf-8-sig")
        self.assertEqual(run_main(self.argv("--import-judgments", bad))[0], 2)
        self.assertEqual(run_main(self.argv("--import-judgments", os.path.join(self.tmp.name, "none.csv")))[0], 2)


class TestIncludeNoAnchor(unittest.TestCase):
    def build_threads(self):
        th = make_theme()
        recs = [
            rec("E1", "通関 candidate subject", "2024-04-10T09:00:00", "Sato Pm", "Owner Taro", cid="a"),     # 通常候補
            rec("E2", "通関 no anchor subject", "2024-04-11T09:00:00", "Zed", "Owner Taro", cid="b"),         # A無し・B・C -> 参考
            rec("E3", "customs clearance no owner", "2024-04-12T09:00:00", "Sato Pm", "Zed", cid="c"),        # B無し -> 対象外
            rec("E4", "Lunch", "2024-04-13T09:00:00", "Zed", "Owner Taro", cid="d"),                          # C無し -> 対象外
        ]
        res, _ = build([cache("S", 1, "\\Inbox", recs)], th)
        return res, th

    def test_default_off_and_on_ordering_and_flag(self):
        res, th = self.build_threads()
        off = T.build_ledger_rows(res["threads"], res["evals"], {}, th)
        self.assertEqual(len(off), 1)
        self.assertEqual(off[0][-1], "")
        on = T.build_ledger_rows(res["threads"], res["evals"], {}, th, include_no_anchor=True)
        self.assertEqual(len(on), 2)
        flag = T.LEDGER_COLUMNS.index(T.NO_ANCHOR_COLUMN)
        self.assertEqual([r[flag] for r in on], ["", "参考"])           # 参考は通常候補の後
        self.assertEqual(T.LEDGER_COLUMNS[-1], T.NO_ANCHOR_COLUMN)
        self.assertTrue(T.is_reference_thread(res["evals"][1]))
        self.assertFalse(T.is_reference_thread(res["evals"][0]))

    def test_cli_flag_end_to_end(self):
        res, th = self.build_threads()
        self.assertTrue(T.parse_args(["--include-no-anchor"]).include_no_anchor)
        self.assertFalse(T.parse_args([]).include_no_anchor)


class TestEstimates(LedgerFlowBase):
    def test_constants_and_stage1_estimate_logged(self):
        self.assertEqual(T.STAGE2_EST_SEC, 0.8)
        self.assertEqual(T.STAGE1_EST_SEC_PER_ITEM, 0.06)
        self.assertEqual(T.STAGE2_MAX_DEFAULT, 5000)
        com, _ = fake_com(self.standard_namespace())
        code, out = run_main(self.argv("--no-stage2"), com)
        self.assertIn("第1段の見積り", out)
        self.assertIn("60ms", out)

    def test_stage2_estimate_uses_0_8_sec(self):
        self.scan_first()
        code, out = run_main(self.argv("--stage2-max", "1"), fake_com(Stage2Helper.namespace())[0])
        self.assertIn("1通0.8秒", out)
        self.assertEqual(code, 3)



class FlakySelectExplorer(FakeExplorer):
    """絞り込みの反映が遅れて、N回目の確認で初めて選択できるようになる Explorer"""

    def __init__(self, ready_on):
        super().__init__()
        self.ready_on = ready_on
        self.checks = 0

    def IsItemSelectableInView(self, item):
        self.checks += 1
        return self.checks >= self.ready_on


class TestExplorerTiming(unittest.TestCase):
    def open_with(self, ex):
        item = FakeItem({"EntryID": "E" * 20, "Subject": "Topic about shipment", "Parent": FakeParentFolder()},
                        display_log=[])
        ns = FakeNamespace([], {("E" * 20, "S" * 20): item})
        with contextlib.redirect_stdout(io.StringIO()):
            return T.open_thread_in_outlook(FakeClient(ns, ex), "E" * 20, "S" * 20, "", T.Reporter(None))

    def test_explorer_wait_defaults(self):
        self.assertTrue(0.3 <= ORIG_WAITS[0] <= 0.5)
        self.assertEqual(ORIG_WAITS[1], 0.3)
        self.assertEqual(ORIG_WAITS[2], 5)

    def test_settle_wait_happens_between_folder_switch_and_search(self):
        import unittest.mock as mock
        order = []

        class Ex(FakeExplorer):
            def __setattr__(self, k, v):
                if k == "CurrentFolder":
                    order.append("folder")
                object.__setattr__(self, k, v)

            def Search(self, q, scope):
                order.append("search")
                return super().Search(q, scope)
        ex = Ex()
        order.clear()
        with mock.patch.object(T.time, "sleep", lambda sec: order.append(("sleep", sec))):
            self.open_with(ex)
        self.assertEqual(order[0], "folder")
        self.assertEqual(order[1][0], "sleep")
        self.assertEqual(order[2], "search")

    def test_selection_retried_until_ready(self):
        ex = FlakySelectExplorer(ready_on=3)
        self.assertEqual(self.open_with(ex), "explorer")
        self.assertEqual(ex.checks, 3)
        ex2 = FlakySelectExplorer(ready_on=99)
        self.assertEqual(self.open_with(ex2), "display")             # 5回で諦めて Display()
        self.assertEqual(ex2.checks, 5)


class TestProtocolExisting(unittest.TestCase):
    def test_register_shows_existing_registration_and_unregister_text(self):
        class WR(FakeWinreg):
            def OpenKey(self, root, sub):
                return ("k", sub)

            def QueryValueEx(self, key, name):
                return ('"C:\\old\\pythonw.exe" "C:\\old.py" --open-url "%1"', 1)
        for flag, expect in (("--register-protocol", "上書きします"), ("--unregister-protocol", "削除対象です")):
            wr = WR()
            buf = io.StringIO()
            with tempfile.TemporaryDirectory() as d, contextlib.redirect_stdout(buf):
                code = T.main([flag, "--yes", "--data-dir", d, "--output-dir", d], winreg_mod=wr)
            self.assertEqual(code, 0)
            self.assertIn("既存の ledger: 登録があります", buf.getvalue())
            self.assertIn("C:\\\\old\\\\pythonw.exe", buf.getvalue())
            self.assertIn(expect, buf.getvalue())
            if flag == "--unregister-protocol":
                self.assertIn("非再帰", buf.getvalue())
                self.assertIn("4キー", buf.getvalue())
                self.assertIn("/4キー", buf.getvalue())

    def test_no_existing_registration_has_no_warning(self):
        wr = FakeWinreg()
        wr.OpenKey = lambda root, sub: (_ for _ in ()).throw(FileNotFoundError(sub))
        buf = io.StringIO()
        with tempfile.TemporaryDirectory() as d, contextlib.redirect_stdout(buf):
            T.main(["--register-protocol", "--yes", "--data-dir", d, "--output-dir", d], winreg_mod=wr)
        self.assertNotIn("既存の ledger", buf.getvalue())


# ============================================================
# S1.6: Table方式・会話ID統一・レジリエンス・診断
# ============================================================
CONV_A = bytes(range(1, 17))
CONV_A_HEX = "0102030405060708090A0B0C0D0E0F10"


class TestConvIdAndClass(unittest.TestCase):
    def test_conv_hex_variants(self):
        self.assertEqual(T.conv_hex(CONV_A), CONV_A_HEX)
        self.assertEqual(T.conv_hex(bytearray(CONV_A)), CONV_A_HEX)
        self.assertEqual(T.conv_hex(memoryview(CONV_A)), CONV_A_HEX)
        self.assertEqual(T.conv_hex(tuple(CONV_A)), CONV_A_HEX)
        self.assertEqual(T.conv_hex(list(CONV_A)), CONV_A_HEX)
        self.assertEqual(T.conv_hex(CONV_A_HEX.lower()), CONV_A_HEX)
        self.assertEqual(T.conv_hex(b""), "")
        for bad in (None, "", "zz-not-hex", 12, (1, "x")):
            self.assertEqual(T.conv_hex(bad), "", repr(bad))

    def test_is_note_class(self):
        for ok in ("IPM.Note", "ipm.note", "IPM.Note.SMIME", "IPM.Note.Exchange.Rules.Reminder"):
            self.assertTrue(T.is_note_class(ok), ok)
        for bad in ("", None, "IPM.Appointment", "IPM.Notefoo", "IPM.Schedule.Meeting.Request", "REPORT.IPM.Note.NDR"):
            self.assertFalse(T.is_note_class(bad), repr(bad))

    def test_same_mail_gives_same_conv_id_by_items_and_table_and_threads_merge(self):
        item = mk_item("A1", "Topic long subject", dt(4, 10), "Sato Pm", "Owner Taro", cid="OUTLOOK-STRING",
                       PropertyAccessor=FakePA({T.PR_CONVERSATION_ID_URL: CONV_A}))
        kind, rec_items = T.read_light_item(item, {})
        self.assertEqual((kind, rec_items["c"], rec_items["c2"]), ("mail", CONV_A_HEX, "OUTLOOK-STRING"))
        f = FakeFolder("F", table_rows=[trow("T1", "Topic long subject", dt(4, 11), conv=CONV_A)])
        st = T.new_scan_state()
        T.scan_folder_table(f, datetime(2024, 1, 1), datetime(2025, 1, 1), st)
        self.assertEqual(st["records"][0]["c"], CONV_A_HEX)
        a = cache("MB", 1, "\\Inbox", [rec_items])
        b = cache("PST", 2, "\\x", [dict(st["records"][0], t="2024-04-11T09:00:00")])
        mails, _ = T.merge_caches_to_mails([a, b], (2024, 1), (2024, 12))
        threads = T.build_threads(mails)
        self.assertEqual(len(threads), 1)                 # Items由来とTable由来のスレッドが統合される
        self.assertEqual(threads[0]["key"], "c:" + CONV_A_HEX)

    def test_c2_used_when_pr_conversation_id_missing(self):
        mails, _ = T.merge_caches_to_mails([cache("MB", 1, "\\Inbox", [rec("E1", "x long subject", "2024-04-10T09:00:00", "A")])],
                                           (2024, 1), (2024, 12))
        r = rec("E2", "y long subject", "2024-04-10T09:00:01", "A")
        r["c2"] = "OLD-FORM"
        mails, _ = T.merge_caches_to_mails([cache("MB", 1, "\\Inbox", [r])], (2024, 1), (2024, 12))
        self.assertEqual((mails[0]["cid"], mails[0]["cid2"]), ("", "OLD-FORM"))      # c2 はスレッドキーにしない（保持のみ）
        self.assertEqual(T.thread_key_of(mails[0])[:2], "s:")                        # 既定: 正規化件名
        self.assertEqual(T.thread_key_of(mails[0], "c2"), "ci:OLD-FORM")             # 選べば接頭辞つきで使う（c由来と衝突しない）
        mails2, _ = T.merge_caches_to_mails([cache("MB", 1, "\\Inbox", [rec("E3", "z long subject", "2024-04-10T09:00:02", "A", cid="OLD-FORM")])],
                                            (2024, 1), (2024, 12))
        self.assertEqual(T.thread_key_of(mails2[0], "c2"), "c:OLD-FORM")             # c由来とは別のキー空間
        self.assertNotEqual(T.thread_key_of(mails[0], "c2"), T.thread_key_of(mails2[0], "c2"))
        self.assertEqual(len(T.build_threads(mails + mails2, "c2")), 2)
    def test_pa_failure_other_than_missing_is_counted(self):
        errs = {}
        item = mk_item("A1", "s", dt(4, 10), "Zed", "Owner Taro",
                       PropertyAccessor=FakePA({}))          # GetProperty が FakeComError
        T.read_light_item(item, errs)
        self.assertIn("ConvID(PA):FakeComError", errs)

    def test_constants_and_schema(self):
        self.assertEqual(T.SCHEMA_VERSION, 2)
        self.assertEqual(ORIG_RETRY_WAITS, (5, 15))
        self.assertEqual((T.RECONNECT_AFTER, T.TABLE_CHUNK), (3, 500))
        self.assertEqual(T.EST_SEC, {"pst": 0.06, "table": 0.010, "items": 1.0})


class TableBase(ComTestBase):
    def argv_t(self, *extra):
        return self.argv("--scan-mode", "table", "--no-stage2", *extra)

    def one_store(self, folder, name="MB", path=""):
        return FakeNamespace([FakeStore(name, FakeFolder("root", [], subfolders=[folder]), path=path)])

    def mails_for(self, n, start_day=1):
        return [trow(f"R{i}", f"s{i}", dt(4, start_day + i), conv=CONV_A) for i in range(n)]


class TestTableScan(TableBase):
    def test_basic_rows_values_and_classes(self):
        rows = [
            trow("R1", "s1", dt(4, 1), smtp="pm@example.com", conv=CONV_A, cc="Cc One"),
            trow("R2", "s2", None, sent=dt(4, 2), semail="/O=EXCH/CN=X"),                 # ReceivedTime 無し -> 送信日時
            trow("R3", "s3", dt(4, 3), cls="IPM.Appointment"),                              # メール以外
            trow("R4", None, dt(4, 4), sender=None, to=None, conv=None),                    # None の値
            trow("R5", "s5", dt(4, 5), cls="IPM.Note.SMIME"),
            trow("R6", "s6", datetime(2023, 1, 1), cls="IPM.Note"),                        # 期間外
            trow("", "s8", dt(4, 8)),                                                      # EntryID なし
        ]
        f = FakeFolder("F", table_rows=rows)
        code, out = run_main(self.argv_t(), fake_com(self.one_store(f))[0])
        self.assertEqual(code, 0)
        c = self.caches()[0]
        by = {r["e"]: r for r in c["records"]}
        self.assertEqual(sorted(by), ["R1", "R2", "R4", "R5"])
        self.assertEqual((by["R1"]["a"], by["R1"]["c"], by["R1"]["cc"]), ("pm@example.com", CONV_A_HEX, "Cc One"))
        self.assertEqual(by["R2"]["t"], "2024-04-02T09:00:00")
        self.assertEqual(by["R2"]["a"], "/O=EXCH/CN=X")
        self.assertEqual((by["R4"]["s"], by["R4"]["n"], by["R4"]["to"], by["R4"]["c"]), ("", "", "", ""))
        self.assertEqual(c["counts"]["non_mail"], 1)
        self.assertEqual((c["schema"], c["date_field"], c["scan_mode"], c["complete"]), (2, "table", "table", True))
        self.assertIn("EntryIDなし", c["counts"]["errors"])
        self.assertEqual(f._items.filters, [])                      # アイテムは開かない（Restrictも呼ばない）
        self.assertEqual(f.table_filters[0].count(" OR "), 1)
        self.assertNotIn(SECRET_ERR, out)

    def test_chunk_boundaries(self):
        old = T.TABLE_CHUNK
        try:
            for n_rows, chunk in ((7, 3), (6, 3), (1, 3), (3, 3), (0, 3), (4, 500)):
                T.TABLE_CHUNK = chunk
                f = FakeFolder("F", table_rows=self.mails_for(n_rows))
                st = T.new_scan_state()
                T.scan_folder_table(f, datetime(2024, 1, 1), datetime(2025, 1, 1), st)
                self.assertEqual([r["e"] for r in st["records"]], [f"R{i}" for i in range(n_rows)], (n_rows, chunk))
                self.assertTrue(st["complete"])
        finally:
            T.TABLE_CHUNK = old

    def test_column_failure_is_tolerated_but_entry_id_failure_is_not(self):
        f = FakeFolder("F", table_rows=[trow("R1", "s1", dt(4, 1), cc="Cc One", conv=CONV_A)],
                       table_opts={"col_fail": ("urn:schemas:httpmail:displaycc", T.PR_CONVERSATION_ID_URL)})
        st = T.new_scan_state()
        T.scan_folder_table(f, datetime(2024, 1, 1), datetime(2025, 1, 1), st)
        self.assertEqual((st["records"][0]["cc"], st["records"][0]["c"], st["records"][0]["s"]), ("", "", "s1"))
        self.assertIn("Column(cc):FakeComError", st["errors"])
        self.assertIn("Column(conv):FakeComError", st["errors"])
        f2 = FakeFolder("F", table_rows=[trow("R1", "s1", dt(4, 1))], table_opts={"col_fail": ("EntryID",)})
        with self.assertRaises(T.TableUnavailable):
            T.scan_folder_table(f2, datetime(2024, 1, 1), datetime(2025, 1, 1), T.new_scan_state())
        f3 = FakeFolder("F", table_rows=[trow("R1", "s1", dt(4, 1))],
                        table_opts={"col_fail": ("urn:schemas:httpmail:datereceived", T.PR_URL + "0x00390040")})
        with self.assertRaises(T.TableUnavailable):                  # 日時の列が全く無い
            T.scan_folder_table(f3, datetime(2024, 1, 1), datetime(2025, 1, 1), T.new_scan_state())

    def test_class_column_failure_accepts_all_rows(self):
        f = FakeFolder("F", table_rows=[trow("R1", "s1", dt(4, 1), cls="IPM.Appointment")],
                       table_opts={"col_fail": (dict(T.TABLE_COLUMNS)["cls"],)})
        st = T.new_scan_state()
        T.scan_folder_table(f, datetime(2024, 1, 1), datetime(2025, 1, 1), st)
        self.assertEqual(len(st["records"]), 1)

    def test_filter_fallback_order(self):
        rows = [trow("IN", "s", dt(4, 1)), trow("OUT", "s", datetime(2023, 5, 1)), trow("SENT", "s", None, sent=dt(5, 1))]
        f = FakeFolder("F", table_rows=rows, table_opts={"fail_or": True})
        st = T.new_scan_state()
        T.scan_folder_table(f, datetime(2024, 1, 1), datetime(2025, 1, 1), st)
        self.assertEqual([x.split()[0] if x else x for x in f.table_filters], ["([ReceivedTime]", "[ReceivedTime]"])
        self.assertEqual([r["e"] for r in st["records"]], ["IN"])             # ReceivedTime のみのフィルタ
        self.assertEqual(st["table_filter"], "ReceivedTime")
        f2 = FakeFolder("F", table_rows=rows, table_opts={"fail_any_filter": True})
        st2 = T.new_scan_state()
        T.scan_folder_table(f2, datetime(2024, 1, 1), datetime(2025, 1, 1), st2)
        self.assertEqual(f2.table_filters[-1], None)
        self.assertEqual(len(f2.table_filters), 3)
        self.assertEqual(sorted(r["e"] for r in st2["records"]), ["IN", "SENT"])   # 全件取得して手元で期間に絞る
        self.assertEqual(st2["table_filter"], "all")
        f_nodate = FakeFolder("F", table_rows=[trow("ND", "s", datetime(4501, 1, 1), sent=None)], table_opts={"fail_any_filter": True})
        st_nd = T.new_scan_state()
        T.scan_folder_table(f_nodate, datetime(2024, 1, 1), datetime(2025, 1, 1), st_nd)
        self.assertEqual((st_nd["records"], st_nd["errors"].get("日時なし")), ([], 1))        # 日時が取れない行はエラー種別だけ集計
        f3 = FakeFolder("F", table_rows=rows)
        st3 = T.new_scan_state()
        T.scan_folder_table(f3, datetime(2024, 1, 1), datetime(2025, 1, 1), st3)
        self.assertEqual(sorted(r["e"] for r in st3["records"]), ["IN", "SENT"])   # OR で送信日時だけのメールも拾う
        self.assertEqual(st3["table_filter"], "or")

    def test_table_failure_falls_back_to_items_with_warning(self):
        f = FakeFolder("F", [mk_item("I1", "s1", dt(4, 1), "Zed", "Owner Taro")], table_rows=None)    # GetTable が失敗
        code, out = run_main(self.argv_t(), fake_com(self.one_store(f))[0])
        self.assertEqual(code, 0)
        self.assertIn("Items方式に切り替えます", out)
        self.assertNotIn(SECRET_ERR, out)
        c = self.caches()[0]
        self.assertEqual([r["e"] for r in c["records"]], ["I1"])
        self.assertEqual(c["scan_mode"], "items")
        self.assertIn("Tableフォールバック", c["counts"]["errors"])

    def test_auto_mode_uses_items_for_pst_and_table_otherwise(self):
        pst_f = FakeFolder("P", [mk_item("P1", "s1", dt(4, 1), "Zed", "Owner Taro")], table_rows=self.mails_for(2))
        mb_f = FakeFolder("M", [mk_item("M1", "s1", dt(4, 1), "Zed", "Owner Taro")], table_rows=self.mails_for(3))
        ns = FakeNamespace([FakeStore("2024_Q2", FakeFolder("root", [], subfolders=[pst_f]), path="C:\\a.pst"),
                            FakeStore("MB", FakeFolder("root", [], subfolders=[mb_f]))])
        code, out = run_main(self.argv("--scan-mode", "auto", "--no-stage2"), fake_com(ns)[0])
        self.assertEqual(code, 0)
        self.assertEqual(pst_f.table_filters, [])               # PST は Items
        self.assertTrue(mb_f.table_filters)                     # PST 以外は Table
        self.assertEqual(mb_f._items.filters, [])
        modes = {c["store_name"]: c["scan_mode"] for c in self.caches()}
        self.assertEqual(modes, {"2024_Q2": "items", "MB": "table"})
        self.assertEqual(T.choose_scan_mode("auto", "PST"), "items")
        self.assertEqual(T.choose_scan_mode("auto", "アーカイブ"), "table")
        self.assertEqual(T.choose_scan_mode("table", "PST"), "table")
        self.assertEqual(T.choose_scan_mode("items", "メールボックス/PST"), "items")

    def test_estimates_by_mode_and_rates_logged(self):
        pst_f = FakeFolder("P", [mk_item("P1", "s1", dt(4, 1), "Zed", "Owner Taro")])
        mb_f = FakeFolder("M", [mk_item("M1", "s1", dt(4, 1), "Zed", "Owner Taro")], table_rows=self.mails_for(3))
        ns = FakeNamespace([FakeStore("2024_Q2", FakeFolder("root", [], subfolders=[pst_f]), path="C:\\a.pst"),
                            FakeStore("MB", FakeFolder("root", [], subfolders=[mb_f]))])
        code, out = run_main(self.argv("--scan-mode", "auto", "--no-stage2"), fake_com(ns)[0])
        self.assertIn("PST(Items方式) 1件 × 60ms", out)
        self.assertIn("Table方式(非PST) 1件 × 10ms(仮置き。実測で補正)", out)
        self.assertIn("行/秒", out)
        self.assertIn("📈", out)
        code, out = run_main(self.argv("--scan-mode", "items", "--no-stage2", "--rescan"),
                             fake_com(FakeNamespace([FakeStore("MB", FakeFolder("root", [], subfolders=[
                                 FakeFolder("M", [mk_item("M1", "s", dt(4, 1), "Zed", "Owner Taro")])]))]))[0])
        self.assertIn("Items方式(非PST) 1件 × 1000ms", out)

    def test_table_mode_has_no_unreadable_items_threshold(self):
        old = T.MAX_ITEM_SKIP
        T.MAX_ITEM_SKIP = 0
        try:
            rows = [trow(f"R{i}", f"s{i}", dt(4, 1 + i % 20)) for i in range(30)] + [trow(None, "bad", dt(4, 1))] * 20
            f = FakeFolder("F", table_rows=rows)
            st = T.new_scan_state()
            T.scan_folder_table(f, datetime(2024, 1, 1), datetime(2025, 1, 1), st)
        finally:
            T.MAX_ITEM_SKIP = old
        self.assertTrue(st["complete"])
        self.assertEqual(len(st["records"]), 30)

    def test_checkpoint_during_table_scan(self):
        saved = []
        st = T.new_scan_state()
        f = FakeFolder("F", table_rows=self.mails_for(7))
        T.scan_folder_table(f, datetime(2024, 1, 1), datetime(2025, 1, 1), st, checkpoint_cb=lambda s_: saved.append(len(s_["records"])),
                            checkpoint_every=3)
        self.assertEqual(saved, [3, 6])


class TestResilience(TableBase):
    def test_transient_error_is_retried_and_resumes_without_duplicates(self):
        old = T.TABLE_CHUNK
        T.TABLE_CHUNK = 2
        try:
            f = FakeFolder("F", table_rows=self.mails_for(7), table_opts={"array_fail_calls": {2}})
            code, out = run_main(self.argv_t(), fake_com(self.one_store(f))[0])
        finally:
            T.TABLE_CHUNK = old
        self.assertEqual(code, 0)
        self.assertIn("リトライ 1/2", out)
        self.assertNotIn(SECRET_ERR, out)
        c = self.caches()[0]
        self.assertTrue(c["complete"])
        eids = [r["e"] for r in c["records"]]
        self.assertEqual(sorted(eids), [f"R{i}" for i in range(7)])
        self.assertEqual(len(eids), len(set(eids)))

    def test_persistent_error_fails_folder_keeps_checkpoint_and_rerun_resumes(self):
        old = T.TABLE_CHUNK
        T.TABLE_CHUNK = 2
        try:
            bad = FakeFolder("F", table_rows=self.mails_for(7), table_opts={"array_fail_calls": set(range(2, 100))})
            code, out = run_main(self.argv_t(), fake_com(self.one_store(bad))[0])
            self.assertIn("リトライ 2/2", out)
            self.assertIn("失敗フォルダ 1個", out)
            self.assertIn("FakeComError=1", out)
            self.assertNotIn(SECRET_ERR, out)
            c = self.caches()[0]
            self.assertFalse(c["complete"])                           # 完了にしない
            self.assertEqual(len(c["records"]), 2)                    # 取得済みの行はチェックポイント保存
            good = FakeFolder("F", table_rows=self.mails_for(7))
            code2, out2 = run_main(self.argv_t(), fake_com(self.one_store(good))[0])
        finally:
            T.TABLE_CHUNK = old
        self.assertIn("再開", out2)
        c2 = self.caches()[0]
        self.assertTrue(c2["complete"])
        eids = [r["e"] for r in c2["records"]]
        self.assertEqual(sorted(eids), [f"R{i}" for i in range(7)])
        self.assertEqual(len(eids), len(set(eids)))

    def test_retry_waits_are_used(self):
        import unittest.mock as mock
        sleeps = []
        f = FakeFolder("F", table_rows=self.mails_for(3), table_opts={"array_fail_calls": set(range(1, 100))})
        with mock.patch.object(T, "RETRY_WAITS", (5, 15)), mock.patch.object(T.time, "sleep", lambda s_: sleeps.append(s_)):
            run_main(self.argv_t(), fake_com(self.one_store(f))[0])
        self.assertEqual([x for x in sleeps if x in (5, 15)], [5, 15])

    def test_reconnect_after_three_consecutive_folder_failures(self):
        bad = [FakeFolder(f"Bad{i}", [mk_item(f"B{i}", "s", dt(4, 1))], items_raise=FakeComError(SECRET_ERR)) for i in range(3)]
        good = FakeFolder("Good", [mk_item("G1", "ok", dt(4, 2), "Zed", "Owner Taro")])
        ns = FakeNamespace([FakeStore("MB", FakeFolder("root", [], subfolders=bad + [good]))])
        com, pc = fake_com(ns)
        client = com[0]
        dispatches = []
        orig = client.Dispatch
        client.Dispatch = lambda name: (dispatches.append(name), orig(name))[1]
        code, out = run_main(self.argv("--no-stage2", "--scan-mode", "items"), com)
        self.assertEqual(code, 0)
        self.assertEqual(len(dispatches), 5)                          # 最初の接続 + 2回目のリトライ前の再接続×3 + 3フォルダ連続失敗の再接続×1
        self.assertIn("接続を作り直します", out)
        self.assertIn("失敗フォルダ 3個", out)
        self.assertIn("再接続 4回", out)
        self.assertEqual([c["folder_path"] for c in self.caches()], ["\\Good"])
        self.assertNotIn(SECRET_ERR, out)

    def test_two_failures_then_success_do_not_reconnect(self):
        bad = [FakeFolder(f"Bad{i}", [], items_raise=FakeComError(SECRET_ERR)) for i in range(2)]
        bad[0]._items.items = bad[1]._items.items = [mk_item("B", "s", dt(4, 1))]
        good = FakeFolder("Good", [mk_item("G1", "ok", dt(4, 2), "Zed", "Owner Taro")])
        com, _ = fake_com(FakeNamespace([FakeStore("MB", FakeFolder("root", [], subfolders=bad + [good]))]))
        client = com[0]
        dispatches = []
        orig = client.Dispatch
        client.Dispatch = lambda name: (dispatches.append(name), orig(name))[1]
        run_main(self.argv("--no-stage2", "--scan-mode", "items"), com)
        self.assertEqual(len(dispatches), 3)          # 最初の接続 + 2回目のリトライ前の再接続×2（連続3失敗の再接続は無し）

    def test_connection_check_failure_triggers_reconnect(self):
        ns = FakeNamespace([FakeStore("MB", FakeFolder("root", [], subfolders=[
            FakeFolder("A", [mk_item("A1", "s", dt(4, 1), "Zed", "Owner Taro")])]))])
        com, _ = fake_com(ns)
        client = com[0]
        dispatches = []
        orig = client.Dispatch
        client.Dispatch = lambda name: (dispatches.append(name), orig(name))[1]
        calls = {"n": 0}
        real_check = T.OutlookConn.check

        def flaky(self_):
            calls["n"] += 1
            return False if calls["n"] == 1 else real_check(self_)
        T.OutlookConn.check = flaky
        try:
            code, out = run_main(self.argv("--no-stage2", "--scan-mode", "items"), com)
        finally:
            T.OutlookConn.check = real_check
        self.assertEqual(code, 0)
        self.assertEqual(len(dispatches), 2)
        self.assertIn("再接続します", out)
        self.assertEqual(len(self.caches()), 1)

    def test_items_mode_unreadable_folder_does_not_give_up_the_store(self):
        old = T.MAX_ITEM_SKIP
        T.MAX_ITEM_SKIP = 1
        try:
            bad = FakeFolder("Bad", [mk_item(f"B{i}", "s", dt(4, 1)) for i in range(3)], next_raises=(1, FakeComError(SECRET_ERR)),
                             item_raises=FakeComError(SECRET_ERR))
            good = FakeFolder("Good", [mk_item("G1", "ok", dt(4, 2), "Zed", "Owner Taro")])
            ns = FakeNamespace([FakeStore("MB", FakeFolder("root", [], subfolders=[bad, good]))])
            code, out = run_main(self.argv("--no-stage2", "--scan-mode", "items"), fake_com(ns)[0])
        finally:
            T.MAX_ITEM_SKIP = old
        self.assertEqual(code, 0)
        paths = {c["folder_path"]: c["complete"] for c in self.caches()}
        self.assertEqual(paths, {"\\Bad": False, "\\Good": True})


class TestSchemaMigration(TableBase):
    def test_old_schema_cache_is_rescanned_and_ignored_in_evaluation(self):
        f = FakeFolder("F", table_rows=self.mails_for(3))
        run_main(self.argv_t(), fake_com(self.one_store(f))[0])
        path = [os.path.join(self.paths()["scan_cache"], x) for x in os.listdir(self.paths()["scan_cache"]) if x.endswith(".json")][0]
        c = _rj(path)
        c["schema"] = 1
        _wj(path, c)
        self.assertEqual(T.load_all_caches(self.paths()["scan_cache"]), [])
        code, out = run_main(["--theme", self.theme_path, "--data-dir", self.data, "--output-dir", self.out, "--evaluate-only"])
        self.assertEqual(code, 2)
        self.assertIn("旧スキーマのキャッシュが1個", out)
        f2 = FakeFolder("F", table_rows=self.mails_for(3))
        code, out = run_main(self.argv_t(), fake_com(self.one_store(f2))[0])
        self.assertIn("新規 1", out)                                 # 旧版は完了済みでもスキップしない
        self.assertEqual(_rj(path)["schema"], 2)


class TestProbe(TableBase):
    SID = "SID-MB"

    def probe_ns(self, n=4, item_fn=None, conv_table=CONV_A, table_time_fn=None):
        """Table の n 行と、同じ EntryID の Item を GetItemFromID で引ける名前空間。"""
        items, rows, by_id = [], [], {}
        for i in range(n):
            t = dt(4, 1 + i)
            it = (item_fn(i, t) if item_fn else
                  mk_item(f"E{i}", SECRET + str(i), t, "Zed", "Owner Taro", cid=f"OLD{i}",
                          PropertyAccessor=FakePA({T.PR_CONVERSATION_ID_URL: CONV_A})))
            if it is not None:
                by_id[(f"E{i}", self.SID)] = it
            items.append(mk_item(f"E{i}", "x", t))
            rows.append(trow(f"E{i}", SECRET + str(i), table_time_fn(t) if table_time_fn else t, conv=conv_table))
        f = FakeFolder("Big", items, table_rows=rows, table_opts={"fail_any_filter": True} if table_time_fn else None)
        small = FakeFolder("Small", [mk_item("S1", "s", dt(4, 1))], table_rows=[trow("S1", "s", dt(4, 1))])
        ns = FakeNamespace([FakeStore("MB", FakeFolder("root", [], subfolders=[small, f]), sid=self.SID)], by_id)
        return ns, f

    def run_probe(self, *extra, ns=None):
        ns = ns or self.probe_ns()[0]
        return run_main(["--data-dir", self.data, "--output-dir", self.out, "--from", "2024-04", "--to", "2024-04", *extra],
                        fake_com(ns)[0])

    def test_probe_table_compares_same_entry_ids_without_subjects(self):
        code, out = self.run_probe("--probe-table")
        self.assertEqual(code, 0)
        self.assertNotIn(SECRET, out)
        self.assertIn("table:", out)
        self.assertIn("items(GetItemFromID): 取得 4件 / 取れず 0件", out)
        self.assertIn("同一メールとして比較できた 4件", out)
        self.assertIn("一致 4 / 不一致 0", out)
        self.assertIn("最頻 +0時間", out)
        self.assertIn("使用行基準 1行あたり", out)
        self.assertIn("形状(行×列)=4×11", out)
        self.assertIn("Big", out)                       # 件数最大のフォルダを選ぶ
        self.assertIn("結果の見方", out)
        self.assertEqual(self.caches(), [])             # 診断はキャッシュを作らない

    def test_item_fetch_failure_is_counted(self):
        ns, _f = self.probe_ns(item_fn=lambda i, t: None if i in (1, 3) else mk_item(
            f"E{i}", "s", t, cid="o", PropertyAccessor=FakePA({T.PR_CONVERSATION_ID_URL: CONV_A})))
        code, out = self.run_probe("--probe-table", ns=ns)
        self.assertIn("取得 2件 / 取れず 2件", out)
        self.assertIn("Itemが取れなかった 2件", out)
        self.assertNotIn(SECRET_ERR, out)

    def test_conversation_id_match_mismatch_and_empty(self):
        other = bytes(range(101, 117))

        def item_fn(i, t):
            if i == 0:
                return mk_item("E0", "s", t, cid="x", PropertyAccessor=FakePA({T.PR_CONVERSATION_ID_URL: CONV_A}))      # 一致
            if i == 1:
                return mk_item("E1", "s", t, cid="x", PropertyAccessor=FakePA({T.PR_CONVERSATION_ID_URL: other}))       # 不一致
            return mk_item(f"E{i}", "s", t, cid="x")                                                                       # Item側が空
        ns, _f = self.probe_ns(n=3, item_fn=item_fn)
        code, out = self.run_probe("--probe-conversation-id", ns=ns)
        self.assertIn("一致 1 / 不一致 1 / 空(Item) 1 / 空(Table) 0", out)
        self.assertIn("#1:", out)
        self.assertIn("PRどうし=一致", out)
        self.assertIn("PRどうし=不一致/空", out)
        self.assertNotIn(SECRET, out)
        ns2, _f = self.probe_ns(n=2, conv_table=None)
        code, out = self.run_probe("--probe-table", ns=ns2)
        self.assertIn("空(Table) 2", out)

    def test_time_type_diagnosis_tz_aware_and_naive(self):
        from datetime import timedelta
        with tokyo_tz():
            # tz付き(UTC)。Itemの受信日時は +9h（ローカル）。auto ならローカルに直して差0
            ns, _f = self.probe_ns(table_time_fn=lambda t: (t - timedelta(hours=9)).replace(tzinfo=UTC))
            code, out = self.run_probe("--probe-table", ns=ns)
            self.assertIn("tz付き 4 / tzなし 0 / utcoffset: +00:00=4", out)
            self.assertIn("例: 2024-04-01T00:", out)
            self.assertIn("最頻 +0時間", out)
            self.assertNotIn("UTC疑い", out)
            # tzなし。UTC値がそのまま入っている環境: 差 -9 -> UTC疑い。--table-time utc で解消
            ns2, _f = self.probe_ns(table_time_fn=lambda t: t - timedelta(hours=9))
            code, out = self.run_probe("--probe-table", ns=ns2)
            self.assertIn("tz付き 0 / tzなし 4", out)
            self.assertIn("-9時間", out)
            self.assertIn("UTC疑い", out)
            self.assertIn("--table-time utc", out)
            code, out = self.run_probe("--probe-table", "--table-time", "utc", ns=ns2)
            self.assertIn("最頻 +0時間", out)
            self.assertNotIn("UTC疑い", out)

    def test_probe_limit_folder_filter_and_errors(self):
        code, out = self.run_probe("--probe-table", "--probe-limit", "2")
        self.assertIn("取得 2件 / 取れず 0件", out)
        self.assertIn("Table 2行", out)
        code, out = self.run_probe("--probe-table", "--probe-folder", "small")
        self.assertIn("Small", out)
        code, out = self.run_probe("--probe-table", "--probe-folder", "nothing-matches")
        self.assertEqual(code, 2)
        code, out = run_main(["--data-dir", self.data, "--output-dir", self.out, "--probe-table", "--stores", "zzz"],
                             fake_com(self.probe_ns()[0])[0])
        self.assertEqual(code, 2)

    def test_probe_reports_table_unavailable(self):
        items = [mk_item("E1", "s", dt(4, 1), "Zed", "Owner Taro")]
        f = FakeFolder("F", items, table_rows=None)
        code, out = run_main(["--data-dir", self.data, "--output-dir", self.out, "--from", "2024-04", "--to", "2024-04",
                              "--probe-table"], fake_com(self.one_store(f))[0])
        self.assertEqual(code, 0)
        self.assertIn("table: 利用できません", out)
        self.assertNotIn(SECRET_ERR, out)

    def test_compare_probe_unit(self):
        table = [{"e": "1", "c": "AA", "t": "2024-04-01T01:00:00"}, {"e": "2", "c": "", "t": "2024-04-02T01:00:00"},
                 {"e": "3", "c": "BB", "t": "2024-04-03T01:00:00"}, {"e": "4", "c": "EE", "t": "2024-04-04T01:00:00"}]
        items = {"1": {"c": "AA", "c2": "aa", "t": "2024-04-01T10:00:00"}, "2": {"c": "CC", "c2": "zz", "t": "2024-04-02T10:00:00"},
                 "3": {"c": "", "c2": "", "t": "2024-04-03T10:00:00"}}
        r = T.compare_probe(table, items)
        self.assertEqual((r["n_table"], r["n_items_ok"], r["n_items_fail"]), (4, 3, 1))
        self.assertEqual((r["conv_match"], r["conv_mismatch"], r["conv_empty_item"], r["conv_empty_table"]), (1, 0, 1, 1))
        self.assertEqual((r["c2_same"], r["c2_diff"], r["c2_na"]), (1, 1, 1))
        self.assertEqual((r["time_diff_mode"], r["time_diff_n"], r["utc_suspect"]), (-9, 3, True))
        self.assertEqual(r["time_diff_counts"], {-9: 3})
        r0 = T.compare_probe(table[:1], {"1": {"c": "AA", "c2": "", "t": "2024-04-01T01:00:00"}})
        self.assertEqual((r0["time_diff_mode"], r0["utc_suspect"]), (0, False))
        far = T.compare_probe(table[:1], {"1": {"c": "", "c2": "", "t": "2024-05-01T01:00:00"}})
        self.assertFalse(far["utc_suspect"])                                  # 差が大きすぎる（数時間でない）
        self.assertIsNone(T.compare_probe([], {})["time_diff_mode"])
        self.assertEqual(T.conv_relation("abcd", "ABCD"), "同一")
        self.assertEqual(T.conv_relation("abcdef", "ABCD"), "先頭一致")
        self.assertEqual(T.conv_relation("zz", "ABCD"), "不一致")
        self.assertEqual(T.conv_relation("", "ABCD"), "比較不可")

    def test_describe_raw_times(self):
        from datetime import timedelta, timezone
        jst = timezone(timedelta(hours=9))
        d = T.describe_raw_times([("a", datetime(2024, 4, 1, tzinfo=UTC), None), ("b", None, datetime(2024, 4, 2, tzinfo=jst)),
                                  ("c", datetime(2024, 4, 3), None), ("d", None, None), ("e", "x", 5)])
        self.assertEqual((d["aware"], d["naive"]), (2, 1))
        self.assertEqual(d["offsets"], {"+00:00": 1, "+09:00": 1})
        self.assertEqual(len(d["samples"]), 3)

    def test_read_item_probe(self):
        ns = FakeNamespace([], {("E", "S" * 20): mk_item("E", "s", dt(4, 1), cid="OLD", PropertyAccessor=FakePA({T.PR_CONVERSATION_ID_URL: CONV_A}))})
        got = T.read_item_probe(ns, "E", "S" * 20)
        self.assertEqual(got, {"t": "2024-04-01T09:00:00", "c": CONV_A_HEX, "c2": "OLD"})
        self.assertIsNone(T.read_item_probe(ns, "missing", "S" * 20))

    def test_probe_does_not_need_theme(self):
        code, out = self.run_probe("--probe-table")
        self.assertEqual(code, 0)


class TestProbeSpeed(TableBase):
    def ns(self, n=8):
        rows = [trow(f"E{i}", SECRET + str(i), dt(4, 1 + i), conv=CONV_A, cc="c", to="t") for i in range(n)]
        f = FakeFolder("Big", [mk_item("x", "s", dt(4, 1))], table_rows=rows)
        return FakeNamespace([FakeStore("MB", FakeFolder("root", [], subfolders=[f]), sid="SID-MB")]), f

    def probe(self, ns, *extra):
        return run_main(["--data-dir", self.data, "--output-dir", self.out, "--from", "2024-04", "--probe-speed", *extra],
                        fake_com(ns)[0])

    def test_runs_all_experiments_without_subjects_or_cache(self):
        ns, f = self.ns()
        code, out = self.probe(ns, "--probe-limit", "5")
        self.assertEqual(code, 0)
        self.assertNotIn(SECRET, out)
        for chunk in (50, 100, 500, 1000, 2000):
            self.assertIn(f"最小列(3)", out)
            self.assertIn(f"chunk={chunk}", out)
        self.assertIn("全列(11)", out)
        for key in ("sent", "sender", "smtp", "semail", "to", "cc", "conv", "cls"):
            self.assertIn(f"+{key}", out)
        self.assertIn("結果の見方", out)
        self.assertIn("--table-chunk", out)
        self.assertEqual(self.caches(), [])
        self.assertEqual(f._items.filters, [])                               # アイテムは開かない

    def test_time_table_fetch_counts_rows_by_chunk(self):
        ns, f = self.ns(8)
        r = T.time_table_fetch(f, datetime(2024, 4, 1), datetime(2024, 5, 1), T.SPEED_BASE, 3, 5)
        self.assertEqual((r["fetched"], r["used"], r["filter"]), (6, 5, "or"))      # チャンク3で2回 = 6行取得、5行使用
        r2 = T.time_table_fetch(f, datetime(2024, 4, 1), datetime(2024, 5, 1), [k for k, _n in T.TABLE_COLUMNS], 500, 5)
        self.assertEqual((r2["fetched"], r2["used"]), (8, 5))                        # 500行チャンクでも実際にある8行まで
        self.assertGreaterEqual(r["setup"], 0)

    def test_slow_column_is_identified(self):
        import unittest.mock as mock
        ns, _f = self.ns()

        def fake_fetch(folder, s, e, keys, chunk, k):
            per_row = 0.001 + (0.060 if "to" in keys else 0) + (0.004 if "cc" in keys else 0)
            return {"setup": 0.01, "fetch": per_row * k, "fetched": k, "used": k, "filter": "or"}
        with mock.patch.object(T, "time_table_fetch", fake_fetch):
            code, out = self.probe(ns)
        self.assertIn("遅い列の候補: to(+60ms/行)", out)
        self.assertNotIn("cc(+", out)                                               # +4ms は閾値未満
        self.assertEqual(T.summarize_speed(10.0, 12.0), (2.0, False))
        self.assertEqual(T.summarize_speed(10.0, 40.0)[1], True)
        self.assertEqual(T.summarize_speed(100.0, 140.0)[1], False)                 # 基準が大きいときは比率も必要

    def test_errors_in_an_experiment_do_not_stop_the_probe(self):
        import unittest.mock as mock
        ns, _f = self.ns()
        calls = {"n": 0}
        real = T.time_table_fetch

        def flaky(folder, s, e, keys, chunk, k):
            calls["n"] += 1
            if calls["n"] == 3:
                raise FakeComError(SECRET_ERR)
            return real(folder, s, e, keys, chunk, k)
        with mock.patch.object(T, "time_table_fetch", flaky):
            code, out = self.probe(ns)
        self.assertEqual(code, 0)
        self.assertIn("エラー FakeComError", out)
        self.assertNotIn(SECRET_ERR, out)
        self.assertIn("結果の見方", out)

    def test_zero_rows_and_theme_not_required(self):
        ns, _f = self.ns()
        code, out = run_main(["--data-dir", self.data, "--output-dir", self.out, "--probe-speed"], fake_com(ns)[0])
        self.assertEqual(code, 2)
        self.assertIn("0件でした。--from YYYY-MM を指定してください", out)


class TestTableChunkOption(TableBase):
    def test_table_chunk_changes_getarray_size(self):
        f = FakeFolder("F", table_rows=[trow(f"R{i}", "s", dt(4, 1 + i % 25)) for i in range(25)])
        code, out = run_main(self.argv_t("--table-chunk", "10"), fake_com(self.one_store(f))[0])
        self.assertEqual(code, 0)
        self.assertEqual(f.array_calls, 3)                 # 10+10+5 行
        self.assertEqual(len(self.caches()[0]["records"]), 25)
        g = FakeFolder("F", table_rows=self.mails_for(7))
        run_main(self.argv_t("--table-chunk", "500", "--rescan"), fake_com(self.one_store(g))[0])
        self.assertEqual(g.array_calls, 1)

    def test_default_and_validation(self):
        self.assertEqual(T.parse_args([]).table_chunk, 0)
        self.assertEqual(T.parse_args(["--table-chunk", "100"]).table_chunk, 100)
        for bad in ("5", "6000", "abc"):
            with contextlib.redirect_stderr(io.StringIO()):
                with self.assertRaises(SystemExit):
                    T.parse_args(["--table-chunk", bad])

    def test_probe_table_uses_table_chunk(self):
        ns, f = TestProbe.probe_ns(TestProbe("test_probe_does_not_need_theme"), n=25)
        code, out = run_main(["--data-dir", self.data, "--output-dir", self.out, "--from", "2024-04", "--probe-table",
                              "--table-chunk", "10", "--probe-limit", "15"], fake_com(ns)[0])
        self.assertIn("チャンク 10", out)
        self.assertIn("実取得 20行基準", out)



class TestRegressionKeepsReadOnly(TableBase):
    def test_table_scan_calls_no_forbidden_methods_and_hides_subjects(self):
        rows = [trow("R1", SECRET, dt(4, 1)), trow("R2", SECRET + "2", dt(4, 2))]
        f = FakeFolder("F", table_rows=rows)
        code, out = run_main(self.argv_t(), fake_com(self.one_store(f))[0])
        self.assertEqual(code, 0)
        self.assertNotIn(SECRET, out)
        self.assertEqual(ACCESSED, [])
        log = sorted(x for x in os.listdir(self.out) if x.endswith(".log"))[-1]
        self.assertNotIn(SECRET, _rd(os.path.join(self.out, log), "utf-8-sig"))

    def test_end_to_end_table_candidates_flow_to_ledger(self):
        rows = [trow("R1", "通関 shipment topic", dt(4, 1), to="Owner Taro", sender="Sato Pm", conv=CONV_A),
                trow("R2", "RE: 通関 shipment topic", dt(4, 2), to="Sato Pm", sender="Owner Taro", conv=CONV_A)]
        f = FakeFolder("F", table_rows=rows)
        code, out = run_main(self.argv_t(), fake_com(self.one_store(f))[0])
        self.assertEqual(code, 0)
        led = sorted(x for x in os.listdir(self.out) if x.startswith("thread_ledger_theme_") and x.endswith(".csv"))[-1]
        got = _rdict(os.path.join(self.out, led))
        self.assertEqual(len(got), 1)
        self.assertEqual(got[0]["メール数"], "2")



# ============================================================
# S1.6 レビュー対応: タイムゾーン・接続系エラー・形状・列値・会話IDカウンタ・診断
# ============================================================
@contextlib.contextmanager
def tokyo_tz():
    import time as _time
    old = os.environ.get("TZ")
    os.environ["TZ"] = "Asia/Tokyo"
    _time.tzset()
    try:
        yield
    finally:
        if old is None:
            os.environ.pop("TZ", None)
        else:
            os.environ["TZ"] = old
        _time.tzset()


UTC = __import__("datetime").timezone.utc


class TestTableTimezone(TableBase):
    def test_to_naive_dt_modes(self):
        aware = datetime(2024, 3, 31, 16, 0, 0, tzinfo=UTC)          # JST では 2024-04-01 01:00
        naive = datetime(2024, 3, 31, 16, 0, 0)
        with tokyo_tz():
            self.assertEqual(T.to_naive_dt(aware), datetime(2024, 3, 31, 16, 0))              # 従来（Items方式）は tzinfo を見ない
            self.assertEqual(T.to_naive_dt(aware, "auto"), datetime(2024, 4, 1, 1, 0))
            self.assertEqual(T.to_naive_dt(naive, "auto"), datetime(2024, 3, 31, 16, 0))       # tzinfo 無しはローカルとみなす
            self.assertEqual(T.to_naive_dt(naive, "utc"), datetime(2024, 4, 1, 1, 0))          # UTCとみなしてローカルに直す
            self.assertEqual(T.to_naive_dt(aware, "utc"), datetime(2024, 4, 1, 1, 0))
            self.assertEqual(T.to_naive_dt(aware, "local"), datetime(2024, 3, 31, 16, 0))
            self.assertEqual(T.to_naive_dt(naive, "local"), datetime(2024, 3, 31, 16, 0))
            for mode in (None, "auto", "utc", "local"):
                self.assertIsNone(T.to_naive_dt(datetime(4501, 1, 1), mode), mode)
                self.assertIsNone(T.to_naive_dt(datetime(4501, 1, 1, tzinfo=UTC), mode), mode)

    def scan_rows(self, rows, mode, start=datetime(2024, 4, 1), end=datetime(2024, 5, 1)):
        f = FakeFolder("F", table_rows=rows, table_opts={"fail_any_filter": True})      # フィルタ無し→手元で期間に絞る
        st = T.new_scan_state()
        st["tz_mode"] = mode
        T.scan_folder_table(f, start, end, st)
        return st

    def test_boundary_is_not_9_hours_off(self):
        rows = [trow("IN", "s", datetime(2024, 3, 31, 16, 0, tzinfo=UTC)),         # JST 04-01 01:00 -> 期間内
                trow("OUT", "s", datetime(2024, 4, 30, 15, 30, tzinfo=UTC))]       # JST 05-01 00:30 -> 期間外
        with tokyo_tz():
            st = self.scan_rows(rows, "auto")
        self.assertEqual([r["e"] for r in st["records"]], ["IN"])
        self.assertEqual(st["records"][0]["t"], "2024-04-01T01:00:00")
        # tzinfo 無し + utc 指定
        rows2 = [trow("IN", "s", datetime(2024, 3, 31, 16, 0)), trow("OUT", "s", datetime(2024, 4, 30, 15, 30))]
        with tokyo_tz():
            self.assertEqual([r["e"] for r in self.scan_rows(rows2, "utc")["records"]], ["IN"])
            self.assertEqual([r["e"] for r in self.scan_rows(rows2, "auto")["records"]], ["OUT"])   # autoは naive をローカル扱い
            self.assertEqual([r["e"] for r in self.scan_rows(rows2, "local")["records"]], ["OUT"])

    def test_cli_default_and_passed_through_state(self):
        self.assertEqual(T.parse_args([]).table_time, "auto")
        self.assertEqual(T.parse_args(["--table-time", "utc"]).table_time, "utc")
        with tokyo_tz():
            rows = [trow("R1", "s", datetime(2024, 3, 31, 16, 0))]
            f = FakeFolder("F", table_rows=rows, table_opts={"fail_any_filter": True})
            code, out = run_main(self.argv_t("--table-time", "utc", "--from", "2024-04", "--to", "2024-04"),
                                 fake_com(self.one_store(f))[0])
        self.assertEqual(self.caches()[0]["records"][0]["t"], "2024-04-01T01:00:00")

    def test_compare_probe_time_difference_and_warning(self):
        table = [{"e": str(i), "c": "X", "t": "2024-04-01T01:00:00"} for i in range(4)] + [{"e": "4", "c": "X", "t": "2024-04-01T10:00:00"}]
        items = {str(i): {"c": "X", "c2": "", "t": "2024-04-01T10:00:00"} for i in range(5)}
        r = T.compare_probe(table, items)
        self.assertEqual((r["time_diff_mode"], r["time_diff_n"], r["utc_suspect"]), (-9, 5, True))
        self.assertEqual(r["time_diff_counts"], {-9: 4, 0: 1})
        r0 = T.compare_probe(table[4:], {"4": items["4"]})
        self.assertEqual((r0["time_diff_mode"], r0["utc_suspect"]), (0, False))
        far = T.compare_probe([{"e": "0", "c": "", "t": "2024-05-01T10:00:00"}], {"0": items["0"]})
        self.assertFalse(far["utc_suspect"])            # 差が大きすぎる（数時間でない）ときは疑わない
        self.assertIsNone(T.compare_probe([], {})["time_diff_mode"])


class TestConnectionErrors(TableBase):
    CONN = -2147221227          # 0x80040115 MAPI_E_NETWORK_ERROR

    def test_is_connection_error(self):
        E = FakeComError
        for code in (0x80040115, 0x80020009, 0x800706BA, 0x800706BE, 0x80070005, 0x80010108, 0x8001010E):
            self.assertTrue(T.is_connection_error(E(code - (1 << 32), "x")), hex(code))
            self.assertTrue(T.is_connection_error(E(code, "x")), hex(code))
        self.assertTrue(T.is_connection_error(E(-2147352567, "x", (0, "src", "d", None, 0, self.CONN))))     # 内側のscodeが接続系
        self.assertFalse(T.is_connection_error(E(-2147352567, "x", (0, "src", "d", None, 0, -2147219196))))   # 内側が構文系ならfalse
        self.assertFalse(T.is_connection_error(E(SECRET_ERR)))
        self.assertFalse(T.is_connection_error(ValueError("x")))
        self.assertFalse(T.is_connection_error(E(0x80040109, "x")))
        self.assertTrue(T.is_disp_error(E(-2147352567, "x")))
        self.assertFalse(T.is_disp_error(E(self.CONN, "x")))

    def test_open_table_reraises_connection_error_but_falls_back_on_syntax_error(self):
        f = FakeFolder("F", table_rows=self.mails_for(2), table_opts={"get_table_exc": FakeComError(self.CONN, "x")})
        with self.assertRaises(FakeComError):
            T.open_table(f, datetime(2024, 1, 1), datetime(2025, 1, 1), {})
        self.assertEqual(len(f.table_filters), 1)                      # 次のフィルタへ進まない
        g = FakeFolder("F", table_rows=self.mails_for(2), table_opts={"fail_or": True})
        tbl, fmode, cols = T.open_table(g, datetime(2024, 1, 1), datetime(2025, 1, 1), {})
        self.assertEqual(fmode, "ReceivedTime")

    def test_column_connection_error_is_not_swallowed(self):
        class Cols:
            def RemoveAll(self):
                pass

            def Add(self, name):
                raise FakeComError(TestConnectionErrors.CONN, "x")

        class Tbl:
            Columns = Cols()
        with self.assertRaises(FakeComError):
            T._setup_table_columns(Tbl(), {})

    def test_transient_connection_error_retries_in_table_mode_without_items_fallback(self):
        f = FakeFolder("F", table_rows=self.mails_for(3),
                       table_opts={"get_table_exc": FakeComError(self.CONN, "x"), "get_table_exc_calls": {1}})
        code, out = run_main(self.argv_t(), fake_com(self.one_store(f))[0])
        self.assertEqual(code, 0)
        self.assertIn("リトライ 1/2", out)
        self.assertNotIn("Items方式に切り替えます", out)
        c = self.caches()[0]
        self.assertEqual((c["scan_mode"], c["complete"], len(c["records"])), ("table", True, 3))
        self.assertEqual(f._items.filters, [])

    def test_persistent_connection_error_fails_folder_and_keeps_table_mode(self):
        f = FakeFolder("F", [mk_item("I1", "s", dt(4, 1))], table_rows=self.mails_for(3),
                       table_opts={"get_table_exc": FakeComError(self.CONN, "x")})
        code, out = run_main(self.argv_t(), fake_com(self.one_store(f))[0])
        self.assertIn("失敗フォルダ 1個", out)
        self.assertNotIn("Items方式に切り替えます", out)
        self.assertEqual(f._items.filters, [])                         # Items（1通1秒の方式）へ落ちない
        self.assertEqual(self.caches(), [])

    def test_disp_error_reconnects_before_first_retry_and_check_probes_a_folder(self):
        bad = FakeFolder("Bad", [mk_item("B", "s", dt(4, 1))], items_raise=FakeComError(-2147352567, "x"))
        com, _ = fake_com(FakeNamespace([FakeStore("MB", FakeFolder("root", [], subfolders=[bad]))]))
        dispatches = []
        orig = com[0].Dispatch
        com[0].Dispatch = lambda name: (dispatches.append(name), orig(name))[1]
        run_main(self.argv("--no-stage2", "--scan-mode", "items"), com)
        self.assertEqual(len(dispatches), 3)             # 最初 + DISP_E_*の1回目リトライ前 + 2回目リトライ前
        conn = T.OutlookConn(FakeClient(FakeNamespace([], {})))
        conn.connect()
        self.assertTrue(conn.check())
        conn.probe_ids = ("ab" * 10, "cd" * 10)
        self.assertFalse(conn.check())                   # GetFolderFromID が失敗 -> 壊れた扱い


class TestShapeAndValues(TableBase):
    def test_transposed_array_falls_back_to_items(self):
        f = FakeFolder("F", [mk_item("I1", "s1", dt(4, 1), "Zed", "Owner Taro")], table_rows=self.mails_for(4),
                       table_opts={"transpose": True})
        with self.assertRaises(T.TableUnavailable) as cm:
            T.scan_folder_table(f, datetime(2024, 1, 1), datetime(2025, 1, 1), T.new_scan_state())
        self.assertIn("形状不正", str(cm.exception))
        code, out = run_main(self.argv_t(), fake_com(self.one_store(f))[0])
        self.assertIn("形状不正", out)
        self.assertIn("Items方式に切り替えます", out)
        self.assertEqual([r["e"] for r in self.caches()[0]["records"]], ["I1"])

    def test_shape_recorded_for_probe(self):
        f = FakeFolder("F", table_rows=self.mails_for(3))
        st = T.new_scan_state()
        T.scan_folder_table(f, datetime(2024, 1, 1), datetime(2025, 1, 1), st)
        self.assertEqual(st["table_shape"], (3, len(T.TABLE_COLUMNS)))

    def test_non_string_values_in_string_columns_become_empty_with_error_count(self):
        rows = [trow("R1", 12345, dt(4, 1), sender=-2147221233, to=FakeComError("x"), cc=3.5, smtp=0x8004010F, conv=7)]
        f = FakeFolder("F", table_rows=rows)
        st = T.new_scan_state()
        T.scan_folder_table(f, datetime(2024, 1, 1), datetime(2025, 1, 1), st)
        r = st["records"][0]
        self.assertEqual((r["s"], r["n"], r["to"], r["cc"], r["a"], r["c"]), ("", "", "", "", "", ""))
        for k in ("subj", "sender", "to", "cc", "smtp", "conv"):
            self.assertIn(f"列値:{k}", st["errors"], k)
        self.assertEqual(T._tstr(None), "")
        self.assertEqual(T._tstr(b"abc"), "abc")
        e = {}
        self.assertEqual(T._tstr(5, e, "x"), "")
        self.assertEqual(e, {"列値:x": 1})

    def test_message_class_uses_unicode_property(self):
        self.assertTrue(dict(T.TABLE_COLUMNS)["cls"].endswith("0x001A001F"))


class TestConvIdCounters(TableBase):
    def test_per_store_counters_and_conv_failures_split_from_error_kinds(self):
        good = mk_item("A1", "s1", dt(4, 1), "Zed", "Owner Taro", cid="OLD1",
                       PropertyAccessor=FakePA({T.PR_CONVERSATION_ID_URL: CONV_A}))
        c2only = mk_item("A2", "s2", dt(4, 2), "Zed", "Owner Taro", cid="OLD2")
        failing = mk_item("A3", "s3", dt(4, 3), "Zed", "Owner Taro", cid="OLD3", PropertyAccessor=FakePA({}))
        f = FakeFolder("F", [good, c2only, failing])
        code, out = run_main(self.argv("--no-stage2"), fake_com(self.one_store(f))[0])
        self.assertIn("メール3件中 PR_CONVERSATION_ID空 2件（うち item.ConversationID のみ 2件）", out)
        self.assertIn("会話ID取得失敗 1件", out)
        self.assertNotIn("ConvID(PA)", out.split("エラー種別")[-1] if "エラー種別" in out else "")
        self.assertIn("--conv-fallback c2", out)
        self.assertEqual(T.conv_counts([{"c": "A"}, {"c": "", "c2": "x"}, {"c": ""}]), (3, 2, 1))
        other, n = T.split_conv_errors({"ConvID(PA):X": 2, "Subject:Y": 1})
        self.assertEqual((other, n), ({"Subject:Y": 1}, 2))

    def test_mode_and_retry_counts_in_summary(self):
        f = FakeFolder("F", table_rows=self.mails_for(2), table_opts={"array_fail_calls": {1}})
        code, out = run_main(self.argv_t(), fake_com(self.one_store(f))[0])
        self.assertIn("方式別フォルダ数: table 1 / items 0 / リトライ 1回", out)

    def test_conv_fallback_option_changes_threading(self):
        r1 = rec("E1", "Short a", "2024-04-10T09:00:00", "Sato Pm", "Owner Taro")
        r2 = rec("E2", "Short b", "2024-04-11T09:00:00", "Owner Taro", "Sato Pm")
        r1["c2"] = r2["c2"] = "OLD"
        th = make_theme()
        c = cache("S", 1, "\\Inbox", [r1, r2])
        res_default = T.evaluate_all([c], th, {}, (2024, 1), (2024, 12))
        res_c2 = T.evaluate_all([c], th, {}, (2024, 1), (2024, 12), conv_fallback="c2")
        self.assertEqual((len(res_default["threads"]), len(res_c2["threads"])), (2, 1))
        self.assertEqual(T.parse_args([]).conv_fallback, "subject")
        self.assertEqual(T.parse_args(["--conv-fallback", "c2"]).conv_fallback, "c2")


class TestProbeDefaults(TableBase):
    def two_stores(self, first_has_mail):
        month_day = dt(4, 5)
        empty = FakeFolder("E", [mk_item("X", "s", datetime(2023, 1, 1))], table_rows=[trow("X", "s", datetime(2023, 1, 1))])
        full_items = [mk_item(f"E{i}", SECRET + str(i), dt(4, 1 + i), "Zed", "Owner Taro", cid="o",
                              PropertyAccessor=FakePA({T.PR_CONVERSATION_ID_URL: CONV_A})) for i in range(3)]
        full = FakeFolder("Full", full_items, table_rows=[trow(f"E{i}", "s", dt(4, 1 + i), conv=CONV_A) for i in range(3)])
        a = FakeStore("StoreA", FakeFolder("root", [], subfolders=[full if first_has_mail else empty]), sid="SA")
        b = FakeStore("StoreB", FakeFolder("root", [], subfolders=[empty if first_has_mail else full]), sid="SB")
        sid = "SA" if first_has_mail else "SB"
        return FakeNamespace([a, b], {(it.EntryID, sid): it for it in full_items})

    def probe(self, ns, *extra):
        return run_main(["--data-dir", self.data, "--output-dir", self.out, *extra], fake_com(ns)[0])

    def test_month_defaults_to_to_month_and_zero_rows_message(self):
        code, out = self.probe(self.two_stores(True), "--probe-table", "--to", "2024-04")
        self.assertEqual(code, 0)
        self.assertIn("対象月: 2024-04", out)
        code, out = self.probe(self.two_stores(True), "--probe-table")          # --from/--to 無し -> 先月（テスト日付では0件）
        self.assertEqual(code, 2)
        self.assertIn("0件でした。--from YYYY-MM を指定してください", out)
        code, out = self.probe(self.two_stores(True), "--probe-table", "--from", "2024-04", "--to", "2024-06")
        self.assertIn("対象月: 2024-04", out)                                    # --from が優先

    def test_tries_stores_in_order_until_one_has_mail(self):
        code, out = self.probe(self.two_stores(False), "--probe-table", "--from", "2024-04")
        self.assertEqual(code, 0)
        self.assertIn("StoreA: 2024-04 に該当メールなし", out)
        self.assertIn("ストア: StoreB", out)
        self.assertNotIn(SECRET, out)

    def test_explanations_and_shape_and_time_warning(self):
        code, out = self.probe(self.two_stores(True), "--probe-table", "--from", "2024-04")
        self.assertIn("結果の見方", out)
        self.assertIn("形状(行×列)=3×", out)
        self.assertIn("受信日時差", out)
        self.assertNotIn("UTC疑い", out)
        # Table が -9 時間ずれて返る（UTC）環境を再現: 行の時刻を 9 時間引く
        ns = self.two_stores(True)
        full = ns.Stores.Item(1).GetRootFolder().Folders.Item(1)
        from datetime import timedelta
        full.table_rows = [dict(r, recv=r["recv"] - timedelta(hours=9)) for r in full.table_rows]
        with tokyo_tz():
            code, out = self.probe(ns, "--probe-table", "--from", "2024-04", "--table-time", "local")
            self.assertIn("UTC疑い", out)
            self.assertIn("-9時間", out)
            self.assertIn("--table-time utc", out)
            code, out = self.probe(ns, "--probe-table", "--from", "2024-04", "--table-time", "utc")
            self.assertNotIn("UTC疑い", out)
            self.assertIn("最頻 +0時間", out)



class TestStage1PropertyGuard(unittest.TestCase):
    def test_guard_detects_non_conversation_id_property_in_stage1(self):
        LIGHT_ACCESSED.clear()
        item = FakeItem({"PropertyAccessor": FakePA({T.PR_CONVERSATION_ID_URL: CONV_A, "http://x/body": "b"})})
        item.PropertyAccessor.GetProperty(T.PR_CONVERSATION_ID_URL)
        self.assertEqual(LIGHT_ACCESSED, [])                                   # 会話IDは許可
        item.PropertyAccessor.GetProperty("http://x/body")
        self.assertEqual(LIGHT_ACCESSED, ["GetProperty:http://x/body"])         # それ以外は検出される
        LIGHT_ACCESSED.clear()

    def test_real_stage1_read_only_requests_conversation_id(self):
        LIGHT_ACCESSED.clear()
        item = mk_item("A1", "s", dt(4, 1), "Zed", "Owner Taro",
                       PropertyAccessor=FakePA({T.PR_CONVERSATION_ID_URL: CONV_A}))
        T.read_light_item(item, {})
        self.assertEqual(LIGHT_ACCESSED, [])

    def test_resume_does_not_double_count_processed_or_non_mail(self):
        rows = [trow(f"R{i}", "s", dt(4, 1 + i)) for i in range(4)] + [trow("N1", "s", dt(4, 9), cls="IPM.Appointment")]
        f = FakeFolder("F", table_rows=rows)
        st = T.new_scan_state()
        T.scan_folder_table(f, datetime(2024, 1, 1), datetime(2025, 1, 1), st)
        self.assertEqual((st["processed"], st["non_mail"], len(st["records"])), (5, 1, 4))
        # 同じ state で再開（リトライ・再実行相当）: 取得済みは飛ばし、メール以外は数え直す。processed は今回の見た行数だけ
        T.scan_folder_table(FakeFolder("F", table_rows=rows), datetime(2024, 1, 1), datetime(2025, 1, 1), st, resume=True)
        self.assertEqual((st["processed"], st["non_mail"], len(st["records"])), (1, 1, 4))


class TestFilterFallbackLogs(TableBase):
    def test_all_filter_and_received_only_are_announced(self):
        f = FakeFolder("F", table_rows=self.mails_for(2), table_opts={"fail_any_filter": True})
        code, out = run_main(self.argv_t(), fake_com(self.one_store(f))[0])
        self.assertIn("フィルタ無し（全件取得して手元で絞る）", out)
        g = FakeFolder("F", table_rows=self.mails_for(2), table_opts={"fail_or": True})
        code, out = run_main(self.argv_t("--rescan"), fake_com(self.one_store(g))[0])
        self.assertIn("ReceivedTime のみで絞りました", out)
        h = FakeFolder("F", table_rows=self.mails_for(2))
        code, out = run_main(self.argv_t("--rescan"), fake_com(self.one_store(h))[0])
        self.assertNotIn("フィルタ無し", out)
        self.assertNotIn("ReceivedTime のみ", out)



class TestAmbiguousDispatchError(TableBase):
    AMB = -2147352567            # 0x80020009 DISP_E_EXCEPTION（内側のscodeなし）

    def amb(self):
        return FakeComError(self.AMB, "x")

    def test_classification(self):
        self.assertTrue(T.is_ambiguous_dispatch_error(self.amb()))
        self.assertTrue(T.is_ambiguous_dispatch_error(FakeComError(self.AMB, "x", (0, "s", "d", None, 0, self.AMB))))
        self.assertFalse(T.is_ambiguous_dispatch_error(FakeComError(-2147221227, "x")))             # 0x80040115 は明確な接続系
        self.assertFalse(T.is_ambiguous_dispatch_error(FakeComError(self.AMB, "x", (0, "s", "d", None, 0, -2147221227))))
        self.assertFalse(T.is_ambiguous_dispatch_error(FakeComError(SECRET_ERR)))
        self.assertTrue(T.is_connection_error(self.amb()))                                        # リトライ側では従来どおり接続系扱い
        self.assertFalse(T.is_definite_connection_error(self.amb()))
        for code in (0x80040115, 0x800706BA, 0x800706BE, 0x80070005, 0x80010108):
            self.assertTrue(T.is_definite_connection_error(FakeComError(code - (1 << 32), "x")), hex(code))

    def test_ambiguous_error_on_or_filter_falls_back_to_received_time_only(self):
        f = FakeFolder("F", table_rows=self.mails_for(3), table_opts={"get_table_exc": self.amb(), "get_table_exc_calls": {1}})
        stats = {}
        tbl, fmode, cols = T.open_table(f, datetime(2024, 1, 1), datetime(2025, 1, 1), {}, alive=lambda: False, stats=stats)
        self.assertEqual(fmode, "ReceivedTime")                    # 接続が死んでいても、後のフィルタが通れば成功（alive は呼ばない）
        self.assertEqual(stats["ambiguous_disp"], 1)
        self.assertEqual(len(f.table_filters), 2)

    def test_all_filters_ambiguous_and_connection_alive_means_table_unavailable(self):
        f = FakeFolder("F", table_rows=self.mails_for(3), table_opts={"get_table_exc": self.amb()})
        stats = {}
        with self.assertRaises(T.TableUnavailable):
            T.open_table(f, datetime(2024, 1, 1), datetime(2025, 1, 1), {}, alive=lambda: True, stats=stats)
        self.assertEqual((len(f.table_filters), stats["ambiguous_disp"]), (3, 3))        # 全フィルタを試した

    def test_all_filters_ambiguous_and_connection_dead_reraises(self):
        f = FakeFolder("F", table_rows=self.mails_for(3), table_opts={"get_table_exc": self.amb()})
        with self.assertRaises(FakeComError):
            T.open_table(f, datetime(2024, 1, 1), datetime(2025, 1, 1), {}, alive=lambda: False)
        self.assertEqual(len(f.table_filters), 3)

    def test_mixed_failures_do_not_check_connection(self):
        f = FakeFolder("F", table_rows=self.mails_for(3), table_opts={"fail_any_filter": True, "fail_or": True})
        calls = []
        tbl, fmode, cols = T.open_table(f, datetime(2024, 1, 1), datetime(2025, 1, 1), {}, alive=lambda: calls.append(1) or False)
        self.assertEqual((fmode, calls), ("all", []))

    def test_definite_connection_error_is_reraised_immediately(self):
        f = FakeFolder("F", table_rows=self.mails_for(3), table_opts={"get_table_exc": FakeComError(-2147221227, "x")})
        calls = []
        with self.assertRaises(FakeComError):
            T.open_table(f, datetime(2024, 1, 1), datetime(2025, 1, 1), {}, alive=lambda: calls.append(1) or True)
        self.assertEqual((len(f.table_filters), calls), (1, []))

    def test_e2e_or_ambiguous_still_scans_in_table_mode_and_counts(self):
        f = FakeFolder("F", table_rows=self.mails_for(3), table_opts={"get_table_exc": self.amb(), "get_table_exc_calls": {1}})
        code, out = run_main(self.argv_t(), fake_com(self.one_store(f))[0])
        self.assertEqual(code, 0)
        c = self.caches()[0]
        self.assertEqual((c["scan_mode"], c["complete"], len(c["records"])), ("table", True, 3))
        self.assertIn("曖昧な 0x80020009（DISP_E_EXCEPTION）を判定した回数: 1回 / table→items フォールバック: 0フォルダ", out)
        self.assertIn("ReceivedTime のみで絞りました", out)
        self.assertEqual(f._items.filters, [])

    def test_e2e_all_ambiguous_with_live_connection_falls_back_to_items(self):
        f = FakeFolder("F", [mk_item("I1", "s", dt(4, 1), "Zed", "Owner Taro")], table_rows=self.mails_for(3),
                       table_opts={"get_table_exc": self.amb()})
        code, out = run_main(self.argv_t(), fake_com(self.one_store(f))[0])
        self.assertEqual(code, 0)
        self.assertIn("Items方式に切り替えます", out)
        self.assertIn("曖昧な 0x80020009（DISP_E_EXCEPTION）を判定した回数: 3回 / table→items フォールバック: 1フォルダ", out)
        self.assertEqual(self.caches()[0]["scan_mode"], "items")

    def test_e2e_all_ambiguous_with_dead_connection_goes_to_retry(self):
        f = FakeFolder("F", [mk_item("I1", "s", dt(4, 1), "Zed", "Owner Taro")], table_rows=self.mails_for(3),
                       table_opts={"get_table_exc": self.amb()})
        real = T.OutlookConn.check
        calls = {"n": 0}

        def dead_after_first(self_):
            calls["n"] += 1
            return calls["n"] == 1          # 走査前の確認は通り、その後（曖昧エラー時）は死んでいる扱い
        T.OutlookConn.check = dead_after_first
        try:
            code, out = run_main(self.argv_t(), fake_com(self.one_store(f))[0])
        finally:
            T.OutlookConn.check = real
        self.assertIn("リトライ 1/2", out)
        self.assertNotIn("Items方式に切り替えます", out)
        self.assertIn("失敗フォルダ 1個", out)
        self.assertEqual(f._items.filters, [])



class TestProbeReviewFixes(TableBase):
    def ns_first_row_non_mail(self):
        # 対象月の先頭の行が会議出席依頼（メール以外）。メールはその後ろにある
        rows = [trow("M0", "x", dt(4, 1), cls="IPM.Schedule.Meeting.Request"), trow("M1", "x", dt(4, 2), cls="IPM.Appointment")] + \
               [trow(f"E{i}", "s", dt(4, 3 + i), conv=CONV_A) for i in range(3)]
        items = {(f"E{i}", "SB"): mk_item(f"E{i}", "s", dt(4, 3 + i), cid="o",
                                         PropertyAccessor=FakePA({T.PR_CONVERSATION_ID_URL: CONV_A})) for i in range(3)}
        empty = FakeFolder("E", [mk_item("X", "s", datetime(2023, 1, 1))], table_rows=[trow("X", "s", datetime(2023, 1, 1))])
        full = FakeFolder("Full", [mk_item("x", "s", dt(4, 1))], table_rows=rows)
        a = FakeStore("StoreA", FakeFolder("root", [], subfolders=[empty]), sid="SA")
        b = FakeStore("StoreB", FakeFolder("root", [], subfolders=[full]), sid="SB")
        return FakeNamespace([a, b], items)

    def test_store_with_non_mail_first_rows_is_not_skipped(self):
        code, out = run_main(["--data-dir", self.data, "--output-dir", self.out, "--from", "2024-04", "--probe-table"],
                             fake_com(self.ns_first_row_non_mail())[0])
        self.assertEqual(code, 0)
        self.assertIn("StoreA: 2024-04 に該当メールなし", out)       # 空のストアはスキップ
        self.assertNotIn("StoreB: 2024-04 に該当メールなし", out)    # 先頭がメール以外でもメールのあるストアはスキップしない
        self.assertIn("ストア: StoreB", out)
        self.assertIn("同一メールとして比較できた 3件", out)

    def test_selection_check_uses_small_chunk_and_up_to_50_rows(self):
        import unittest.mock as mock
        seen = []
        real = T.scan_folder_table

        def spy(folder, s_, e_, state, **kw):
            seen.append(kw.get("max_rows"))
            seen.append(kw.get("chunk"))
            return real(folder, s_, e_, state, **kw)
        with mock.patch.object(T, "scan_folder_table", spy):
            run_main(["--data-dir", self.data, "--output-dir", self.out, "--from", "2024-04", "--probe-table"],
                     fake_com(self.ns_first_row_non_mail())[0])
        self.assertEqual(seen[:2], [50, 50])

    def test_recommended_chunk_is_the_one_with_min_per_row_time(self):
        import unittest.mock as mock
        rows = [trow(f"E{i}", "s", dt(4, 1 + i % 25)) for i in range(30)]
        f = FakeFolder("Big", [mk_item("x", "s", dt(4, 1))], table_rows=rows)
        ns = FakeNamespace([FakeStore("MB", FakeFolder("root", [], subfolders=[f]), sid="SID-MB")])
        per_row = {50: 0.003, 100: 0.002, 500: 0.001, 1000: 0.0015, 2000: 0.004}      # 秒/行
        seen_k = []

        def fake_fetch(folder, s_, e_, keys, chunk, k):
            seen_k.append(k)
            fetched = max(chunk, k)             # K(500)<チャンク(1000,2000)なら、チャンク分だけ取得する
            return {"setup": 0.01, "fetch": per_row[chunk] * fetched, "fetched": fetched, "used": k, "filter": "or"}
        with mock.patch.object(T, "time_table_fetch", fake_fetch):
            code, out = run_main(["--data-dir", self.data, "--output-dir", self.out, "--from", "2024-04", "--probe-speed"],
                                 fake_com(ns)[0])
        self.assertEqual(code, 0)
        self.assertIn("--table-chunk 500", out)                    # 総時間（K行を得る時間）ではなく1行あたり時間が最小のチャンク
        self.assertIn("1.0ms/行", out)
        self.assertEqual(set(seen_k), {500})                       # 既定 K=500
        self.assertEqual(T.SPEED_CHUNKS, (50, 100, 500, 1000, 2000))
        with mock.patch.object(T, "time_table_fetch", fake_fetch):
            run_main(["--data-dir", self.data, "--output-dir", self.out, "--from", "2024-04", "--probe-speed", "--probe-limit", "800"],
                     fake_com(ns)[0])
        self.assertIn(800, seen_k)

    def test_each_condition_is_measured_twice_and_the_faster_is_used(self):
        import unittest.mock as mock
        ns, _f = TestProbeSpeed.ns(TestProbeSpeed("test_zero_rows_and_theme_not_required"))
        calls = []

        def fake_fetch(folder, s_, e_, keys, chunk, k):
            calls.append((len(keys), chunk))
            n = calls.count((len(keys), chunk))
            return {"setup": 0.0, "fetch": 0.5 if n == 1 else 0.05, "fetched": 100, "used": 100, "filter": "or"}
        with mock.patch.object(T, "time_table_fetch", fake_fetch):
            code, out = run_main(["--data-dir", self.data, "--output-dir", self.out, "--from", "2024-04", "--probe-speed"],
                                 fake_com(ns)[0])
        self.assertEqual(code, 0)
        self.assertEqual(calls.count((11, 500)), 2)
        self.assertIn("1行あたり    0.5ms", out)                   # 0.05秒/100行 = 0.5ms（遅い方の5msではない）

    def test_help_and_notes(self):
        buf = io.StringIO()
        with contextlib.redirect_stdout(buf):
            with self.assertRaises(SystemExit):
                T.parse_args(["--help"])
        text = "".join(buf.getvalue().split())          # ヘルプは折り返されるので空白を除いて照合する
        self.assertIn("同じEntryIDのItem", text)
        self.assertIn("送信済みフォルダ", text)
        code, out = run_main(["--data-dir", self.data, "--output-dir", self.out, "--from", "2024-04", "--probe-table"],
                             fake_com(self.ns_first_row_non_mail())[0])
        self.assertIn("送信済みフォルダでは", out)



if __name__ == "__main__":
    unittest.main()
