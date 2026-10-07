# -*- coding: utf-8 -*-
"""サイト名のみを検索する — v20261007_01 で追加

document_search_manager: 「1. SharePoint」タブの「サイト名のみを検索する」を
オンにすると、文書ではなく、キーワードに名前が合う**サイトそのもの**を
結果の表に出す機能を検証する（越智さんの要望、2026-10-07。A案）。
  - サイトの行（SearchResult.is_site）への変換
  - SearchManager.search_site_list（件数の上限・失敗時・0件）
  - /api/search の site_only（SharePointタブだけ・キーワード必須・サイト指定は無視）
  - サイトの行を ZIP・要約・Excel/CSV に混ぜない（誤操作の防止）
  - 画面（チェックボックス・列構成・[hidden]対策・キャッシュキー・無効化の戻し忘れ）
Graph API には一切アクセスせず、スタブで検証する。
実行: python tests/test_32_site_only_search.py
（まとめて実行する場合は python tests/run_tests.py）
"""
import csv
import io
import sys
import tempfile
from dataclasses import asdict
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from _harness import dsm, DummyAuth  # noqa: E402

ok, ng = 0, 0


def check(label, cond, detail=""):
    global ok, ng
    if cond:
        ok += 1
        print(f"  OK   {label}")
    else:
        ng += 1
        print(f"  NG   {label}  {detail}")


CFG = dict(dsm.DEFAULT_CFG, mcas_host="nexperia.sharepoint.com.mcas.ms")
PO = "https://nexperia.sharepoint.com/sites/PO"
LORRY = "https://nexperia.sharepoint.com/sites/P010024_Lorry"
JEEP = "https://nexperia.sharepoint.com/sites/P010008-Jeeppkg-spin"

_ORIG = {"post": dsm.http_req.post, "get": dsm.http_req.get,
         "cfg": dsm._cfg, "mgr": dsm._manager, "export": dsm.EXPORT_DIR,
         "last": dsm._last_results}
WORK = Path(tempfile.mkdtemp())


class FakeResp:
    def __init__(self, status, body=None, text=""):
        self.status_code, self._body, self.text = status, body, text

    def json(self):
        if self._body is None:
            raise ValueError("not json")
        return self._body


def site_hits(*sites):
    return {"value": [{"hitsContainers": [{
        "total": len(sites),
        "hits": [{"resource": {"displayName": n, "name": n, "webUrl": u,
                               "description": d}} for n, u, d in sites],
    }]}]}


try:
    # ── T1: サイトの行への変換 ────────────────────────────────
    print("\n[T1] search_site_results（サイト → 検索結果の行）")
    sp = dsm.SharePointProvider(CFG, DummyAuth())
    seen = []

    def fake_post(url, **kw):
        seen.append(kw.get("json"))
        return FakeResp(200, site_hits(("P010024_Lorry", LORRY, "Lorry project"),
                                       ("PO", PO, "")))
    dsm.http_req.post = fake_post
    r = sp.search_site_results("PO", 50)
    check("成功し、サイトの行を返す", r.get("ok") and len(r["results"]) == 2, r)
    row = r["results"][0]
    check("is_site が True（ZIP・要約・探索の対象から外すための印）", row.is_site is True)
    check("種別は「サイト」", row.doc_type == "サイト", row.doc_type)
    check("タイトルとサイト名はサイトの表示名", row.title == "P010024_Lorry" and row.site == "P010024_Lorry")
    check("URL・サイトURLはサイトのトップ", row.url == LORRY and row.site_url == LORRY)
    check("場所は /sites/名前", row.path_text == "/sites/P010024_Lorry", row.path_text)
    check("説明を持つ", row.description == "Lorry project")
    check("フォルダではない（フォルダ探索の対象にならない）", row.is_folder is False)
    check("ソースは SharePoint", row.source == "SharePoint")
    check("順位（rank）は1から", [x.rank for x in r["results"]] == [1, 2])
    check("要求は entityTypes=site（実機で確認済みの方式）",
          seen[0]["requests"][0]["entityTypes"] == ["site"], seen[0])
    check("要求の件数は渡した上限", seen[0]["requests"][0]["size"] == 50, seen[0])

    cfg_mcas = dict(CFG, rewrite_host_to_mcas=True)
    sp_m = dsm.SharePointProvider(cfg_mcas, DummyAuth())
    rm = sp_m.search_site_results("PO", 10)
    check("rewrite_host_to_mcas がオンなら、リンクは .mcas.ms 経由（文書行と同じ扱い）",
          "mcas.ms" in rm["results"][0].url and "mcas.ms" in rm["results"][0].site_url,
          rm["results"][0].url)

    # 0件のとき先頭の語で探し直す
    calls = []

    def fake_post2(url, **kw):
        term = kw["json"]["requests"][0]["query"]["queryString"]
        calls.append(term)
        if term == "P010008":
            return FakeResp(200, site_hits(("P010008 - Jeep pkg-spin", JEEP, "")))
        return FakeResp(200, site_hits())
    dsm.http_req.post = fake_post2
    r = sp.search_site_results("P010008 - Jeep pkg-spin", 10)
    check("0件なら先頭の語で探し直す（9/25の実機で確認した方法）",
          calls == ["P010008 - Jeep pkg-spin", "P010008"] and len(r["results"]) == 1, calls)
    check("探し直した事実を note に残す（黙って別の語にしない）",
          "P010008" in r.get("note", "") and "探し直し" in r.get("note", ""), r.get("note"))

    dsm.http_req.post = lambda url, **kw: FakeResp(403, None, "denied")
    dsm.http_req.get = lambda url, **kw: FakeResp(500, None, "error")
    r = sp.search_site_results("PO", 10)
    check("両方式が失敗したら ok=False と理由（例外にしない）",
          r.get("ok") is False and "HTTP" in r.get("message", ""), r)
    check("空の語は ok=False", sp.search_site_results('  "" ', 10).get("ok") is False)

    # ── T2: SearchManager.search_site_list ────────────────────
    print("\n[T2] SearchManager.search_site_list")
    mgr = dsm.SearchManager(dict(CFG, site_only_max_results=20), DummyAuth())
    sizes = []

    def fake_post3(url, **kw):
        sizes.append(kw["json"]["requests"][0]["size"])
        return FakeResp(200, site_hits(("PO", PO, "")))
    dsm.http_req.post = fake_post3
    out = mgr.search_site_list("PO", 100)
    check("応答の形は search() と同じ（results / statuses / excluded_nexus）",
          set(out) == {"results", "statuses", "excluded_nexus"}, set(out))
    check("画面の件数が大きくても、サイト用の上限（既定50→ここでは20）に抑える",
          sizes == [20], sizes)
    check("画面の件数が小さければそちらを使う", (mgr.search_site_list("PO", 5), sizes[-1])[1] == 5, sizes)
    st = out["statuses"][0]
    check("状態は ok・件数あり・系統は sharepoint",
          st["key"] == "sharepoint" and st["state"] == "ok" and st["count"] == 1, st)
    dsm.http_req.post = lambda url, **kw: FakeResp(200, site_hits())
    out = mgr.search_site_list("zzz", 10)
    check("0件は state=empty", out["statuses"][0]["state"] == "empty" and out["results"] == [])
    dsm.http_req.post = lambda url, **kw: FakeResp(403, None, "denied")
    dsm.http_req.get = lambda url, **kw: FakeResp(500, None, "error")
    out = mgr.search_site_list("PO", 10)
    check("失敗は state=error と理由（例外を素通ししない）",
          out["statuses"][0]["state"] == "error" and out["statuses"][0]["message"], out["statuses"])
    check("既定の上限は config に定義されている",
          dsm.DEFAULT_CFG.get("site_only_max_results") == 50)

    # ── T3: /api/search（site_only） ──────────────────────────
    print("\n[T3] /api/search（site_only）")
    dsm._cfg = CFG
    mgr3 = dsm.SearchManager(CFG, DummyAuth())
    doc_calls = []

    def sp_search(keyword, max_results):
        doc_calls.append((keyword, mgr3.providers[dsm.TARGET_SHAREPOINT].site_scope_url))
        return {"results": [], "total": 0, "note": ""}
    mgr3.providers[dsm.TARGET_SHAREPOINT].search = sp_search
    for k in (dsm.TARGET_NEXUS, dsm.TARGET_ENOVIA):
        mgr3.providers[k].search = lambda kw, mx: {"results": [], "total": 0, "note": ""}
    dsm._manager = mgr3
    client = dsm.flask_app.test_client()
    dsm.http_req.post = lambda url, **kw: FakeResp(200, site_hits(("PO", PO, "po site"),
                                                                  ("P010024_Lorry", LORRY, "")))

    resp = client.post("/api/search", json={"keyword": "PO", "target": "sharepoint",
                                            "site_only": True})
    body = resp.get_json()
    check("サイトの一覧を返す（文書検索は呼ばない）",
          resp.status_code == 200 and len(body["results"]) == 2 and not doc_calls, (resp.status_code, doc_calls))
    check("行は is_site=True で、画面が列構成を切り替えられるよう site_only を返す",
          body["results"][0]["is_site"] is True and body["site_only"] is True, body.get("site_only"))
    check("target は sharepoint のまま", body["target"] == "sharepoint")
    resp = client.post("/api/search", json={"keyword": "", "target": "sharepoint", "site_only": True})
    check("キーワードが空なら 400（サイト名が要る）",
          resp.status_code == 400 and "サイト名" in resp.get_json().get("error", ""), resp.get_json())
    resp = client.post("/api/search", json={"keyword": "", "target": "sharepoint", "site_only": True,
                                            "site_scope": LORRY})
    check("サイトを選んでいても、サイト名のみはキーワードが必要（サイト内の一覧にならない）",
          resp.status_code == 400)
    doc_calls.clear()
    resp = client.post("/api/search", json={"keyword": "PO", "target": "sharepoint", "site_only": True,
                                            "site_scope": LORRY})
    check("サイト名のみのときは、対象サイトの指定を無視する（文書を探さない）",
          resp.status_code == 200 and not doc_calls and resp.get_json()["site_scope"] == "")
    for tgt in ("all", "nexus", "enovia"):
        doc_calls.clear()
        resp = client.post("/api/search", json={"keyword": "PO", "target": tgt, "site_only": True})
        b = resp.get_json()
        check(f"{tgt}タブでは site_only を無視する（従来どおりの検索）",
              resp.status_code == 200 and b["site_only"] is False
              and not any(r.get("is_site") for r in b["results"]), (tgt, b.get("site_only")))
    doc_calls.clear()
    resp = client.post("/api/search", json={"keyword": "PO", "target": "sharepoint"})
    check("site_only を付けなければ従来どおり文書検索",
          resp.status_code == 200 and doc_calls == [("PO", "")] and resp.get_json()["site_only"] is False,
          doc_calls)

    # ── T4: サイトの行を ZIP・要約・Excel/CSVに混ぜない ─────────
    print("\n[T4] サイトの行の扱い（ZIP・要約・出力）")
    client.post("/api/search", json={"keyword": "PO", "target": "sharepoint", "site_only": True})
    check("最後の結果がサイトの行になっている", dsm._last_results and dsm._last_results[0].is_site)

    downloaded = []
    orig_get = dsm.http_req.get
    dsm.http_req.get = lambda url, **kw: downloaded.append(url) or FakeResp(200, None, "")
    resp = client.get("/api/download?idx=0,1")
    check("サイトの行は ZIP の対象にならない（400。サイトのトップページを取りに行かない）",
          resp.status_code == 400 and not downloaded, (resp.status_code, downloaded))
    dsm.http_req.get = orig_get
    ok_s, reason = dsm._summarizable_reason(dsm._last_results[0])
    check("サイトの行は要約できない（理由つき）", ok_s is False and "サイト" in reason, reason)
    resp = client.post("/api/summarize", json={"idx": "0"})
    check("要約APIも 400 で断る", resp.status_code == 400 and "サイト" in resp.get_json().get("error", ""),
          resp.get_json())

    dsm.EXPORT_DIR = WORK / "exports"
    resp = client.get("/api/export?format=csv&idx=0,1")
    text = resp.data.decode("utf-8-sig")
    rows = list(csv.reader(io.StringIO(text)))
    check("CSVはサイト用の列（サイト名・場所・説明・サイトURL）", rows[0] == ["サイト名", "場所", "説明", "サイトURL"], rows[0])
    check("CSVの内容は、サイト名・場所・説明・URL",
          rows[1] == ["PO", "/sites/PO", "po site", PO], rows[1])
    resp = client.get("/api/export?format=xlsx&idx=0,1")
    check("Excelも出力できる", resp.status_code == 200 and resp.data[:2] == b"PK", resp.status_code)
    try:
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(resp.data))
        ws = wb.active
        check("Excelの見出しはサイト用で、サイト名のセルにはサイトへのリンクが付く",
              [c.value for c in ws[1]] == ["サイト名", "場所", "説明"]
              and ws["A2"].hyperlink is not None and ws["A2"].hyperlink.target == PO,
              [c.value for c in ws[1]])
    except ModuleNotFoundError:
        check("（openpyxl 未導入のためExcelの中身の検査は省略）", True)

    # ── T5: 画面 ──────────────────────────────────────────────
    print("\n[T5] 画面（チェックボックス・列構成・無効化・キャッシュ）")
    html = dsm.INDEX_HTML
    check("チェックボックスは SharePointタブ専用の行（folderOnlyRow）の中にある",
          html.index('id="folderOnlyRow"') < html.index('id="siteOnly"') < html.index('id="siteScopeBox"'))
    check("「サイト名のみを検索する（SharePointタブのみ）」と表示する",
          "サイト名のみを検索する（SharePointタブのみ）" in html)
    check("注記の span は既定で隠れている", 'id="siteOnlyNote" hidden' in html)
    check("サイト用の列構成を持つ（サイト名・場所・説明・サイト内検索）",
          'sites: [' in html and '{ key: "__sitego"' in html and 'label: "サイト名"' in html)
    sites_block = html[html.index("sites: ["):html.index("explore: [")]
    check("サイト用の列に 選択・探索・要約 は含めない",
          "__select" not in sites_block and "__explore" not in sites_block
          and "__summary" not in sites_block, sites_block)
    check("列構成は応答の site_only で切り替える", "applyColumnSet(actualTarget, !!data.site_only)" in html)
    check("サイト一覧のときは探索ボタンを出さない",
          'btn.hidden = (shownTarget !== "sharepoint" || shownSiteOnly)' in html)
    check("サイト一覧のときは「フォルダのみ」の絞り込みを自動設定しない",
          '&& !data.site_only) {\n    filters.doc_type' in html)
    check("要約ボタンの判定もサイトを断る", 'if (r.is_site) { return "サイトは要約できません。"; }' in html)
    check("一括ZIPの選択対象からサイトを外す", "!r.is_folder && !r.is_site" in html)
    check("キャッシュのキーに「サイト名のみ」を含める（文書の結果とサイトの結果を取り違えない）",
          "typeScope, siteScope, siteOnly)" in html and 'siteOnly ? "1" : ""' in html)
    check("サイト名のみはSharePointタブでだけ有効（他のタブでは無視）",
          'return currentTarget() === "sharepoint" && document.getElementById("siteOnly").checked;' in html)
    check("リクエストに site_only を載せる", "site_only: p.siteOnly" in html)
    check("キーワードが空ならログで案内して送らない", "「サイト名のみ」はキーワードが必要です" in html)
    check("オンの間は、文書検索用の指定を無効にする（タイトル限定・前方一致・フォルダのみ・サイト）",
          all(x in html for x in ('"titleOnly", "prefixSearch", "folderOnly", "siteSearchInput"',
                                  '"favoriteSites", "btnFavToggle", "btnClearSite"')))
    check("タブを切り替えるたびに無効化を見直す（Nexus等で無効のまま残さない）",
          "  updateSiteOnlyUi();\n  updateTabHint();\n}" in html)
    check("busy のときもサイトを探すボタンの状態と矛盾させない",
          'btnSiteSearch").disabled = busy || isSiteOnlyActive()' in html)
    check("サイトの行の🔍でサイト内検索に進むとき、サイト名のみをオフに戻す",
          'document.getElementById("siteOnly").checked = false;' in html
          and "selectSite({ name: r.site || r.title" in html)
    check("チェックを変えたのに再検索していないときは、案内を出す",
          "「サイト名のみ」の指定を変えました" in html)
    check("サイト名のみは次回起動時に引き継がない（保存対象に含めない）",
          'site_only' not in html[html.index("function saveState()"):html.index("// ── ヘッダー描画")])
    check("バージョン表記", "v20261007_01" in html and "サイト名のみ検索" in html)

finally:
    dsm.http_req.post, dsm.http_req.get = _ORIG["post"], _ORIG["get"]
    dsm._cfg, dsm._manager = _ORIG["cfg"], _ORIG["mgr"]
    dsm.EXPORT_DIR, dsm._last_results = _ORIG["export"], _ORIG["last"]


print(f"\n{'=' * 46}")
print(f"  成功 {ok} 件 / 失敗 {ng} 件")
print(f"{'=' * 46}")
sys.exit(1 if ng else 0)
