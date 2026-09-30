r"""
VERSION 2026.0930.01
PowerPoint Gemini 翻訳ツール（書式完全保持版）

■ このバージョンでの変更（画像だけのスライドの翻訳: メモ欄版）
文字データが1つも無く、全面に画像が貼られているだけのスライド（他のツールで
書き出した資料など）は、従来の run 単位の翻訳では何も拾えず「翻訳対象のテキストが
見つかりません」で終了していた。**文字が1つも見つからなかった場合に限り**、画像そのものを
Gemini へ送って画像内の文字と訳文を読み取り、各スライド下のメモ欄(ノート)へ追記する。
  - 対象スライド: 画面の70%以上を覆う画像が1枚だけあり、翻訳対象の文字データが無いスライド。
    複数の画像・回転した画像・70%未満の画像のスライドは対象外（件数を結果に出す）。
  - 確認ダイアログで「画像そのものを外部AI(Gemini)へ送信する」ことを明示し、処理する
    スライドの範囲（例: 6 / 1-3,5。空欄=全部）を指定できる。まず1枚だけ試すための入力。
  - 元のPPTXは上書きしない。既存のメモは消さず、末尾に区切り線つきで追記する。
  - 失敗したスライドは無加工のまま残し、番号を結果に出す。3枚連続で失敗したら中断する
    （それまでに終わったスライドは保存する）。
  - 数字・%だけのブロックは返させない。ロゴ・アイコンは除外、固有名詞は訳さない。
  - 送信する画像は幅1600pxのJPEG(q85)。ログには原文を書かない（件数とスライド番号のみ）。
  - 画像を載せるため、シム(_CommonGeminiModels.generate_content)に「contents が list のとき
    パーツ列として送る」分岐と、generationConfig の maxOutputTokens / thinkingBudget を追加した。
    contents が str のとき（従来の翻訳）の payload は変わらない（テストで完全一致を確認）。
  - 文字ありスライドの翻訳経路は変更していない。ただし、文字のあるスライドと画像だけの
    スライドが混ざった資料では、完了メッセージに「画像だけのスライドN枚は未対応」と1行出す。
  - 【仕様】メモ欄に翻訳対象の文章があると「文字あり」の資料として従来の経路になる
    （画像の文字は翻訳されず、上の1行で知らせる）。このツールが画像の訳をメモ欄へ書いた
    ファイルを再度読み込んだ場合も同じ（メモの訳文が文字として扱われる）。翻訳は元のファイル
    から行うこと。

■ 20260911_01 での変更（出力ファイル名の言語コードを2文字へ短縮）
翻訳後のファイル名に付く言語部分を、翻訳先言語そのまま（`_gemini_japanese`）から
2文字の言語コード（`_ja`）へ短縮した。pdf_translator と同じ考え方で、
LANGUAGE_SUFFIX_MAP / lang_to_suffix() を用意して変換している。

  旧: 資料_gemini_japanese.pptx / 資料_gemini_english.pptx / 資料_gemini_chinese.pptx
  新: 資料_ja.pptx            / 資料_en.pptx            / 資料_cn.pptx

変更したのは `translate_ppt_document_thread` の出力パス組み立て2行と、
新規追加した `lang_to_suffix()` だけ。翻訳処理・書式保持・Gemini呼び出しには
一切手を入れていない。

【申し送り】**出力ファイル名が変わる。** 旧版で作った `_gemini_japanese.pptx` は
そのまま残るため、同じ資料を新版で翻訳すると `_ja.pptx` が別ファイルとして
できる（上書きはされない）。
【申し送り】中国語簡体字は `cn` にしている。pdf_translator は同じ言語に `zh`
（ISO 639-1）を使っているため、ツール間で綴りが揃っていない。揃えたくなったら
下の LANGUAGE_SUFFIX_MAP を直せばよい。

■ 20260812_01 での変更（Gemini APIプロキシ対応）
会社PCからGemini APIへの直接アクセスが遮断された事象（2026-08-10頃）への対応。
旧SDK google.generativeai を直接呼ぶ方式をやめ、共通モジュール
gemini_client.py の generate_advanced() 経由（直接呼び出しが失敗したら
自宅PCのプロキシへ自動フォールバック）へ移行した。
rtocs_organizer / analog_ic_se_strategy_organizer / outlook_total_organizer /
excel_translation / pdf_translator と同じ方式。

変更した関数は次の3つだけで、PowerPoint処理側（スライド走査・run単位の書き戻し・
進捗表示・ファイル選択）は一切変更していない。
  `check_dependencies` … google-generativeai のチェックを共通モジュールのチェックへ
  `init_gemini`        … genai.list_models() による自動モデル検出を廃止（下記）
  `translate_batch_gemini` … 互換シム経由の呼び出しへ置換

【重要】自動モデル検出を廃止した理由:
旧版は init_gemini() の中で genai.list_models() を呼んで使用可能モデルを
自動検出していたが、これはネットワークアクセスを伴うため、直接アクセスが
遮断された環境では必ず例外になり「API初期化エラー」で sys.exit(1) していた
（＝プロキシ経由なら翻訳できるのに、ツールが起動すらできない）。
共通モジュール・プロキシのどちらにも list_models 相当が無いため、固定モデル名
（環境変数 GEMINI_MODEL で上書き可。既定 gemini-2.5-flash）を使う方式に変更した。
副作用として、指定モデルが使えない環境では404が出るようになる。

【挙動差】旧版の request_options={"timeout": 40} は共通モジュールに同等機能が
無いため削除した。タイムアウトは gemini_client.py 側の固定値（直接15秒 /
プロキシ60秒）になる。遮断下では最初のバッチだけ直接呼び出しの15秒を待つぶん
遅くなるが、一度失敗すると以降はプロキシ直行になるため2バッチ目以降は影響しない。
これは仕様どおりの挙動であり、不具合ではない。

■ 必要な環境変数（どちらか一方以上）
  GEMINI_API_KEY    … 直接呼び出し用
  GEMINI_PROXY_URL  … 自宅PCプロキシのURL（直接呼び出し失敗時のフォールバック先）
  GEMINI_MODEL      … 使用モデルを変えたい場合のみ（任意。既定 gemini-2.5-flash）
  GEMINI_COMMON_DIR … gemini_client.py の置き場所を明示したい場合のみ（任意）

■ 使い方
  1. run_ppt_translator.bat を実行（または python ppt_translation_20260911_01.py）
  2. 翻訳したい .pptx を選ぶ
  3. 翻訳先言語を選んで「翻訳開始」を押す
  4. 完了すると同じフォルダに `元ファイル名_ja.pptx`（英語なら `_en`、
     中国語簡体字なら `_cn`）のように保存される

■ 20260309_03 までの経緯（旧docstringより）
  - 進捗の可視化とUIハングアップ対策: 並列処理エンジンがバッチ（10項目）を処理する
    ごとにプログレスバーを更新するコールバック関数を導入。
  - APIタイムアウトとリトライ機構: 失敗時は「3回・3秒間隔」で再試行。
  - フェイルファスト（即時撤退）と強制切断: 3バッチ連続でエラーが発生した場合、
    残りの処理を即座にキャンセルして終了する安全装置。
  - デバッグ用ログ（LOG）出力: 処理の足跡と通信エラーの詳細を translation_debug.log
    に記録し、コンソールにも出力。
  - ファイルロックの事前検知: 翻訳処理を開始する前に出力先ファイルの書き込み権限を
    チェックし、ロックされている場合は即座に警告。
"""
import tkinter as tk
from tkinter import filedialog, messagebox
import os
import sys
import shutil
import re
import time
import threading
from concurrent.futures import ThreadPoolExecutor, as_completed
import logging
import traceback
import base64
import io
import json

# ============================================================
# Gemini 共通クライアント(gemini_client.py)への互換シム
# ============================================================
# 会社PCからGemini APIへの直接アクセスが遮断される事象(2026-08-10頃)を受け、
# rtocs_organizer / analog_ic_se_strategy_organizer / outlook_total_organizer /
# excel_translation / pdf_translator と同様に、共通モジュール gemini_client.py の
# generate_advanced() 経由(直接呼び出しが失敗したら自宅PCプロキシへ自動フォール
# バック)へ移行した。
#
# 本ツールは旧SDK(google.generativeai)の genai.GenerativeModel(...).generate_content()
# を1箇所で使っていただけだが、他ツールと実装を揃えるため、同じ形の薄い互換シム
# (_CommonGeminiClient)を用意し、そこを経由する方式にした。これにより、レスポンスを
# 読む側(response.parts で空判定 → response.text を正規表現で番号付きリストへ戻す処理)や
# リトライ・フェイルファスト・進捗表示のロジックは一切変更しないで済んでいる。
# 旧SDK(google-generativeai)への依存は本バージョンで無くなった。
#
# 必要な環境変数(会社PC):
#   GEMINI_API_KEY   … 直接呼び出し用(gemini_client.py 側が読む)
#   GEMINI_PROXY_URL … 自宅PCプロキシのURL(直接呼び出し失敗時のフォールバック先)
#   GEMINI_MODEL     … 使用モデルを変えたい場合のみ(任意。既定 gemini-2.5-flash)
#   GEMINI_COMMON_DIR… gemini_client.py の置き場所を明示したい場合のみ(任意)

_GEMINI_COMMON_DIR_ENV = os.environ.get("GEMINI_COMMON_DIR")
if _GEMINI_COMMON_DIR_ENV:
    _COMMON_DIR_CANDIDATES = [_GEMINI_COMMON_DIR_ENV]
else:
    # 会社PCでは本スクリプトが PythonScripts\Powerpoint\ppt_translator\ に、
    # gemini_client.py が PythonScripts\common\ に置かれるため、正解は「1つ上」では
    # なく「2つ上」の common になる。他ツール(1つ上が common)と同じ配置に置かれた
    # 場合でも動くよう、上位ディレクトリを順に探して最初に見つかったものを使う。
    _SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
    _COMMON_DIR_CANDIDATES = [
        os.path.abspath(os.path.join(_SCRIPT_DIR, *([os.pardir] * _n + ["common"])))
        for _n in (1, 2, 3)
    ]

_COMMON_DIR = next(
    (_d for _d in _COMMON_DIR_CANDIDATES
     if os.path.isfile(os.path.join(_d, "gemini_client.py"))),
    _COMMON_DIR_CANDIDATES[0])

if _COMMON_DIR not in sys.path:
    sys.path.insert(0, _COMMON_DIR)

# gemini_client のインポートはここで試みるが、失敗しても import 時点では落とさない
# (原因が分かるメッセージを起動時チェックで出すため)。
try:
    from gemini_client import generate_advanced as _generate_advanced
    _GEMINI_CLIENT_IMPORT_ERROR = None
except Exception as _e:
    _generate_advanced = None
    _GEMINI_CLIENT_IMPORT_ERROR = _e

# 本ツールは全機能が翻訳(AI呼び出し)のため、共通モジュールを読み込めたかどうかが
# そのまま「Gemini が使えるか」になる。
HAS_GEMINI = _generate_advanced is not None
if not HAS_GEMINI:
    print(f"警告: Gemini共通モジュール(gemini_client.py)を読み込めませんでした: "
          f"{_GEMINI_CLIENT_IMPORT_ERROR}")

# 旧版は genai.list_models() で使用可能モデルを自動検出していたが、これはネットワーク
# アクセスを伴うため、直接アクセスが遮断された環境では必ず失敗し、起動時の
# init_gemini() が False を返して sys.exit(1) していた(＝ツールが起動できない)。
# 共通モジュール・プロキシのどちらにも list_models 相当が無いため、自動検出は廃止し、
# 固定モデル名(環境変数 GEMINI_MODEL で上書き可)を使う方式に変更した。
GEMINI_MODEL_NAME = os.environ.get("GEMINI_MODEL", "gemini-2.5-flash")


def _gemini_common_module_error_message():
    """共通モジュールを読み込めなかったときの、原因が分かる案内文を組み立てる。"""
    return ("Gemini共通モジュール(gemini_client.py)を読み込めませんでした。\n"
            f"探索したパス: {' / '.join(_COMMON_DIR_CANDIDATES)}\n"
            f"元のエラー: {_GEMINI_CLIENT_IMPORT_ERROR}\n\n"
            "gemini-common-tools を配置し、必要なら環境変数 GEMINI_COMMON_DIR で\n"
            "gemini_client.py のあるフォルダを指定してください。")


def _schema_to_jsonable(value):
    """REST APIのpayloadへそのまま載せられる素のdict/listへ変換する。
    本ツールの safety_settings は既に素のdictのリストなのでそのまま返るが、
    他ツールのシムと実装を揃えるために残してある(pydanticモデル等が渡された
    場合にJSON化できなくなるのを防ぐ保険)。"""
    if value is None or isinstance(value, (dict, list, str, int, float, bool)):
        return value
    # pydantic v2 (model_dump) / v1 (dict) の両方に対応。REST APIのフィールド名は
    # camelCase なので by_alias=True で別名を使う。
    for attr, kwargs in (("model_dump", {"mode": "json", "exclude_none": True, "by_alias": True}),
                         ("dict", {"exclude_none": True, "by_alias": True})):
        fn = getattr(value, attr, None)
        if callable(fn):
            try:
                return fn(**kwargs)
            except Exception:
                try:
                    return fn()
                except Exception:
                    pass
    return value


class _GeminiGenerateConfig:
    """旧SDKの genai.types.GenerationConfig(...) ＋ safety_settings 相当の設定
    オブジェクト。旧SDKへの依存を断つため、同等の入れ物をここに置く
    (シム側は getattr で属性を読むだけなので実装差の影響を受けない)。"""
    def __init__(self, temperature=None, safety_settings=None,
                 system_instruction=None, response_mime_type=None, response_schema=None,
                 max_output_tokens=None, thinking_budget=None):
        self.temperature = temperature
        self.safety_settings = safety_settings
        self.system_instruction = system_instruction
        self.response_mime_type = response_mime_type
        self.response_schema = response_schema
        # 20260930_01 で追加（画像スライドの翻訳だけが使う。未指定なら payload に載せない）
        self.max_output_tokens = max_output_tokens
        self.thinking_budget = thinking_budget


class _CommonUsageMetadata:
    """response.usage_metadata 互換(トークン計測用)。本ツールは現時点で参照して
    いないが、他ツールのシムと契約を揃えておく。"""
    def __init__(self, usage):
        usage = usage if isinstance(usage, dict) else {}
        self.prompt_token_count = usage.get("promptTokenCount", 0)
        self.candidates_token_count = usage.get("candidatesTokenCount", 0)


class _CommonGeminiResponse:
    """client.models.generate_content(...) の戻り値互換。
    本ツールは response.parts で空応答を判定してから response.text を読むため、
    parts も提供する(空応答なら空リストになるので、呼び出し側の
    「空なら ValueError を投げてリトライ」という既存の挙動がそのまま保たれる)。
    レスポンスが想定外の形でも例外を投げず、text は空文字にする。"""
    def __init__(self, raw):
        try:
            self.text = raw["candidates"][0]["content"]["parts"][0]["text"]
        except (KeyError, IndexError, TypeError):
            self.text = ""
        try:
            self.parts = raw["candidates"][0]["content"]["parts"] or []
        except (KeyError, IndexError, TypeError):
            self.parts = []
        self.usage_metadata = _CommonUsageMetadata(
            raw.get("usageMetadata", {}) if isinstance(raw, dict) else {})


class _CommonGeminiModels:
    """client.models 互換。"""
    def generate_content(self, model=None, contents=None, config=None):
        if _generate_advanced is None:
            raise RuntimeError(_gemini_common_module_error_message())

        if isinstance(contents, (list, tuple)):
            # 20260930_01 で追加: 画像などのパーツ列（str は {"text": ...} に包む）。
            # contents が str のとき（従来の翻訳）は下の else の従来どおりの形になる。
            parts = [{"text": c} if isinstance(c, str) else c for c in contents]
        else:
            parts = [{"text": contents}]
        payload = {"contents": [{"parts": parts}]}
        if config is not None:
            # safety_settings は旧SDKへ渡していた時点で既にREST形式
            # (category / threshold のdictリスト)なので、そのまま載せればよい。
            # 載せ忘れると BLOCK_NONE 指定が消え、資料の内容によっては応答が空になり
            # 「一部のバッチだけ翻訳されない」という切り分けにくい症状になる。
            safety = getattr(config, "safety_settings", None)
            if safety:
                payload["safetySettings"] = _schema_to_jsonable(safety)

            system_instruction = getattr(config, "system_instruction", None)
            if isinstance(system_instruction, str) and system_instruction:
                payload["systemInstruction"] = {"parts": [{"text": system_instruction}]}
            elif isinstance(system_instruction, dict):
                payload["systemInstruction"] = system_instruction

            gen_cfg = {}
            mime = getattr(config, "response_mime_type", None)
            if mime:
                gen_cfg["responseMimeType"] = mime
            schema = getattr(config, "response_schema", None)
            if schema is not None:
                gen_cfg["responseSchema"] = _schema_to_jsonable(schema)
            temp = getattr(config, "temperature", None)
            if temp is not None:
                gen_cfg["temperature"] = temp
            max_tokens = getattr(config, "max_output_tokens", None)
            if max_tokens is not None:
                gen_cfg["maxOutputTokens"] = max_tokens
            thinking = getattr(config, "thinking_budget", None)
            if thinking is not None:
                gen_cfg["thinkingConfig"] = {"thinkingBudget": thinking}
            if gen_cfg:
                payload["generationConfig"] = gen_cfg

        # model は明示的に渡す(共通モジュール側の既定モデルへ勝手にフォールバック
        # されると、意図したモデルと実際に使われるモデルが食い違うため)。
        raw = _generate_advanced(payload, model=model)
        return _CommonGeminiResponse(raw)


class _CommonGeminiClient:
    """genai.Client(api_key=...) 相当。api_key は gemini_client.py 側が環境変数
    GEMINI_API_KEY から読むため、ここでは互換性のために受け取るだけで使用しない。"""
    def __init__(self, api_key=None):
        self.models = _CommonGeminiModels()


def gemini_credentials_available():
    """AI呼び出しが行える見込みがあるかどうかの事前チェック。
    直接呼び出しが遮断されていてもプロキシ経由なら成功しうるため、
    GEMINI_API_KEY / GEMINI_PROXY_URL のどちらか一方でも設定されていれば通す
    (プロキシ専用構成を誤って弾かないため)。"""
    return bool(os.environ.get("GEMINI_API_KEY") or os.environ.get("GEMINI_PROXY_URL"))


# --- 依存関係の確認 ---
try:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    HAS_PPTX = True
except ImportError:
    HAS_PPTX = False

def check_dependencies(root_window):
    """起動時の依存関係チェック"""
    missing_libs = []
    if not HAS_PPTX:
        missing_libs.append("python-pptx")

    if missing_libs:
        error_msg = "以下のライブラリがインストールされていません:\n"
        for lib in missing_libs:
            error_msg += f"- {lib}\n"
        error_msg += "\n以下のコマンドでインストールしてください:\n"
        error_msg += f"pip install {' '.join(missing_libs)}"

        messagebox.showerror("依存関係エラー", error_msg, parent=root_window)
        return False

    # 本ツールは全機能が翻訳(AI呼び出し)のため、共通モジュールが読めない場合は
    # 起動を続けても何もできない。原因が分かる形で案内して終了する。
    if not HAS_GEMINI:
        messagebox.showerror("依存関係エラー", _gemini_common_module_error_message(),
                             parent=root_window)
        return False

    return True

# --- グローバル変数（Gemini互換シムのクライアント） ---
gemini_client = None

# --- 新規追加: ロガーの初期化 ---
def get_logger():
    """デバッグ用ロガーの初期化（コンソールとファイル両方に出力）"""
    logger = logging.getLogger("PPT_Translation")
    if not logger.handlers:
        logger.setLevel(logging.DEBUG)
        fh = logging.FileHandler("translation_debug.log", encoding="utf-8")
        ch = logging.StreamHandler()
        formatter = logging.Formatter("[%(asctime)s] %(levelname)s - %(message)s", "%H:%M:%S")
        fh.setFormatter(formatter)
        ch.setFormatter(formatter)
        logger.addHandler(fh)
        logger.addHandler(ch)
    return logger

def init_gemini(root_window):
    """Gemini呼び出しの事前チェックと互換シムのクライアント生成。

    旧版はここで genai.configure() の後に genai.list_models() を呼び、使用可能な
    モデルを自動検出していた。しかしこれはネットワークアクセスを伴うため、直接
    アクセスが遮断された環境では必ず例外になり、「API初期化エラー」ダイアログを
    出して sys.exit(1) していた（＝ツールが起動できない）。プロキシ経由なら翻訳
    自体は可能なのに起動段階で止まってしまうため、自動モデル検出は廃止し、固定
    モデル名（環境変数 GEMINI_MODEL で上書き可）を使う方式に変更した。
    この関数はネットワークアクセスを一切行わない。
    """
    global gemini_client
    if _generate_advanced is None:
        messagebox.showerror("エラー", _gemini_common_module_error_message(), parent=root_window)
        return False

    if not gemini_credentials_available():
        messagebox.showerror("エラー",
                             "Gemini認証情報が設定されていません。\n"
                             "以下のいずれかを設定してください:\n"
                             "- 環境変数 GEMINI_API_KEY （直接接続用）\n"
                             "- 環境変数 GEMINI_PROXY_URL （自宅PCプロキシ経由用）\n\n"
                             "※ setx で設定した場合は、コマンドプロンプトを\n"
                             "　 開き直してから起動してください。", parent=root_window)
        return False

    gemini_client = _CommonGeminiClient()
    print(f"使用モデル: {GEMINI_MODEL_NAME}")
    return True

def is_translatable(text):
    """翻訳が必要なテキストかどうかを判定"""
    if not text or str(text).strip() == "":
        return False
    
    text_str = str(text).strip()
    
    if text_str in ["", "#", "-", "N/A", "NULL", "•", "◦", "▪", "**", "*", ":", "：", 
                    "I.", "II.", "III.", "IV.", "V.", "VI.", "***"]:
        return False
    if text_str.replace(".", "").replace("-", "").isdigit():
        return False
    if len(text_str) <= 2:
        return False
    return True

def translate_batch_gemini(texts, target_language="Japanese", batch_idx=0, logger=None):
    """Gemini APIを使用した小バッチ翻訳（リトライ・タイムアウト機構付き）"""
    if not gemini_client or not texts:
        return texts, False
    
    batch_input = "\n".join([f"[{i+1}] {t}" for i, t in enumerate(texts)])
    
    prompt = f"""
    Task: Translate the following text into {target_language}.
    
    Guidelines:
    1. Maintain the exact format [number] for each translated line.
    2. Output ONLY the numbered list. No extra explanations.
    3. Keep technical terms natural.
    
    Source Text:
    {batch_input}
    """

    safety_settings = [
        {"category": "HARM_CATEGORY_HARASSMENT", "threshold": "BLOCK_NONE"},
        {"category": "HARM_CATEGORY_HATE_SPEECH", "threshold": "BLOCK_NONE"},
        {"category": "HARM_CATEGORY_SEXUALLY_EXPLICIT", "threshold": "BLOCK_NONE"},
        {"category": "HARM_CATEGORY_DANGEROUS_CONTENT", "threshold": "BLOCK_NONE"},
    ]

    for attempt in range(1, 4):  # 最大3回リトライ
        try:
            if logger: logger.info(f"バッチ {batch_idx+1} 通信開始 (試行 {attempt}/3)")
            
            # 旧: gemini_model.generate_content(prompt, generation_config=...,
            #         safety_settings=..., request_options={"timeout": 40})
            # 新: 共通モジュール(gemini_client.py)経由の互換シムで同じ内容を送る。
            # request_options={"timeout": 40} は共通モジュールに同等機能が無いため削除した。
            # タイムアウトは gemini_client.py 側の固定値（直接15秒 / プロキシ60秒）になる。
            # 遮断下では最初のバッチだけ直接呼び出しの15秒を待つぶん遅くなるが、一度失敗
            # すると以降はプロキシ直行になるため2バッチ目以降は影響しない（仕様どおり）。
            response = gemini_client.models.generate_content(
                model=GEMINI_MODEL_NAME,
                contents=prompt,
                config=_GeminiGenerateConfig(temperature=0.1, safety_settings=safety_settings),
            )
            time.sleep(1.5)

            if not response.parts:
                if logger: logger.warning(f"バッチ {batch_idx+1} 空のレスポンスを受信")
                raise ValueError("Empty response from API")

            response_text = response.text.strip()
            results = [None] * len(texts)
            lines = response_text.split('\n')
            
            for line in lines:
                match = re.match(r'^\[(\d+)\]\s*(.*)', line.strip())
                if match:
                    idx = int(match.group(1)) - 1
                    if 0 <= idx < len(texts):
                        results[idx] = match.group(2).strip()
            
            for i in range(len(results)):
                if results[i] is None:
                    results[i] = texts[i]
                    
            if logger: logger.info(f"バッチ {batch_idx+1} 成功！")
            return results, False  # 成功（エラーフラグFalse）

        except Exception as e:
            if logger: logger.error(f"バッチ {batch_idx+1} エラー発生: {str(e)}")
            if attempt < 3:
                time.sleep(3)  # エラー時は3秒間隔で待機
            else:
                if logger: logger.error(f"バッチ {batch_idx+1} は3回失敗したためスキップします。")
    
    return texts, True  # 失敗（原文を返し、エラーフラグTrueを通知）

def translate_super_fast_parallel(all_texts, target_language="Japanese", max_workers=3, progress_callback=None, logger=None):
    """並列処理エンジン（コールバックと強制切断機能付き）"""
    if not all_texts:
        return []
    
    batch_size = 10
    chunks = [all_texts[i:i + batch_size] for i in range(0, len(all_texts), batch_size)]
    results = [None] * len(chunks)
    
    abort_event = threading.Event()
    consecutive_errors = 0
    processed_items = 0
    
    def translate_chunk(chunk_idx, chunk_texts):
        if abort_event.is_set():
            return chunk_idx, (chunk_texts, False)  # 中断フラグが立っていればスルー
        return chunk_idx, translate_batch_gemini(chunk_texts, target_language, chunk_idx, logger)

    with ThreadPoolExecutor(max_workers=max_workers) as executor:
        futures = [executor.submit(translate_chunk, i, chunk) for i, chunk in enumerate(chunks)]
        
        for future in as_completed(futures):
            try:
                chunk_idx, (translated_chunk, is_error) = future.result()
                results[chunk_idx] = translated_chunk
                
                # エラーカウントの判定
                if is_error:
                    consecutive_errors += 1
                else:
                    consecutive_errors = 0  # 1つでも成功すればリセット
                    
                # 3回連続エラーで即時撤退（フェイルファスト）
                if consecutive_errors >= 3:
                    abort_event.set()
                    if logger: logger.critical("【致命的エラー】3バッチ連続で通信エラー発生。処理を強制中断します。")
                    raise RuntimeError("Gemini APIへの通信が3回連続で失敗しました。\nネットワーク接続かAPI制限をご確認ください。")
                
                # 進捗UIの更新
                processed_items += len(chunks[chunk_idx])
                if progress_callback:
                    progress_callback(processed_items)
                    
            except RuntimeError as e:
                raise e  # 致命的エラーはそのまま投げる
            except Exception as e:
                if logger: logger.error(f"チャンク結果取得エラー: {str(e)}")
    
    final_results = []
    for chunk_result in results:
        if chunk_result:
            final_results.extend(chunk_result)
        else:
            final_results.extend([""] * batch_size)
            
    return final_results

class WordProgressWindow:
    """進捗表示用ウィンドウ（スレッドセーフ版）"""
    def __init__(self, parent):
        self.window = tk.Toplevel(parent)
        self.window.title("Gemini 翻訳進捗")
        self.window.geometry("450x180")
        self.window.resizable(False, False)
        
        try:
            self.window.transient(parent)
            self.window.grab_set()
        except:
            pass
        
        self.progress_label = tk.Label(self.window, text="Gemini AI 翻訳を準備中...", font=("Arial", 11, "bold"))
        self.progress_label.pack(pady=15)
        
        self.status_label = tk.Label(self.window, text="処理を開始します...", font=("Arial", 9))
        self.status_label.pack(pady=5)
        
        self.progress_frame = tk.Frame(self.window, width=350, height=20, bg="white", relief="sunken")
        self.progress_frame.pack(pady=10)
        
        self.progress_bar = tk.Frame(self.progress_frame, height=18, bg="#0078D4")
        self.progress_bar.place(x=1, y=1)
        
        self.time_label = tk.Label(self.window, text="", font=("Arial", 8), fg="blue")
        self.time_label.pack(pady=2)
        
        self.start_time = time.time()
        
    def update_progress(self, current, total, status=""):
        # 別スレッドから安全にGUIを更新するため after を使用
        try:
            self.window.after(0, self._update_gui, current, total, status)
        except Exception:
            pass
            
    def _update_gui(self, current, total, status):
        try:
            percentage = int((current / total) * 100) if total > 0 else 0
            bar_width = int((current / total) * 348) if total > 0 else 0
            self.progress_bar.config(width=bar_width)
            
            elapsed_time = time.time() - self.start_time
            
            self.progress_label.config(text=f"翻訳進捗: {current}/{total} ({percentage}%)")
            if status:
                self.status_label.config(text=status)
            self.time_label.config(text=f"経過時間: {elapsed_time:.1f}s")
        except:
            pass
        
    def close(self):
        try:
            self.window.after(0, self.window.destroy)
        except:
            pass

LANGUAGE_SUFFIX_MAP = {
    "Japanese": "ja",
    "English": "en",
    "Chinese Simplified": "cn",
    "Korean": "ko",
}


def lang_to_suffix(target_language):
    """出力ファイル名の末尾に付ける2文字言語コードを返す（例: Japanese -> ja）。

    20260812_01 までは `target_language.split()[0].lower()` をそのまま使っていたため
    `_gemini_japanese.pptx` のように言語名がまるごと入っていた。ファイル名が長くなり
    他ツール（pdf_translator は `_ja.pdf`）とも揃っていなかったため、2文字コードへ
    短縮した。未知の言語が来た場合は先頭2文字を小文字にして使う（pdf_translator の
    lang_to_suffix と同じ考え方）。
    """
    if target_language in LANGUAGE_SUFFIX_MAP:
        return LANGUAGE_SUFFIX_MAP[target_language]
    return target_language.strip()[:2].lower()


# ============================================================
# 画像だけのスライドの翻訳（20260930_01〜）
# ============================================================
# 文字データが1つも見つからなかった場合に限り、全面に画像が貼られているだけの
# スライドの画像を Gemini へ送り、画像内の文字と訳文を読み取ってメモ欄へ追記する。
# 文字ありスライドの経路（従来の翻訳）には一切影響しない。

IMAGE_COVERAGE_MIN = 0.70        # 画面のこの割合以上を覆う画像を「全面画像」とみなす
IMAGE_MAX_WIDTH = 1600           # 送信前に縮小する幅(px)。小さい文字が読める範囲
IMAGE_JPEG_QUALITY = 85
IMAGE_WORKERS = 2                # 画像は1回の応答が重いので既存(3)より控えめ
IMAGE_MAX_OUTPUT_TOKENS = 16384
IMAGE_NOTE_HEADER = "【画像内テキストの翻訳】"
IMAGE_NOTE_SEPARATOR = "----------"
# 画像のスライド確認ダイアログに追記する文言（重ね貼り版で上書きされる）
IMAGE_CONFIRM_EXTRA = ""

IMAGE_SAFETY_SETTINGS = [
    {"category": "HARM_CATEGORY_HARASSMENT", "threshold": "BLOCK_NONE"},
    {"category": "HARM_CATEGORY_HATE_SPEECH", "threshold": "BLOCK_NONE"},
    {"category": "HARM_CATEGORY_SEXUALLY_EXPLICIT", "threshold": "BLOCK_NONE"},
    {"category": "HARM_CATEGORY_DANGEROUS_CONTENT", "threshold": "BLOCK_NONE"},
]

# 項目の順序は box_2d → 原文 → 訳文 → 行数（位置を先に確定させてから訳させる）
IMAGE_RESPONSE_SCHEMA = {
    "type": "ARRAY",
    "items": {
        "type": "OBJECT",
        "properties": {
            "box_2d": {"type": "ARRAY", "items": {"type": "INTEGER"}},
            "text": {"type": "STRING"},
            "translation": {"type": "STRING"},
            "lines": {"type": "INTEGER"},
        },
        "required": ["box_2d", "text", "translation", "lines"],
        "propertyOrdering": ["box_2d", "text", "translation", "lines"],
    },
}


def parse_slide_spec(spec_str, total_slides):
    """スライド指定文字列（例: "1-3,5"、1始まり）を0始まりの番号集合に変換する。
    空文字列・空白のみの場合は None（全スライド対象）を返す。不正な指定はValueErrorを送出する。
    （pdf_translator の parse_page_spec と同じ書式）"""
    spec_str = (spec_str or "").strip()
    if not spec_str:
        return None

    slides = set()
    for token in spec_str.split(","):
        token = token.strip()
        if not token:
            continue
        if "-" in token:
            parts = token.split("-")
            if len(parts) != 2 or not parts[0].strip().isdigit() or not parts[1].strip().isdigit():
                raise ValueError(f"スライド指定の形式が正しくありません: '{token}'")
            start, end = int(parts[0].strip()), int(parts[1].strip())
            if start > end:
                raise ValueError(f"スライドの範囲が逆になっています: '{token}'")
            slides.update(range(start, end + 1))
        else:
            if not token.isdigit():
                raise ValueError(f"スライド指定の形式が正しくありません: '{token}'")
            slides.add(int(token))

    if not slides:
        raise ValueError("スライドが1つも指定されていません。")

    for s in slides:
        if s < 1 or s > total_slides:
            raise ValueError(f"スライド番号 {s} はこの資料（全{total_slides}枚）の範囲外です。")

    return {s - 1 for s in slides}


def _shape_is_picture(shape):
    try:
        return shape.shape_type == MSO_SHAPE_TYPE.PICTURE
    except Exception:
        return False  # 種類を判定できない図形は画像として扱わない


def _slide_has_translatable_text(slide):
    """図形・表に、翻訳対象の文字データ(run)があるか（スピーカーノートは見ない）。
    ページ番号だけのテキストボックスは is_translatable が False なので「文字なし」扱いになる。"""
    for shape in slide.shapes:
        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    if is_translatable(run.text):
                        return True
        if shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    if cell.text_frame:
                        for para in cell.text_frame.paragraphs:
                            for run in para.runs:
                                if is_translatable(run.text):
                                    return True
    return False


def _picture_coverage(pic, prs):
    """画像がスライド面積に占める割合（スライドの外にはみ出した分は数えない）"""
    sw, sh = prs.slide_width, prs.slide_height
    if not sw or not sh:
        return 0.0
    left, top = pic.left or 0, pic.top or 0
    width, height = pic.width or 0, pic.height or 0
    ix = max(0, min(left + width, sw) - max(left, 0))
    iy = max(0, min(top + height, sh) - max(top, 0))
    return (ix * iy) / float(sw * sh)


def find_image_slides(prs):
    """画像だけのスライド（対象）を [(スライド番号(0始まり), 画像図形), ...] で返す。
    条件: 翻訳対象の文字データが無い / 画面の70%以上を覆う画像がちょうど1枚 / 回転していない。"""
    targets = []
    for idx, slide in enumerate(prs.slides):
        if _slide_has_translatable_text(slide):
            continue
        big = [sh for sh in slide.shapes
               if _shape_is_picture(sh) and _picture_coverage(sh, prs) >= IMAGE_COVERAGE_MIN]
        if len(big) != 1:
            continue
        try:
            if big[0].rotation:
                continue
        except Exception:
            pass
        targets.append((idx, big[0]))
    return targets


def _prepare_slide_image(blob, crop, max_width=IMAGE_MAX_WIDTH):
    """画像を（トリミングを反映して）縮小し、(JPEGのバイト列, RGBのPIL画像) を返す。
    トリミング後の画像を送ることで、Gemini が返す座標(0-1000)がそのまま画像図形の
    表示枠に対する割合になり、座標変換が単純になる。"""
    from PIL import Image
    im = Image.open(io.BytesIO(blob))
    im.load()
    if im.mode in ("RGBA", "LA") or (im.mode == "P" and "transparency" in im.info):
        rgba = im.convert("RGBA")
        flat = Image.new("RGB", rgba.size, (255, 255, 255))
        flat.paste(rgba, mask=rgba.split()[3])
        im = flat
    else:
        im = im.convert("RGB")

    left, top, right, bottom = [max(0.0, float(c or 0.0)) for c in crop]
    w, h = im.size
    box = (round(left * w), round(top * h), w - round(right * w), h - round(bottom * h))
    if box != (0, 0, w, h) and box[2] > box[0] and box[3] > box[1]:
        im = im.crop(box)

    if im.width > max_width:
        new_h = max(1, round(im.height * max_width / float(im.width)))
        im = im.resize((max_width, new_h), Image.LANCZOS)

    buf = io.BytesIO()
    im.save(buf, "JPEG", quality=IMAGE_JPEG_QUALITY)
    return buf.getvalue(), im


def _build_image_prompt(target_language):
    return f"""You are given one slide rendered as an image. Read every piece of text in it and translate it into {target_language}.

Return a JSON array. Each element is one text block (a heading, a table cell, a bullet line, a short label, or a paragraph) with:
- box_2d: [ymin, xmin, ymax, xmax] of the block, integers from 0 to 1000 (fractions of the image height and width, origin at the top-left). Make the box tight around the text.
- text: the original text exactly as written.
- translation: the translation into {target_language}.
- lines: the number of text lines the original block occupies in the image.

Rules:
1. List the blocks in natural reading order.
2. Do NOT return blocks that contain only numbers, percentages, symbols or punctuation (for example "80%", "( n=5 )", "1").
3. Do NOT return logos, icons, or text that is part of a picture or graphic (for example a company logo).
4. Keep proper nouns, company names and product names untranslated.
5. Use consistent terminology for terms that repeat.
6. Translate faithfully. Do not add, omit or summarize anything. Keep numbers that appear inside a sentence unchanged.
7. If the image contains no text, return an empty array [].
"""


def _validate_image_blocks(raw_text):
    """Gemini の JSON 応答を検証し、使える文字ブロックのリストを返す。
    形式が不正なら ValueError（呼び出し側が再試行する）。座標が不正なブロックは捨てる。
    元の応答に1つ以上の要素があるのに使えるブロックが0個のときも ValueError にする
    （全ブロックの位置が0で返る等の壊れた応答を「文字なし」と誤認しないため）。"""
    data = json.loads(raw_text)
    if not isinstance(data, list):
        raise ValueError("応答のJSONが配列ではありません")

    blocks = []
    for item in data:
        if not isinstance(item, dict):
            continue
        box = item.get("box_2d")
        text = item.get("text")
        translation = item.get("translation")
        if not (isinstance(box, (list, tuple)) and len(box) == 4
                and all(isinstance(v, (int, float)) and not isinstance(v, bool) for v in box)):
            continue
        ymin, xmin, ymax, xmax = [max(0.0, min(1000.0, float(v))) for v in box]
        if ymax <= ymin or xmax <= xmin:
            continue
        if not isinstance(text, str) or not isinstance(translation, str) or not text.strip():
            continue
        lines = item.get("lines")
        if not (isinstance(lines, (int, float)) and not isinstance(lines, bool) and 1 <= lines <= 20):
            lines = 1
        blocks.append({"box_2d": [ymin, xmin, ymax, xmax], "text": text,
                       "translation": translation, "lines": int(lines)})

    if data and not blocks:
        raise ValueError("使える文字ブロックがありませんでした（座標や形式が不正）")
    return blocks


def _has_letters(text):
    """文字（英字・かな・漢字など）を2つ以上含むか。数字・%・記号だけのブロックを除く専用の判定。
    既存の is_translatable は "80%" や "( n=5 )" を True にしてしまうため、画像には使わない。"""
    return sum(1 for c in text if c.isalpha()) >= 2


def _translate_slide_image(jpeg_bytes, target_language, slide_no, logger=None, should_stop=None):
    """1スライド分の画像を Gemini へ送り、(ブロックのリスト, エラー内容) を返す。
    成功時は (blocks, "")。blocks が空リストなら「文字が見つからなかった」。
    3回リトライしても失敗したら (None, エラーの種類)。ログには原文を書かない。
    should_stop() が True を返したら、それ以上リトライせず (None, "中断") で戻る
    （3枚連続失敗で処理を中断したあと、実行中のスライドが無駄に待ち続けないため）。"""
    if not gemini_client:
        return None, "Geminiが初期化されていません"

    contents = [
        {"inlineData": {"mimeType": "image/jpeg",
                        "data": base64.b64encode(jpeg_bytes).decode("ascii")}},
        _build_image_prompt(target_language),
    ]
    # thinking はモデルによっては無効化できない（gemini-2.5-pro 等）ため、2.5-flash 系のときだけ0にする。
    # 有効なままだと出力トークン枠を消費し、JSONが途中で切れやすい。
    thinking = 0 if str(GEMINI_MODEL_NAME).startswith("gemini-2.5-flash") else None

    last_error = ""
    for attempt in range(1, 4):
        if should_stop is not None and should_stop():
            return None, "中断"
        try:
            if logger: logger.info(f"スライド {slide_no} 通信開始 (試行 {attempt}/3)")
            response = gemini_client.models.generate_content(
                model=GEMINI_MODEL_NAME,
                contents=contents,
                config=_GeminiGenerateConfig(
                    temperature=0.1,
                    safety_settings=IMAGE_SAFETY_SETTINGS,
                    response_mime_type="application/json",
                    response_schema=IMAGE_RESPONSE_SCHEMA,
                    max_output_tokens=IMAGE_MAX_OUTPUT_TOKENS,
                    thinking_budget=thinking),
            )
            time.sleep(1.5)

            if not response.parts:
                if logger: logger.warning(f"スライド {slide_no} 空のレスポンスを受信")
                raise ValueError("Empty response from API")

            blocks = _validate_image_blocks(response.text)
            if logger: logger.info(f"スライド {slide_no} 成功（文字ブロック {len(blocks)} 個）")
            return blocks, ""
        except Exception as e:
            last_error = type(e).__name__
            if logger: logger.error(f"スライド {slide_no} エラー発生: {str(e)}")
            if attempt < 3:
                time.sleep(3)
            else:
                if logger: logger.error(f"スライド {slide_no} は3回失敗したためスキップします。")
    return None, last_error


_XML_INVALID_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _clean_note_text(s):
    """メモ欄に書ける形にする（XMLで使えない制御文字を除き、改行は空白にする）"""
    return _XML_INVALID_CHARS.sub("", s).replace("\r", " ").replace("\n", " ").strip()


def _image_note_lines(blocks):
    """メモ欄に書く行（訳文と原文）。文字を含まないブロック・訳文が空/原文と同じブロックは除く。"""
    lines = []
    for b in blocks:
        translation = _clean_note_text(b["translation"])
        original = _clean_note_text(b["text"])
        if not _has_letters(original) or not translation or translation == original:
            continue
        lines.append(f"{translation}（原文: {original}）")
    return lines


def _append_image_notes(slide, blocks):
    """メモ欄の末尾に、画像内テキストの翻訳を追記する。既存のメモは消さない。
    追記した行数を返す（追記しなかったら0）。"""
    lines = _image_note_lines(blocks)
    if not lines:
        return 0
    tf = slide.notes_slide.notes_text_frame
    if tf is None:
        return 0
    existing = tf.text.strip()
    if IMAGE_NOTE_HEADER in existing:
        return 0  # 既に追記済みなら二重に書かない

    entries = ([IMAGE_NOTE_SEPARATOR] if existing else []) + [IMAGE_NOTE_HEADER] + lines
    for i, entry in enumerate(entries):
        para = tf.paragraphs[0] if (i == 0 and not existing) else tf.add_paragraph()
        para.text = entry
    return len(lines)


def _run_on_main_thread(widget, func):
    """tkinter の操作は主スレッドで行う必要があるため、ワーカースレッドから
    主スレッドに処理を依頼し、結果が出るまで待って返す。"""
    done = threading.Event()
    holder = {}

    def _call():
        try:
            holder["result"] = func()
        except Exception as e:
            holder["error"] = e
        finally:
            done.set()

    widget.after(0, _call)
    done.wait()
    if "error" in holder:
        raise holder["error"]
    return holder.get("result")


def _show_image_confirm_dialog(parent, message, validate):
    """画像スライドの翻訳を実行するかの確認ダイアログ。処理するスライドの範囲を入力できる。
    実行なら入力文字列（空欄=全部）、キャンセルなら None を返す。"""
    result = {"spec": None}
    win = tk.Toplevel(parent)
    win.title("画像スライドの翻訳")
    win.resizable(False, False)
    try:
        win.transient(parent)
        win.grab_set()
    except Exception:
        pass

    tk.Label(win, text=message, justify="left", wraplength=480,
             font=("Arial", 10)).pack(padx=20, pady=(16, 8))

    row = tk.Frame(win)
    row.pack(padx=20, pady=4)
    tk.Label(row, text="処理するスライド（空欄=全部。例: 6 / 1-3,5）:",
             font=("Arial", 10)).pack(side=tk.LEFT)
    var = tk.StringVar(win)
    entry = tk.Entry(row, textvariable=var, width=14)
    entry.pack(side=tk.LEFT, padx=6)

    err = tk.Label(win, text="", fg="red", font=("Arial", 9))
    err.pack()

    def on_ok():
        try:
            validate(var.get())
        except ValueError as e:
            err.config(text=str(e))
            return
        result["spec"] = var.get()
        win.destroy()

    buttons = tk.Frame(win)
    buttons.pack(pady=(6, 16))
    tk.Button(buttons, text="実行", command=on_ok, bg="#0078D4", fg="white",
              padx=20, pady=6, font=("Arial", 10, "bold")).pack(side=tk.LEFT, padx=8)
    tk.Button(buttons, text="キャンセル", command=win.destroy,
              padx=20, pady=6, font=("Arial", 10)).pack(side=tk.LEFT, padx=8)

    win.protocol("WM_DELETE_WINDOW", win.destroy)
    entry.focus_set()
    win.wait_window()
    return result["spec"]


def _ask_image_confirmation(progress_window, message, validate):
    """確認ダイアログを主スレッドで出して結果を返す。テストではここを差し替える。"""
    parent = progress_window.window
    return _run_on_main_thread(parent, lambda: _show_image_confirm_dialog(parent, message, validate))


def _remove_quietly(path):
    try:
        os.remove(path)
    except OSError:
        pass


def _translate_image_slides(prs, targets, output_path, target_language, progress_window,
                            logger, start_total_time):
    """画像だけのスライドを翻訳し、メモ欄へ追記して保存する。"""
    total_slides = len(prs.slides)
    message = (f"画像だけのスライドが {len(targets)} 枚見つかりました（文字データがありません）。\n"
               "画像内の文字を読み取って翻訳します。\n\n"
               "・スライド画像そのものを外部AI（Gemini）へ送信します。\n"
               "・機密・個人情報を含む資料では実行しないでください。\n"
               "・訳文は各スライド下のメモ欄（ノート）に追記されます。\n"
               + IMAGE_CONFIRM_EXTRA +
               "・元のファイルは変更しません（新しいファイルに保存します）。")

    spec = _ask_image_confirmation(progress_window, message,
                                   lambda s: parse_slide_spec(s, total_slides))
    if spec is None:
        logger.info("画像スライドの翻訳: ユーザーがキャンセルしました")
        progress_window.close()
        _remove_quietly(output_path)  # 先に作った未翻訳のコピーを残さない
        messagebox.showinfo("キャンセル", "画像スライドの翻訳をキャンセルしました。")
        return

    selected = parse_slide_spec(spec, total_slides)
    work = [(i, p) for i, p in targets if selected is None or i in selected]
    if not work:
        progress_window.close()
        _remove_quietly(output_path)
        messagebox.showinfo("完了", "処理するスライドがありませんでした。\n"
                            "（指定した範囲に、画像だけのスライドがありません）")
        return

    # python-pptx の読み取りは主のワーカースレッドで済ませ、通信だけを並列に行う
    jobs = []
    for idx, pic in work:
        try:
            crop = (pic.crop_left, pic.crop_top, pic.crop_right, pic.crop_bottom)
            jobs.append((idx, pic.image.blob, crop))
        except Exception as e:
            logger.error(f"スライド {idx + 1} の画像を読み込めません: {type(e).__name__}")

    total_jobs = len(jobs)
    progress_window.update_progress(0, total_jobs, f"画像内の文字を翻訳中... (0/{total_jobs}枚)")
    logger.info(f"画像スライドの翻訳を開始: 対象{total_jobs}枚")

    abort_event = threading.Event()

    def run_job(idx, blob, crop):
        if abort_event.is_set():
            return idx, None, "中断"
        try:
            jpeg, _ = _prepare_slide_image(blob, crop)
        except Exception as e:
            return idx, None, type(e).__name__
        blocks, err = _translate_slide_image(jpeg, target_language, idx + 1, logger, abort_event.is_set)
        return idx, blocks, err

    results = {}
    failed = []
    consecutive_errors = 0
    aborted = False
    processed = 0
    with ThreadPoolExecutor(max_workers=IMAGE_WORKERS) as executor:
        futures = [executor.submit(run_job, *job) for job in jobs]
        for future in as_completed(futures):
            try:
                idx, blocks, err = future.result()
            except Exception as e:
                logger.error(f"スライド結果の取得エラー: {type(e).__name__}")
                continue
            if err == "中断":
                continue
            processed += 1
            if blocks is None:
                failed.append(idx + 1)
                consecutive_errors += 1
            else:
                results[idx] = blocks
                consecutive_errors = 0
            progress_window.update_progress(
                processed, total_jobs, f"画像内の文字を翻訳中... ({processed}/{total_jobs}枚)")
            if consecutive_errors >= 3 and not aborted:
                aborted = True
                abort_event.set()
                logger.critical("【致命的エラー】3枚連続で通信エラー発生。処理を中断します。")
                for f in futures:
                    f.cancel()

    progress_window.update_progress(total_jobs, total_jobs, "翻訳結果をメモ欄に書き込み中...")
    extra = _apply_image_results(prs, results, work, target_language, logger)

    progress_window.update_progress(total_jobs, total_jobs, "保存中...")
    try:
        prs.save(output_path)
    except PermissionError:
        progress_window.close()
        messagebox.showerror("保存エラー", "ファイルが他のプログラム（PowerPointなど）で開かれています。\n閉じてから再度実行してください。")
        return

    progress_window.close()
    total_time = time.time() - start_total_time
    logger.info(f"=== 画像スライドの処理完了 ({total_time:.1f}秒) 翻訳{extra['translated']}枚 "
                f"文字なし{extra['no_text']}枚 失敗{len(failed)}枚 ===")

    lines = ["画像スライドの翻訳が完了しました！",
             f"保存先: {output_path}",
             f"翻訳したスライド: {extra['translated']}枚"]
    if extra["no_text"]:
        lines.append(f"文字が見つからなかったスライド: {extra['no_text']}枚")
    lines.extend(extra["detail_lines"])
    if failed:
        lines.append(f"失敗したスライド: {len(failed)}枚（番号: {', '.join(str(n) for n in sorted(failed))}）")
        lines.append("  → もう一度実行し、範囲にその番号を入力すると再試行できます。")
    if aborted:
        lines.append("※ 3枚連続で通信に失敗したため、途中で中断しました。")
    not_target = total_slides - len(targets)
    if not_target > 0:
        lines.append(f"対象外のスライド: {not_target}枚")
    lines.append(f"処理時間: {total_time:.1f}秒")
    lines.append("")
    lines.append("※ AIによる翻訳です。内容は必ず確認してください。")
    messagebox.showinfo("完了", "\n".join(lines))


def _apply_image_results(prs, results, work, target_language, logger):
    """翻訳結果を書き込む（ノート版: メモ欄への追記のみ）。
    戻り値: {"translated": 翻訳したスライド数, "no_text": 文字が無かったスライド数,
             "detail_lines": 完了メッセージへ足す行のリスト}"""
    translated = 0
    no_text = 0
    for idx, _pic in work:
        if idx not in results:
            continue
        blocks = results[idx]
        if not blocks:
            no_text += 1
            continue
        if _append_image_notes(prs.slides[idx], blocks) > 0:
            translated += 1
        else:
            no_text += 1
    return {"translated": translated, "no_text": no_text, "detail_lines": []}


def translate_ppt_document_thread(file_path, target_language, progress_window):
    """バックグラウンドで実行されるメイン処理"""
    logger = get_logger()
    try:
        start_total_time = time.time()
        lang_suffix = lang_to_suffix(target_language)
        output_path = os.path.splitext(file_path)[0] + f"_{lang_suffix}.pptx"
        
        logger.info(f"=== PPT翻訳開始: {os.path.basename(file_path)} ===")
        
        # WinError 32 (ファイルロック) の事前チェック
        try:
            with open(file_path, 'a'): pass
        except PermissionError:
            logger.error(f"[WinError 32事前検知] 読み込み元ファイルがロックされています: {file_path}")
            progress_window.close()
            messagebox.showerror("ファイルエラー", "対象のPowerPointファイルが別のアプリで開かれています。\nファイルを閉じてから再度実行してください。")
            return
            
        if os.path.exists(output_path):
            try:
                with open(output_path, 'a'): pass
            except PermissionError:
                logger.error(f"[WinError 32事前検知] 保存先ファイルがロックされています: {output_path}")
                progress_window.close()
                messagebox.showerror("ファイルエラー", "以前に作成した翻訳ファイルが開かれています。\nファイルを閉じてから再度実行してください。")
                return
        
        shutil.copy2(file_path, output_path)
        prs = Presentation(output_path)
        
        translatable_items = [] 
        
        # PPTのスライドごとの処理
        for slide in prs.slides:
            # 1. 通常の図形とテーブル
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        for run in para.runs:
                            if is_translatable(run.text):
                                translatable_items.append(run)
                if shape.has_table:
                    for row in shape.table.rows:
                        for cell in row.cells:
                            if cell.text_frame:
                                for para in cell.text_frame.paragraphs:
                                    for run in para.runs:
                                        if is_translatable(run.text):
                                            translatable_items.append(run)
            
            # 2. スピーカーノート
            if slide.has_notes_slide:
                for shape in slide.notes_slide.shapes:
                    if shape.has_text_frame:
                        for para in shape.text_frame.paragraphs:
                            for run in para.runs:
                                if is_translatable(run.text):
                                    translatable_items.append(run)

        if not translatable_items:
            # 文字が1つも見つからなかったときだけ、画像だけのスライドの翻訳へ進む
            # （文字ありの資料の動きは変えない）。
            image_targets = find_image_slides(prs)
            if image_targets:
                _translate_image_slides(prs, image_targets, output_path, target_language,
                                        progress_window, logger, start_total_time)
                return
            progress_window.close()
            hint = ""
            if any(_shape_is_picture(sh) for s in prs.slides for sh in s.shapes):
                hint = ("\n（画像内の文字の翻訳は、画面の7割以上を覆う画像が1枚だけのスライドが対象です。"
                        "\n　そのようなスライドは見つかりませんでした）")
            messagebox.showinfo("完了", "翻訳対象のテキストが見つかりませんでした。" + hint)
            return

        texts_only = [run.text for run in translatable_items]
        total_items = len(texts_only)
        progress_window.update_progress(0, total_items, f"Gemini APIで並列翻訳中... (0/{total_items}項目)")
        
        # UI更新用のコールバック関数
        def update_ui_callback(processed_count):
            progress_window.update_progress(processed_count, total_items, f"Gemini APIで並列翻訳中... ({processed_count}/{total_items}項目)")
        
        # ※バックグラウンドスレッドで重い通信処理を実行
        translated_texts = translate_super_fast_parallel(texts_only, target_language, max_workers=3, progress_callback=update_ui_callback, logger=logger)
        
        progress_window.update_progress(len(translatable_items), len(translatable_items), "翻訳結果をPowerPointに適用中...")
        for i, run in enumerate(translatable_items):
            if i < len(translated_texts) and translated_texts[i]:
                run.text = translated_texts[i]
                if "Japanese" in target_language or "日本" in target_language:
                    run.font.name = '游ゴシック'
        
        progress_window.update_progress(len(translatable_items), len(translatable_items), "保存中...")
        
        try:
            prs.save(output_path)
        except PermissionError:
            progress_window.close()
            messagebox.showerror("保存エラー", "ファイルが他のプログラム（PowerPointなど）で開かれています。\n閉じてから再度実行してください。")
            return

        progress_window.close()
        total_time = time.time() - start_total_time
        logger.info(f"=== 処理完了: 成功 ({total_time:.1f}秒) ===")
        
        image_only_note = ""
        image_only_count = len(find_image_slides(prs))
        if image_only_count:
            image_only_note = (f"\n\n※ 画像だけのスライドが {image_only_count} 枚あります（未翻訳）。"
                               f"\n　文字のあるスライドと混ざった資料では、画像内の文字は翻訳されません。")
        messagebox.showinfo("完了", 
                          f"書式保持翻訳完了！\n"
                          f"保存先: {output_path}\n"
                          f"翻訳項目数: {len(translatable_items)}\n"
                          f"処理時間: {total_time:.1f}秒" + image_only_note)
                          
    except RuntimeError as e:
        progress_window.close()
        messagebox.showerror("通信エラー強制終了", str(e))
        
    except Exception as e:
        logger.error(f"予期せぬエラー: {traceback.format_exc()}")
        progress_window.close()
        messagebox.showerror("エラー", f"翻訳処理中にエラーが発生しました:\n{str(e)}")

def select_file():
    path = filedialog.askopenfilename(
        title="翻訳するPowerPointファイルを選択してください",
        filetypes=[("PowerPoint files", "*.pptx")]
    )
    
    if not path:
        return
    
    lang_win = tk.Toplevel(root)
    lang_win.title("PPT翻訳設定")
    lang_win.geometry("400x250")
    lang_win.resizable(False, False)
    lang_win.transient(root)
    lang_win.grab_set()
    
    tk.Label(lang_win, text="翻訳先言語を選択してください", font=("Arial", 12, "bold")).pack(padx=20, pady=20)
    
    languages = {
        "日本語 (Japanese)": "Japanese",
        "英語 (English)": "English",
        "中国語簡体字 (Chinese)": "Chinese Simplified"
    }
    
    lang_var = tk.StringVar(lang_win)
    lang_var.set("日本語 (Japanese)")
    
    lang_menu = tk.OptionMenu(lang_win, lang_var, *languages.keys())
    lang_menu.config(font=("Arial", 10), width=25)
    lang_menu.pack(padx=20, pady=10)
    
    def start_translation():
        selected_language = languages[lang_var.get()]
        lang_win.destroy()
        
        # プログレスウィンドウを作成
        progress_window = WordProgressWindow(root)
        
        # 画面をフリーズさせないために、別スレッドで翻訳処理を開始！
        thread = threading.Thread(
            target=translate_ppt_document_thread, 
            args=(path, selected_language, progress_window)
        )
        thread.daemon = True
        thread.start()
    
    button_frame = tk.Frame(lang_win)
    button_frame.pack(pady=25)
    
    tk.Button(button_frame, text="翻訳開始", command=start_translation, 
             bg="#0078D4", fg="white", padx=20, pady=8, font=("Arial", 11, "bold")).pack(side=tk.LEFT, padx=10)
    tk.Button(button_frame, text="キャンセル", command=lang_win.destroy, 
             padx=20, pady=8, font=("Arial", 11)).pack(side=tk.LEFT, padx=10)

# --- GUI初期設定 ---
if __name__ == "__main__":
    root = tk.Tk()
    root.withdraw()
    
    if not check_dependencies(root):
        sys.exit(1)
        
    if not init_gemini(root):
        sys.exit(1)
        
    root.deiconify()
    root.title("PowerPoint Gemini 翻訳ツール")
    root.geometry("500x280")
    root.resizable(False, False)
    
    main_frame = tk.Frame(root)
    main_frame.pack(expand=True, fill='both', padx=20, pady=20)
    
    title_label = tk.Label(main_frame, text="PowerPoint Gemini 翻訳ツール", font=("Arial", 16, "bold"))
    title_label.pack(pady=8)
    
    subtitle_label = tk.Label(main_frame, text="書式完全保持版 (Gemini API)", font=("Arial", 12), fg="#0078D4")
    subtitle_label.pack(pady=2)
    
    desc_label = tk.Label(main_frame, 
                         text="PowerPointファイル(.pptx)を選択して翻訳します\n"
                              "フォント、色、配置、テーブル、ノート書式を完全保持",
                         font=("Arial", 10))
    desc_label.pack(pady=8)
    
    select_button = tk.Button(main_frame, text="ファイル選択", command=select_file,
                             font=("Arial", 12), bg="#0078D4", fg="white", padx=20, pady=10)
    select_button.pack(pady=15)
    
    root.mainloop()