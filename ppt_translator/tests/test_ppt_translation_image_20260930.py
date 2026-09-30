"""ppt_translation の「画像だけのスライドの翻訳」(20260930_01 / 20260930_02) の検証テスト。

実際の Gemini・プロキシ・社内資料には一切依存しない。
- 偽の gemini_client を **対象モジュールのロード前に** sys.modules へ注入する。
- 「画像だけのスライド」は、このテストが PIL で描く合成スライド（表・帯・段落・ロゴ風の図形）。
  文字の位置（正解）が分かっているので、偽 Gemini はその正解を box_2d で返せる。
  ※ 実在の資料（社名入り等）はリポジトリに含めない。
- python-pptx は本物を使う。tkinter だけスタブ化する（GUIは別ファイル
  test_ppt_translation_image_gui_20260930.py で本物の tkinter を使って確認する）。

実行方法:
    pip install python-pptx
    python3 tests/test_ppt_translation_image_20260930.py            # 最新版(_02)を検証
    python3 tests/test_ppt_translation_image_20260930.py 01         # ノート版(_01)を検証
_01 には重ね貼りが無いので、重ね貼りのテストは自動的にスキップされる。
"""

import copy
import hashlib
import importlib.util
import io
import json
import logging
import os
import shutil
import sys
import tempfile
import types as pytypes

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.join(HERE, "..")
VERSION = sys.argv[1] if len(sys.argv) > 1 else "02"
TARGET = os.path.join(ROOT, f"ppt_translation_20260930_{VERSION}.py")
PREV_TARGET = os.path.join(ROOT, "ppt_translation_20260911_01.py")

RESULTS = []


def check(name, ok, detail=""):
    RESULTS.append((name, bool(ok), detail))
    print(f"{'PASS' if ok else 'FAIL'}: {name}" + (f"  ({detail})" if detail else ""))


def skip(name, why):
    print(f"SKIP: {name}  ({why})")


# ------------------------------------------------------------
# tkinter スタブ
# ------------------------------------------------------------
def _install_env_stubs():
    class _Anything:
        def __init__(self, *a, **k):
            pass

        def __call__(self, *a, **k):
            return _Anything()

        def __getattr__(self, _n):
            return _Anything()

    tk = pytypes.ModuleType("tkinter")
    for attr in ("Tk", "Toplevel", "Frame", "Label", "Button", "Entry", "Checkbutton",
                 "OptionMenu", "StringVar", "BooleanVar", "LEFT", "END"):
        setattr(tk, attr, _Anything())
    filedialog = pytypes.ModuleType("tkinter.filedialog")
    filedialog.askopenfilename = lambda *a, **k: ""
    messagebox = pytypes.ModuleType("tkinter.messagebox")
    messagebox.CALLS = []

    def _record(kind):
        def _fn(title="", message="", *a, **k):
            messagebox.CALLS.append((kind, title, message))
            return True
        return _fn

    messagebox.showerror = _record("error")
    messagebox.showwarning = _record("warning")
    messagebox.showinfo = _record("info")
    tk.filedialog = filedialog
    tk.messagebox = messagebox
    sys.modules["tkinter"] = tk
    sys.modules["tkinter.filedialog"] = filedialog
    sys.modules["tkinter.messagebox"] = messagebox
    return messagebox


MESSAGEBOX = _install_env_stubs()

from PIL import Image, ImageDraw, ImageFont  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.util import Emu, Pt  # noqa: E402

# ------------------------------------------------------------
# 偽 gemini_client
# ------------------------------------------------------------
CAPTURED = {"calls": []}
RESPONDER = {"fn": None}


def _fake_generate_advanced(payload, model=None, **kwargs):
    CAPTURED["calls"].append({"payload": payload, "model": model})
    return RESPONDER["fn"](payload)


def _load(module_name, target=TARGET):
    sys.modules.pop("gemini_client", None)
    fake = pytypes.ModuleType("gemini_client")
    fake.generate_advanced = _fake_generate_advanced
    sys.modules["gemini_client"] = fake
    spec = importlib.util.spec_from_file_location(module_name, target)
    mod = importlib.util.module_from_spec(spec)
    sys.modules[module_name] = mod
    spec.loader.exec_module(mod)
    stub = pytypes.ModuleType("time")
    stub.sleep = lambda _s: None
    stub.time = __import__("time").time
    mod.time = stub
    mod.gemini_client = mod._CommonGeminiClient()
    return mod


mod = _load("target_main")
HAS_OVERLAY = hasattr(mod, "_numbers_preserved")


class FakeProgress:
    class _W:
        def after(self, _ms, fn, *a):
            fn(*a)

    def __init__(self):
        self.window = self._W()
        self.updates = []
        self.closed = 0

    def update_progress(self, cur, total, status=""):
        self.updates.append((cur, total, status))

    def close(self):
        self.closed += 1


# ------------------------------------------------------------
# 合成スライド（正解の位置つき）
# ------------------------------------------------------------
def _font(size):
    for p in ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
              "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"):
        if os.path.exists(p):
            return ImageFont.truetype(p, size)
    return ImageFont.load_default(size)


TR = {  # 訳文（日本語）
    "Executive Summary": "エグゼクティブサマリー",
    "Employee Engagement": "従業員エンゲージメント",
    "The Engagement Index is comprised of 4 questions.": "エンゲージメント指数は4つの質問で構成されています。",
    "Wellbeing": "ウェルビーイング",
    "Teamwork and Collaboration": "チームワークと協働",
    "Performance Management": "パフォーマンスマネジメント",
    "My Group": "自分のグループ",
    "Nexperia": "Nexperia",
}


def _tr(text):
    return TR.get(text, "訳文:" + text)


def make_slide_image(size=(1920, 1080), with_junk_left=False, marker=0):
    """合成スライドと、正解のブロック(px座標)を返す。
    blocks: 翻訳対象の文字。numbers: 数字だけのブロック（触ってはいけない）。"""
    W, H = size
    im = Image.new("RGB", size, (255, 255, 255))
    d = ImageDraw.Draw(im)
    blocks, numbers = [], []

    def put(text, xy, font, fill, store, anchor_center_x=None):
        x, y = xy
        if anchor_center_x is not None:
            w = d.textlength(text, font=font)
            x = anchor_center_x - w / 2
        d.text((x, y), text, font=font, fill=fill)
        x0, y0, x1, y1 = d.textbbox((x, y), text, font=font)
        store.append({"text": text, "box": [x0, y0, x1, y1]})

    # 濃い帯(ヘッダー) — 白文字。帯の色は文字の色と大きく違う
    d.rectangle([0, 0, W, int(H * 0.09)], fill=(0, 112, 122))
    put("Executive Summary", (int(W * 0.02), int(H * 0.02)), _font(int(H * 0.05)), (255, 255, 255), blocks)

    # 見出し(青) と 段落(2行)
    put("Employee Engagement", (int(W * 0.03), int(H * 0.14)), _font(int(H * 0.028)), (20, 60, 160), blocks)
    put("The Engagement Index is comprised of 4 questions.",
        (int(W * 0.03), int(H * 0.20)), _font(int(H * 0.024)), (30, 30, 30), blocks)

    # 縞々の表: 行ごとに背景色が違う。左=ラベル、右=数字(触らない)
    row_h = int(H * 0.055)
    top = int(H * 0.32)
    labels = ["Wellbeing", "Teamwork and Collaboration", "Performance Management"]
    values = ["80%", "70%", "( n=5 )"]
    for i, (lab, val) in enumerate(zip(labels, values)):
        y0 = top + i * row_h
        d.rectangle([int(W * 0.03), y0, int(W * 0.47), y0 + row_h], fill=(249, 250, 249) if i % 2 == 0 else (255, 255, 255))
        put(lab, (int(W * 0.04), y0 + int(row_h * 0.2)), _font(int(H * 0.022)), (30, 30, 30), blocks)
        put(val, (int(W * 0.40), y0 + int(row_h * 0.2)), _font(int(H * 0.022)), (30, 30, 30), numbers)

    # 濃灰の帯の上の小さい見出し（中央寄せ）
    bx0, bx1 = int(W * 0.53), int(W * 0.97)
    d.rectangle([bx0, top, bx1, top + row_h], fill=(106, 106, 106))
    put("My Group", (0, top + int(row_h * 0.25)), _font(int(H * 0.02)), (20, 20, 40), blocks,
        anchor_center_x=(bx0 + bx1) / 2)

    # ロゴ風の図形(文字なし) と 固有名詞(訳さない)
    d.ellipse([int(W * 0.85), int(H * 0.15), int(W * 0.93), int(H * 0.15) + int(W * 0.08)], fill=(230, 70, 40))
    put("Nexperia", (int(W * 0.86), int(H * 0.30)), _font(int(H * 0.03)), (0, 112, 122), blocks)

    # スライドごとに違う画像にする目印（偽Geminiが送信画像のmd5でスライドを見分けるため）
    d.rectangle([W - 30, H - 30, W - 10, H - 10], fill=((marker * 37) % 256, (marker * 91) % 256, 200))

    if with_junk_left:  # トリミングで切り落とされる左端の文字
        put("JUNK", (int(W * 0.01), int(H * 0.60)), _font(int(H * 0.03)), (0, 0, 0), [])
    return im, blocks, numbers


def png_bytes(im):
    b = io.BytesIO()
    im.save(b, "PNG")
    return b.getvalue()


def expected_response(blocks, numbers, size, crop=(0, 0, 0, 0), include_numbers=True, jitter=0, lines=None):
    """正解の位置を、トリミング後の画像に対する 0-1000 の座標に直して、Gemini の応答JSONにする。"""
    W, H = size
    l, t, r, b = crop
    cx0, cy0 = l * W, t * H
    cw, ch = W * (1 - l - r), H * (1 - t - b)
    out = []
    for blk in blocks + (numbers if include_numbers else []):
        x0, y0, x1, y1 = blk["box"]
        if x1 <= cx0 or x0 >= cx0 + cw:
            continue
        box = [round((y0 - cy0) / ch * 1000) - jitter, round((x0 - cx0) / cw * 1000) - jitter,
               round((y1 - cy0) / ch * 1000) + jitter, round((x1 - cx0) / cw * 1000) + jitter]
        is_num = blk in numbers
        out.append({"box_2d": box, "text": blk["text"],
                    "translation": blk["text"] if is_num else _tr(blk["text"]),
                    "lines": (lines or {}).get(blk["text"], 1)})
    return out


def gemini_reply(blocks_json):
    return {"candidates": [{"content": {"parts": [{"text": json.dumps(blocks_json, ensure_ascii=False)}]},
                            "finishReason": "STOP"}]}


def add_picture_slide(prs, im, left=0, top=0, width=None, height=None, crop=(0, 0, 0, 0), rotation=0):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    pic = slide.shapes.add_picture(io.BytesIO(png_bytes(im)), Emu(left), Emu(top),
                                   Emu(width if width is not None else prs.slide_width),
                                   Emu(height if height is not None else prs.slide_height))
    pic.crop_left, pic.crop_top, pic.crop_right, pic.crop_bottom = crop
    if rotation:
        pic.rotation = rotation
    return slide, pic


def new_prs():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(11430000), Emu(6429375)
    return prs


def picture_md5(path, slide_idx=0):
    prs = Presentation(path)
    for sh in prs.slides[slide_idx].shapes:
        if sh.shape_type == 13:
            return hashlib.md5(sh.image.blob).hexdigest(), (sh.left, sh.top, sh.width, sh.height)
    return None, None


def notes_text(path, slide_idx):
    prs = Presentation(path)
    s = prs.slides[slide_idx]
    return s.notes_slide.notes_text_frame.text if s.has_notes_slide else ""


def build_responder(deck_specs):
    """スライドの画像(送信JPEG)のmd5から、そのスライドの応答を引く偽Geminiを作る。"""
    table = {}
    for spec in deck_specs:
        jpeg, _ = mod._prepare_slide_image(spec["blob"], spec["crop"])
        table[hashlib.md5(jpeg).hexdigest()] = spec["reply"]

    def responder(payload):
        parts = payload["contents"][0]["parts"]
        inline = next(p["inlineData"] for p in parts if "inlineData" in p)
        import base64
        key = hashlib.md5(base64.b64decode(inline["data"])).hexdigest()
        reply = table[key]
        if callable(reply):
            return reply()
        return gemini_reply(reply)
    return responder


def run_flow(path, spec_value="", language="Japanese", m=None):
    """translate_ppt_document_thread を、確認ダイアログの答えを固定して実行する。"""
    m = m or mod
    MESSAGEBOX.CALLS.clear()
    captured = {}

    def fake_ask(pw, message, validate):
        captured["message"] = message
        captured["validate"] = validate
        return spec_value
    m._ask_image_confirmation = fake_ask
    pw = FakeProgress()
    CAPTURED["calls"].clear()
    m.translate_ppt_document_thread(path, language, pw)
    return pw, captured


# ============================================================
# 1. 部品の単体検証
# ============================================================
# --- スライド指定 ---
check("parse_slide_spec: 空欄は None(全部)", mod.parse_slide_spec("  ", 24) is None)
check("parse_slide_spec: 単一", mod.parse_slide_spec("6", 24) == {5})
check("parse_slide_spec: 範囲と複数", mod.parse_slide_spec("1-3,5", 24) == {0, 1, 2, 4})
for bad, label in (("25", "範囲外"), ("0", "0番"), ("3-1", "逆順"), ("a", "文字"), ("1-2-3", "形式不正"), (",", "空要素のみ")):
    try:
        mod.parse_slide_spec(bad, 24)
        check(f"parse_slide_spec: {label}を弾く", False, "例外が出なかった")
    except ValueError:
        check(f"parse_slide_spec: {label}を弾く", True)

# --- 数字・記号だけのブロックを除く専用判定 ---
for t in ("80%", "60.0%", "( n=5 )", "( n=6,797 )", "65.8% (-15.8)", "1", "•", "12"):
    check(f"_has_letters: {t!r} は文字なし(除外)", mod._has_letters(t) is False)
for t in ("Favorability %", "Nexperia", "Top 5 Scoring Questions", "日本語のテキスト", "My Group"):
    check(f"_has_letters: {t!r} は文字あり(対象)", mod._has_letters(t) is True)
check("(参考)既存の is_translatable は '80%' を True にする=専用判定が必要な根拠",
      mod.is_translatable("80%") is True)

# --- 応答JSONの検証 ---
ok = mod._validate_image_blocks(json.dumps([
    {"box_2d": [10, 20, 30, 400], "text": "Hello", "translation": "こんにちは", "lines": 2}]))
check("_validate_image_blocks: 正常な応答", len(ok) == 1 and ok[0]["lines"] == 2 and ok[0]["box_2d"] == [10.0, 20.0, 30.0, 400.0])
check("_validate_image_blocks: 空配列は「文字なし」で例外にしない", mod._validate_image_blocks("[]") == [])
for label, raw in (("壊れたJSON", "[{"), ("配列でない", '{"a":1}'), ("文字列", '"x"')):
    try:
        mod._validate_image_blocks(raw)
        check(f"_validate_image_blocks: {label}は例外", False)
    except ValueError:
        check(f"_validate_image_blocks: {label}は例外", True)
zero = json.dumps([{"box_2d": [0, 0, 0, 0], "text": "A", "translation": "あ", "lines": 1}] * 3)
try:
    mod._validate_image_blocks(zero)
    check("_validate_image_blocks: 全ブロックの位置が0なら例外(文字なしと誤認しない)", False)
except ValueError:
    check("_validate_image_blocks: 全ブロックの位置が0なら例外(文字なしと誤認しない)", True)
mixed = mod._validate_image_blocks(json.dumps([
    {"box_2d": [0, 0, 0, 0], "text": "bad", "translation": "x", "lines": 1},
    {"box_2d": [10, 10, 50, 500], "text": "good", "translation": "良い", "lines": 1},
    {"box_2d": [10, 10], "text": "short", "translation": "x", "lines": 1},
    {"box_2d": [10, 10, 50, 500], "text": "", "translation": "x", "lines": 1}]))
check("_validate_image_blocks: 不正なブロックだけ捨てて残りを使う", [b["text"] for b in mixed] == ["good"])
clamped = mod._validate_image_blocks(json.dumps([
    {"box_2d": [-5, -5, 1200, 1300], "text": "edge", "translation": "端", "lines": 99}]))
check("_validate_image_blocks: 範囲外の座標は0-1000へ丸め、行数の異常値は1にする",
      clamped[0]["box_2d"] == [0.0, 0.0, 1000.0, 1000.0] and clamped[0]["lines"] == 1)

# --- 画像の準備 ---
big, _, _ = make_slide_image((3200, 1800))
jpeg, prepared = mod._prepare_slide_image(png_bytes(big), (0, 0, 0, 0))
check("_prepare_slide_image: 幅1600pxへ縮小しJPEGにする",
      prepared.width == 1600 and prepared.height == 900 and jpeg[:2] == b"\xff\xd8",
      f"size={prepared.size}, bytes={len(jpeg)}")
small_im, _, _ = make_slide_image((1000, 600))
_, prepared_small = mod._prepare_slide_image(png_bytes(small_im), (0, 0, 0, 0))
check("_prepare_slide_image: 小さい画像は拡大しない", prepared_small.size == (1000, 600))
_, prepared_crop = mod._prepare_slide_image(png_bytes(small_im), (0.1, 0.0, 0.2, 0.5))
check("_prepare_slide_image: トリミングを反映する(幅70%・高さ50%)", prepared_crop.size == (700, 300),
      f"size={prepared_crop.size}")
rgba = Image.new("RGBA", (200, 100), (0, 0, 0, 0))
_, prepared_alpha = mod._prepare_slide_image(png_bytes(rgba), (0, 0, 0, 0))
check("_prepare_slide_image: 透明部分は黒でなく白にする", prepared_alpha.getpixel((10, 10)) == (255, 255, 255))

# ============================================================
# 2. シム拡張(画像の受け口)
# ============================================================
prev = _load("target_prev", PREV_TARGET)
seen = {}


def _capture_into(name):
    def f(payload, model=None):
        seen[name] = json.dumps({"payload": payload, "model": model}, sort_keys=True, ensure_ascii=False)
        return {"candidates": [{"content": {"parts": [{"text": "x"}]}}]}
    return f


mod._generate_advanced = _capture_into("new")
prev._generate_advanced = _capture_into("old")
safety = [{"category": "HARM_CATEGORY_HARASSMENT", "threshold": "BLOCK_NONE"}]
same = True
for kwargs in (dict(temperature=0.1, safety_settings=safety),
               dict(temperature=0.1, system_instruction="s", response_mime_type="application/json",
                    response_schema={"type": "ARRAY"}),
               dict()):
    prev._CommonGeminiClient().models.generate_content(model="m", contents="hello", config=prev._GeminiGenerateConfig(**kwargs))
    mod._CommonGeminiClient().models.generate_content(model="m", contents="hello", config=mod._GeminiGenerateConfig(**kwargs))
    same &= seen["old"] == seen["new"]
check("シム: contents が str のとき payload が直前版と完全一致(従来の翻訳に影響しない)", same)
prev._CommonGeminiClient().models.generate_content(model="m", contents="hello", config=None)
mod._CommonGeminiClient().models.generate_content(model="m", contents="hello", config=None)
check("シム: config が None でも従来と完全一致", seen["old"] == seen["new"])

mod._CommonGeminiClient().models.generate_content(
    model="gemini-2.5-flash",
    contents=[{"inlineData": {"mimeType": "image/jpeg", "data": "AAAA"}}, "translate"],
    config=mod._GeminiGenerateConfig(temperature=0, max_output_tokens=123, thinking_budget=0,
                                     response_mime_type="application/json"))
p = json.loads(seen["new"])["payload"]
check("シム: contents が list のとき画像パーツとテキストパーツになる",
      p["contents"][0]["parts"] == [{"inlineData": {"mimeType": "image/jpeg", "data": "AAAA"}}, {"text": "translate"}])
check("シム: maxOutputTokens が generationConfig に載る", p["generationConfig"]["maxOutputTokens"] == 123)
check("シム: thinkingBudget=0 が thinkingConfig に載る(0も落とさない)",
      p["generationConfig"]["thinkingConfig"] == {"thinkingBudget": 0})
mod._CommonGeminiClient().models.generate_content(model="m", contents="x", config=mod._GeminiGenerateConfig(temperature=0.1))
p2 = json.loads(seen["new"])["payload"]
check("シム: 未指定なら maxOutputTokens / thinkingConfig を載せない",
      "maxOutputTokens" not in p2["generationConfig"] and "thinkingConfig" not in p2["generationConfig"])
mod._generate_advanced = _fake_generate_advanced

# ============================================================
# 3. 対象スライドの判定
# ============================================================
prs = new_prs()
im1, blocks1, nums1 = make_slide_image()
add_picture_slide(prs, im1)                                                     # 0: 対象
s1, _ = add_picture_slide(prs, im1)                                             # 1: 対象(ページ番号だけのテキスト付き)
tb = s1.shapes.add_textbox(Emu(100), Emu(100), Emu(500000), Emu(300000)); tb.text_frame.text = "12"
add_picture_slide(prs, im1, width=int(prs.slide_width * 0.6), height=int(prs.slide_height * 0.6))   # 2: 70%未満
s3, _ = add_picture_slide(prs, im1)                                             # 3: 大きい画像が2枚
s3.shapes.add_picture(io.BytesIO(png_bytes(im1)), Emu(0), Emu(0), prs.slide_width, prs.slide_height)
add_picture_slide(prs, im1, rotation=15)                                        # 4: 回転
s5, _ = add_picture_slide(prs, im1)                                             # 5: 文字あり(英文)
tb = s5.shapes.add_textbox(Emu(100), Emu(100), Emu(5000000), Emu(300000)); tb.text_frame.text = "This slide has real text"
add_picture_slide(prs, im1, left=-int(prs.slide_width * 0.05), width=int(prs.slide_width * 1.1))    # 6: はみ出しても70%以上
found = mod.find_image_slides(prs)
check("find_image_slides: 対象は 0,1,6 の3枚", [i for i, _ in found] == [0, 1, 6], f"{[i for i, _ in found]}")
check("find_image_slides: ページ番号だけの文字は「文字なし」扱い(スライド2が対象)", 1 in [i for i, _ in found])
check("find_image_slides: 70%未満・画像2枚・回転・文字ありは対象外", not ({2, 3, 4, 5} & {i for i, _ in found}))

# ============================================================
# 4. エンドツーエンド(ノート追記)
# ============================================================
def build_deck(path, n_slides=3, with_notes_on=None, crop_slide=None):
    """n_slides枚の画像スライドの資料を作り、偽Geminiの表を返す。"""
    prs = new_prs()
    specs, truth = [], []
    for i in range(n_slides):
        crop = (0.1, 0.0, 0.0, 0.0) if i == crop_slide else (0, 0, 0, 0)
        im, blocks, nums = make_slide_image(with_junk_left=(i == crop_slide), marker=i + 1)
        slide, pic = add_picture_slide(prs, im, crop=crop)
        if with_notes_on is not None and i == with_notes_on:
            slide.notes_slide.notes_text_frame.text = "備考"   # 2文字=翻訳対象にならない短いメモ
        reply = expected_response(blocks, nums, im.size, crop)
        specs.append({"blob": png_bytes(im), "crop": crop, "reply": reply})
        truth.append((blocks, nums, im.size, crop))
    prs.save(path)
    return specs, truth


tmp = tempfile.mkdtemp()
cwd0 = os.getcwd()
os.chdir(tmp)
try:
    src = os.path.join(tmp, "deck.pptx")
    specs, truth = build_deck(src, 3, with_notes_on=1)
    RESPONDER["fn"] = build_responder(specs)
    orig_md5 = [picture_md5(src, i) for i in range(3)]
    orig_bytes = open(src, "rb").read()

    pw, cap = run_flow(src, "")
    out = os.path.join(tmp, "deck_ja.pptx")
    check("出力ファイル名が 元_ja.pptx", os.path.isfile(out))
    check("元のPPTXは変更されない", open(src, "rb").read() == orig_bytes)
    check("エラーダイアログが出ていない", not any(c[0] == "error" for c in MESSAGEBOX.CALLS), f"{MESSAGEBOX.CALLS}")
    check("確認ダイアログの文言に「画像そのものを外部AIへ送信」が明記される",
          "画像そのもの" in cap["message"] and "外部AI" in cap["message"] and "Gemini" in cap["message"])
    check("確認ダイアログの文言に「機密・個人情報」の注意がある", "機密" in cap["message"])
    check("確認ダイアログの文言に対象枚数(3枚)が出る", "3 枚" in cap["message"])
    check("Gemini呼び出しはスライド数と同じ3回(1枚1回)", len(CAPTURED["calls"]) == 3, f"{len(CAPTURED['calls'])}")

    call = CAPTURED["calls"][0]
    pl = call["payload"]
    parts = pl["contents"][0]["parts"]
    check("payload: model が明示的に渡る", call["model"] == "gemini-2.5-flash", f"{call['model']!r}")
    check("payload: 画像は image/jpeg のinlineData", parts[0]["inlineData"]["mimeType"] == "image/jpeg")
    check("payload: 画像のあとにプロンプトのテキストが続く", "text" in parts[1] and "translate" in parts[1]["text"].lower())
    check("payload: 翻訳先言語がプロンプトに入る", "Japanese" in parts[1]["text"])
    gc = pl["generationConfig"]
    check("payload: JSONモードとスキーマが載る(項目順 box_2d→text→translation→lines)",
          gc["responseMimeType"] == "application/json"
          and gc["responseSchema"]["items"]["propertyOrdering"] == ["box_2d", "text", "translation", "lines"])
    check("payload: maxOutputTokens と thinkingBudget=0(2.5-flash)", gc["maxOutputTokens"] > 0
          and gc["thinkingConfig"] == {"thinkingBudget": 0})
    check("payload: safetySettings が4カテゴリ載る", len(pl["safetySettings"]) == 4)
    check("payload: json.dumps 可能で、送る画像は約1MB未満(base64)",
          len(json.dumps(pl)) < 1_000_000, f"{len(json.dumps(pl))} bytes")
    check("プロンプトに「数字だけは返さない・ロゴ除外・固有名詞は訳さない」が入る",
          all(k in parts[1]["text"] for k in ("only numbers", "logos", "proper nouns")))

    n0 = notes_text(out, 0)
    check("スライド1: メモ欄に見出しと訳文が入る",
          "【画像内テキストの翻訳】" in n0 and "エグゼクティブサマリー" in n0 and "ウェルビーイング" in n0)
    check("メモ欄に原文も併記される", "（原文: Executive Summary）" in n0)
    check("数字だけのブロック(80%)はメモ欄に入らない", "80%" not in n0 and "n=5" not in n0)
    check("訳文が原文と同じ固有名詞(Nexperia)はメモ欄に入らない", "Nexperia" not in n0)
    n1 = notes_text(out, 1)
    check("既存のメモは消えず、区切り線のあとに追記される",
          n1.startswith("備考") and "----------" in n1 and "【画像内テキストの翻訳】" in n1)
    check("元の画像は3枚とも位置・バイト列が完全に同じ",
          [picture_md5(out, i) for i in range(3)] == orig_md5)
    check("スライド枚数は変わらない", len(Presentation(out).slides) == 3)
    check("完了ダイアログ: 翻訳したスライド3枚と「AIによる翻訳です」の注意",
          any("翻訳したスライド: 3枚" in c[2] and "確認してください" in c[2] for c in MESSAGEBOX.CALLS if c[0] == "info"),
          f"{MESSAGEBOX.CALLS}")

    try:
        cap["validate"]("99")
        check("確認ダイアログの入力検証: 範囲外を弾く", False)
    except ValueError:
        check("確認ダイアログの入力検証: 範囲外を弾く", True)

    # --- ログに原文を書かない ---
    handlers = logging.getLogger("PPT_Translation").handlers
    log_path = next(h.baseFilename for h in handlers if hasattr(h, "baseFilename"))
    log_text = open(log_path, encoding="utf-8").read()
    secrets = ["Executive Summary", "Employee Engagement", "Wellbeing", "Performance Management", "エグゼクティブ"]
    check("ログに原文・訳文を書かない(件数とスライド番号だけ)",
          not any(s in log_text for s in secrets) and "スライド 1 成功" in log_text,
          "漏れ: " + str([s for s in secrets if s in log_text]))

    # --- 範囲指定: 2枚目だけ ---
    if os.path.exists(out):
        os.remove(out)
    pw, cap = run_flow(src, "2")
    check("範囲指定「2」: Gemini呼び出しは1回だけ", len(CAPTURED["calls"]) == 1, f"{len(CAPTURED['calls'])}")
    check("範囲指定「2」: 2枚目だけメモが付き、1枚目・3枚目には付かない",
          "【画像内テキストの翻訳】" in notes_text(out, 1)
          and "【画像内テキストの翻訳】" not in notes_text(out, 0)
          and "【画像内テキストの翻訳】" not in notes_text(out, 2))

    # --- キャンセル ---
    if os.path.exists(out):
        os.remove(out)
    pw, cap = run_flow(src, None)
    check("キャンセル: Geminiを呼ばない", len(CAPTURED["calls"]) == 0)
    check("キャンセル: 未翻訳のコピー(_ja.pptx)を残さない", not os.path.exists(out))
    check("キャンセル: 元ファイルは無傷", open(src, "rb").read() == orig_bytes)
    check("キャンセル: 進捗ウィンドウを閉じる", pw.closed >= 1)

    # --- 翻訳済みファイルを再度読み込んだ場合 ---
    # メモ欄に訳文(文字)が入っているため「文字あり」の資料として従来経路になる。
    # 画像は送らず、確認ダイアログも出ない。画像の訳が二重に書かれることはない。
    run_flow(src, "")
    twice_src = os.path.join(tmp, "twice.pptx")
    shutil.copy(out, twice_src)
    pw, cap = run_flow(twice_src, "")
    check("翻訳済みのファイルを再度読み込んでも、画像は送らず確認ダイアログも出ない(従来経路)",
          "message" not in cap
          and all("inlineData" not in json.dumps(c["payload"]) for c in CAPTURED["calls"]),
          f"{len(CAPTURED['calls'])}回")
    check("翻訳済みのファイルを再度読み込んでも、画像の訳が二重に書かれない",
          notes_text(os.path.join(tmp, "twice_ja.pptx"), 0).count("【画像内テキストの翻訳】") <= 1)
    dup_prs = Presentation(out)
    check("メモ追記の二重防止: 見出しが既にあれば追記しない",
          mod._append_image_notes(dup_prs.slides[0], mod._validate_image_blocks(json.dumps(specs[0]["reply"]))) == 0)

    # --- トリミングされた画像 ---
    crop_src = os.path.join(tmp, "crop.pptx")
    cspecs, ctruth = build_deck(crop_src, 2, crop_slide=1)
    RESPONDER["fn"] = build_responder(cspecs)
    run_flow(crop_src, "")
    crop_out = os.path.join(tmp, "crop_ja.pptx")
    n_crop = notes_text(crop_out, 1)
    check("トリミングされた画像: 切り落とされた左端の文字(JUNK)はメモに入らず、見えている文字は入る",
          "JUNK" not in n_crop and "エグゼクティブサマリー" in n_crop)

    # --- 失敗の扱い ---
    fdeck = os.path.join(tmp, "fail.pptx")
    fspecs, _ = build_deck(fdeck, 6)
    calls_by_slide = {"n": 0}

    def make_fail_from(k):
        results = []
        for i, spec in enumerate(fspecs):
            if i < k:
                results.append(spec["reply"])
            else:
                results.append(lambda: (_ for _ in ()).throw(RuntimeError("proxy down")))
        for spec, r in zip(fspecs, results):
            spec["reply"] = r
    orig_replies = [s["reply"] for s in fspecs]
    make_fail_from(2)
    RESPONDER["fn"] = build_responder(fspecs)
    mod.IMAGE_WORKERS = 1
    run_flow(fdeck, "")
    fout = os.path.join(tmp, "fail_ja.pptx")
    info = next((c[2] for c in MESSAGEBOX.CALLS if c[0] == "info"), "")
    check("3枚連続失敗: 中断して、それまでに終わったスライド(1,2枚目)は保存する",
          os.path.exists(fout) and "【画像内テキストの翻訳】" in notes_text(fout, 0)
          and "【画像内テキストの翻訳】" in notes_text(fout, 1), info)
    check("3枚連続失敗: 失敗した3枚の番号(3,4,5)を結果に出し、中断したと知らせる",
          "3, 4, 5" in info and "中断" in info, info)
    check("3枚連続失敗: 4枚目以降はもう呼ばない(3枚目までで各3回リトライ=計2+9回)",
          len(CAPTURED["calls"]) == 2 + 9, f"{len(CAPTURED['calls'])}")
    check("失敗したスライドは無加工のまま", "【画像内テキストの翻訳】" not in notes_text(fout, 2))
    check("失敗の案内に「範囲に番号を入力して再試行」が書かれる", "再試行" in info)
    mod.IMAGE_WORKERS = 2

    # 途中だけ失敗(連続ではない) → 中断せず続行
    for spec, r in zip(fspecs, orig_replies):
        spec["reply"] = r
    fspecs[2]["reply"] = lambda: (_ for _ in ()).throw(RuntimeError("one bad slide"))
    RESPONDER["fn"] = build_responder(fspecs)
    mod.IMAGE_WORKERS = 1
    run_flow(fdeck, "")
    info = next((c[2] for c in MESSAGEBOX.CALLS if c[0] == "info"), "")
    check("1枚だけ失敗: 中断せず残りを続け、失敗した番号(3)を結果に出す",
          "番号: 3）" in info and "中断" not in info
          and "【画像内テキストの翻訳】" in notes_text(fout, 5), info)
    mod.IMAGE_WORKERS = 2

    # --- 応答の異常 ---
    sdeck = os.path.join(tmp, "one.pptx")
    sspecs, _ = build_deck(sdeck, 1)

    def run_single(reply):
        sspecs[0]["reply"] = reply
        RESPONDER["fn"] = build_responder(sspecs)
        CAPTURED["calls"].clear()
        run_flow(sdeck, "")
        return next((c[2] for c in MESSAGEBOX.CALLS if c[0] == "info"), "")

    info = run_single([])
    check("応答が空配列: 「文字が見つからなかった」と通知し、失敗扱いにしない",
          "文字が見つからなかったスライド: 1枚" in info and "失敗" not in info, info)
    check("応答が空配列: メモ欄は作らない", "【画像内テキストの翻訳】" not in notes_text(os.path.join(tmp, "one_ja.pptx"), 0))

    info = run_single([{"box_2d": [0, 0, 0, 0], "text": "A", "translation": "あ", "lines": 1}] * 3)
    check("全ブロックの位置が0: 3回リトライして失敗として通知", len(CAPTURED["calls"]) == 3 and "失敗したスライド: 1枚" in info, info)

    seq = {"n": 0}

    def flaky():
        seq["n"] += 1
        if seq["n"] == 1:
            return {"candidates": [{"content": {"parts": [{"text": "[{\"box"}]}, "finishReason": "MAX_TOKENS"}]}
        return gemini_reply(expected_response(*truth[0][:2], truth[0][2], (0, 0, 0, 0)))
    info = run_single(flaky)
    check("途中で切れたJSON(MAX_TOKENS)は再試行して成功する", seq["n"] == 2 and "翻訳したスライド: 1枚" in info, info)

    info = run_single(lambda: {"candidates": [{"content": {"parts": []}}]})
    check("空応答(セーフティ等): 3回リトライして失敗として通知", len(CAPTURED["calls"]) == 3 and "失敗したスライド: 1枚" in info, info)

    # --- 混在資料・画像だけでない資料 ---
    mixed = new_prs()
    ms, _ = add_picture_slide(mixed, im1)
    tb = ms.shapes.add_textbox(Emu(100), Emu(100), Emu(5000000), Emu(300000)); tb.text_frame.text = "Quarterly results overview"
    add_picture_slide(mixed, im1)
    mixed_path = os.path.join(tmp, "mixed.pptx")
    mixed.save(mixed_path)

    def text_responder(payload):
        text = payload["contents"][0]["parts"][0]["text"]
        n = text.count("[")
        return {"candidates": [{"content": {"parts": [{"text": "\n".join(f"[{i}] 訳文{i}" for i in range(1, n + 1))}]}}]}
    RESPONDER["fn"] = text_responder
    pw, cap = run_flow(mixed_path, "")
    info = next((c[2] for c in MESSAGEBOX.CALLS if c[0] == "info"), "")
    check("混在資料: 従来の翻訳が動き、確認ダイアログ(画像送信)は出ない", "message" not in cap and "翻訳項目数: 1" in info, info)
    check("混在資料: 完了メッセージに「画像だけのスライドが1枚あります(未翻訳)」と出る",
          "画像だけのスライドが 1 枚" in info and "未翻訳" in info, info)
    check("混在資料: 画像はGeminiへ送らない(文字のテキストだけ)",
          all("inlineData" not in json.dumps(c["payload"]) for c in CAPTURED["calls"]))
    check("混在資料: 従来どおり文字が書き戻される",
          Presentation(os.path.join(tmp, "mixed_ja.pptx")).slides[0].shapes[-1].text_frame.text == "訳文1")

    textonly = new_prs()
    ts = textonly.slides.add_slide(textonly.slide_layouts[6])
    tb = ts.shapes.add_textbox(Emu(100), Emu(100), Emu(5000000), Emu(300000)); tb.text_frame.text = "Only text here"
    tpath = os.path.join(tmp, "textonly.pptx"); textonly.save(tpath)
    pw, cap = run_flow(tpath, "")
    info = next((c[2] for c in MESSAGEBOX.CALLS if c[0] == "info"), "")
    check("文字だけの資料: 完了メッセージに画像の注意書きが出ない(従来と同じ)", "画像だけ" not in info and "翻訳項目数: 1" in info, info)

    empty = new_prs(); empty.slides.add_slide(empty.slide_layouts[6])
    epath = os.path.join(tmp, "empty.pptx"); empty.save(epath)
    pw, cap = run_flow(epath, "")
    info = next((c[2] for c in MESSAGEBOX.CALLS if c[0] == "info"), "")
    check("文字も画像も無い資料: 従来のメッセージのまま(余計な案内を出さない)",
          info == "翻訳対象のテキストが見つかりませんでした。", info)

    part = new_prs(); add_picture_slide(part, im1, width=int(part.slide_width * 0.5), height=int(part.slide_height * 0.5))
    ppath = os.path.join(tmp, "partial.pptx"); part.save(ppath)
    pw, cap = run_flow(ppath, "")
    info = next((c[2] for c in MESSAGEBOX.CALLS if c[0] == "info"), "")
    check("画像が全面でない資料: 「対象は画面の7割以上を覆う画像」と案内して終了し、Geminiを呼ばない",
          "7割以上" in info and len(CAPTURED["calls"]) == 0, info)

    # --- 画像スライドのメモに「翻訳対象になる文章」がある資料 ---
    notes_deck = new_prs()
    ns_, _ = add_picture_slide(notes_deck, im1)
    ns_.notes_slide.notes_text_frame.text = "Please mention the survey response rate"
    add_picture_slide(notes_deck, im1)
    notes_path = os.path.join(tmp, "notesdeck.pptx"); notes_deck.save(notes_path)
    RESPONDER["fn"] = text_responder
    pw, cap = run_flow(notes_path, "")
    info = next((c[2] for c in MESSAGEBOX.CALLS if c[0] == "info"), "")
    check("メモに文章がある画像スライド: 従来経路でメモを翻訳し、画像は送らない(確認ダイアログも出ない)",
          "message" not in cap and "翻訳項目数: 1" in info
          and all("inlineData" not in json.dumps(c["payload"]) for c in CAPTURED["calls"]), info)
    check("メモに文章がある画像スライド: 「画像だけのスライドが2枚あります(未翻訳)」と知らせる",
          "画像だけのスライドが 2 枚" in info and "未翻訳" in info, info)

    # --- 文字ありの経路が直前版と同一挙動 ---
    RESPONDER["fn"] = text_responder
    tdeck = new_prs()
    tsl = tdeck.slides.add_slide(tdeck.slide_layouts[6])
    tb = tsl.shapes.add_textbox(Emu(100), Emu(100), Emu(5000000), Emu(900000))
    tb.text_frame.text = "First paragraph of text"
    r = tb.text_frame.paragraphs[0].runs[0]; r.font.size = Pt(18); r.font.bold = True
    tsl.notes_slide.notes_text_frame.text = "Speaker notes are here"
    tdpath = os.path.join(tmp, "textdeck.pptx"); tdeck.save(tdpath)

    def text_signature(path):
        prs_ = Presentation(path)
        out_ = []
        for s in prs_.slides:
            for sh in s.shapes:
                if sh.has_text_frame:
                    for para in sh.text_frame.paragraphs:
                        for run in para.runs:
                            out_.append((run.text, run.font.size.pt if run.font.size else None, run.font.bold, run.font.name))
            if s.has_notes_slide:
                out_.append(("notes", s.notes_slide.notes_text_frame.text))
        return out_
    prev._ask_image_confirmation = lambda *a: ""
    prev.gemini_client = prev._CommonGeminiClient()
    prev._generate_advanced = _fake_generate_advanced
    shutil.copy(tdpath, os.path.join(tmp, "textdeck_prev.pptx"))
    prev.translate_ppt_document_thread(os.path.join(tmp, "textdeck_prev.pptx"), "Japanese", FakeProgress())
    mod.translate_ppt_document_thread(tdpath, "Japanese", FakeProgress())
    check("文字ありの資料: 直前版(20260911_01)と出力(文字・書式・ノート・フォント)が完全一致",
          text_signature(os.path.join(tmp, "textdeck_prev_ja.pptx")) == text_signature(os.path.join(tmp, "textdeck_ja.pptx")),
          f"{text_signature(os.path.join(tmp, 'textdeck_ja.pptx'))}")
finally:
    os.chdir(cwd0)
    shutil.rmtree(tmp, ignore_errors=True)

# ============================================================
# 5. 変更範囲(AST): 文字ありの翻訳経路が未変更であること
# ============================================================
import ast


def qualified_hashes(path):
    src = open(path, encoding="utf-8").read()
    lines = src.splitlines()
    out = {}

    def walk(node, prefix=""):
        for n in ast.iter_child_nodes(node):
            if isinstance(n, (ast.FunctionDef, ast.AsyncFunctionDef, ast.ClassDef)):
                q = prefix + n.name
                out[q] = hashlib.md5("\n".join(lines[n.lineno - 1:n.end_lineno]).encode()).hexdigest()
                if isinstance(n, ast.ClassDef):
                    walk(n, q + ".")
    walk(ast.parse(src))
    return out


a, b = qualified_hashes(PREV_TARGET), qualified_hashes(TARGET)
changed = {k for k in a if k in b and a[k] != b[k]}
allowed = {"_CommonGeminiModels", "_CommonGeminiModels.generate_content", "_GeminiGenerateConfig",
           "_GeminiGenerateConfig.__init__", "translate_ppt_document_thread"}
check("AST: 直前版から変わった既存の定義は、シム(画像の受け口)と分岐点だけ", changed <= allowed, f"changed={sorted(changed)}")
for keep in ("translate_batch_gemini", "translate_super_fast_parallel", "is_translatable", "lang_to_suffix",
             "init_gemini", "check_dependencies", "select_file", "WordProgressWindow"):
    check(f"AST: {keep} は直前版と完全一致", keep in a and a[keep] == b.get(keep))

# ============================================================
# 6. 重ね貼り(_02 のみ)
# ============================================================
if HAS_OVERLAY:
    exec(open(os.path.join(HERE, "test_ppt_translation_image_overlay_20260930.py"), encoding="utf-8").read())
else:
    skip("重ね貼りのテスト", "このバージョン(_01)には重ね貼りが無い")

print("\n" + "=" * 60)
failed = [r for r in RESULTS if not r[1]]
print(f"対象: {os.path.basename(TARGET)}")
print(f"合計 {len(RESULTS)} 項目 / 合格 {len(RESULTS) - len(failed)} / 失敗 {len(failed)}")
if failed:
    for name, _ok, detail in failed:
        print(f"  FAILED: {name}  {detail}")
sys.exit(1 if failed else 0)
